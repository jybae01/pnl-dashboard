from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from forecast.bff.pnl_reporting_template import (
    PNL_REPORTING_TEMPLATE_FILENAME,
    PNL_REPORTING_TEMPLATE_RESOURCE,
)
from forecast.reporting import (
    DatasetType,
    ReportingState,
    ValidationErrorCode,
    build_pnl_reporting_read_model,
    build_reporting_state,
    parse_pnl_reporting_workbook,
)
from forecast.reporting.registry import (
    MANUFACTURING_COGS_ROWS,
    MONTHLY_PNL_ROWS,
    PRODUCT_PNL_ROWS,
    SGA_ROWS,
    SHEET_MANUFACTURING_COGS,
    SHEET_MONTHLY_PNL,
    SHEET_PRODUCT_PNL,
    SHEET_SGA,
)


_AUTHORITATIVE_INPUT_SERIES = 52
_AUTHORITATIVE_PLAN_INPUTS = 624
_METADATA_MISMATCH_CODES = (
    ValidationErrorCode.LABEL_MISMATCH,
    ValidationErrorCode.UNIT_MISMATCH,
    ValidationErrorCode.CATEGORY_MISMATCH,
    ValidationErrorCode.PRODUCT_GROUP_MISMATCH,
)
_REGULAR_INPUT_LAYOUTS = (
    (SHEET_MONTHLY_PNL, MONTHLY_PNL_ROWS, 15, 3),
    (SHEET_MANUFACTURING_COGS, MANUFACTURING_COGS_ROWS, 15, 3),
    (SHEET_SGA, SGA_ROWS, 16, 4),
)


def _authoritative_input_coordinates(workbook):
    coordinates = []
    for sheet_name, definitions, key_column, month_start_column in _REGULAR_INPUT_LAYOUTS:
        worksheet = workbook[sheet_name]
        rows_by_key = {
            worksheet.cell(row, key_column).value: row
            for row in range(2, worksheet.max_row + 1)
        }
        for definition in definitions:
            if not definition.is_input:
                continue
            row = rows_by_key[definition.key]
            coordinates.extend(
                (sheet_name, row, month_start_column + month - 1, month)
                for month in range(1, 13)
            )

    worksheet = workbook[SHEET_PRODUCT_PNL]
    rows_by_key = {
        (worksheet.cell(row, 17).value, worksheet.cell(row, 18).value): row
        for row in range(2, worksheet.max_row + 1)
    }
    for definition in PRODUCT_PNL_ROWS:
        if not definition.is_input:
            continue
        row = rows_by_key[(definition.product_group_key, definition.key)]
        coordinates.extend(
            (SHEET_PRODUCT_PNL, row, 5 + month - 1, month)
            for month in range(1, 13)
        )

    assert len(coordinates) == _AUTHORITATIVE_PLAN_INPUTS
    assert len(set(coordinates)) == _AUTHORITATIVE_PLAN_INPUTS
    return tuple(coordinates)


def _cell_value_snapshot(workbook):
    return {
        (worksheet.title, cell.coordinate): cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    }


def _populated_production_template(
    dataset_type: DatasetType,
    *,
    actual_through_month: int | None = None,
    zero_values: bool = False,
) -> tuple[bytes, int]:
    workbook = load_workbook(
        BytesIO(PNL_REPORTING_TEMPLATE_RESOURCE.read_bytes()),
        read_only=False,
        data_only=False,
        keep_links=False,
    )
    try:
        coordinates = _authoritative_input_coordinates(workbook)
        assert all(
            workbook[sheet_name].cell(row, column).value is None
            for sheet_name, row, column, _month in coordinates
        )
        before = _cell_value_snapshot(workbook)
        populated = []
        for index, (sheet_name, row, column, month) in enumerate(coordinates):
            applicable = (
                dataset_type is DatasetType.PLAN
                or (
                    actual_through_month is not None
                    and month <= actual_through_month
                )
            )
            if not applicable:
                continue
            value = 0 if zero_values else ((index // 12) + 1) * 100 + month
            workbook[sheet_name].cell(row, column).value = value
            populated.append((sheet_name, row, column))

        after = _cell_value_snapshot(workbook)
        changed = {
            key
            for key, value in before.items()
            if after[key] != value
        }
        expected_changed = {
            (sheet_name, workbook[sheet_name].cell(row, column).coordinate)
            for sheet_name, row, column in populated
        }
        assert changed == expected_changed

        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), len(populated)
    finally:
        workbook.close()


def _parse(source: bytes, dataset_type: DatasetType, actual_through_month=None):
    return parse_pnl_reporting_workbook(
        source,
        dataset_type=dataset_type,
        reporting_year=2026,
        actual_through_month=actual_through_month,
        source_filename=PNL_REPORTING_TEMPLATE_FILENAME,
    )


def _canonical_input_series(result):
    assert result.canonical_payload is not None
    for sheet in result.canonical_payload.sheets:
        yield from sheet.rows
        for product_group in sheet.product_groups:
            yield from product_group.metrics


def _assert_parser_pass_with_no_metadata_mismatch(result):
    assert result.valid is True
    assert result.error_count == 0
    mismatch_counts = {
        code: sum(issue.error_code is code for issue in result.issues)
        for code in _METADATA_MISMATCH_CODES
    }
    assert mismatch_counts == {code: 0 for code in _METADATA_MISMATCH_CODES}


def test_bundled_production_template_plan_parses_with_exact_624_inputs():
    source, populated_count = _populated_production_template(DatasetType.PLAN)
    result = _parse(source, DatasetType.PLAN)

    assert populated_count == _AUTHORITATIVE_PLAN_INPUTS
    _assert_parser_pass_with_no_metadata_mismatch(result)
    series = tuple(_canonical_input_series(result))
    assert len(series) == _AUTHORITATIVE_INPUT_SERIES
    assert sum(len(row.values) for row in series) == _AUTHORITATIVE_PLAN_INPUTS
    assert all(value is not None for row in series for value in row.values)


@pytest.mark.parametrize("actual_through_month", [1, 6, 12])
def test_bundled_production_template_actual_matrix_preserves_future_blanks(
    actual_through_month,
):
    source, populated_count = _populated_production_template(
        DatasetType.ACTUAL,
        actual_through_month=actual_through_month,
    )
    result = _parse(source, DatasetType.ACTUAL, actual_through_month)

    assert populated_count == _AUTHORITATIVE_INPUT_SERIES * actual_through_month
    _assert_parser_pass_with_no_metadata_mismatch(result)
    series = tuple(_canonical_input_series(result))
    assert len(series) == _AUTHORITATIVE_INPUT_SERIES
    for row in series:
        assert all(value is not None for value in row.values[:actual_through_month])
        assert row.values[actual_through_month:] == (None,) * (12 - actual_through_month)


def test_bundled_production_template_numeric_zero_reaches_ready_read_model():
    plan_source, plan_populated = _populated_production_template(
        DatasetType.PLAN,
        zero_values=True,
    )
    actual_source, actual_populated = _populated_production_template(
        DatasetType.ACTUAL,
        actual_through_month=6,
        zero_values=True,
    )
    plan = _parse(plan_source, DatasetType.PLAN)
    actual = _parse(actual_source, DatasetType.ACTUAL, 6)

    assert plan_populated == _AUTHORITATIVE_PLAN_INPUTS
    assert actual_populated == _AUTHORITATIVE_INPUT_SERIES * 6
    _assert_parser_pass_with_no_metadata_mismatch(plan)
    _assert_parser_pass_with_no_metadata_mismatch(actual)
    assert all(
        value == 0
        for row in _canonical_input_series(plan)
        for value in row.values
    )
    for row in _canonical_input_series(actual):
        assert row.values[:6] == (0,) * 6
        assert row.values[6:] == (None,) * 6

    state = build_reporting_state(plan.canonical_payload, actual.canonical_payload)
    assert state.state is ReportingState.READY
    assert state.pair is not None
    report = build_pnl_reporting_read_model(state.pair)

    assert len(report.kpis) == 3
    assert len(report.monthly_trends) == 12
    assert len(report.monthly_data_rows) == 8
    assert len(report.pnl_rows) == 25
    assert len(report.cogs_rows) == 5
    assert all(len(row.months) == 12 for row in report.cogs_rows)
    assert len(report.sga_rows) == 23
    assert {segment.key: len(segment.rows) for segment in report.product_segments} == {
        "SW": 10,
        "BW": 10,
        "LC": 10,
        "FS": 10,
        "NEW_BUSINESS": 8,
    }
    assert report.monthly_trends[0].actual_operating_margin is None
    operating_margin = next(row for row in report.pnl_rows if row.key == "operating_margin")
    assert operating_margin.comparison_by_period["2026-01"].actual is None
    assert report.cogs_rows[0].months[0].revenue_share is None
