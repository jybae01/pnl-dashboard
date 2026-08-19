from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from forecast.reporting import (
    DatasetType,
    ValidationErrorCode,
    ValidationSeverity,
    parse_pnl_reporting_workbook,
)
from forecast.reporting.parser import _finite_number
from forecast.reporting.registry import (
    MANUFACTURING_COGS_HEADERS,
    MANUFACTURING_COGS_ROWS,
    MONTHLY_PNL_HEADERS,
    MONTHLY_PNL_ROWS,
    PRODUCT_GROUPS,
    PRODUCT_PNL_HEADERS,
    PRODUCT_PNL_ROWS,
    REQUIRED_SHEETS,
    SGA_HEADERS,
    SGA_ROWS,
    SHEET_MANUFACTURING_COGS,
    SHEET_MONTHLY_PNL,
    SHEET_PRODUCT_PNL,
    SHEET_SGA,
    TEMPLATE_VERSION,
    authoritative_input_count,
)


def _workbook_bytes(*, dataset_type: DatasetType = DatasetType.PLAN, actual_through: int = 6) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    monthly = workbook.create_sheet(SHEET_MONTHLY_PNL)
    monthly.append(MONTHLY_PNL_HEADERS)
    for definition in MONTHLY_PNL_ROWS:
        row = [definition.label, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.extend([definition.key, None])
        monthly.append(row)
    monthly["P2"] = TEMPLATE_VERSION

    cogs = workbook.create_sheet(SHEET_MANUFACTURING_COGS)
    cogs.append(MANUFACTURING_COGS_HEADERS)
    for definition in MANUFACTURING_COGS_ROWS:
        row = [definition.label, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.append(definition.key)
        cogs.append(row)

    sga = workbook.create_sheet(SHEET_SGA)
    sga.append(SGA_HEADERS)
    for definition in SGA_ROWS:
        row = [definition.label, definition.category, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.append(definition.key)
        sga.append(row)

    product = workbook.create_sheet(SHEET_PRODUCT_PNL)
    product.append(PRODUCT_PNL_HEADERS)
    group_by_key = {group.key: group for group in PRODUCT_GROUPS}
    for definition in PRODUCT_PNL_ROWS:
        group = group_by_key[definition.product_group_key]
        row = [group.display, group.dimension, definition.label, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.extend([definition.product_group_key, definition.key])
        product.append(row)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _month_values(is_input: bool, dataset_type: DatasetType, actual_through: int) -> list[int | str | None]:
    if not is_input:
        return ["=0"] * 12
    return [
        month * 100 if dataset_type is DatasetType.PLAN or month <= actual_through else None
        for month in range(1, 13)
    ]


def _edit(source: bytes, mutator) -> bytes:
    workbook = load_workbook(BytesIO(source), read_only=False, data_only=False)
    mutator(workbook)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _row_for_key(worksheet, key: str, key_column: int) -> int:
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row, key_column).value == key:
            return row
    raise AssertionError(f"missing test key {key}")


def _product_row(worksheet, group_key: str, row_key: str) -> int:
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row, 17).value == group_key and worksheet.cell(row, 18).value == row_key:
            return row
    raise AssertionError(f"missing test product key {group_key}/{row_key}")


def _set_cell(worksheet, row: int, column: int, value):
    worksheet.cell(row, column).value = value


def _parse(source: bytes, *, dataset_type: DatasetType = DatasetType.PLAN, actual_through=None):
    return parse_pnl_reporting_workbook(
        source,
        dataset_type=dataset_type,
        reporting_year=2026,
        actual_through_month=actual_through,
    )


def _codes(result) -> set[ValidationErrorCode]:
    return {issue.error_code for issue in result.issues}


def _canonical_input_series(result):
    assert result.canonical_payload is not None
    for sheet in result.canonical_payload.sheets:
        yield from sheet.rows
        for group in sheet.product_groups:
            yield from group.metrics


def _rewrite_zip(source: bytes, transform) -> bytes:
    destination = BytesIO()
    with zipfile.ZipFile(BytesIO(source)) as original, zipfile.ZipFile(destination, "w") as output:
        for entry in original.infolist():
            payload = original.read(entry.filename)
            payload = transform(entry.filename.replace("\\", "/"), payload)
            output.writestr(entry, payload)
    return destination.getvalue()


def test_registry_has_exact_contract_counts_and_product_safety():
    assert len(MONTHLY_PNL_ROWS) == 25
    assert sum(row.is_input for row in MONTHLY_PNL_ROWS) == 11
    assert len(MANUFACTURING_COGS_ROWS) == 5
    assert sum(row.is_input for row in MANUFACTURING_COGS_ROWS) == 4
    assert len(SGA_ROWS) == 23
    assert sum(row.is_input for row in SGA_ROWS) == 18
    assert len(PRODUCT_PNL_ROWS) == 48
    assert sum(row.is_input for row in PRODUCT_PNL_ROWS) == 19
    assert authoritative_input_count() == 52

    group_by_key = {group.key: group for group in PRODUCT_GROUPS}
    assert (group_by_key["LC"].display, group_by_key["LC"].dimension, group_by_key["LC"].quantity_unit) == ("4인치 LC", "4-inch", "PCS")
    assert (group_by_key["FS"].display, group_by_key["FS"].dimension, group_by_key["FS"].quantity_unit) == ("FS", "LENGTH", "m")
    new_metrics = {row.key for row in PRODUCT_PNL_ROWS if row.product_group_key == "NEW_BUSINESS"}
    assert "volume" not in new_metrics
    assert "asp" not in new_metrics


def test_valid_plan_requires_and_preserves_624_numeric_cells():
    result = _parse(_workbook_bytes())
    assert result.valid is True
    assert result.error_count == 0
    assert result.warning_count == 0
    series = tuple(_canonical_input_series(result))
    assert len(series) == 52
    assert sum(len(row.values) for row in series) == 624
    assert all(value is not None for row in series for value in row.values)
    assert len(result.canonical_payload.to_json()) > 0


def test_parser_accepts_a_path_source(tmp_path):
    source = tmp_path / "pnl-reporting.xlsx"
    source.write_bytes(_workbook_bytes())
    result = parse_pnl_reporting_workbook(
        source,
        dataset_type=DatasetType.PLAN,
        reporting_year=2026,
    )
    assert result.valid is True


def test_valid_actual_through_six_preserves_zero_and_future_null():
    source = _workbook_bytes(dataset_type=DatasetType.ACTUAL, actual_through=6)
    source = _edit(source, lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 7, 0))
    result = _parse(source, dataset_type=DatasetType.ACTUAL, actual_through=6)
    assert result.valid is True
    payload = result.canonical_payload.to_dict()
    revenue = payload["sheets"][SHEET_MONTHLY_PNL]["rev_product"]
    assert revenue["05"] == 0
    assert tuple(revenue[key] for key in ("07", "08", "09", "10", "11", "12")) == (None,) * 6


def test_negative_numeric_input_is_valid():
    source = _edit(
        _workbook_bytes(),
        lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(
            _row_for_key(workbook[SHEET_MONTHLY_PNL], "rev_rebate", 15), 3, -25
        ),
    )
    result = _parse(source)
    assert result.valid is True
    assert result.canonical_payload.to_dict()["sheets"][SHEET_MONTHLY_PNL]["rev_rebate"]["01"] == -25


@pytest.mark.parametrize(
    ("dataset_type", "actual_through", "mutator", "expected"),
    [
        (
            DatasetType.PLAN,
            None,
            lambda workbook: _set_cell(workbook[SHEET_MONTHLY_PNL], 3, 3, None),
            ValidationErrorCode.PLAN_MONTH_MISSING,
        ),
        (
            DatasetType.ACTUAL,
            6,
            lambda workbook: _set_cell(workbook[SHEET_MONTHLY_PNL], 3, 4, None),
            ValidationErrorCode.ACTUAL_MONTH_MISSING,
        ),
        (
            DatasetType.ACTUAL,
            6,
            lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 9, 0),
            ValidationErrorCode.ACTUAL_FUTURE_VALUE_PRESENT,
        ),
    ],
)
def test_availability_failures(dataset_type, actual_through, mutator, expected):
    source = _workbook_bytes(dataset_type=dataset_type, actual_through=actual_through or 6)
    result = _parse(_edit(source, mutator), dataset_type=dataset_type, actual_through=actual_through)
    assert result.valid is False
    assert result.canonical_payload is None
    assert expected in _codes(result)


def test_actual_future_formula_reports_type_and_availability_violations():
    source = _workbook_bytes(dataset_type=DatasetType.ACTUAL, actual_through=6)
    source = _edit(
        source,
        lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 9, "=1+1"),
    )

    result = _parse(source, dataset_type=DatasetType.ACTUAL, actual_through=6)

    assert result.valid is False
    assert {
        ValidationErrorCode.INVALID_VALUE,
        ValidationErrorCode.ACTUAL_FUTURE_VALUE_PRESENT,
    }.issubset(_codes(result))


@pytest.mark.parametrize("actual_through", [None, 0, 13, "6", True])
def test_actual_through_must_be_integer_one_to_twelve(actual_through):
    result = parse_pnl_reporting_workbook(
        b"not-opened",
        dataset_type=DatasetType.ACTUAL,
        reporting_year=2026,
        actual_through_month=actual_through,
    )
    assert result.valid is False
    assert _codes(result) == {ValidationErrorCode.INVALID_ACTUAL_THROUGH}


def test_plan_rejects_actual_through():
    result = parse_pnl_reporting_workbook(
        b"not-opened",
        dataset_type=DatasetType.PLAN,
        reporting_year=2026,
        actual_through_month=6,
    )
    assert _codes(result) == {ValidationErrorCode.INVALID_ACTUAL_THROUGH}


def test_unsupported_source_type_returns_safe_invalid_template_issue():
    result = parse_pnl_reporting_workbook(
        object(),  # type: ignore[arg-type]
        dataset_type=DatasetType.PLAN,
        reporting_year=2026,
    )

    assert result.valid is False
    assert _codes(result) == {ValidationErrorCode.INVALID_TEMPLATE}


@pytest.mark.parametrize(
    "invalid_value",
    ["=1+1", "0", True, date(2026, 1, 1), "#DIV/0!", "N/A", "arbitrary"],
)
def test_invalid_authoritative_input_types_are_rejected(invalid_value):
    def mutate(workbook):
        workbook[SHEET_MONTHLY_PNL].cell(3, 3, invalid_value)

    result = _parse(_edit(_workbook_bytes(), mutate))
    assert result.valid is False
    assert ValidationErrorCode.INVALID_VALUE in _codes(result)


@pytest.mark.parametrize("invalid_value", [float("nan"), float("inf"), float("-inf")])
def test_non_finite_numeric_primitive_is_rejected(invalid_value):
    workbook = Workbook()
    cell = workbook.active["A1"]
    cell.value = invalid_value
    assert _finite_number(cell) is None
    workbook.close()


@pytest.mark.parametrize(
    ("mutator", "expected"),
    [
        (lambda workbook: workbook.remove(workbook[SHEET_SGA]), ValidationErrorCode.MISSING_SHEET),
        (lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(1, 3, "01월"), ValidationErrorCode.INVALID_HEADER),
        (lambda workbook: workbook[SHEET_MONTHLY_PNL].__setitem__("P2", "PNL_REPORTING_V2"), ValidationErrorCode.UNSUPPORTED_TEMPLATE_VERSION),
        (lambda workbook: _set_cell(workbook[SHEET_MANUFACTURING_COGS], 2, 15, None), ValidationErrorCode.MISSING_KEY),
        (lambda workbook: workbook[SHEET_MANUFACTURING_COGS].cell(3, 15, "mfg_material"), ValidationErrorCode.DUPLICATE_KEY),
        (lambda workbook: workbook[SHEET_MANUFACTURING_COGS].cell(2, 15, "mfg_unknown"), ValidationErrorCode.UNKNOWN_KEY),
        (lambda workbook: _set_cell(workbook[SHEET_MANUFACTURING_COGS], 6, 15, None), ValidationErrorCode.MISSING_KEY),
        (lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 1, "wrong label"), ValidationErrorCode.LABEL_MISMATCH),
        (lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 2, "KRW"), ValidationErrorCode.UNIT_MISMATCH),
        (lambda workbook: workbook[SHEET_SGA].cell(3, 2, "wrong category"), ValidationErrorCode.CATEGORY_MISMATCH),
        (lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(2, 17, "UNKNOWN"), ValidationErrorCode.PRODUCT_GROUP_MISMATCH),
        (
            lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(
                _product_row(workbook[SHEET_PRODUCT_PNL], "LC", "volume"), 2, "16-inch"
            ),
            ValidationErrorCode.PRODUCT_GROUP_MISMATCH,
        ),
        (
            lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(
                _product_row(workbook[SHEET_PRODUCT_PNL], "LC", "volume"), 1, "16인치 LC"
            ),
            ValidationErrorCode.PRODUCT_GROUP_MISMATCH,
        ),
        (
            lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(
                _product_row(workbook[SHEET_PRODUCT_PNL], "LC", "volume"), 4, "m"
            ),
            ValidationErrorCode.UNIT_MISMATCH,
        ),
        (
            lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(
                _product_row(workbook[SHEET_PRODUCT_PNL], "FS", "volume"), 4, "PCS"
            ),
            ValidationErrorCode.UNIT_MISMATCH,
        ),
        (
            lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(
                _product_row(workbook[SHEET_PRODUCT_PNL], "SW", "revenue"), 3, "wrong metric"
            ),
            ValidationErrorCode.LABEL_MISMATCH,
        ),
        (
            lambda workbook: workbook[SHEET_PRODUCT_PNL].cell(
                _product_row(workbook[SHEET_PRODUCT_PNL], "BW", "revenue"), 17, "SW"
            ),
            ValidationErrorCode.DUPLICATE_KEY,
        ),
        (lambda workbook: workbook[SHEET_MONTHLY_PNL].merge_cells("A2:B2"), ValidationErrorCode.INVALID_TEMPLATE),
    ],
)
def test_structure_and_metadata_failures_have_exact_codes(mutator, expected):
    result = _parse(_edit(_workbook_bytes(), mutator))
    assert result.valid is False
    assert result.canonical_payload is None
    assert expected in _codes(result)


@pytest.mark.parametrize("unknown_metric", ["volume", "asp"])
def test_new_business_volume_and_asp_are_unknown_composite_keys(unknown_metric):
    def mutate(workbook):
        sheet = workbook[SHEET_PRODUCT_PNL]
        row = _product_row(sheet, "NEW_BUSINESS", "cogs")
        sheet.cell(row, 18, unknown_metric)

    result = _parse(_edit(_workbook_bytes(), mutate))
    assert result.valid is False
    assert ValidationErrorCode.UNKNOWN_KEY in _codes(result)


def test_product_quantity_units_are_exact_and_no_mixed_total_is_created():
    result = _parse(_workbook_bytes())
    assert result.valid
    product_payload = result.canonical_payload.to_dict()["sheets"][SHEET_PRODUCT_PNL]
    assert set(product_payload) == {"SW", "BW", "LC", "FS", "NEW_BUSINESS"}
    assert set(product_payload["SW"]) == {"revenue", "volume", "cogs", "sga"}
    assert set(product_payload["FS"]) == {"revenue", "volume", "cogs", "sga"}
    assert set(product_payload["NEW_BUSINESS"]) == {"revenue", "cogs", "sga"}
    assert "volume" not in product_payload["NEW_BUSINESS"]
    assert "asp" not in product_payload["NEW_BUSINESS"]
    assert "total_volume" not in product_payload


@pytest.mark.parametrize("replacement", [None, 12345, "#N/A"])
def test_deleted_literal_and_error_helper_cells_warn_but_do_not_change_canonical(replacement):
    baseline = _parse(_workbook_bytes())

    def mutate(workbook):
        _set_cell(workbook[SHEET_MANUFACTURING_COGS], 6, 3, replacement)

    changed = _parse(_edit(_workbook_bytes(), mutate))
    assert changed.valid is True
    assert ValidationErrorCode.HELPER_VALUE_CHANGED in _codes(changed)
    assert changed.canonical_payload.to_json() == baseline.canonical_payload.to_json()


def test_altered_helper_formula_is_ignored_and_canonical_is_unchanged():
    baseline = _parse(_workbook_bytes())
    changed = _parse(
        _edit(
            _workbook_bytes(),
            lambda workbook: workbook[SHEET_MANUFACTURING_COGS].cell(6, 3, "=999999"),
        )
    )
    assert changed.valid is True
    assert changed.canonical_payload.to_json() == baseline.canonical_payload.to_json()


def test_helper_formula_cached_value_never_enters_canonical_payload():
    source = _workbook_bytes()

    def add_cached_value(name: str, payload: bytes) -> bytes:
        if name == "xl/worksheets/sheet2.xml":
            payload = payload.replace(b'<c r="C6"><f>0</f><v></v></c>', b'<c r="C6"><f>0</f><v>999999</v></c>')
        return payload

    baseline = _parse(source)
    cached = _parse(_rewrite_zip(source, add_cached_value))
    assert cached.valid is True
    assert cached.canonical_payload.to_json() == baseline.canonical_payload.to_json()
    assert "999999" not in cached.canonical_payload.to_json()


def test_input_formula_is_rejected_even_when_package_contains_a_cached_number():
    source = _edit(_workbook_bytes(), lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 3, "=1+1"))

    def add_cached_value(name: str, payload: bytes) -> bytes:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(b'<c r="C3"><f>1+1</f><v></v></c>', b'<c r="C3"><f>1+1</f><v>2</v></c>')
        return payload

    result = _parse(_rewrite_zip(source, add_cached_value))
    assert result.valid is False
    assert ValidationErrorCode.INVALID_VALUE in _codes(result)


def test_row_order_change_warns_but_produces_identical_canonical_serialization():
    baseline = _parse(_workbook_bytes())

    def swap_rows(workbook):
        sheet = workbook[SHEET_MANUFACTURING_COGS]
        first = [sheet.cell(2, column).value for column in range(1, 16)]
        second = [sheet.cell(3, column).value for column in range(1, 16)]
        for column, value in enumerate(second, 1):
            sheet.cell(2, column, value)
        for column, value in enumerate(first, 1):
            sheet.cell(3, column, value)

    reordered = _parse(_edit(_workbook_bytes(), swap_rows))
    assert reordered.valid is True
    assert ValidationErrorCode.ROW_ORDER_CHANGED in _codes(reordered)
    assert reordered.canonical_payload.to_json() == baseline.canonical_payload.to_json()


def test_same_input_has_deterministic_serialization():
    source = _workbook_bytes(dataset_type=DatasetType.ACTUAL, actual_through=6)
    first = _parse(source, dataset_type=DatasetType.ACTUAL, actual_through=6)
    second = _parse(source, dataset_type=DatasetType.ACTUAL, actual_through=6)
    assert first.valid and second.valid
    assert first.canonical_payload.to_json() == second.canonical_payload.to_json()


def test_extra_sheet_and_sheet_order_are_warnings_not_blockers():
    def mutate(workbook):
        workbook.create_sheet("00_입력안내", 0)
        sheet = workbook._sheets.pop(workbook._sheets.index(workbook[SHEET_SGA]))
        workbook._sheets.insert(1, sheet)

    result = _parse(_edit(_workbook_bytes(), mutate))
    assert result.valid is True
    assert ValidationErrorCode.UNEXPECTED_SHEET in _codes(result)
    assert ValidationErrorCode.SHEET_ORDER_CHANGED in _codes(result)


def test_merge_outside_contract_block_is_ignored():
    result = _parse(
        _edit(
            _workbook_bytes(),
            lambda workbook: workbook[SHEET_MONTHLY_PNL].merge_cells("A30:B30"),
        )
    )
    assert result.valid is True
    assert ValidationErrorCode.INVALID_TEMPLATE not in _codes(result)


def test_calc_chain_is_allowed_for_helper_formula_workbooks():
    source = _workbook_bytes()
    destination = BytesIO()
    with zipfile.ZipFile(BytesIO(source)) as original, zipfile.ZipFile(destination, "w") as output:
        for entry in original.infolist():
            output.writestr(entry, original.read(entry.filename))
        output.writestr(
            "xl/calcChain.xml",
            b'<?xml version="1.0" encoding="UTF-8"?><calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
        )
    result = _parse(destination.getvalue())
    assert result.valid is True


def test_external_relationship_is_rejected_as_invalid_template():
    source = _workbook_bytes()

    def add_external(name: str, payload: bytes) -> bytes:
        if name == "_rels/.rels":
            relation = (
                b'<Relationship Id="unsafe" Type="urn:test" '
                b'Target="https://example.invalid/" TargetMode="External"/>'
            )
            payload = payload.replace(b"</Relationships>", relation + b"</Relationships>")
        return payload

    result = _parse(_rewrite_zip(source, add_external))
    assert result.valid is False
    assert _codes(result) == {ValidationErrorCode.INVALID_TEMPLATE}
    assert all("example.invalid" not in issue.safe_message for issue in result.issues)


def test_malformed_and_compression_bomb_packages_are_rejected_safely():
    malformed = _parse(b"not an xlsx")
    assert _codes(malformed) == {ValidationErrorCode.INVALID_TEMPLATE}

    bomb = BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"0" * (2 * 1024 * 1024))
        archive.writestr("_rels/.rels", b"x")
        archive.writestr("xl/workbook.xml", b"x")
    bomb_result = _parse(bomb.getvalue())
    assert _codes(bomb_result) == {ValidationErrorCode.INVALID_TEMPLATE}


def test_issue_model_is_structured_and_safe():
    result = _parse(
        _edit(
            _workbook_bytes(dataset_type=DatasetType.ACTUAL, actual_through=6),
            lambda workbook: _set_cell(workbook[SHEET_SGA], 3, 4, None),
        ),
        dataset_type=DatasetType.ACTUAL,
        actual_through=6,
    )
    issue = next(item for item in result.errors if item.error_code is ValidationErrorCode.ACTUAL_MONTH_MISSING)
    assert issue.severity is ValidationSeverity.BLOCKING
    assert issue.sheet == SHEET_SGA
    assert issue.row_key == "admin_labor"
    assert issue.month == 1
    assert issue.field == "D3"
    assert "Traceback" not in issue.safe_message
    assert ":\\" not in issue.safe_message


def test_reporting_parser_has_no_business_or_infrastructure_dependency_leakage():
    root = Path(__file__).resolve().parents[1] / "forecast" / "reporting"
    production_source = "\n".join(path.read_text(encoding="utf-8").casefold() for path in root.glob("*.py"))
    forbidden = (
        "sales_effects",
        "inventory_timing",
        "residual_rca",
        "golden_adapter",
        "legacy pnl_dashboard",
        "forecast.engine",
        "forecast.analysis",
        "forecast.bff",
        "supabase",
        "fastapi",
        "starlette",
    )
    assert all(token not in production_source for token in forbidden)
