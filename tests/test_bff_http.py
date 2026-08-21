from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
import tempfile
from pathlib import Path
from io import BytesIO
from openpyxl import Workbook, load_workbook

import pytest
from fastapi.testclient import TestClient

from forecast.bff.application import (
    AnalysisModelListService,
    AnalysisSubmissionService,
    JobQueryService,
    ResultQueryService,
    TrustedBffApplication,
    WorkerAdministrationService,
)
from forecast.bff.auth import AccessCodeSessionService
from forecast.bff.gateway import (
    GatewayIdempotencyConflictError,
    GatewayValidationError,
    SubmissionRecord,
)
from forecast.bff.errors import ApiErrorCode, BffError
from forecast.bff.http import HttpBffSettings, InMemoryLoginRateLimiter, create_http_bff
from forecast.bff.production import AuditSink
from forecast.bff.evidence_history import EvidenceArtifact
from forecast.analysis_export import MIME_XLSX
from forecast.provenance import ResultProvenance
from forecast.bff.persistent_delete import (
    PersistentDeleteBatchResult,
    PersistentDeleteItemResult,
)


BASE = "11111111-1111-4111-8111-111111111111"
COMP = "22222222-2222-4222-8222-222222222222"
JOB = "33333333-3333-4333-8333-333333333333"
RESULT = "44444444-4444-4444-8444-444444444444"
PROVENANCE = ResultProvenance("engine", "mapping", "a" * 64, "1")


class Gateway:
    def __init__(self):
        self.requests = {}
        self.job_status = "pending"
        self.available = True
        self.model_published = True

    def submit_analysis(self, **values):
        if not self.model_published:
            raise GatewayValidationError("analysis models must be published")
        key = (values["idempotency_actor"], values["idempotency_key"])
        fingerprint = (
            values["baseline_model_id"], values["comparison_model_id"],
            values["start_month"], values["end_month"],
            values["baseline_sales_fx"], values["comparison_sales_fx"],
            tuple(sorted((values["baseline_sales_fx_monthly"] or {}).items())),
            tuple(sorted((values["comparison_sales_fx_monthly"] or {}).items())),
        )
        if key in self.requests:
            if self.requests[key] != fingerprint:
                raise GatewayIdempotencyConflictError("collision")
            return SubmissionRecord(JOB, self.job_status, True)
        self.requests[key] = fingerprint
        return SubmissionRecord(JOB, self.job_status, False)

    def get_job_status(self, _job_id):
        return {
            "job_id": JOB, "status": self.job_status,
            "baseline_model_id": BASE, "comparison_model_id": COMP,
            "start_month": 1, "end_month": 6, "attempt": 1, "max_attempts": 3,
            "created_at": "2026-08-11T00:00:00Z", "heartbeat_at": None,
            "completed_at": "2026-08-11T00:01:00Z" if self.job_status == "completed" else None,
            "result_id": RESULT if self.job_status == "completed" else None,
            "error_code": "worker_execution_failed" if self.job_status == "failed" else None,
            "error_message": "SECRET postgres traceback / storage/path.xlsx",
            "claim_token": "must-never-leak", "queue_message_id": 99,
        }

    def get_admin_result_preview(self, _result_id):
        return result_row(is_published=False, published_at=None)

    def get_viewer_result(self, _result_id, **_kwargs):
        return result_row() if self.available else None

    def validate_result_availability(self, _result_id, **_kwargs):
        return self.available


def result_row(**overrides):
    row = {
        "result_id": RESULT, "job_id": JOB, "analysis_view": {"summary": {"status": "PASS"}},
        "baseline_model_id": BASE, "comparison_model_id": COMP,
        "baseline_workbook_sha256": "b" * 64, "comparison_workbook_sha256": "c" * 64,
        "engine_version": "engine", "mapping_version": "mapping", "mapping_hash": "a" * 64,
        "result_schema_version": "1", "is_published": True, "is_default": False,
        "published_at": "2026-08-11T00:01:00Z", "created_at": "2026-08-11T00:01:00Z",
    }
    row.update(overrides)
    return row


@dataclass
class Fixture:
    client: TestClient
    gateway: Gateway
    forecast: object
    forecast_metadata: object

    def login(self, code="admin-code"):
        response = self.client.post("/api/session/login", json={"access_code": code})
        assert response.status_code == 200
        return response

    @property
    def csrf(self):
        return self.client.cookies.get("pnl_csrf")


def make_fixture(*, limiter=None, clock=None, worker_control=None, audit_sink=None, persistent_delete=None) -> Fixture:
    sessions = AccessCodeSessionService(
        viewer_code="viewer-code", admin_code="admin-code",
        actor_namespace_secret="actor-namespace-secret-at-least-32-chars", ttl_seconds=3600,
        clock=clock,
    )
    gateway = Gateway()
    repository = SimpleNamespace(list=lambda: [
        SimpleNamespace(id=BASE, name="Base", model_type="PLAN", year=2026,
                        start_month=1, end_month=12, is_published=True,
                        is_default=True, workbook_sha256="b" * 64),
        SimpleNamespace(id=COMP, name="Comparison", model_type="ACTUAL", year=2026,
                        start_month=1, end_month=12, is_published=True,
                        is_default=False, workbook_sha256="c" * 64),
    ])
    forecast = FakeForecast()
    forecast_metadata = FakeForecastMetadata()
    app_service = TrustedBffApplication(
        sessions,
        AnalysisSubmissionService(sessions, gateway, PROVENANCE),
        JobQueryService(sessions, gateway),
        ResultQueryService(sessions, gateway, supported_result_schema_versions=("1",)),
        AnalysisModelListService(sessions, repository),
        evidence=FakeEvidence(),
        history=FakeHistory(),
        presentation=FakePresentation(),
        pnl_dashboard=FakePnlDashboard(),
        forecast_generation=forecast,
        forecast_input_metadata=forecast_metadata,
        worker_administration=(
            WorkerAdministrationService(sessions, worker_control)
            if worker_control is not None else None
        ),
        persistent_delete=persistent_delete,
    )
    app = create_http_bff(
        app_service,
        settings=HttpBffSettings(environment="test", csrf_secret="csrf-secret-at-least-32-characters"),
        rate_limiter=limiter,
        audit_sink=audit_sink,
    )
    return Fixture(TestClient(app), gateway, forecast, forecast_metadata)


class FakeEvidence:
    last_root = None
    def admin_download(self, _session, result_id):
        return self._artifact(result_id)

    def viewer_download(self, _session, result_id):
        return self._artifact(result_id)

    @staticmethod
    def _artifact(result_id):
        root = Path(tempfile.mkdtemp(prefix="test-evidence-"))
        FakeEvidence.last_root = root
        path = root / "evidence.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = result_id
        workbook.save(path)
        return EvidenceArtifact(path, f"손익분석_근거_{result_id[:8]}.xlsx", MIME_XLSX, root)


class FakePnlDashboard:
    def viewer_read(self, _session):
        return {"result_id": RESULT, "job_id": JOB, "dto_version": "1"}


class FakeForecast:
    def __init__(self):
        self.last_request = None

    def generate(self, _session, request):
        self.last_request = request
        return {
            "generation_id": JOB, "model_id": COMP, "display_name": request.name,
            "model_year": request.model_year, "start_month": request.start_month,
            "end_month": request.end_month, "is_published": False, "is_default": False,
            "workbook_sha256": "d" * 64, "idempotency_replayed": False,
            "execution_mode": "SYNCHRONOUS", "dto_version": "1",
        }


class FakeForecastMetadata:
    def get(self, _session, base_model_id):
        return {
            "base_model_id": base_model_id,
            "manufacturing": [{
                "adjustment_key": "manufacturing:000", "display_name": "노무비",
                "category": "manufacturing", "section": None, "unit": "KRW",
            }],
            "sga": [{
                "adjustment_key": "sga:000", "display_name": "운반비",
                "category": "sga", "section": "selling", "unit": "KRW",
            }],
            "reason_max_length": 500,
            "dto_version": "1",
        }

    def resolve_adjustment_keys(self, _session, _base_model_id, category, keys):
        expected = {"manufacturing": "manufacturing:000", "sga": "sga:000"}[category]
        if any(key != expected for key in keys):
            raise BffError(ApiErrorCode.VALIDATION_ERROR, "invalid key")
        row = {"manufacturing": 290, "sga": 1168}[category]
        return {key: row for key in keys}


class FakeHistory:
    def list_admin(self, _session, **_kwargs):
        return {
            "items": [{
                "job_id": JOB, "result_id": RESULT, "status": "COMPLETED",
                "baseline_model_id": BASE, "baseline_model_name": "Base",
                "comparison_model_id": COMP, "comparison_model_name": "Comparison",
                "start_month": 1, "end_month": 6, "attempt": 1, "max_attempts": 3,
                "created_at": "2026-08-11T00:00:00Z", "completed_at": "2026-08-11T00:01:00Z",
                "error_code": None, "error_message": None, "is_published": False,
            }],
            "next_before_created_at": None, "next_before_job_id": None, "dto_version": "1",
        }


class FakePresentation:
    def admin_read(self, _session, result_id):
        return {"identity": {"result_id": result_id}, "scope": "admin", "dto_version": "1"}

    def viewer_read(self, _session, result_id):
        return {"identity": {"result_id": result_id}, "scope": "viewer", "dto_version": "1"}

    def list_viewer(self, _session):
        return {
            "results": [{
                "result_id": RESULT,
                "label": "2026 계획 대비 12월 실적 · 12월",
                "completed_at": "2026-08-11T00:01:00Z",
                "published_at": "2026-08-11T00:02:00Z",
            }],
            "dto_version": "1",
        }


class FakeWorkerControl:
    def __init__(self):
        self.wakes = 0
        self.stops = 0

    @staticmethod
    def _status(desired):
        return {
            "desired_instance_count": desired, "configured_instance_count": desired,
            "actual_instance_count": desired, "queue_depth": 0,
            "claimable_count": 0, "pending_count": 0, "processing_count": 0,
            "active_lease_count": 0, "active_heartbeat_count": 0,
            "recovery_pending_count": 0, "work_exists": False,
            "idle_seconds": 0, "last_worker_activity_at": "2026-08-12T00:00:00Z",
            "last_scaling_result": "scaled", "platform_reconciling": False,
            "platform_ready": True, "operating_policy": "DEMAND_ONLY",
            "idle_policy_seconds": 1800, "dto_version": "1",
        }

    def status(self): return self._status(0)
    def emergency_wake(self):
        self.wakes += 1
        return self._status(1)
    def safe_stop(self):
        self.stops += 1
        return self._status(0)


class RecordingAudit(AuditSink):
    def __init__(self): self.events = []
    def record(self, **event): self.events.append(event)


def test_admin_worker_actions_return_success_and_audit_non_uuid_resource_once():
    control = FakeWorkerControl()
    audit = RecordingAudit()
    fx = make_fixture(worker_control=control, audit_sink=audit)
    fx.login()
    headers = {"X-CSRF-Token": fx.csrf}

    wake = fx.client.post("/api/admin/worker/emergency-wake", headers=headers)
    stop = fx.client.post("/api/admin/worker/safe-stop", headers=headers)

    assert wake.status_code == stop.status_code == 200
    assert control.wakes == control.stops == 1
    worker_events = [event for event in audit.events if event["event_type"].startswith("worker_")]
    assert [event["event_type"] for event in worker_events] == [
        "worker_emergency_wake", "worker_safe_stop",
    ]
    assert all(len(event["operation_id"]) == 36 for event in worker_events)


@pytest.mark.parametrize("code,role", [("viewer-code", "viewer"), ("admin-code", "admin")])
def test_login_issues_opaque_cookie_without_body_secret(code, role):
    fx = make_fixture()
    response = fx.login(code)
    assert response.json()["role"] == role and response.json()["authenticated"] is True
    assert "session_id" not in response.text and code not in response.text
    cookies = response.headers.get_list("set-cookie")
    assert any("pnl_session=" in value and "HttpOnly" in value and "SameSite=strict" in value for value in cookies)
    assert any("pnl_csrf=" in value and "HttpOnly" not in value for value in cookies)


def test_invalid_and_malformed_session_are_rejected_and_rate_limited_generically():
    limiter = InMemoryLoginRateLimiter(max_attempts=2, window_seconds=60)
    fx = make_fixture(limiter=limiter)
    assert fx.client.post("/api/session/login", json={"access_code": "bad"}).status_code == 401
    locked = fx.client.post("/api/session/login", json={"access_code": "bad2"})
    assert locked.status_code == 429 and "Retry-After" in locked.headers
    limited = fx.client.post("/api/session/login", json={"access_code": "admin-code"})
    assert limited.status_code == 429 and "Retry-After" in limited.headers
    fx.client.cookies.set("pnl_session", "short")
    assert fx.client.get("/api/session").status_code == 401


def test_expired_http_session_is_rejected():
    now = [datetime(2026, 8, 11, tzinfo=timezone.utc)]
    fx = make_fixture(clock=lambda: now[0]); fx.login()
    now[0] += timedelta(hours=2)
    assert fx.client.get("/api/session").status_code == 401


def test_csrf_logout_and_viewer_admin_capability_boundary():
    fx = make_fixture()
    fx.login("viewer-code")
    assert fx.client.get("/api/models").status_code == 403
    assert fx.client.post("/api/session/logout").status_code == 403
    assert fx.client.post("/api/session/logout", headers={"X-CSRF-Token": "wrong"}).status_code == 403
    assert fx.client.post("/api/session/logout", headers={"X-CSRF-Token": fx.csrf}).status_code == 200
    assert fx.client.get("/api/session").status_code == 401


def test_model_list_submit_idempotency_collision_and_unpublished_rejection():
    fx = make_fixture(); fx.login()
    assert len(fx.client.get("/api/models").json()["models"]) == 2
    body = {
        "baseline_model_id": BASE, "comparison_model_id": COMP,
        "start_month": 1, "end_month": 6,
        "baseline_sales_fx_monthly": {f"2026-{month:02d}": 1480 for month in range(1, 7)},
        "comparison_sales_fx_monthly": {f"2026-{month:02d}": 1500 for month in range(1, 7)},
        "idempotency_key": "logical-action-1",
    }
    assert fx.client.post("/api/analyses", json=body).status_code == 403
    assert fx.client.post("/api/analyses", json=body, headers={"X-CSRF-Token": "wrong"}).status_code == 403
    first = fx.client.post("/api/analyses", json=body, headers={"X-CSRF-Token": fx.csrf})
    replay = fx.client.post("/api/analyses", json=body, headers={"X-CSRF-Token": fx.csrf})
    assert first.status_code == replay.status_code == 200
    assert replay.json()["idempotency_replayed"] is True
    changed = {**body["comparison_sales_fx_monthly"], "2026-03": 1501}
    conflict = fx.client.post("/api/analyses", json={**body, "comparison_sales_fx_monthly": changed}, headers={"X-CSRF-Token": fx.csrf})
    assert conflict.status_code == 409 and conflict.json()["error"]["code"] == "IDEMPOTENCY_CONFLICT"
    fx.gateway.model_published = False
    denied = fx.client.post("/api/analyses", json={**body, "idempotency_key": "new"}, headers={"X-CSRF-Token": fx.csrf})
    assert denied.status_code == 422


@pytest.mark.parametrize("field,value", [("start_month", True), ("end_month", 13), ("baseline_sales_fx", True), ("comparison_sales_fx", 0)])
def test_submit_strict_validation(field, value):
    fx = make_fixture(); fx.login()
    body = {"baseline_model_id": BASE, "comparison_model_id": COMP, "start_month": 1,
            "end_month": 6, "baseline_sales_fx": 1480, "comparison_sales_fx": 1500,
            "idempotency_key": "validation"}
    response = fx.client.post("/api/analyses", json={**body, field: value}, headers={"X-CSRF-Token": fx.csrf})
    assert response.status_code == 422 and response.json()["error"]["code"] == "VALIDATION_ERROR"


@pytest.mark.parametrize("status", ["pending", "processing", "completed", "failed"])
def test_job_by_id_states_and_internal_fields_never_leak(status):
    fx = make_fixture(); fx.login(); fx.gateway.job_status = status
    response = fx.client.get(f"/api/jobs/{JOB}")
    assert response.status_code == 200 and response.json()["status"] == status.upper()
    assert "claim_token" not in response.text and "queue_message_id" not in response.text
    assert "postgres" not in response.text and "storage/path" not in response.text


def test_admin_preview_and_strict_viewer_result_are_separate():
    admin = make_fixture(); admin.login()
    preview = admin.client.get(f"/api/admin/results/{RESULT}")
    assert preview.status_code == 200 and preview.json()["is_published"] is False

    viewer = make_fixture(); viewer.login("viewer-code")
    assert viewer.client.get(f"/api/admin/results/{RESULT}").status_code == 403
    assert viewer.client.get(f"/api/viewer/results/{RESULT}").status_code == 200
    viewer.gateway.available = False
    hidden = viewer.client.get(f"/api/viewer/results/{RESULT}")
    assert hidden.status_code == 404 and hidden.json()["error"]["code"] == "RESULT_NOT_AVAILABLE"


def test_presentation_http_routes_keep_admin_and_viewer_capabilities_separate():
    admin = make_fixture(); admin.login()
    preview = admin.client.get(f"/api/admin/results/{RESULT}/presentation")
    assert preview.status_code == 200 and preview.json()["scope"] == "admin"
    assert "claim_token" not in preview.text and "workbook_path" not in preview.text

    viewer = make_fixture(); viewer.login("viewer-code")
    assert viewer.client.get(f"/api/admin/results/{RESULT}/presentation").status_code == 403
    visible = viewer.client.get(f"/api/viewer/results/{RESULT}/presentation")
    assert visible.status_code == 200 and visible.json()["scope"] == "viewer"


def test_viewer_analysis_result_list_is_minimal_no_store_and_not_admin_history():
    anonymous = make_fixture()
    assert anonymous.client.get("/api/viewer/analysis-results").status_code == 401

    viewer = make_fixture(); viewer.login("viewer-code")
    response = viewer.client.get("/api/viewer/analysis-results")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    assert set(response.json()) == {"results", "dto_version"}
    assert set(response.json()["results"][0]) == {
        "result_id", "label", "completed_at", "published_at",
    }
    assert viewer.client.get("/api/admin/calculation-history").status_code == 403
    assert viewer.client.get(f"/api/viewer/results/{RESULT}/presentation").status_code == 200


def test_production_requires_secure_cookie_and_shared_rate_limiter():
    fx = make_fixture()
    with pytest.raises(ValueError, match="Secure"):
        HttpBffSettings(environment="production", csrf_secret="csrf-secret-at-least-32-characters")
    with pytest.raises(ValueError, match="shared login rate limiter"):
        create_http_bff(
            fx.client.app.state if False else make_application_only(),
            settings=HttpBffSettings(environment="production", cookie_secure=True, csrf_secret="csrf-secret-at-least-32-characters"),
        )


def make_application_only():
    sessions = AccessCodeSessionService(viewer_code="v", admin_code="a", actor_namespace_secret="x" * 32)
    gateway = Gateway()
    return TrustedBffApplication(
        sessions, AnalysisSubmissionService(sessions, gateway, PROVENANCE),
        JobQueryService(sessions, gateway), ResultQueryService(sessions, gateway, supported_result_schema_versions=("1",)),
    )


def test_evidence_http_streams_xlsx_and_history_is_admin_only():
    admin = make_fixture(); admin.login()
    evidence = admin.client.get(f"/api/admin/results/{RESULT}/evidence")
    assert evidence.status_code == 200
    assert evidence.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in evidence.headers["content-disposition"]
    assert evidence.content.startswith(b"PK")
    assert load_workbook(BytesIO(evidence.content)).active["A1"].value == RESULT
    assert FakeEvidence.last_root is not None and not FakeEvidence.last_root.exists()
    history = admin.client.get("/api/admin/calculation-history?limit=20")
    assert history.status_code == 200
    assert history.json()["items"][0]["result_id"] == RESULT
    assert "claim_token" not in history.text and "queue_message_id" not in history.text

    viewer = make_fixture(); viewer.login("viewer-code")
    assert viewer.client.get("/api/admin/calculation-history").status_code == 403
    assert viewer.client.get(f"/api/admin/results/{RESULT}/evidence").status_code == 403
    assert viewer.client.get(f"/api/viewer/results/{RESULT}/evidence").status_code == 200


def test_pnl_dashboard_is_viewer_capability_with_no_store_and_no_internal_fields():
    anonymous = make_fixture()
    assert anonymous.client.get("/api/viewer/pnl-dashboard").status_code == 401
    viewer = make_fixture(); viewer.login("viewer-code")
    response = viewer.client.get("/api/viewer/pnl-dashboard")
    assert response.status_code == 200
    assert response.headers["cache-control"].startswith("no-store")
    assert response.json() == {"result_id": RESULT, "job_id": JOB, "dto_version": "1"}
    assert "claim_token" not in response.text and "storage_path" not in response.text


def test_forecast_http_is_admin_csrf_protected_and_returns_only_safe_draft_dto():
    body = {
        "base_model_id": BASE, "name": "2026 Forecast", "model_year": 2026,
        "version": "V1", "start_month": 7, "end_month": 7,
        "months": [{"month": 7, "sales": [], "production": []}],
        "idempotency_key": "forecast-key",
    }
    anonymous = make_fixture()
    assert anonymous.client.post("/api/admin/forecasts", json=body).status_code == 401
    viewer = make_fixture(); viewer.login("viewer-code")
    assert viewer.client.post("/api/admin/forecasts", json=body,
        headers={"X-CSRF-Token": viewer.csrf}).status_code == 403
    admin = make_fixture(); admin.login()
    assert admin.client.post("/api/admin/forecasts", json=body).status_code == 403
    response = admin.client.post("/api/admin/forecasts", json=body,
        headers={"X-CSRF-Token": admin.csrf})
    assert response.status_code == 200
    assert response.json()["is_published"] is False
    assert response.json()["is_default"] is False
    assert "workbook_path" not in response.text and "service_role" not in response.text


def test_forecast_metadata_and_opaque_adjustment_keys_preserve_admin_boundary():
    anonymous = make_fixture()
    assert anonymous.client.get(
        f"/api/admin/forecasts/input-metadata?base_model_id={BASE}"
    ).status_code == 401
    viewer = make_fixture(); viewer.login("viewer-code")
    assert viewer.client.get(
        f"/api/admin/forecasts/input-metadata?base_model_id={BASE}"
    ).status_code == 403

    admin = make_fixture(); admin.login()
    metadata = admin.client.get(
        f"/api/admin/forecasts/input-metadata?base_model_id={BASE}"
    )
    assert metadata.status_code == 200
    assert metadata.json()["manufacturing"][0]["adjustment_key"] == "manufacturing:000"
    assert "row" not in metadata.text and "Data!" not in metadata.text

    body = {
        "base_model_id": BASE, "name": "Advanced Forecast", "model_year": 2026,
        "version": "V1", "start_month": 7, "end_month": 7,
        "months": [{
            "month": 7, "sales": [], "production": [],
            "manufacturing_adjustments": [{
                "adjustment_key": "manufacturing:000", "amount": -150000000,
                "reason": "비가동 조정",
            }],
            "sga_adjustments": [{
                "adjustment_key": "sga:000", "amount": 25000000,
                "reason": "운반비 조정",
            }],
        }],
        "idempotency_key": "forecast-advanced",
    }
    response = admin.client.post(
        "/api/admin/forecasts", json=body,
        headers={"X-CSRF-Token": admin.csrf},
    )
    assert response.status_code == 200
    month = admin.forecast.last_request.months[0]
    assert month.manufacturing_adjustments[0].row == 290
    assert month.manufacturing_adjustments[0].amount == -150000000
    assert month.manufacturing_adjustments[0].reason == "비가동 조정"
    assert month.sga_adjustments[0].row == 1168


def test_forecast_disabled_mode_is_admin_authenticated_and_fails_closed():
    application = make_application_only()
    app = create_http_bff(
        application,
        settings=HttpBffSettings(environment="test", csrf_secret="csrf-secret-at-least-32-characters"),
    )
    client = TestClient(app)
    login = client.post("/api/session/login", json={"access_code": "a"})
    assert login.status_code == 200
    body = {
        "base_model_id": BASE, "name": "Disabled Forecast", "model_year": 2026,
        "version": "V1", "start_month": 7, "end_month": 7,
        "months": [{"month": 7, "sales": [], "production": []}],
        "idempotency_key": "forecast-disabled",
    }
    denied = client.post(
        "/api/admin/forecasts", json=body,
        headers={"X-CSRF-Token": client.cookies.get("pnl_csrf")},
    )
    assert denied.status_code == 403
    assert denied.json()["error"]["code"] == "FORECAST_SCOPE_NOT_APPROVED"


class FakePersistentDelete:
    def __init__(self):
        self.calls = []

    def delete_models(self, session_id, ids):
        self.calls.append(("model", session_id, tuple(ids)))
        return PersistentDeleteBatchResult(
            resource_type="model", requested_count=2, deleted_count=1,
            blocked_count=1, failed_count=0,
            items=(
                PersistentDeleteItemResult(ids[0], "DELETED", "DELETED", {}, False),
                PersistentDeleteItemResult(ids[1], "BLOCKED_IN_USE", "MODEL_IN_USE", {"analysis_jobs": 1}, False),
            ),
        )

    def delete_analyses(self, session_id, ids):
        self.calls.append(("analysis", session_id, tuple(ids)))
        return PersistentDeleteBatchResult(
            resource_type="analysis", requested_count=1, deleted_count=0,
            blocked_count=0, failed_count=0, cleanup_required_count=1,
            items=(PersistentDeleteItemResult(
                ids[0], "CLEANUP_REQUIRED",
                "DB_DELETED_STORAGE_CLEANUP_REQUIRED", {}, True,
            ),),
        )

    def list_model_recoveries(self, session_id):
        self.calls.append(("model-recovery", session_id))
        return PersistentDeleteBatchResult(
            resource_type="model", requested_count=1, deleted_count=0,
            blocked_count=0, failed_count=0, cleanup_required_count=1,
            items=(PersistentDeleteItemResult(
                BASE, "CLEANUP_REQUIRED",
                "DB_DELETED_STORAGE_CLEANUP_REQUIRED", {}, True,
            ),),
        )

    def list_analysis_recoveries(self, session_id):
        self.calls.append(("analysis-recovery", session_id))
        return PersistentDeleteBatchResult(
            resource_type="analysis", requested_count=1, deleted_count=0,
            blocked_count=0, failed_count=0, cleanup_required_count=1,
            items=(PersistentDeleteItemResult(
                JOB, "CLEANUP_REQUIRED",
                "DB_DELETED_STORAGE_CLEANUP_REQUIRED", {}, True,
            ),),
        )

    def retry_model_cleanup(self, session_id, ids):
        self.calls.append(("model-retry", session_id, tuple(ids)))
        return PersistentDeleteBatchResult(
            resource_type="model", requested_count=1, deleted_count=1,
            blocked_count=0, failed_count=0,
            items=(PersistentDeleteItemResult(ids[0], "DELETED", "DELETED", {}, True),),
        )

    def retry_analysis_cleanup(self, session_id, ids):
        self.calls.append(("analysis-retry", session_id, tuple(ids)))
        return PersistentDeleteBatchResult(
            resource_type="analysis", requested_count=1, deleted_count=1,
            blocked_count=0, failed_count=0,
            items=(PersistentDeleteItemResult(ids[0], "DELETED", "DELETED", {}, True),),
        )


def test_persistent_delete_routes_require_admin_csrf_and_return_per_item_results():
    capability = FakePersistentDelete()
    anonymous = make_fixture(persistent_delete=capability)
    assert anonymous.client.post("/api/admin/models/delete", json={"ids": [BASE, COMP]}).status_code == 401

    viewer = make_fixture(persistent_delete=capability); viewer.login("viewer-code")
    assert viewer.client.post(
        "/api/admin/models/delete", json={"ids": [BASE, COMP]},
        headers={"X-CSRF-Token": viewer.csrf},
    ).status_code == 403

    audit = RecordingAudit()
    admin = make_fixture(persistent_delete=capability, audit_sink=audit); admin.login()
    assert admin.client.post("/api/admin/models/delete", json={"ids": [BASE, COMP]}).status_code == 403
    response = admin.client.post(
        "/api/admin/models/delete", json={"ids": [BASE, COMP]},
        headers={"X-CSRF-Token": admin.csrf},
    )
    assert response.status_code == 200
    assert response.json()["deleted_count"] == 1
    assert response.json()["items"][1]["reason"] == "MODEL_IN_USE"
    delete_events = [event for event in audit.events if event["event_type"] == "model_persistent_delete"]
    assert [event["outcome"] for event in delete_events] == ["success", "denied"]
    assert "storage_path" not in response.text and "service_role" not in response.text


def test_analysis_delete_route_reports_storage_cleanup_and_rejects_arbitrary_fields():
    capability = FakePersistentDelete()
    audit = RecordingAudit()
    admin = make_fixture(persistent_delete=capability, audit_sink=audit); admin.login()
    headers = {"X-CSRF-Token": admin.csrf}
    invalid = admin.client.post(
        "/api/admin/calculation-history/delete",
        json={"ids": [JOB], "storage_path": "models/other/source.xlsx"}, headers=headers,
    )
    assert invalid.status_code == 422
    assert capability.calls == []

    response = admin.client.post(
        "/api/admin/calculation-history/delete", json={"ids": [JOB]}, headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["cleanup_required_count"] == 1
    assert response.json()["items"][0]["status"] == "CLEANUP_REQUIRED"
    event = next(event for event in audit.events if event["event_type"] == "analysis_persistent_delete")
    assert event["outcome"] == "cleanup_required"
    assert event["error_code"] == "DB_DELETED_STORAGE_CLEANUP_REQUIRED"


def test_persistent_delete_recovery_status_and_retry_are_admin_only_and_path_free():
    capability = FakePersistentDelete()
    admin = make_fixture(persistent_delete=capability); admin.login()
    headers = {"X-CSRF-Token": admin.csrf}

    status = admin.client.get("/api/admin/models/delete/recovery")
    assert status.status_code == 200
    assert status.json()["items"][0]["status"] == "CLEANUP_REQUIRED"
    assert "storage_path" not in status.text and "storage_bucket" not in status.text

    no_csrf = admin.client.post(
        "/api/admin/models/delete/retry", json={"ids": [BASE]},
    )
    assert no_csrf.status_code == 403
    retry = admin.client.post(
        "/api/admin/models/delete/retry", json={"ids": [BASE]}, headers=headers,
    )
    assert retry.status_code == 200
    assert retry.json()["items"][0]["status"] == "DELETED"
    assert ("model-retry", admin.client.cookies.get("pnl_session"), (BASE,)) in capability.calls

    viewer = make_fixture(persistent_delete=capability); viewer.login("viewer-code")
    assert viewer.client.get("/api/admin/calculation-history/delete/recovery").status_code == 403
