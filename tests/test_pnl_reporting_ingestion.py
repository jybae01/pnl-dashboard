from __future__ import annotations

import copy
import hashlib
import threading
import uuid
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from forecast.bff.auth import AccessCodeSessionService
from forecast.bff.errors import ApiErrorCode, BffError
from forecast.bff.pnl_reporting_ingestion import (
    PnlReportingFinalizeError,
    PnlReportingFinalizeUncertainError,
    PnlReportingIdempotencyConflictError,
    PnlReportingInProgressError,
    PnlReportingIngestionService,
    PnlReportingReservation,
    PnlReportingSourceIntegrityError,
    PnlReportingStorageError,
    PnlReportingUploadRequest,
    PnlReportingValidationFailure,
    pnl_reporting_request_fingerprint,
    validation_failure_payload,
)
from forecast.reporting import (
    DatasetType,
    PnlReportingValidationResult,
    TEMPLATE_VERSION,
    ValidationErrorCode,
    ValidationIssue,
    ValidationSeverity,
)
from forecast.reporting.registry import (
    MANUFACTURING_COGS_HEADERS,
    MANUFACTURING_COGS_ROWS,
    MONTHLY_PNL_HEADERS,
    MONTHLY_PNL_ROWS,
    PRODUCT_GROUPS,
    PRODUCT_PNL_HEADERS,
    PRODUCT_PNL_ROWS,
    SGA_HEADERS,
    SGA_ROWS,
    SHEET_MANUFACTURING_COGS,
    SHEET_MONTHLY_PNL,
    SHEET_PRODUCT_PNL,
    SHEET_SGA,
)


def workbook_bytes(
    *,
    dataset_type: DatasetType = DatasetType.PLAN,
    actual_through: int = 6,
    zero_value: bool = False,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    monthly = workbook.create_sheet(SHEET_MONTHLY_PNL)
    monthly.append(MONTHLY_PNL_HEADERS)
    for definition in MONTHLY_PNL_ROWS:
        row = [definition.label, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.extend([definition.key, None])
        monthly.append(row)
    monthly["P2"] = TEMPLATE_VERSION

    cogs = workbook.create_sheet(SHEET_MANUFACTURING_COGS)
    cogs.append(MANUFACTURING_COGS_HEADERS)
    for definition in MANUFACTURING_COGS_ROWS:
        row = [definition.label, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.append(definition.key)
        cogs.append(row)

    sga = workbook.create_sheet(SHEET_SGA)
    sga.append(SGA_HEADERS)
    for definition in SGA_ROWS:
        row = [definition.label, definition.category, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.append(definition.key)
        sga.append(row)

    product = workbook.create_sheet(SHEET_PRODUCT_PNL)
    product.append(PRODUCT_PNL_HEADERS)
    groups = {group.key: group for group in PRODUCT_GROUPS}
    for definition in PRODUCT_PNL_ROWS:
        group = groups[definition.product_group_key]
        row = [group.display, group.dimension, definition.label, definition.unit]
        row.extend(_month_values(definition.is_input, dataset_type, actual_through))
        row.extend([definition.product_group_key, definition.key])
        product.append(row)

    if zero_value:
        monthly.cell(3, 3, 0)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _month_values(
    is_input: bool,
    dataset_type: DatasetType,
    actual_through: int,
) -> list[int | str | None]:
    if not is_input:
        return ["=0"] * 12
    return [
        month * 100 if dataset_type is DatasetType.PLAN or month <= actual_through else None
        for month in range(1, 13)
    ]


def edit_workbook(source: bytes, mutator) -> bytes:
    workbook = load_workbook(BytesIO(source), read_only=False, data_only=False)
    mutator(workbook)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def clear_cell(workbook, sheet: str, row: int, column: int) -> None:
    workbook[sheet].cell(row, column).value = None


class MemoryGateway:
    def __init__(self) -> None:
        self.lock = threading.Lock()
        self.requests: dict[tuple[str, str], dict] = {}
        self.objects: dict[str, bytes] = {}
        self.datasets: dict[str, dict] = {}
        self.active: dict[tuple[int, str], str] = {}
        self.superseded: dict[str, str] = {}
        self.fail_upload = False
        self.fail_finalize = False
        self.force_mismatch = False
        self.fail_cleanup = False
        self.finalize_uncertain = False
        self.fail_record = False
        self.upload_entered: threading.Event | None = None
        self.upload_release: threading.Event | None = None

    def reserve(self, *, actor, request, source_sha256, request_fingerprint):
        key = (actor, request.idempotency_key)
        with self.lock:
            state = self.requests.get(key)
            if state is not None:
                if state["fingerprint"] != request_fingerprint:
                    raise PnlReportingIdempotencyConflictError
                if state["status"] == "COMPLETED":
                    return PnlReportingReservation(
                        state["ingestion_id"], state["dataset_id"], "COMPLETED", None,
                        source_sha256,
                    )
                if state["status"] == "CLEANUP_REQUIRED":
                    from forecast.bff.pnl_reporting_ingestion import PnlReportingCleanupRequiredError

                    raise PnlReportingCleanupRequiredError
                if state["status"] == "FAILED":
                    state["status"] = "RESERVED"
                    return PnlReportingReservation(
                        state["ingestion_id"], state["dataset_id"], "RESERVED",
                        str(uuid.uuid4()), source_sha256,
                    )
                raise PnlReportingInProgressError
            state = {
                "ingestion_id": str(uuid.uuid4()),
                "dataset_id": str(uuid.uuid4()),
                "fingerprint": request_fingerprint,
                "status": "RESERVED",
                "request": request,
                "response": None,
                "failures": [],
            }
            self.requests[key] = state
            return PnlReportingReservation(
                state["ingestion_id"], state["dataset_id"], "RESERVED",
                str(uuid.uuid4()), source_sha256,
            )

    def upload_source(self, reservation, source):
        if self.upload_entered is not None:
            self.upload_entered.set()
        if self.upload_release is not None:
            assert self.upload_release.wait(timeout=10)
        if self.fail_upload:
            raise PnlReportingStorageError("upload failed")
        self.objects[reservation.dataset_id] = Path(source).read_bytes()

    def verify_source(self, reservation):
        if self.force_mismatch:
            raise PnlReportingSourceIntegrityError("mismatch")
        payload = self.objects.get(reservation.dataset_id)
        if payload is None:
            raise PnlReportingStorageError("missing")
        if hashlib.sha256(payload).hexdigest() != reservation.source_sha256:
            raise PnlReportingSourceIntegrityError("mismatch")

    def finalize(
        self,
        reservation,
        *,
        canonical_payload,
        original_filename,
        validation_summary,
    ):
        if self.finalize_uncertain:
            raise PnlReportingFinalizeUncertainError
        if self.fail_finalize:
            raise PnlReportingFinalizeError
        with self.lock:
            state = self._state(reservation.ingestion_id)
            request = state["request"]
            identity = (request.reporting_year, request.dataset_type.value)
            previous = self.active.get(identity)
            response = {
                "datasetId": reservation.dataset_id,
                "datasetType": request.dataset_type.value,
                "reportingYear": request.reporting_year,
                "actualThroughMonth": request.actual_through_month,
                "templateVersion": TEMPLATE_VERSION,
                "sourceSha256": reservation.source_sha256,
                "uploadedAt": "2026-08-19T00:00:00+00:00",
                "warnings": copy.deepcopy(validation_summary["warnings"]),
                "supersededDatasetId": previous,
                "replayed": False,
            }
            self.datasets[reservation.dataset_id] = {
                "canonical": copy.deepcopy(canonical_payload),
                "filename": original_filename,
                "response": copy.deepcopy(response),
            }
            self.active[identity] = reservation.dataset_id
            if previous is not None:
                self.superseded[previous] = reservation.dataset_id
            state["status"] = "COMPLETED"
            state["response"] = copy.deepcopy(response)
            return response

    def completed_response(self, reservation):
        return copy.deepcopy(self._state(reservation.ingestion_id)["response"])

    def remove_source(self, reservation):
        if self.fail_cleanup:
            raise PnlReportingStorageError("cleanup uncertain")
        self.objects.pop(reservation.dataset_id, None)

    def record_failure(self, reservation, **values):
        if self.fail_record:
            raise RuntimeError("record unavailable")
        with self.lock:
            state = self._state(reservation.ingestion_id)
            if state["status"] == "COMPLETED":
                raise RuntimeError("completed is immutable")
            state["status"] = "FAILED" if values["cleanup_succeeded"] else "CLEANUP_REQUIRED"
            state["failures"].append(copy.deepcopy(values))

    def heartbeat(self, reservation):
        assert reservation.lease_token

    def _state(self, ingestion_id):
        return next(
            state for state in self.requests.values() if state["ingestion_id"] == ingestion_id
        )


def fixture():
    sessions = AccessCodeSessionService(
        viewer_code="viewer",
        admin_code="admin",
        actor_namespace_secret="stable-actor-namespace-secret-32chars",
        ttl_seconds=3600,
    )
    admin = sessions.login("admin").session_id
    viewer = sessions.login("viewer").session_id
    gateway = MemoryGateway()
    return admin, viewer, gateway, PnlReportingIngestionService(sessions, gateway)


def request(dataset_type=DatasetType.PLAN, *, key="upload-1", through=None, year=2026):
    return PnlReportingUploadRequest(
        dataset_type=dataset_type,
        reporting_year=year,
        actual_through_month=through,
        original_filename="report.xlsx",
        idempotency_key=key,
    )


def write(tmp_path: Path, name: str, payload: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(payload)
    return path


def test_first_plan_and_actual_upload_create_independent_active_datasets(tmp_path):
    plan = write(tmp_path, "plan.xlsx", workbook_bytes())
    actual = write(
        tmp_path,
        "actual.xlsx",
        workbook_bytes(dataset_type=DatasetType.ACTUAL, actual_through=6, zero_value=True),
    )
    admin, _, gateway, service = fixture()
    plan_result = service.ingest(admin, request(), plan)
    actual_result = service.ingest(
        admin, request(DatasetType.ACTUAL, key="actual-1", through=6), actual
    )
    assert gateway.active[(2026, "PLAN")] == plan_result.dataset_id
    assert gateway.active[(2026, "ACTUAL")] == actual_result.dataset_id
    actual_payload = gateway.datasets[actual_result.dataset_id]["canonical"]
    assert actual_payload["sheets"][SHEET_MONTHLY_PNL]["rev_product"]["01"] == 0
    assert actual_payload["sheets"][SHEET_MONTHLY_PNL]["rev_product"]["07"] is None


@pytest.mark.parametrize(
    ("dataset_type", "through"),
    [(DatasetType.PLAN, None), (DatasetType.ACTUAL, 6)],
)
def test_valid_replacement_is_immutable_and_supersedes_previous(tmp_path, dataset_type, through):
    source = write(
        tmp_path,
        "source.xlsx",
        workbook_bytes(dataset_type=dataset_type, actual_through=through or 6),
    )
    admin, _, gateway, service = fixture()
    first = service.ingest(admin, request(dataset_type, key="first", through=through), source)
    old_snapshot = copy.deepcopy(gateway.datasets[first.dataset_id])
    second = service.ingest(admin, request(dataset_type, key="second", through=through), source)
    assert first.dataset_id != second.dataset_id
    assert gateway.active[(2026, dataset_type.value)] == second.dataset_id
    assert gateway.superseded[first.dataset_id] == second.dataset_id
    assert gateway.datasets[first.dataset_id] == old_snapshot
    assert len(gateway.datasets) == 2


def test_validation_failure_creates_no_reservation_storage_or_dataset_and_keeps_active(tmp_path):
    valid = write(tmp_path, "valid.xlsx", workbook_bytes())
    invalid_bytes = edit_workbook(
        workbook_bytes(), lambda workbook: clear_cell(workbook, SHEET_MONTHLY_PNL, 3, 3)
    )
    invalid = write(tmp_path, "invalid.xlsx", invalid_bytes)
    admin, _, gateway, service = fixture()
    first = service.ingest(admin, request(key="first"), valid)
    counts = (len(gateway.requests), len(gateway.objects), len(gateway.datasets))
    with pytest.raises(PnlReportingValidationFailure) as raised:
        service.ingest(admin, request(key="invalid"), invalid)
    assert raised.value.payload["status"] == "INVALID"
    assert raised.value.payload["errorCount"] > 0
    assert len(raised.value.payload["errors"]) <= 100
    assert gateway.active[(2026, "PLAN")] == first.dataset_id
    assert (len(gateway.requests), len(gateway.objects), len(gateway.datasets)) == counts


def test_validation_response_counts_before_truncation_and_bounds_workbook_text():
    issues = tuple(
        ValidationIssue(
            severity=ValidationSeverity.BLOCKING,
            error_code=ValidationErrorCode.INVALID_VALUE,
            safe_message="m" * 800,
            sheet="s" * 300,
            row_key="r" * 300,
            product_group_key="p" * 300,
            display_label="d" * 300,
            month=13,
            field="f" * 300,
        )
        for _ in range(105)
    ) + tuple(
        ValidationIssue(
            severity=ValidationSeverity.WARNING,
            error_code=ValidationErrorCode.ROW_ORDER_CHANGED,
            safe_message="warning",
        )
        for _ in range(7)
    )
    payload = validation_failure_payload(
        PnlReportingValidationResult(canonical_payload=None, issues=issues)
    )

    assert payload["errorCount"] == 105
    assert payload["warningCount"] == 7
    assert payload["truncated"] is True
    assert len(payload["errors"]) == 100
    assert payload["warnings"] == []
    assert len(payload["errors"][0]["message"]) == 500
    for field in ("sheet", "rowKey", "productGroupKey", "displayLabel", "field"):
        assert len(payload["errors"][0][field]) == 200
    assert "month" not in payload["errors"][0]


@pytest.mark.parametrize(
    ("failure", "expected_code"),
    [
        ("fail_upload", ApiErrorCode.TRANSIENT_SYSTEM_ERROR),
        ("fail_finalize", ApiErrorCode.TRANSIENT_SYSTEM_ERROR),
        ("force_mismatch", ApiErrorCode.INPUT_INTEGRITY_MISMATCH),
    ],
)
def test_failed_replacement_keeps_old_active_and_removes_partial_source(
    tmp_path, failure, expected_code
):
    source = write(tmp_path, "source.xlsx", workbook_bytes())
    admin, _, gateway, service = fixture()
    first = service.ingest(admin, request(key="first"), source)
    setattr(gateway, failure, True)
    with pytest.raises(BffError) as raised:
        service.ingest(admin, request(key="second"), source)
    assert raised.value.code == expected_code
    assert gateway.active[(2026, "PLAN")] == first.dataset_id
    assert len(gateway.datasets) == 1
    assert set(gateway.objects) == {first.dataset_id}


def test_cleanup_uncertain_and_finalize_uncertain_are_durable_and_never_swap_active(tmp_path):
    source = write(tmp_path, "source.xlsx", workbook_bytes())
    admin, _, gateway, service = fixture()
    first = service.ingest(admin, request(key="first"), source)

    gateway.fail_finalize = gateway.fail_cleanup = True
    with pytest.raises(BffError) as cleanup:
        service.ingest(admin, request(key="cleanup"), source)
    assert cleanup.value.code == ApiErrorCode.INGESTION_CLEANUP_REQUIRED
    cleanup_state = next(
        value for key, value in gateway.requests.items() if key[1] == "cleanup"
    )
    assert cleanup_state["status"] == "CLEANUP_REQUIRED"
    assert gateway.active[(2026, "PLAN")] == first.dataset_id

    gateway.fail_finalize = gateway.fail_cleanup = False
    gateway.finalize_uncertain = True
    with pytest.raises(BffError) as uncertain:
        service.ingest(admin, request(key="uncertain"), source)
    assert uncertain.value.code == ApiErrorCode.INGESTION_CLEANUP_REQUIRED
    uncertain_state = next(
        value for key, value in gateway.requests.items() if key[1] == "uncertain"
    )
    assert uncertain_state["status"] == "CLEANUP_REQUIRED"
    assert uncertain_state["dataset_id"] in gateway.objects
    assert gateway.active[(2026, "PLAN")] == first.dataset_id


def test_idempotency_replay_conflict_and_same_source_new_key(tmp_path):
    source = write(tmp_path, "source.xlsx", workbook_bytes())
    changed = write(
        tmp_path,
        "changed.xlsx",
        edit_workbook(
            workbook_bytes(), lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 3, 101)
        ),
    )
    admin, _, gateway, service = fixture()
    first = service.ingest(admin, request(), source)
    replay = service.ingest(admin, request(), source)
    assert replay.dataset_id == first.dataset_id and replay.replayed is True
    assert len(gateway.datasets) == 1
    with pytest.raises(BffError) as conflict:
        service.ingest(admin, request(), changed)
    assert conflict.value.code == ApiErrorCode.IDEMPOTENCY_CONFLICT
    second = service.ingest(admin, request(key="upload-2"), source)
    assert second.dataset_id != first.dataset_id
    assert gateway.active[(2026, "PLAN")] == second.dataset_id
    assert len(gateway.datasets) == 2


def test_same_key_can_retry_after_verified_cleanup_without_becoming_a_replay(tmp_path):
    source = write(tmp_path, "source.xlsx", workbook_bytes())
    admin, _, gateway, service = fixture()
    gateway.fail_upload = True
    with pytest.raises(BffError) as failed:
        service.ingest(admin, request(), source)
    assert failed.value.code == ApiErrorCode.TRANSIENT_SYSTEM_ERROR
    state = next(iter(gateway.requests.values()))
    assert state["status"] == "FAILED" and not gateway.objects and not gateway.datasets

    gateway.fail_upload = False
    result = service.ingest(admin, request(), source)
    assert result.replayed is False
    assert result.dataset_id == state["dataset_id"]
    assert len(gateway.datasets) == 1


def test_concurrent_duplicate_submit_has_one_logical_dataset(tmp_path):
    source = write(tmp_path, "source.xlsx", workbook_bytes())
    admin, _, gateway, service = fixture()
    gateway.upload_entered = threading.Event()
    gateway.upload_release = threading.Event()
    outcomes: list[object] = []

    def run():
        try:
            outcomes.append(service.ingest(admin, request(), source))
        except BffError as exc:
            outcomes.append(exc.code)

    first = threading.Thread(target=run)
    first.start()
    assert gateway.upload_entered.wait(timeout=10)
    second = threading.Thread(target=run)
    second.start()
    second.join(timeout=10)
    gateway.upload_release.set()
    first.join(timeout=10)
    assert len(gateway.datasets) == 1
    assert len(outcomes) == 2
    assert ApiErrorCode.TRANSIENT_SYSTEM_ERROR in outcomes
    assert sum(hasattr(value, "dataset_id") for value in outcomes) == 1


def test_concurrent_same_year_type_activation_keeps_one_pointer_and_coherent_chain(tmp_path):
    source = write(tmp_path, "source.xlsx", workbook_bytes())
    admin, _, gateway, service = fixture()
    initial = service.ingest(admin, request(key="initial"), source)
    results: list = []

    def run(key):
        results.append(service.ingest(admin, request(key=key), source))

    threads = [threading.Thread(target=run, args=(key,)) for key in ("next-1", "next-2")]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=20)
    assert len(results) == 2 and len(gateway.datasets) == 3
    active = gateway.active[(2026, "PLAN")]
    result_ids = {item.dataset_id for item in results}
    assert active in result_ids
    earlier = (result_ids - {active}).pop()
    assert gateway.superseded[initial.dataset_id] == earlier
    assert gateway.superseded[earlier] == active


def test_actual_future_value_blocks_before_reservation_and_viewer_is_forbidden(tmp_path):
    valid = workbook_bytes(dataset_type=DatasetType.ACTUAL, actual_through=6)
    future = edit_workbook(
        valid, lambda workbook: workbook[SHEET_MONTHLY_PNL].cell(3, 9, 0)
    )
    source = write(tmp_path, "actual.xlsx", future)
    admin, viewer, gateway, service = fixture()
    with pytest.raises(PnlReportingValidationFailure):
        service.ingest(
            admin, request(DatasetType.ACTUAL, key="actual", through=6), source
        )
    assert not gateway.requests
    with pytest.raises(BffError) as denied:
        service.ingest(viewer, request(DatasetType.ACTUAL, through=6), source)
    assert denied.value.code == ApiErrorCode.FORBIDDEN


def test_fingerprint_is_exact_canonical_sha_and_changes_with_contract_fields():
    source_sha = "a" * 64
    fingerprint = pnl_reporting_request_fingerprint(
        dataset_type=DatasetType.PLAN,
        reporting_year=2026,
        actual_through_month=None,
        template_version=TEMPLATE_VERSION,
        source_sha256=source_sha,
    )
    expected = hashlib.sha256(
        b'{"actual_through_month":null,"dataset_type":"PLAN","reporting_year":2026,'
        b'"source_sha256":"aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa",'
        b'"template_version":"PNL_REPORTING_V1"}'
    ).hexdigest()
    assert fingerprint == expected
    assert fingerprint != pnl_reporting_request_fingerprint(
        dataset_type=DatasetType.ACTUAL,
        reporting_year=2026,
        actual_through_month=6,
        template_version=TEMPLATE_VERSION,
        source_sha256=source_sha,
    )
