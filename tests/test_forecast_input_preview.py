from __future__ import annotations

from io import BytesIO
from pathlib import Path
import zipfile
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from forecast.bff.application import TrustedBffApplication
from forecast.bff.auth import AccessCodeSessionService
from forecast.bff.forecast_input_preview import (
    MIME_XLSX,
    PRODUCTION_HEADERS,
    SALES_HEADERS,
    ForecastInputPreviewService,
    build_input_template,
    parse_input_workbook,
)
from forecast.bff.http import HttpBffSettings, create_http_bff


def _write_workbook(path: Path, *, extra_sheet: bool = False) -> None:
    workbook = load_workbook(BytesIO(build_input_template()))
    sales, production = workbook.worksheets
    sales["A2"] = 7
    sales["E2"] = 12
    sales["F2"] = 99_000
    production["A2"] = 7
    production["D2"] = 100
    if extra_sheet:
        workbook.create_sheet("extra")
    workbook.save(path)
    workbook.close()


def test_template_is_exact_two_sheet_data_only_workbook():
    workbook = load_workbook(BytesIO(build_input_template()), read_only=False, data_only=False)
    assert workbook.sheetnames == ["판매계획", "생산계획"]
    assert tuple(cell.value for cell in next(workbook.worksheets[0].iter_rows(max_row=1))) == SALES_HEADERS
    assert tuple(cell.value for cell in next(workbook.worksheets[1].iter_rows(max_row=1))) == PRODUCTION_HEADERS
    assert workbook.worksheets[0].max_row == 12
    assert workbook.worksheets[1].max_row == 7
    assert workbook.worksheets[0].freeze_panes == "A2"
    assert workbook.worksheets[0]["A1"].font.bold is True
    assert workbook.worksheets[0].column_dimensions["A"].width >= 10
    workbook.close()
    with zipfile.ZipFile(BytesIO(build_input_template())) as archive:
        assert archive.testzip() is None
        assert "xl/calcChain.xml" not in archive.namelist()


def test_preview_returns_source_rows_and_unit_separated_summaries():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path)
        value = parse_input_workbook(path, start_month=7, end_month=12, source_filename="input.xlsx")
        assert value.valid is True
        assert value.source_filename == "input.xlsx"
        assert value.sales_rows[0].source_sheet == "판매계획"
        assert value.sales_rows[0].source_row == 2
        assert value.business_production_rows[0].source_sheet == "생산계획"
        assert value.sales_summary[0].unit == "PCS"
        assert value.sales_summary[0].quantity_total == 12
        assert value.production_summary[0].unit == "m"
        assert value.production_summary[0].quantity_total == 100
    finally:
        path.unlink(missing_ok=True)


def test_metadata_only_template_is_structured_blocking_empty_input():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        path.write_bytes(build_input_template())
        value = parse_input_workbook(path, start_month=7, end_month=7, source_filename="input.xlsx")
        assert value.valid is False and value.blocking is True
        assert value.sales_rows == () and value.business_production_rows == ()
        assert value.issues[0].code == "empty_input"
    finally:
        path.unlink(missing_ok=True)


def test_preview_rejects_unknown_metadata_duplicate_and_invalid_numbers():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path)
        workbook = load_workbook(path)
        sales = workbook.worksheets[0]
        sales["B2"] = "UNKNOWN"
        sales["A3"] = 7
        sales["B3"] = "SW400"
        sales["C3"] = "SW400"
        sales["D3"] = "SW"
        sales["E3"] = -1
        sales["F3"] = "NaN"
        workbook.save(path)
        workbook.close()
        value = parse_input_workbook(path, start_month=7, end_month=12, source_filename="input.xlsx")
        codes = {item.code for item in value.issues}
        assert value.valid is False and value.blocking is True
        assert {"unknown_product_code", "negative_numeric", "invalid_numeric"} <= codes
    finally:
        path.unlink(missing_ok=True)


def test_preview_rejects_formulas_and_wrong_sheet_shape():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path, extra_sheet=True)
        with pytest.raises(Exception):
            parse_input_workbook(path, start_month=7, end_month=12, source_filename="input.xlsx")

        workbook = load_workbook(BytesIO(build_input_template()))
        workbook.worksheets[0]["E2"] = "=1+1"
        workbook.save(path)
        workbook.close()
        with pytest.raises(Exception):
            parse_input_workbook(path, start_month=7, end_month=12, source_filename="input.xlsx")
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize("variant", ["missing_sheet", "wrong_header", "calc_chain", "macro", "external_link", "malformed"])
def test_preview_rejects_structurally_unsafe_ooxml_variants(variant):
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        if variant == "malformed":
            path.write_bytes(b"not-a-zip")
        elif variant in {"missing_sheet", "wrong_header"}:
            workbook = load_workbook(BytesIO(build_input_template()))
            if variant == "missing_sheet":
                workbook.remove(workbook.worksheets[1])
            else:
                workbook.worksheets[0]["A1"] = "wrong"
            workbook.save(path)
            workbook.close()
        elif variant in {"calc_chain", "macro"}:
            source = BytesIO(build_input_template())
            with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as output:
                for item in original.infolist():
                    output.writestr(item, original.read(item.filename))
                output.writestr("xl/calcChain.xml" if variant == "calc_chain" else "xl/vbaProject.bin", b"unsafe")
        else:
            source = BytesIO(build_input_template())
            relationship_part = "xl/_rels/workbook.xml.rels"
            marker = b"</Relationships>"
            external = (
                b"<Relationship Id='rIdExternal' "
                b"Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink' "
                b"Target='https://example.invalid/' TargetMode=' External '/>"
            )
            with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as output:
                for item in original.infolist():
                    payload = original.read(item.filename)
                    if item.filename == relationship_part:
                        payload = payload.replace(marker, external + marker)
                    output.writestr(item, payload)
        with pytest.raises(Exception):
            parse_input_workbook(path, start_month=7, end_month=7, source_filename="input.xlsx")
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize("sheet_name", ["판매계획", "생산계획"])
def test_preview_rejects_each_missing_required_sheet(sheet_name):
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        workbook = load_workbook(BytesIO(build_input_template()))
        workbook.remove(workbook[sheet_name])
        workbook.save(path)
        workbook.close()
        with pytest.raises(Exception):
            parse_input_workbook(path, start_month=7, end_month=7, source_filename="input.xlsx")
    finally:
        path.unlink(missing_ok=True)


def test_preview_period_is_bounded_to_six_months():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path)
        with pytest.raises(Exception, match="too long"):
            parse_input_workbook(path, start_month=1, end_month=7, source_filename="input.xlsx")
    finally:
        path.unlink(missing_ok=True)


def test_month_text_08_is_normalized_but_numeric_text_is_not():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path)
        workbook = load_workbook(path)
        sales = workbook.worksheets[0]
        sales["A2"] = "08"
        sales["E2"] = 12
        sales["F2"] = 99
        workbook.worksheets[1]["A2"] = 8
        workbook.save(path)
        workbook.close()
        value = parse_input_workbook(path, start_month=8, end_month=8, source_filename="input.xlsx")
        assert value.valid is True
        assert value.sales_rows[0].month == 8
        workbook = load_workbook(path)
        workbook.worksheets[0]["E2"] = "12"
        workbook.save(path)
        workbook.close()
        invalid = parse_input_workbook(path, start_month=8, end_month=8, source_filename="input.xlsx")
        assert invalid.valid is False
        assert invalid.sales_rows == ()
        assert any(item.code == "invalid_numeric" for item in invalid.issues)
    finally:
        path.unlink(missing_ok=True)


def test_http_preview_is_admin_csrf_and_exact_origin_protected():
    sessions = AccessCodeSessionService(
        viewer_code="viewer-code",
        admin_code="admin-code",
        actor_namespace_secret="actor-namespace-secret-at-least-32-chars",
        ttl_seconds=3600,
    )
    application = TrustedBffApplication(
        sessions, None, None, None,
        forecast_input_preview=ForecastInputPreviewService(sessions),
    )
    app = create_http_bff(
        application,
        settings=HttpBffSettings(
            environment="test",
            csrf_secret="csrf-secret-at-least-32-characters",
            allowed_origins=("https://frontend.test",),
        ),
    )
    client = TestClient(app)
    assert client.get("/api/admin/forecasts/input-template").status_code == 401
    client.post("/api/session/login", json={"access_code": "admin-code"})
    template = build_input_template()
    response = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "12"},
        files={"file": ("input.xlsx", template, MIME_XLSX)},
        headers={"X-CSRF-Token": client.cookies.get("pnl_csrf")},
    )
    assert response.status_code == 403
    response = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "12"},
        files={"file": ("input.xlsx", template, MIME_XLSX)},
        headers={
            "X-CSRF-Token": client.cookies.get("pnl_csrf"),
            "Origin": "https://frontend.test",
        },
    )
    assert response.status_code == 200
    assert response.json()["valid"] is False
    assert response.json()["blocking"] is True
    assert response.json()["issues"][0]["code"] == "empty_input"
    assert response.json()["business_production_rows"] == []


def _http_client(*, origins=("https://frontend.test",), max_bytes=1024 * 1024):
    sessions = AccessCodeSessionService(
        viewer_code="viewer-code",
        admin_code="admin-code",
        actor_namespace_secret="actor-namespace-secret-at-least-32-chars",
        ttl_seconds=3600,
    )
    application = TrustedBffApplication(
        sessions, None, None, None,
        forecast_input_preview=ForecastInputPreviewService(sessions),
    )
    app = create_http_bff(
        application,
        settings=HttpBffSettings(
            environment="test",
            csrf_secret="csrf-secret-at-least-32-characters",
            allowed_origins=origins,
            forecast_request_max_bytes=max_bytes,
        ),
    )
    return TestClient(app)


def test_template_admin_download_headers_and_viewer_boundary():
    anonymous = _http_client()
    assert anonymous.get("/api/admin/forecasts/input-template").status_code == 401
    viewer = _http_client()
    viewer.post("/api/session/login", json={"access_code": "viewer-code"})
    assert viewer.get("/api/admin/forecasts/input-template").status_code == 403
    admin = _http_client()
    admin.post("/api/session/login", json={"access_code": "admin-code"})
    response = admin.get("/api/admin/forecasts/input-template")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(MIME_XLSX)
    assert "attachment" in response.headers["content-disposition"]
    assert "forecast_input_template.xlsx" in response.headers["content-disposition"]
    assert "단가" not in response.content.decode("latin-1", errors="ignore")


@pytest.mark.parametrize(
    "filename, expected",
    [("input.xlsm", 422), ("input.csv", 422), ("input.xlsx", 422)],
)
def test_preview_rejects_extension_or_malformed_package(filename, expected):
    client = _http_client()
    client.post("/api/session/login", json={"access_code": "admin-code"})
    content = b"not-an-ooxml-package" if filename.endswith(".xlsx") else build_input_template()
    response = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "7"},
        files={"file": (filename, content, MIME_XLSX)},
        headers={"X-CSRF-Token": client.cookies.get("pnl_csrf"), "Origin": "https://frontend.test"},
    )
    assert response.status_code == expected


def test_preview_missing_csrf_invalid_origin_and_oversize_are_rejected():
    client = _http_client(max_bytes=16 * 1024)
    client.post("/api/session/login", json={"access_code": "admin-code"})
    payload = build_input_template()
    missing_csrf = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "7"},
        files={"file": ("input.xlsx", payload, MIME_XLSX)},
        headers={"Origin": "https://frontend.test"},
    )
    assert missing_csrf.status_code == 403
    invalid_origin = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "7"},
        files={"file": ("input.xlsx", payload, MIME_XLSX)},
        headers={"X-CSRF-Token": client.cookies.get("pnl_csrf"), "Origin": "https://evil.test"},
    )
    assert invalid_origin.status_code == 403
    oversized = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "7"},
        files={"file": ("input.xlsx", b"x" * 20_000, MIME_XLSX)},
        headers={"X-CSRF-Token": client.cookies.get("pnl_csrf"), "Origin": "https://frontend.test"},
    )
    assert oversized.status_code == 413


def test_preview_invalid_dimensions_metadata_duplicates_and_missing_values_are_blocking():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path)
        workbook = load_workbook(path)
        sales, production = workbook.worksheets
        sales["A2"] = 6  # out of requested period
        sales["C2"] = "wrong name"
        sales["D2"] = "wrong group"
        sales["A3"] = 7
        sales["B3"] = "SW400"  # duplicate code/month
        sales["C3"] = "SW400"
        sales["D3"] = "SW"
        sales["E3"] = None
        sales["F3"] = None
        production["A2"] = 7
        production["B2"] = "bad process"
        production["C2"] = "SW"
        production["D2"] = -1
        production["E2"] = "PCS"
        workbook.save(path)
        workbook.close()
        value = parse_input_workbook(path, start_month=7, end_month=7, source_filename="input.xlsx")
        codes = {item.code for item in value.issues}
        assert value.blocking is True and value.valid is False
        assert {"month_out_of_period", "product_name_mismatch", "product_group_mismatch", "missing_numeric", "invalid_production_dimension", "negative_numeric"} <= codes
    finally:
        path.unlink(missing_ok=True)


def test_preview_reports_invalid_month_unit_and_sales_and_production_duplicates():
    path = Path.cwd() / f".forecast-input-{uuid4().hex}.xlsx"
    try:
        _write_workbook(path)
        workbook = load_workbook(path)
        sales, production = workbook.worksheets

        # Keep the first template rows valid, then append exact duplicate keys.
        sales.append((7, "SW400", "SW400", "SW", 2, 3))
        production.append((7, "전공정", "SW", 2, "m"))
        # Independent malformed rows prove strict month and unit handling.
        sales.append((13, "SW440", "SW440", "SW", 2, 3))
        production.append((7, "후공정", "BW", 2, "m"))
        workbook.save(path)
        workbook.close()

        value = parse_input_workbook(path, start_month=7, end_month=7, source_filename="input.xlsx")
        codes = {item.code for item in value.issues}
        assert {
            "duplicate_sales_row",
            "duplicate_production_row",
            "invalid_month",
            "production_unit_mismatch",
        } <= codes
        assert value.blocking is True
    finally:
        path.unlink(missing_ok=True)


def test_preview_post_rejects_viewer():
    client = _http_client()
    client.post("/api/session/login", json={"access_code": "viewer-code"})
    response = client.post(
        "/api/admin/forecasts/input-preview",
        data={"start_month": "7", "end_month": "7"},
        files={"file": ("input.xlsx", build_input_template(), MIME_XLSX)},
        headers={"X-CSRF-Token": client.cookies.get("pnl_csrf"), "Origin": "https://frontend.test"},
    )
    assert response.status_code == 403
