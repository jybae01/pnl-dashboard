from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from forecast.analysis_export import build_comparison_audit_workbook
from forecast.comparison import GenericComparisonEngine, PeriodOption
from forecast.evidence_reporting import audit_reporting_workbook, collect_reporting_sources

try:
    from tests.test_golden_analysis_adapter import _build_workbook, _meta
except ModuleNotFoundError:
    from test_golden_analysis_adapter import _build_workbook, _meta


ROOT = Path(__file__).resolve().parents[1]
OFFICIAL_EFFECTS = (
    "sales_quantity", "sales_mix", "sales_price", "sales_fx", "tariff",
    "material_total", "manufacturing_realized", "inventory_timing",
    "sga_variable", "sga_fixed",
)


def _n(value):
    return float(value or 0.0)


def _sales_projection(result) -> dict[str, float]:
    sales = result.sales_analysis
    trace = list(sales["trace_rows"])
    pools: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in trace:
        pools[(str(row["period"]), str(row["pool"]))].append(row)
    quantity = 0.0
    mix = 0.0
    displayed_price = 0.0
    sales_fx = 0.0
    for rows in pools.values():
        base_total = sum(_n(row["base_quantity"]) for row in rows)
        comparison_total = sum(_n(row["comparison_quantity"]) for row in rows)
        weighted_gp = 0.0
        mix_component = 0.0
        for row in rows:
            base_quantity = _n(row["base_quantity"])
            comparison_quantity = _n(row["comparison_quantity"])
            base_gp = (
                (_n(row["base_revenue"]) - _n(row["base_cogs"])) / base_quantity
                if base_quantity else 0.0
            )
            base_mix = base_quantity / base_total if base_total else 0.0
            comparison_mix = (
                comparison_quantity / comparison_total if comparison_total else 0.0
            )
            weighted_gp += base_mix * base_gp
            mix_component += (comparison_mix - base_mix) * base_gp
            base_fx = _n(row["base_fx"])
            comparison_fx = _n(row["comparison_fx"])
            base_price = _n(row["base_revenue"]) / base_quantity if base_quantity else 0.0
            comparison_price = (
                _n(row["comparison_revenue"]) / comparison_quantity
                if comparison_quantity else 0.0
            )
            base_foreign = base_price / base_fx if base_fx else 0.0
            comparison_foreign = comparison_price / comparison_fx if comparison_fx else 0.0
            displayed_price += (
                comparison_quantity
                * (comparison_foreign - base_foreign)
                * (base_fx + comparison_fx)
                / 2
            )
            sales_fx += (
                comparison_quantity
                * (comparison_fx - base_fx)
                * (base_foreign + comparison_foreign)
                / 2
            )
        quantity += (comparison_total - base_total) * weighted_gp
        mix += comparison_total * mix_component

    new_business = list(sales.get("new_business_trace_rows") or [])
    if not new_business:
        new_business = list(result.sales_cogs_scope_analysis.get("new_business_rows") or [])
    for row in new_business:
        base_revenue = _n(row["base_revenue"])
        comparison_revenue = _n(row["comparison_revenue"])
        base_gp = base_revenue - _n(row["base_cogs"])
        comparison_gp = comparison_revenue - _n(row["comparison_cogs"])
        base_rate = base_gp / base_revenue if base_revenue else 0.0
        comparison_rate = comparison_gp / comparison_revenue if comparison_revenue else 0.0
        quantity += (comparison_revenue - base_revenue) * base_rate
        displayed_price += comparison_revenue * (comparison_rate - base_rate)

    freight = 0.0
    tariff = 0.0
    for row in sales["freight_trace_rows"]:
        base_ex = _n(row["base_freight_including_tariff"]) - _n(row["base_tariff"])
        comparison_ex = (
            _n(row["comparison_freight_including_tariff"])
            - _n(row["comparison_tariff"])
        )
        base_equivalent = (
            _n(row["base_sw_pcs"]) + _n(row["base_bw_pcs"])
            + _n(row["base_lc_pcs"]) + _n(row["base_fs_length"]) / 45
        )
        comparison_equivalent = (
            _n(row["comparison_sw_pcs"]) + _n(row["comparison_bw_pcs"])
            + _n(row["comparison_lc_pcs"]) + _n(row["comparison_fs_length"]) / 45
        )
        base_unit = base_ex / base_equivalent if base_equivalent else 0.0
        comparison_unit = (
            comparison_ex / comparison_equivalent if comparison_equivalent else 0.0
        )
        freight += (base_unit - comparison_unit) * comparison_equivalent
        tariff += _n(row["base_tariff"]) - _n(row["comparison_tariff"])
    return {
        "sales_quantity": quantity,
        "sales_mix": mix,
        "sales_price": displayed_price + freight,
        "sales_fx": sales_fx,
        "tariff": tariff,
    }


def _material_projection(result) -> float:
    total = 0.0
    for row in result.material_analysis["trace_rows"]:
        base_cost = _n(row["base_cost"])
        comparison_cost = _n(row["comparison_cost"])
        base_output = _n(row["base_output"])
        comparison_output = _n(row["comparison_output"])
        if (base_cost and not base_output) or (comparison_cost and not comparison_output):
            continue
        base_unit = base_cost / base_output if base_output else 0.0
        comparison_unit = comparison_cost / comparison_output if comparison_output else 0.0
        total += (base_unit - comparison_unit) * _n(row["comparison_sales"])
    return total


def _manufacturing_projection(result) -> float:
    total = 0.0
    for row in result.manufacturing_analysis["trace_rows"]:
        base = _n(row["baseline_amount"])
        comparison = _n(row["comparison_amount"])
        base_ratio = _n(row["front_ratio_base"])
        comparison_ratio = _n(row["front_ratio_comparison"])
        base_front = base * base_ratio
        comparison_front = comparison * comparison_ratio
        base_back = base - base_front
        comparison_back = comparison - comparison_front
        if row["classification"] == "fixed":
            total += base - comparison
            continue
        base_front_activity = _n(row.get("base_front_activity"))
        comparison_front_activity = _n(row.get("comparison_front_activity"))
        base_back_activity = _n(row.get("base_back_activity"))
        comparison_back_activity = _n(row.get("comparison_back_activity"))
        base_front_unit = base_front / base_front_activity if base_front_activity else 0.0
        comparison_front_unit = (
            comparison_front / comparison_front_activity
            if comparison_front_activity else 0.0
        )
        base_back_unit = base_back / base_back_activity if base_back_activity else 0.0
        comparison_back_unit = (
            comparison_back / comparison_back_activity
            if comparison_back_activity else 0.0
        )
        front_activity = (
            (base_front_activity - comparison_front_activity) * base_front_unit
            if base_front_activity and comparison_front_activity else 0.0
        )
        back_activity = (
            (base_back_activity - comparison_back_activity) * base_back_unit
            if base_back_activity and comparison_back_activity else 0.0
        )
        front_unit = (
            comparison_front_activity * (base_front_unit - comparison_front_unit)
            if base_front_activity and comparison_front_activity
            else base_front - comparison_front
        )
        back_unit = (
            comparison_back_activity * (base_back_unit - comparison_back_unit)
            if base_back_activity and comparison_back_activity
            else base_back - comparison_back
        )
        total += front_activity + back_activity + front_unit + back_unit
    return total


def _inventory_projection(result) -> float:
    source_rows = result.inventory_analysis["source_details"]
    def source_total(side: str, canonical: str) -> float:
        return sum(
            _n(row["value"]) for row in source_rows
            if row["side"] == side and row["canonical_field"] == canonical
        )
    base_manufactured = (
        source_total("BASE", "finished_goods_cogs")
        + source_total("BASE", "semi_finished_goods_cogs")
    )
    comparison_manufactured = (
        source_total("COMPARISON", "finished_goods_cogs")
        + source_total("COMPARISON", "semi_finished_goods_cogs")
    )
    current_effect = (
        source_total("BASE", "current_manufacturing_cost")
        - source_total("COMPARISON", "current_manufacturing_cost")
    )
    gross = base_manufactured - comparison_manufactured - current_effect
    core = 0.0
    for row in result.inventory_analysis["core_overlap_details"]:
        if not row["selected"]:
            continue
        base_quantity = _n(row["base_quantity"])
        comparison_quantity = _n(row["comparison_quantity"])
        pool_base = _n(row["pool_base_quantity"])
        pool_comparison = _n(row["pool_comparison_quantity"])
        base_cogs_per_unit = (
            _n(row["base_core_manufactured_cogs"]) / base_quantity
            if base_quantity else 0.0
        )
        base_mix = base_quantity / pool_base if pool_base else 0.0
        comparison_mix = comparison_quantity / pool_comparison if pool_comparison else 0.0
        quantity_expense = (pool_comparison - pool_base) * base_mix * base_cogs_per_unit
        mix_expense = pool_comparison * (comparison_mix - base_mix) * base_cogs_per_unit
        core += -(quantity_expense + mix_expense)
    return gross - core


def _sga_projection(result) -> dict[str, float]:
    totals = {"variable": 0.0, "fixed": 0.0}
    for row in result.sga_monthly_trace:
        classification = str(row["classification"])
        if classification not in totals:
            continue
        if str(row["bridge_position"]) in {"판매효과", "외부효과/관세"}:
            continue
        totals[classification] += _n(row["base_amount"]) - _n(row["comparison_amount"])
    return {"sga_variable": totals["variable"], "sga_fixed": totals["fixed"]}


def test_formula_projection_matches_backend_authoritative_result(tmp_path):
    baseline = tmp_path / "baseline.xlsx"
    comparison = tmp_path / "comparison.xlsx"
    _build_workbook(baseline, comparison=False)
    _build_workbook(comparison, comparison=True)
    result = GenericComparisonEngine(ROOT / "config" / "model_mapping.json").compare(
        _meta("base"), baseline, _meta("comparison"), comparison,
        PeriodOption("M2026_01", "2026-01", (1,), "사용자정의"),
        baseline_sales_fx=1_400.0, comparison_sales_fx=1_450.0,
    )
    projected = {
        **_sales_projection(result),
        "material_total": _material_projection(result),
        "manufacturing_realized": _manufacturing_projection(result),
        "inventory_timing": _inventory_projection(result),
        **_sga_projection(result),
    }
    authoritative = {
        row["code"]: float(row["profit_effect"]) for row in result.effects
    }
    assert tuple(projected) == OFFICIAL_EFFECTS
    for code in OFFICIAL_EFFECTS:
        assert projected[code] == pytest.approx(authoritative[code], abs=1e-6)
        assert projected[code] / 1000 == pytest.approx(authoritative[code] / 1000)
    assert sum(projected.values()) == pytest.approx(result.effects_total, abs=1e-6)
    assert result.effects_total + result.residual == pytest.approx(
        result.operating_profit_delta, abs=1e-6
    )

    payload = build_comparison_audit_workbook(
        result=asdict(result),
        sales_rows=result.sales_analysis["rows"],
        sales_totals=result.sales_analysis["totals"],
        baseline_fx=1_400.0,
        comparison_fx=1_450.0,
        mapping_path=ROOT / "config" / "model_mapping.json",
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    audit = audit_reporting_workbook(workbook)
    assert audit["hard_coded_derived_duplicates"] == 0
    assert audit["formula_error_count"] == 0

    effect_sheet = workbook["02_손익영향"]
    for code in OFFICIAL_EFFECTS:
        row = next(
            row for row in range(1, effect_sheet.max_row + 1)
            if effect_sheet[f"G{row}"].value == code
        )
        formula = effect_sheet[f"C{row}"].value
        assert formula.startswith("='03_판매근거'!") or formula.startswith("='04_원가근거'!")
    summary = workbook["01_보고요약"]
    assert all(
        not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool)
        for row in summary.iter_rows()
        for cell in row
    )
    cost = workbook["04_원가근거"]
    manufacturing_section = next(
        row for row in range(1, cost.max_row + 1)
        if cost[f"A{row}"].value == "C. 제조경비"
    )
    activity_row = manufacturing_section + 2
    assert cost[f"D{activity_row}"].value.startswith("=")
    assert cost[f"F{activity_row}"].value.startswith("=MAX(")
    normal = next(
        row for row in result.manufacturing_analysis["trace_rows"]
        if row.get("back_activity_basis") == "BACK"
        and row.get("classification") == "variable"
    )
    outsourcing = next(
        row for row in result.manufacturing_analysis["trace_rows"]
        if row.get("back_activity_basis") == "OUTSOURCING_BACK"
    )
    normal_row = next(
        row for row in range(1, cost.max_row + 1)
        if cost[f"B{row}"].value == normal["account"]
    )
    outsourcing_row = next(
        row for row in range(1, cost.max_row + 1)
        if cost[f"B{row}"].value == outsourcing["account"]
    )
    assert cost[f"O{normal_row}"].value == f"=D{activity_row}"
    assert cost[f"P{normal_row}"].value == f"=E{activity_row}"
    assert cost[f"O{outsourcing_row}"].value == f"=F{activity_row}"
    assert cost[f"P{outsourcing_row}"].value == f"=G{activity_row}"
    assert normal["base_back_activity"] != outsourcing["base_back_activity"]


def test_production_weighted_formula_uses_amount_sum_over_quantity_sum(tmp_path):
    baseline = tmp_path / "baseline.xlsx"
    comparison = tmp_path / "comparison.xlsx"
    _build_workbook(baseline, comparison=False)
    _build_workbook(comparison, comparison=True)
    result = GenericComparisonEngine(ROOT / "config" / "model_mapping.json").compare(
        _meta("base"), baseline, _meta("comparison"), comparison,
        PeriodOption("M2026_01", "2026-01", (1,), "사용자정의"),
        baseline_sales_fx=1_400.0, comparison_sales_fx=1_450.0,
    )
    by_group = {row["production_basis"]: row for row in result.production_evidence}
    for group in ("FS", "SW", "BW", "LC"):
        row = by_group[group]
        components = row["source_components"]
        base_quantity = sum(
            _n(item["baseline_value"]) for item in components if item["kind"] == "quantity"
        )
        base_amount = sum(
            _n(item["baseline_value"]) for item in components if item["kind"] == "amount"
        )
        expected = base_amount / base_quantity if base_quantity else None
        assert row["baseline_weighted_unit_cost"] == pytest.approx(expected) if expected is not None else row["baseline_weighted_unit_cost"] is None
    back = by_group["SW+BW+LC"]
    assert back["baseline_quantity"] == sum(by_group[group]["baseline_quantity"] for group in ("SW", "BW", "LC"))
    assert back["baseline_amount"] == sum(by_group[group]["baseline_amount"] for group in ("SW", "BW", "LC"))
    assert back["baseline_weighted_unit_cost"] == pytest.approx(
        back["baseline_amount"] / back["baseline_quantity"]
    )
    assert by_group["FS"]["baseline_quantity"] not in {
        back["baseline_quantity"],
        sum(by_group[group]["baseline_quantity"] for group in ("FS", "SW", "BW", "LC")),
    }


@pytest.mark.parametrize("invalid_cogs", [None, float("nan"), float("inf"), "not-a-number"])
def test_legacy_missing_raw_cogs_fails_closed(invalid_cogs):
    result = {
        "period": {"label": "2026-01", "months": [1]},
        "pnl": [],
        "sales_analysis": {},
        "sales_groups": [{
            "product_group": "SW",
            "baseline_quantity": 10,
            "comparison_quantity": 11,
            "baseline_amount": 100,
            "comparison_amount": 120,
            "baseline_cogs": invalid_cogs,
            "comparison_cogs": invalid_cogs,
        }],
    }
    with pytest.raises(ValueError, match="legacy raw COGS is unavailable"):
        collect_reporting_sources(result, [], 1_400.0, 1_450.0)
