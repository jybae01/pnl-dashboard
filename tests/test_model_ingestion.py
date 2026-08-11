from __future__ import annotations

import hashlib
import threading
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from forecast.bff.auth import AccessCodeSessionService
from forecast.bff.dto import ModelUploadRequest
from forecast.bff.errors import ApiErrorCode, BffError
from forecast.bff.model_ingestion import (
    MAX_WORKBOOK_BYTES,
    ModelIngestionConflictError,
    ModelIngestionFinalizeUncertainError,
    ModelIngestionInProgressError,
    ModelIngestionReservation,
    ModelIngestionService,
    ModelManagementService,
    ModelPublicationService,
    _validate_xlsx_package,
)
from forecast.provenance import ResultProvenance


MODEL_ID = "11111111-1111-4111-8111-111111111111"
PROVENANCE = ResultProvenance("1.1.0", "analysis-v1.0.0", "a" * 64, "1")


def workbook(path: Path, *, year: int = 2026) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    for month, column in enumerate(range(5, 17), 1):
        sheet.cell(2, column, f"{year}년 {month}월")
        sheet.cell(3, column, "실적")
    book.save(path)
    book.close()
    return path.read_bytes()


class Validator:
    def require(self, path, *, expected_year=None):
        assert Path(path).is_file() and expected_year == 2026
        return SimpleNamespace(passed=True)


class Repository:
    def __init__(self):
        self.rows = {}

    def list(self):
        return list(self.rows.values())

    def get(self, model_id):
        if model_id not in self.rows:
            raise KeyError(model_id)
        return self.rows[model_id]

    def set_publication(self, model_id, *, is_published, is_default=False):
        row = self.get(model_id)
        if is_default and not is_published:
            raise ValueError("default requires published")
        row.is_published = is_published
        row.confirmed = is_published
        row.is_default = is_default
        return row


class Gateway:
    def __init__(self, repository):
        self.repository = repository
        self.lock = threading.Lock()
        self.requests = {}
        self.objects = {}
        self.finalized = 0
        self.fail_upload = False
        self.fail_finalize = False
        self.fail_cleanup = False
        self.fail_diagnostic = False
        self.malformed_finalize = False
        self.finalize_response_lost = False
        self.wrong_finalize_sha = False
        self.failures = []

    def reserve(self, *, actor, request, workbook_sha256, period_types, provenance):
        key = (actor, request.idempotency_key)
        fingerprint = (request, workbook_sha256, tuple(sorted(period_types.items())), provenance)
        with self.lock:
            existing = self.requests.get(key)
            if existing:
                if existing[0] != fingerprint:
                    raise ModelIngestionConflictError
                if existing[1] == "completed":
                    return ModelIngestionReservation("ingestion", MODEL_ID, "completed", None, True)
                raise ModelIngestionInProgressError
            reservation = ModelIngestionReservation("ingestion", MODEL_ID, "reserved", "lease", False)
            self.requests[key] = [fingerprint, "reserved"]
            return reservation

    def upload_source(self, model_id, source, workbook_sha256):
        if self.fail_upload:
            raise OSError("storage unavailable")
        self.objects[model_id] = Path(source).read_bytes()

    def verify_source(self, model_id, workbook_sha256):
        if hashlib.sha256(self.objects[model_id]).hexdigest() != workbook_sha256:
            raise ValueError("mismatch")

    def finalize(self, reservation):
        if self.fail_finalize:
            raise RuntimeError("database failure")
        payload = self.objects[reservation.model_id]
        row = SimpleNamespace(
            id=reservation.model_id, name="2026 Actual", model_type="ACTUAL", year=2026,
            start_month=1, end_month=12, version="V1", file_name="actual.xlsx",
            workbook_sha256=hashlib.sha256(payload).hexdigest(), confirmed=False,
            is_published=False, is_default=False, uploaded_at="2026-08-11T00:00:00Z",
        )
        self.repository.rows[row.id] = row
        if self.wrong_finalize_sha:
            row.workbook_sha256 = "f" * 64
        self.finalized += 1
        for value in self.requests.values():
            value[1] = "completed"
        if self.finalize_response_lost:
            raise ModelIngestionFinalizeUncertainError
        if self.malformed_finalize:
            return {"id": row.id}
        return row

    def remove_source(self, model_id):
        if self.fail_cleanup:
            raise OSError("cleanup failure")
        self.objects.pop(model_id, None)

    def record_failure(self, reservation, **values):
        if self.fail_diagnostic:
            raise OSError("diagnostic failure")
        if any(value[1] == "completed" for value in self.requests.values()):
            raise RuntimeError("completed ingestion is immutable")
        self.failures.append(values)
        for value in self.requests.values():
            value[1] = "failed" if values["cleanup_succeeded"] else "cleanup_required"


def fixture():
    sessions = AccessCodeSessionService(
        viewer_code="viewer", admin_code="admin",
        actor_namespace_secret="stable-actor-namespace-secret-32chars", ttl_seconds=3600,
    )
    admin = sessions.login("admin").session_id
    viewer = sessions.login("viewer").session_id
    repository = Repository()
    gateway = Gateway(repository)
    service = ModelIngestionService(sessions, repository, gateway, Validator(), PROVENANCE)
    return sessions, admin, viewer, repository, gateway, service


def request(**overrides):
    values = dict(name="2026 Actual", model_type="ACTUAL", model_year=2026,
                  version="V1", file_name="actual.xlsx", idempotency_key="upload-1")
    values.update(overrides)
    return ModelUploadRequest(**values)


def test_valid_upload_stores_exact_bytes_sha_and_creates_unpublished_model(tmp_path):
    path = tmp_path / "actual.xlsx"
    exact = workbook(path)
    _, admin, _, _, gateway, service = fixture()
    result = service.ingest(admin, request(), path)
    assert gateway.objects[MODEL_ID] == exact
    assert result.model.workbook_sha256 == hashlib.sha256(exact).hexdigest()
    assert result.model.is_published is False and result.model.is_default is False
    assert gateway.finalized == 1


def test_invalid_package_error_never_leaks_server_path(tmp_path):
    path = tmp_path / "private-server-name.xlsx"
    path.write_bytes(b"not-a-zip")
    with pytest.raises(BffError) as raised:
        _validate_xlsx_package(path, "upload.xlsx")
    assert raised.value.error.field_errors == {"file": "invalid_xlsx"}
    assert str(tmp_path) not in str(raised.value.error.field_errors)


def test_response_mapping_failure_after_finalize_does_not_remove_durable_source(tmp_path):
    path = tmp_path / "actual.xlsx"
    exact = workbook(path)
    _, admin, _, repository, gateway, service = fixture()
    gateway.malformed_finalize = True

    with pytest.raises(BffError) as raised:
        service.ingest(admin, request(), path)

    assert raised.value.code == ApiErrorCode.INPUT_INTEGRITY_MISMATCH
    assert repository.get(MODEL_ID).workbook_sha256 == hashlib.sha256(exact).hexdigest()
    assert gateway.objects[MODEL_ID] == exact
    assert gateway.failures == []


def test_finalize_response_must_match_reserved_identity_and_exact_sha(tmp_path):
    path = tmp_path / "actual.xlsx"
    exact = workbook(path)
    _, admin, _, _, gateway, service = fixture()
    gateway.wrong_finalize_sha = True
    with pytest.raises(BffError) as raised:
        service.ingest(admin, request(), path)
    assert raised.value.code == ApiErrorCode.INPUT_INTEGRITY_MISMATCH
    assert gateway.objects[MODEL_ID] == exact


def test_lost_finalize_response_never_deletes_a_possibly_committed_source(tmp_path):
    path = tmp_path / "actual.xlsx"
    exact = workbook(path)
    _, admin, _, repository, gateway, service = fixture()
    gateway.finalize_response_lost = True

    with pytest.raises(BffError) as raised:
        service.ingest(admin, request(), path)

    assert raised.value.code == ApiErrorCode.INGESTION_CLEANUP_REQUIRED
    assert repository.get(MODEL_ID).workbook_sha256 == hashlib.sha256(exact).hexdigest()
    assert gateway.objects[MODEL_ID] == exact
    gateway.finalize_response_lost = False
    replay = service.ingest(admin, request(), path)
    assert replay.idempotency_replayed is True
    assert replay.model.model_id == MODEL_ID


def test_idempotency_replay_conflicts_and_same_sha_new_key_is_new_business_identity(tmp_path):
    path = tmp_path / "actual.xlsx"
    workbook(path)
    _, admin, _, _, gateway, service = fixture()
    first = service.ingest(admin, request(), path)
    replay = service.ingest(admin, request(), path)
    assert replay.model.model_id == first.model.model_id and replay.idempotency_replayed is True
    changed = tmp_path / "changed.xlsx"
    workbook(changed, year=2027)
    with pytest.raises(BffError) as collision:
        service.ingest(admin, request(), changed)
    assert collision.value.code == ApiErrorCode.IDEMPOTENCY_CONFLICT
    with pytest.raises(BffError) as metadata_collision:
        service.ingest(admin, request(name="Other"), path)
    assert metadata_collision.value.code == ApiErrorCode.IDEMPOTENCY_CONFLICT
    # A different key is not deduplicated by SHA; the fake's fixed UUID models
    # only the fact that a second reservation/finalization is attempted.
    with gateway.lock:
        gateway.requests.clear()
    service.ingest(admin, request(idempotency_key="upload-2"), path)
    assert gateway.finalized == 2


def test_concurrent_same_key_has_one_finalizer(tmp_path):
    path = tmp_path / "actual.xlsx"
    workbook(path)
    _, admin, _, _, gateway, service = fixture()
    barrier = threading.Barrier(2)
    original = gateway.upload_source

    def slow_upload(*args):
        barrier.wait(timeout=5)
        original(*args)

    gateway.upload_source = slow_upload
    outcomes = []

    def run():
        try:
            outcomes.append(service.ingest(admin, request(), path))
        except BffError as exc:
            outcomes.append(exc.code)

    first = threading.Thread(target=run)
    first.start()
    # The second request observes the active reservation; release the first
    # from its barrier after the duplicate has been rejected.
    second = threading.Thread(target=run)
    second.start()
    try:
        barrier.wait(timeout=5)
    except threading.BrokenBarrierError:
        pass
    first.join(5); second.join(5)
    assert gateway.finalized == 1
    assert len(outcomes) == 2
    assert all(
        value == ApiErrorCode.TRANSIENT_SYSTEM_ERROR
        or (getattr(value, "model", None) is not None and value.model.model_id == MODEL_ID)
        for value in outcomes
    )


def test_storage_and_database_failure_compensation_and_cleanup_diagnostic(tmp_path):
    path = tmp_path / "actual.xlsx"
    workbook(path)
    _, admin, _, repository, gateway, service = fixture()
    gateway.fail_upload = True
    with pytest.raises(BffError):
        service.ingest(admin, request(), path)
    assert not repository.rows and not gateway.objects and gateway.failures[-1]["cleanup_succeeded"] is True

    _, admin, _, repository, gateway, service = fixture()
    gateway.fail_finalize = True
    with pytest.raises(BffError):
        service.ingest(admin, request(), path)
    assert not repository.rows and not gateway.objects and gateway.failures[-1]["cleanup_succeeded"] is True

    _, admin, _, _, gateway, service = fixture()
    gateway.fail_finalize = gateway.fail_cleanup = True
    with pytest.raises(BffError) as cleanup:
        service.ingest(admin, request(), path)
    assert cleanup.value.code == ApiErrorCode.INGESTION_CLEANUP_REQUIRED
    assert gateway.failures[-1]["cleanup_succeeded"] is False

    _, admin, _, _, gateway, service = fixture()
    gateway.fail_finalize = gateway.fail_diagnostic = True
    with pytest.raises(BffError) as diagnostic:
        service.ingest(admin, request(), path)
    assert diagnostic.value.code == ApiErrorCode.INGESTION_CLEANUP_REQUIRED
    assert not gateway.objects


def test_viewer_cannot_ingest_publish_or_list_management(tmp_path):
    path = tmp_path / "actual.xlsx"
    workbook(path)
    sessions, _, viewer, repository, gateway, service = fixture()
    with pytest.raises(BffError) as denied:
        service.ingest(viewer, request(), path)
    assert denied.value.code == ApiErrorCode.FORBIDDEN
    with pytest.raises(BffError):
        ModelManagementService(sessions, repository).list_models(viewer)
    with pytest.raises(BffError):
        ModelPublicationService(sessions, repository, gateway).set_publication(
            viewer, MODEL_ID, is_published=True
        )


def test_publication_verifies_source_and_default_requires_published(tmp_path):
    path = tmp_path / "actual.xlsx"
    workbook(path)
    sessions, admin, _, repository, gateway, ingestion = fixture()
    model = ingestion.ingest(admin, request(), path).model
    service = ModelPublicationService(sessions, repository, gateway)
    published = service.set_publication(admin, model.model_id, is_published=True)
    assert published.model.is_published is True
    with pytest.raises(BffError) as invalid_default:
        service.set_publication(admin, model.model_id, is_published=False, is_default=True)
    assert invalid_default.value.code == ApiErrorCode.VALIDATION_ERROR
    gateway.objects[MODEL_ID] += b"tampered"
    with pytest.raises(BffError) as mismatch:
        service.set_publication(admin, model.model_id, is_published=True)
    assert mismatch.value.code == ApiErrorCode.INPUT_INTEGRITY_MISMATCH


@pytest.mark.parametrize("name", ["model.xls", "model.csv", "../model.xlsx", "folder/model.xlsx"])
def test_extension_and_unsafe_filename_are_rejected(tmp_path, name):
    path = tmp_path / "payload"
    workbook(path)
    with pytest.raises(BffError):
        _validate_xlsx_package(path, name)


def test_oversize_malformed_fake_and_zip_bomb_are_rejected(tmp_path):
    oversized = tmp_path / "oversized.xlsx"
    with oversized.open("wb") as stream:
        stream.truncate(MAX_WORKBOOK_BYTES + 1)
    malformed = tmp_path / "malformed.xlsx"
    malformed.write_bytes(b"not a zip")
    fake = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(fake, "w") as archive:
        archive.writestr("hello.txt", "not OOXML")
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"0" * (2 * 1024 * 1024))
        archive.writestr("_rels/.rels", b"x")
        archive.writestr("xl/workbook.xml", b"x")
    for path in (oversized, malformed, fake, bomb):
        with pytest.raises(BffError):
            _validate_xlsx_package(path, path.name)
