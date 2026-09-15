from __future__ import annotations

import copy
import json
from pathlib import Path
import pytest

from forecast.analysis.configuration import AnalysisConfig
from forecast.analysis.golden_adapter import GoldenAnalysisAdapter
from forecast.bff.analysis_presentation import (
    AnalysisPresentationService,
    build_analysis_presentation,
)
from forecast.storage import ModelMeta
from forecast.workbook import GoldenWorkbook


def _meta(identifier: str) -> ModelMeta:
    return ModelMeta(
        identifier, identifier, "계획", 2026, 1, 12,
        "2026-01-01", "V1", True, f"{identifier}.xlsx", "now",
    )


def _load_adapter():
    mapping = json.loads(Path("config/model_mapping.json").read_text(encoding="utf-8"))
    pe_map = json.loads(
        Path("config/analysis_production_evidence_sources.json").read_text(encoding="utf-8")
    )
    cfg = AnalysisConfig.load("config/analysis_v1.json")
    return GoldenAnalysisAdapter(mapping, cfg, pe_map)


def test_case_a_lc_raw_material_mapping():
    """CASE A: LC raw-material mapping in config/model_mapping.json."""
    mapping = json.loads(Path("config/model_mapping.json").read_text(encoding="utf-8"))
    material = mapping["analysis_adapter"]["material"]
    lc_component = material["raw_material_sources"]["finished_product_components"]["LC"]
    assert lc_component["back_total_rows"] == [927]

    lc_group = material["groups"]["LC"]
    assert lc_group["back_material_terms"] == [
        {"product_code": "LC", "back_total_rows": [927]}
    ]


def test_case_b_row_928_exclusion_and_927_inclusion():
    """CASE B: LC raw_material_components contain row 927 and do NOT contain row 928."""
    adapter = _load_adapter()
    lc_spec = adapter.adapter["material"]["groups"]["LC"]
    source_map = adapter.adapter["material"]["raw_material_sources"]

    source_rows = adapter._material_source_rows(lc_spec, source_map)
    assert 927 in source_rows
    assert 928 not in source_rows


def test_case_c_and_d_numeric_invariance_on_928_and_sensitivity_on_927(tmp_path):
    """CASE C & D: Changing row 928 alone must NOT change LC material cost/effect.

    Changing row 927 must change LC material cost and effect.
    """
    adapter = _load_adapter()

    # Use data/base.xlsx and create modified copies
    wb_base_path = Path("data/base.xlsx")
    assert wb_base_path.exists()

    meta_base = _meta("base")
    meta_comp = _meta("comp")
    months = (1,)

    base_scenario = adapter.build(GoldenWorkbook(wb_base_path), meta_base, months)

    # 1. Modify row 928 in comparison workbook
    import openpyxl
    wb_928_mod_path = tmp_path / "comp_928_mod.xlsx"
    wb_928 = openpyxl.load_workbook(wb_base_path)
    # change row 928 from ~107M to 999M
    wb_928["Data"]["E928"] = 999_999_999.0
    wb_928.save(wb_928_mod_path)

    comp_928_scenario = adapter.build(GoldenWorkbook(wb_928_mod_path), meta_comp, months)
    res_928 = adapter.material_analysis(base_scenario, comp_928_scenario)
    lc_928 = next(r for r in res_928["product_groups"] if r["product_group"] == "LC")

    # BASE scenario LC cost vs comp with modified 928
    base_lc_prod = next(p for p in base_scenario.scenario.products if p.product_group == "LC")
    comp_lc_prod = next(p for p in comp_928_scenario.scenario.products if p.product_group == "LC")

    # CASE C: Changing row 928 alone does NOT change LC raw material cost or unit cost
    assert comp_lc_prod.raw_material_cost == base_lc_prod.raw_material_cost
    assert lc_928["comparison_unit_cost"] == lc_928["baseline_unit_cost"]
    assert lc_928["total"] == 0.0

    # 2. Modify row 927 in comparison workbook
    wb_927_mod_path = tmp_path / "comp_927_mod.xlsx"
    wb_927 = openpyxl.load_workbook(wb_base_path)
    # change row 927 by +10,000,000
    wb_927["Data"]["E927"] = 80_000_000.0
    wb_927.save(wb_927_mod_path)

    comp_927_scenario = adapter.build(GoldenWorkbook(wb_927_mod_path), meta_comp, months)
    res_927 = adapter.material_analysis(base_scenario, comp_927_scenario)
    lc_927 = next(r for r in res_927["product_groups"] if r["product_group"] == "LC")
    comp_927_lc_prod = next(p for p in comp_927_scenario.scenario.products if p.product_group == "LC")

    # CASE D: Changing row 927 MUST change LC raw material cost, unit cost, and effect
    assert comp_927_lc_prod.raw_material_cost != base_lc_prod.raw_material_cost
    assert lc_927["comparison_unit_cost"] != lc_927["baseline_unit_cost"]
    assert lc_927["total"] != 0.0


def test_case_e_other_product_groups_unchanged():
    """CASE E: SW/BW/FS source mappings remain unchanged."""
    mapping = json.loads(Path("config/model_mapping.json").read_text(encoding="utf-8"))
    comps = mapping["analysis_adapter"]["material"]["raw_material_sources"]["finished_product_components"]
    groups = mapping["analysis_adapter"]["material"]["groups"]

    assert comps["SW400"]["back_total_rows"] == [899, 900]
    assert comps["SW440"]["back_total_rows"] == [906, 907]
    assert comps["BW400"]["back_total_rows"] == [913, 914]
    assert comps["BW440"]["back_total_rows"] == [920, 921]

    assert groups["SW"]["back_material_terms"] == [
        {"product_code": "SW400", "back_total_rows": [899, 900]},
        {"product_code": "SW440", "back_total_rows": [906, 907]},
    ]
    assert groups["BW"]["back_material_terms"] == [
        {"product_code": "BW400", "back_total_rows": [913, 914]},
        {"product_code": "BW440", "back_total_rows": [920, 921]},
    ]


def test_case_f_lc_mcm_remains_none():
    """CASE F: LC MCM remains none."""
    mapping = json.loads(Path("config/model_mapping.json").read_text(encoding="utf-8"))
    lc_group = mapping["analysis_adapter"]["material"]["groups"]["LC"]
    assert lc_group["mcm_rows"] == []
    assert lc_group["mcm_material_rows"] == []


def test_case_g_presentation_hotfix_regression():
    """CASE G: The previously failing Result 97876183 must still Presentation PASS."""
    from forecast.provenance import load_registered_provenance
    fixture_path = Path("C:/Users/bjy01/.gemini/antigravity/brain/f318710c-955b-4808-aba3-ecbd4230a673/scratch/STAGING_row_97876183-91b4-41d3-b99a-5f2bac105436.json")
    if not fixture_path.exists():
        pytest.skip("Staging incident row fixture not found")

    data = json.loads(fixture_path.read_text(encoding="utf-8"))
    provenance = load_registered_provenance(
        "config/model_mapping.json",
        "config/mapping_registry.json",
        "config/release.json",
    )
    res = build_analysis_presentation(
        data["result_id"],
        data,
        provenance,
        ["1"],
    )
    assert res is not None
    assert res.identity.result_id == "97876183-91b4-41d3-b99a-5f2bac105436"
    effects = {eff.code: eff for eff in res.effects}
    assert effects["sales_price"].profit_effect == pytest.approx(-312_293_072.30)
    assert effects["sga_variable"].profit_effect == pytest.approx(30_116_673.28)


def test_section_11_evidence_source_lineage():
    """Section 11: LC raw material evidence lineage contains Data!E927 and does not contain Data!E928."""
    adapter = _load_adapter()
    wb_base_path = Path("data/base.xlsx")
    meta_base = _meta("base")
    months = (1,)

    base_scenario = adapter.build(GoldenWorkbook(wb_base_path), meta_base, months)
    res = adapter.material_analysis(base_scenario, base_scenario)

    lc_trace = next(r for r in res["trace_rows"] if r["product_group"] == "LC")
    assert "Data!E927" in lc_trace["base_source_reference"]
    assert "Data!E928" not in lc_trace["base_source_reference"]
