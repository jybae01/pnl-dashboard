from __future__ import annotations

from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from forecast.analysis_export import build_comparison_audit_workbook
from forecast.comparison import GenericComparisonEngine, PeriodOption
from forecast.evidence_reporting import REPORT_SHEETS, audit_reporting_workbook

try:
    from tests.test_golden_analysis_adapter import _build_workbook, _meta
except ModuleNotFoundError:
    from test_golden_analysis_adapter import _build_workbook, _meta


ROOT = Path(__file__).resolve().parents[1]


def _row_with_value(ws, column: str, value: object) -> int:
    return next(
        row for row in range(1, ws.max_row + 1)
        if ws[f"{column}{row}"].value == value
    )


def _build_three_month_sentinel_workbook(path: Path, *, comparison: bool) -> None:
    _build_workbook(path, comparison=comparison)
    workbook = load_workbook(path)
    sheet = workbook["Data"]
    ratio_rows = {273, 274, 275, 345, 346, 347, 788, 789, 790, 791, 792, 956, 957}
    for column, factor in zip(("K", "L", "M"), (1.01, 2.02, 3.03), strict=True):
        for row in range(1, sheet.max_row + 1):
            value = sheet[f"E{row}"].value
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                sheet[f"{column}{row}"] = value if row in ratio_rows else value * factor
        month = {"K": 7, "L": 8, "M": 9}[column]
        sentinel = {7: 101.0, 8: 202.0, 9: 303.0}[month]
        revenue = sentinel * (110.0 if comparison else 100.0)
        cogs = sentinel * (55.0 if comparison else 60.0)
        sheet[f"{column}1594"] = sentinel * (1_100.0 if comparison else 1_000.0)
        sheet[f"{column}113"] = 0.0
        sheet[f"{column}114"] = revenue
        sheet[f"{column}1733"] = revenue
        sheet[f"{column}1734"] = cogs
        sheet[f"{column}9"] = 12.0 + month / 100.0
    workbook.save(path)


class EvidenceWorkbookTraceabilityTests(unittest.TestCase):
    def test_three_month_formula_lineage_preserves_monthly_sources_and_results(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            baseline = root / "base.xlsx"
            comparison = root / "comparison.xlsx"
            _build_three_month_sentinel_workbook(baseline, comparison=False)
            _build_three_month_sentinel_workbook(comparison, comparison=True)
            base_meta = _meta("base")
            comparison_meta = _meta("comparison")
            base_meta.regional_sales_monthly = {"7": 101_000, "8": 202_000, "9": 303_000}
            comparison_meta.regional_sales_monthly = {
                "7": 202_000, "8": 404_000, "9": 606_000,
            }
            for meta in (base_meta, comparison_meta):
                meta.tariff_adjustment_monthly = {}
                meta.tariff_applicable_rate = 0.85
                meta.tariff_rate = 0.10
            engine = GenericComparisonEngine(ROOT / "config" / "model_mapping.json")
            period = PeriodOption(
                "R2026_07_09", "2026-07 ~ 2026-09", (7, 8, 9), "사용자정의"
            )
            result = engine.compare(
                base_meta,
                baseline,
                comparison_meta,
                comparison,
                period,
                baseline_sales_fx=None,
                comparison_sales_fx=None,
                baseline_sales_fx_monthly={
                    "2026-07": 1_480.0,
                    "2026-08": 1_480.0,
                    "2026-09": 1_480.0,
                },
                comparison_sales_fx_monthly={
                    "2026-07": 1_380.0,
                    "2026-08": 1_420.0,
                    "2026-09": 1_500.0,
                },
            )
            before = {
                baseline: sha256(baseline.read_bytes()).hexdigest(),
                comparison: sha256(comparison.read_bytes()).hexdigest(),
            }
            payload = build_comparison_audit_workbook(
                result=asdict(result),
                sales_rows=result.sales_analysis["rows"],
                sales_totals=result.sales_analysis["totals"],
                baseline_fx=None,
                comparison_fx=None,
                baseline_path=baseline,
                comparison_path=comparison,
                mapping_path=ROOT / "config" / "model_mapping.json",
            )
            after = {
                baseline: sha256(baseline.read_bytes()).hexdigest(),
                comparison: sha256(comparison.read_bytes()).hexdigest(),
            }

        self.assertEqual(before, after)
        self.assertAlmostEqual(
            result.effects_total + result.residual, result.operating_profit_delta
        )
        expected_periods = {"2026-07", "2026-08", "2026-09"}
        self.assertEqual(
            {row["period"] for row in result.sales_analysis["monthly_effects"]},
            expected_periods,
        )
        self.assertEqual(
            {row["period"] for row in result.material_analysis["trace_rows"]},
            expected_periods,
        )
        self.assertEqual(
            {row["month"] for row in result.manufacturing_analysis["trace_rows"]},
            expected_periods,
        )
        self.assertEqual(
            {row["period"] for row in result.inventory_analysis["selected_monthly_details"]},
            expected_periods,
        )

        workbook = load_workbook(BytesIO(payload), data_only=False)
        audit = audit_reporting_workbook(workbook)
        self.assertEqual(tuple(workbook.sheetnames), REPORT_SHEETS)
        self.assertGreater(audit["formula_count"], 1_000)
        self.assertEqual(audit["hard_coded_derived_duplicates"], 0)

        source = workbook["90_원본값"]
        source_rows = [
            {
                "key": source.cell(row, 1).value,
                "period": source.cell(row, 5).value,
                "unit": source.cell(row, 6).value,
                "base_source": source.cell(row, 7).value,
                "base_value": source.cell(row, 8).value,
                "comparison_value": source.cell(row, 10).value,
            }
            for row in range(6, source.max_row + 1)
        ]
        self.assertEqual(
            {row["period"] for row in source_rows if str(row["key"]).startswith("sales_fx:")},
            expected_periods,
        )
        self.assertEqual(
            [row["base_value"] for row in source_rows if str(row["key"]).startswith("sales_fx:")],
            [1_480.0, 1_480.0, 1_480.0],
        )
        self.assertEqual(
            [row["comparison_value"] for row in source_rows if str(row["key"]).startswith("sales_fx:")],
            [1_380.0, 1_420.0, 1_500.0],
        )
        keys = [row["key"] for row in source_rows]
        self.assertEqual(len(keys), len(set(keys)))
        self.assertTrue(any("Data!K1594" in str(row["base_source"]) for row in source_rows))
        self.assertTrue(any("Data!L1594" in str(row["base_source"]) for row in source_rows))
        self.assertTrue(any("Data!M1594" in str(row["base_source"]) for row in source_rows))

        sales = workbook["03_판매근거"]
        monthly_rows = [
            row for row in range(1, sales.max_row + 1)
            if sales[f"A{row}"].value in expected_periods
        ]
        self.assertTrue(monthly_rows)
        self.assertTrue(all(
            any(
                isinstance(sales.cell(row, column).value, str)
                and sales.cell(row, column).value.startswith("=")
                for column in range(4, min(sales.max_column, 16) + 1)
            )
            for row in monthly_rows
        ))
        new_business_rows = [
            row for row in range(1, sales.max_row + 1)
            if sales[f"A{row}"].value == "신사업"
        ]
        self.assertEqual(len(new_business_rows), 3)
        self.assertTrue(all(
            sales[f"E{row}"].value is None
            and sales[f"F{row}"].value is None
            and sales[f"G{row}"].value is None
            for row in new_business_rows
        ))

        effects = workbook["02_손익영향"]
        effect_rows = [
            row for row in range(1, effects.max_row + 1)
            if effects[f"B{row}"].value in {
                "sales_quantity", "sales_mix", "sales_price", "sales_fx", "tariff",
                "material_total", "manufacturing_realized", "inventory_timing",
                "sga_variable", "sga_fixed",
            }
        ]
        self.assertEqual(len(effect_rows), 10)
        self.assertTrue(all(
            str(effects[f"D{row}"].value).startswith("='0")
            for row in effect_rows
        ))
        summary = workbook["01_보고요약"]
        self.assertTrue(all(
            isinstance(cell.value, str) and cell.value.startswith("='02_손익영향'!")
            for row in summary.iter_rows()
            for cell in row
            if cell.number_format == '#,##0;[Red](#,##0);-'
        ))

    def test_explicit_monthly_fx_matches_scalar_business_result(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            baseline = root / "base.xlsx"
            comparison = root / "comparison.xlsx"
            _build_workbook(baseline, comparison=False)
            _build_workbook(comparison, comparison=True)
            engine = GenericComparisonEngine(ROOT / "config" / "model_mapping.json")
            period = PeriodOption("M2026_01", "2026-01", (1,), "사용자정의")
            scalar = engine.compare(
                _meta("base"), baseline, _meta("comparison"), comparison, period,
                baseline_sales_fx=1_000.0, comparison_sales_fx=1_100.0,
            )
            monthly = engine.compare(
                _meta("base"), baseline, _meta("comparison"), comparison, period,
                baseline_sales_fx=None,
                comparison_sales_fx=None,
                baseline_sales_fx_monthly={"2026-01": 1_000.0},
                comparison_sales_fx_monthly={"2026-01": 1_100.0},
            )
            payload = build_comparison_audit_workbook(
                result=asdict(monthly),
                sales_rows=monthly.sales_analysis["rows"],
                sales_totals=monthly.sales_analysis["totals"],
                baseline_fx=1_000.0,
                comparison_fx=1_100.0,
                mapping_path=ROOT / "config" / "model_mapping.json",
            )

        self.assertEqual(scalar.effects, monthly.effects)
        self.assertEqual(scalar.effects_total, monthly.effects_total)
        self.assertEqual(scalar.residual, monthly.residual)
        workbook = load_workbook(BytesIO(payload), data_only=False)
        source = workbook["90_원본값"]
        fx_row = _row_with_value(source, "A", "sales_fx:2026-01")
        self.assertEqual(source[f"H{fx_row}"].value, 1_000.0)
        self.assertEqual(source[f"J{fx_row}"].value, 1_100.0)
        self.assertEqual(source[f"H{fx_row}"].number_format, "#,##0.00")

    def test_formula_workbook_has_no_value_only_regression(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            baseline = root / "base.xlsx"
            comparison = root / "comparison.xlsx"
            _build_workbook(baseline, comparison=False)
            _build_workbook(comparison, comparison=True)
            result = GenericComparisonEngine(
                ROOT / "config" / "model_mapping.json"
            ).compare(
                _meta("base"), baseline, _meta("comparison"), comparison,
                PeriodOption("M2026_01", "2026-01", (1,), "사용자정의"),
                baseline_sales_fx=1_400.0, comparison_sales_fx=1_450.0,
            )
            payload = build_comparison_audit_workbook(
                result=asdict(result),
                sales_rows=result.sales_analysis["rows"],
                sales_totals=result.sales_analysis["totals"],
                baseline_fx=1_400.0,
                comparison_fx=1_450.0,
                mapping_path=ROOT / "config" / "model_mapping.json",
            )

        formulas = load_workbook(BytesIO(payload), data_only=False)
        values = load_workbook(BytesIO(payload), data_only=True)
        self.assertGreater(
            sum(
                1 for sheet in formulas.worksheets for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ),
            400,
        )
        # openpyxl does not evaluate formulas. A value-only rewrite would populate
        # all formula coordinates instead of preserving formula cells for Excel.
        formula_coordinates = [
            (sheet.title, cell.coordinate)
            for sheet in formulas.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        self.assertTrue(formula_coordinates)
        self.assertTrue(any(
            values[sheet][coordinate].value is None
            for sheet, coordinate in formula_coordinates
        ))
        self.assertEqual(
            audit_reporting_workbook(formulas)["hard_coded_derived_duplicates"], 0
        )


if __name__ == "__main__":
    unittest.main()
