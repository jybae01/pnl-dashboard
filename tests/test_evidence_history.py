from __future__ import annotations

import hashlib
import uuid
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from forecast.bff.auth import AccessCodeSessionService
from forecast.bff.errors import ApiErrorCode, BffError
from forecast.bff.evidence_history import CalculationHistoryService, EvidenceDeliveryService
from forecast.provenance import ResultProvenance


PROVENANCE = ResultProvenance("engine-1", "mapping-1", "a" * 64, "1")
BASE_ID = str(uuid.uuid4())
COMP_ID = str(uuid.uuid4())
JOB_ID = str(uuid.uuid4())
RESULT_ID = str(uuid.uuid4())
SOURCE = b"exact workbook bytes"
SOURCE_SHA = hashlib.sha256(SOURCE).hexdigest()


class Gateway:
    def __init__(self) -> None:
        self.admin_row = evidence_row()
        self.viewer_row = evidence_row()
        self.history_rows = []
        self.available = True

    def get_admin_evidence_payload(self, *_args, **_kwargs):
        return self.admin_row

    def get_viewer_evidence_payload(self, *_args, **_kwargs):
        return self.viewer_row

    def download_model_source(self, bucket, path):
        assert bucket == "pnl-models"
        assert path in {f"models/{BASE_ID}/source.xlsx", f"models/{COMP_ID}/source.xlsx"}
        return SOURCE

    def list_calculation_history(self, **_kwargs):
        return list(self.history_rows)

    def validate_result_availability(self, *_args, **_kwargs):
        return self.available


def evidence_row():
    return {
        "result_id": RESULT_ID,
        "job_id": JOB_ID,
        "result_payload": {"comparison_result": {
            "period": {"label": "2026-01_2026-12", "months": [1]},
            "sales_analysis": {
                "baseline_fx_krw_per_usd": 1450,
                "comparison_fx_krw_per_usd": 1500,
                "rows": [],
                "totals": {},
            },
        }},
        "analysis_request": {"start_month": 1, "end_month": 12},
        "baseline_model_id": BASE_ID,
        "comparison_model_id": COMP_ID,
        "baseline_workbook_sha256": SOURCE_SHA,
        "comparison_workbook_sha256": SOURCE_SHA,
        "baseline_workbook_bucket": "pnl-models",
        "baseline_workbook_path": f"models/{BASE_ID}/source.xlsx",
        "comparison_workbook_bucket": "pnl-models",
        "comparison_workbook_path": f"models/{COMP_ID}/source.xlsx",
        "engine_version": PROVENANCE.engine_version,
        "mapping_version": PROVENANCE.mapping_version,
        "mapping_hash": PROVENANCE.mapping_hash,
        "result_schema_version": PROVENANCE.result_schema_version,
    }


@pytest.fixture
def auth():
    sessions = AccessCodeSessionService(
        viewer_code="viewer", admin_code="admin", actor_namespace_secret="secret" * 8
    )
    return sessions, sessions.login("admin").session_id, sessions.login("viewer").session_id


@pytest.fixture(autouse=True)
def registered_mapping_hash(monkeypatch):
    monkeypatch.setattr(
        "forecast.bff.evidence_history.mapping_hash",
        lambda _path: PROVENANCE.mapping_hash,
    )


def test_admin_and_viewer_evidence_use_stored_result_and_cleanup(monkeypatch, tmp_path, auth):
    sessions, admin, viewer = auth
    gateway = Gateway()
    calls = []

    def write(path, **kwargs):
        calls.append(kwargs)
        workbook = Workbook()
        workbook.active["A1"] = kwargs["result"]["evidence_provenance"]["result_id"]
        workbook.save(path)

    monkeypatch.setattr("forecast.bff.evidence_history.write_comparison_audit_workbook", write)
    mapping = tmp_path / "mapping.json"
    mapping.write_text("{}", encoding="utf-8")
    service = EvidenceDeliveryService(
        sessions, gateway, PROVENANCE, mapping_path=mapping,
        supported_result_schema_versions=("1",),
    )
    artifact = service.admin_download(admin, RESULT_ID)
    assert artifact.path.exists()
    assert load_workbook(artifact.path).active["A1"].value == RESULT_ID
    assert calls[0]["result"]["sales_analysis"]["rows"] == []
    root = artifact.path.parent
    artifact.cleanup()
    assert not root.exists()

    viewer_artifact = service.viewer_download(viewer, RESULT_ID)
    viewer_artifact.cleanup()


def test_viewer_unavailable_and_sha_mismatch_are_safe(auth, tmp_path):
    sessions, admin, viewer = auth
    gateway = Gateway()
    gateway.viewer_row = None
    service = EvidenceDeliveryService(
        sessions, gateway, PROVENANCE, mapping_path=tmp_path / "mapping.json",
        supported_result_schema_versions=("1",),
    )
    with pytest.raises(BffError) as denied:
        service.viewer_download(viewer, RESULT_ID)
    assert denied.value.code == ApiErrorCode.RESULT_NOT_AVAILABLE

    gateway.admin_row = evidence_row()
    gateway.admin_row["baseline_workbook_sha256"] = "b" * 64
    with pytest.raises(BffError) as mismatch:
        service.admin_download(admin, RESULT_ID)
    assert mismatch.value.code == ApiErrorCode.INPUT_INTEGRITY_MISMATCH


def test_generation_failure_has_distinct_safe_code(monkeypatch, auth, tmp_path):
    sessions, admin, _viewer = auth
    gateway = Gateway()
    monkeypatch.setattr(
        "forecast.bff.evidence_history.write_comparison_audit_workbook",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("internal path")),
    )
    service = EvidenceDeliveryService(
        sessions, gateway, PROVENANCE, mapping_path=tmp_path / "mapping.json",
        supported_result_schema_versions=("1",),
    )
    with pytest.raises(BffError) as caught:
        service.admin_download(admin, RESULT_ID)
    assert caught.value.code == ApiErrorCode.EVIDENCE_GENERATION_FAILED
    assert "internal path" not in str(caught.value)


def test_actual_mapping_file_hash_must_match_pinned_provenance(monkeypatch, auth, tmp_path):
    sessions, admin, _viewer = auth
    monkeypatch.setattr("forecast.bff.evidence_history.mapping_hash", lambda _path: "b" * 64)
    service = EvidenceDeliveryService(
        sessions, Gateway(), PROVENANCE, mapping_path=tmp_path / "mapping.json",
        supported_result_schema_versions=("1",),
    )
    with pytest.raises(BffError) as caught:
        service.admin_download(admin, RESULT_ID)
    assert caught.value.code == ApiErrorCode.INPUT_INTEGRITY_MISMATCH


def test_viewer_rechecks_availability_after_generation(monkeypatch, auth, tmp_path):
    sessions, _admin, viewer = auth
    gateway = Gateway()
    gateway.available = False
    monkeypatch.setattr(
        "forecast.bff.evidence_history.write_comparison_audit_workbook",
        lambda path, **_kwargs: Workbook().save(path),
    )
    service = EvidenceDeliveryService(
        sessions, gateway, PROVENANCE, mapping_path=tmp_path / "mapping.json",
        supported_result_schema_versions=("1",),
    )
    with pytest.raises(BffError) as caught:
        service.viewer_download(viewer, RESULT_ID)
    assert caught.value.code == ApiErrorCode.RESULT_NOT_AVAILABLE


def test_history_is_bounded_stable_and_never_invents_result_ids(auth):
    sessions, admin, viewer = auth
    gateway = Gateway()
    gateway.history_rows = [history_row("completed", result_id=RESULT_ID), history_row("pending")]
    service = CalculationHistoryService(sessions, gateway)
    page = service.list_admin(admin, limit=1)
    assert len(page.items) == 1
    assert page.items[0].result_id == RESULT_ID
    assert page.next_before_created_at == page.items[0].created_at
    assert page.next_before_job_id == page.items[0].job_id
    with pytest.raises(BffError) as denied:
        service.list_admin(viewer)
    assert denied.value.code == ApiErrorCode.FORBIDDEN


def history_row(status, result_id=None):
    return {
        "job_id": JOB_ID,
        "result_id": result_id,
        "status": status,
        "baseline_model_id": BASE_ID,
        "baseline_model_name": "Base",
        "comparison_model_id": COMP_ID,
        "comparison_model_name": "Comparison",
        "start_month": 1,
        "end_month": 12,
        "attempt": 0,
        "max_attempts": 3,
        "created_at": "2026-08-11T00:00:00+00:00",
        "completed_at": "2026-08-11T01:00:00+00:00" if status == "completed" else None,
        "error_code": None,
        "error_message": None,
        "is_published": False,
    }


def test_migration_008_is_narrow_and_additive():
    path = Path("supabase/migrations/202608090008_evidence_history_vertical_slice.sql")
    sql = path.read_text(encoding="utf-8").lower()
    assert "get_calculation_result_evidence_admin_by_id" in sql
    assert "get_calculation_result_evidence_viewer_by_id" in sql
    assert "validate_calculation_result_availability" in sql
    assert "list_calculation_history_admin" in sql
    assert "order by job.created_at desc, job.id desc" in sql
    assert "limit p_limit + 1" in sql
    assert "claim_token" not in sql
    assert "queue_receipt" not in sql
    assert "from public, anon, authenticated" in sql
    assert "to service_role" in sql
