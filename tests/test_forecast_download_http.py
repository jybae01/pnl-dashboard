from __future__ import annotations

from io import BytesIO
import hashlib
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink

from forecast.bff.application import TrustedBffApplication
from forecast.bff.auth import AccessCodeSessionService
from forecast.bff.errors import ApiErrorCode
from forecast.bff.forecast_download import ForecastWorkbookDownloadService
from forecast.bff.gateway import GatewayTransientError
from forecast.bff.http import HttpBffSettings, create_http_bff


MODEL_ID = "11111111-1111-4111-8111-111111111111"
GENERATION_ID = "22222222-2222-4222-8222-222222222222"
SOURCE_MODEL_ID = "33333333-3333-4333-8333-333333333333"
SHA = "a" * 64


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "입력반영내역"
    sheet["A1"] = "=1+1"
    data = workbook.create_sheet("Data")
    # Representative direct-reflection categories from a generated Forecast.
    # Internal links and reasons must survive byte-for-byte delivery without
    # an openpyxl save/rewrite in the BFF.
    markers = (
        ("판매", "K32", 123456, ""),
        ("생산", "K119", 234567, ""),
        ("MCM", "K568", 345678, ""),
        ("제조경비", "K290", 456789, "TEST_MFG_REASON"),
        ("판관비", "K1168", 567890, "TEST_SGA_REASON"),
        ("매출원가", "K1273", 678901, "TEST_COGS_REASON"),
        ("신사업 상품원가", "K1289", 789012, ""),
        ("원재료 관세 환급", "K1294", -890123, "TEST_REFUND_REASON"),
    )
    for row_number, (category, target, value, reason) in enumerate(markers, start=2):
        sheet[f"A{row_number}"] = category
        sheet[f"I{row_number}"] = reason
        sheet[f"J{row_number}"] = "이동"
        sheet[f"J{row_number}"].hyperlink = Hyperlink(
            ref=f"J{row_number}", location=f"'Data'!{target}", display="이동"
        )
        data[target] = value
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


class DownloadGateway:
    def __init__(self, row=None, payload: bytes | None = None, error: Exception | None = None):
        self.row = row
        self.payload = payload
        self.error = error
        self.download_calls: list[tuple[str, str]] = []

    def get_forecast_model(self, model_id: str):
        assert model_id == MODEL_ID
        return self.row

    def download_model_source(self, bucket: str, path: str) -> bytes:
        self.download_calls.append((bucket, path))
        if self.error is not None:
            raise self.error
        return self.payload


def _row(payload: bytes, **overrides):
    values = {
        "id": MODEL_ID,
        "name": "Generated Forecast",
        "model_type": "FORECAST",
        "source_kind": "forecast_generated",
        "source_model_id": SOURCE_MODEL_ID,
        "forecast_generation_id": GENERATION_ID,
        "workbook_bucket": "pnl-models",
        "workbook_path": f"models/{MODEL_ID}/source.xlsx",
        "workbook_sha256": hashlib.sha256(payload).hexdigest(),
        "file_name": "forecast_2026_07_07.xlsx",
    }
    values.update(overrides)
    return values


def _fixture(gateway: DownloadGateway):
    sessions = AccessCodeSessionService(
        viewer_code="viewer-code",
        admin_code="admin-code",
        actor_namespace_secret="stable-server-only-actor-namespace",
    )
    application = TrustedBffApplication(
        sessions=sessions,
        submissions=object(),
        jobs=object(),
        results=object(),
        forecast_download=ForecastWorkbookDownloadService(sessions, gateway),
    )
    app = create_http_bff(
        application,
        settings=HttpBffSettings(
            environment="test",
            csrf_secret="csrf-secret-at-least-32-characters",
        ),
    )
    return TestClient(app)


def _login(client: TestClient, code: str = "admin-code") -> None:
    response = client.post("/api/session/login", json={"access_code": code})
    assert response.status_code == 200


def test_admin_download_returns_exact_private_bytes_and_safe_xlsx_headers():
    payload = _workbook_bytes()
    gateway = DownloadGateway(_row(payload), payload)
    client = _fixture(gateway)
    _login(client)

    response = client.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")

    assert response.status_code == 200
    assert response.content == payload
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="forecast_2026_07_07.xlsx"; '
        "filename*=utf-8''forecast_2026_07_07.xlsx"
    )
    assert response.headers["x-content-type-options"] == "nosniff"
    assert b"source.xlsx" not in response.content
    assert gateway.download_calls == [("pnl-models", f"models/{MODEL_ID}/source.xlsx")]

    # The service never opens/saves the workbook, so formulas, reason text, and
    # hyperlinks remain exactly as stored.
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
    sheet = workbook.active
    assert sheet["A1"].value == "=1+1"
    expected = (
        ("판매", "'Data'!K32", 123456, None),
        ("생산", "'Data'!K119", 234567, None),
        ("MCM", "'Data'!K568", 345678, None),
        ("제조경비", "'Data'!K290", 456789, "TEST_MFG_REASON"),
        ("판관비", "'Data'!K1168", 567890, "TEST_SGA_REASON"),
        ("매출원가", "'Data'!K1273", 678901, "TEST_COGS_REASON"),
        ("신사업 상품원가", "'Data'!K1289", 789012, None),
        ("원재료 관세 환급", "'Data'!K1294", -890123, "TEST_REFUND_REASON"),
    )
    for row_number, (category, location, value, reason) in enumerate(expected, start=2):
        assert sheet[f"A{row_number}"].value == category
        assert sheet[f"I{row_number}"].value == reason
        assert sheet[f"J{row_number}"].value == "이동"
        assert workbook["Data"][location.split("!")[1]].value == value

    # Some openpyxl read-only versions omit hyperlink objects on streaming
    # cells. Reopen the same bytes without mutation only to inspect links.
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    sheet = workbook.active
    for row_number, (_category, location, _value, _reason) in enumerate(expected, start=2):
        assert sheet[f"J{row_number}"].hyperlink.location == location


def test_download_requires_admin_session():
    payload = b"private-workbook-bytes"
    gateway = DownloadGateway(_row(payload), payload)

    anonymous = _fixture(gateway)
    assert anonymous.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook").status_code == 401

    viewer = _fixture(gateway)
    _login(viewer, "viewer-code")
    denied = viewer.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")
    assert denied.status_code == 403
    assert denied.json()["error"]["code"] == ApiErrorCode.FORBIDDEN.value


def test_unknown_or_wrong_type_models_are_rejected_without_storage_access():
    payload = b"private-workbook-bytes"

    unknown_gateway = DownloadGateway(None, payload)
    unknown = _fixture(unknown_gateway)
    _login(unknown)
    missing = unknown.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == ApiErrorCode.MODEL_NOT_FOUND.value
    assert not unknown_gateway.download_calls

    wrong_gateway = DownloadGateway(_row(payload, model_type="ACTUAL"), payload)
    wrong = _fixture(wrong_gateway)
    _login(wrong)
    rejected = wrong.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")
    assert rejected.status_code == 422
    assert rejected.json()["error"]["code"] == ApiErrorCode.VALIDATION_ERROR.value
    assert not wrong_gateway.download_calls


def test_missing_storage_and_sha_mismatch_fail_closed():
    payload = b"private-workbook-bytes"

    missing_storage_gateway = DownloadGateway(
        _row(payload), payload=None, error=GatewayTransientError("missing storage")
    )
    missing_storage = _fixture(missing_storage_gateway)
    _login(missing_storage)
    unavailable = missing_storage.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")
    assert unavailable.status_code == 503
    assert unavailable.json()["error"]["code"] == ApiErrorCode.TRANSIENT_SYSTEM_ERROR.value

    mismatch_gateway = DownloadGateway(_row(payload, workbook_sha256=SHA), payload)
    mismatch = _fixture(mismatch_gateway)
    _login(mismatch)
    rejected = mismatch.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")
    assert rejected.status_code == 409
    assert rejected.json()["error"]["code"] == ApiErrorCode.INPUT_INTEGRITY_MISMATCH.value
    assert rejected.content.startswith(b"{")


@pytest.mark.parametrize(
    "row_overrides",
    [
        {"id": None},
        {"source_kind": None},
        {"forecast_generation_id": None},
        {"forecast_generation_id": "not-a-uuid"},
        {"source_model_id": None},
        {"source_model_id": "not-a-uuid"},
        {"workbook_sha256": None},
        {"workbook_path": "models/other/source.xlsx"},
        {"source_kind": "uploaded"},
    ],
)
def test_invalid_recorded_provenance_is_rejected_before_storage(row_overrides):
    payload = b"private-workbook-bytes"
    gateway = DownloadGateway(_row(payload, **row_overrides), payload)
    client = _fixture(gateway)
    _login(client)

    response = client.get(f"/api/admin/forecast-models/{MODEL_ID}/workbook")

    assert response.status_code in {409, 422}
    assert response.json()["error"]["code"] in {
        ApiErrorCode.INPUT_INTEGRITY_MISMATCH.value,
        ApiErrorCode.VALIDATION_ERROR.value,
    }
    assert not gateway.download_calls
