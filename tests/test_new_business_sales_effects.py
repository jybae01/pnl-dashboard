from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from forecast.analysis.configuration import AnalysisConfig
from forecast.analysis.sales_effects import calculate_sales_effects
from forecast.analysis.schema import AnalysisScenario, ProductRecord, ScenarioMeta
from forecast.evidence_traceability import write_sales_evidence
from forecast.sales_comparison import calculate_sales_effect_rows


ROOT = Path(__file__).resolve().parents[1]
CONFIG = AnalysisConfig.load(ROOT / "config" / "analysis_v1.json")


def _scenario(identifier: str, products: list[ProductRecord]) -> AnalysisScenario:
    return AnalysisScenario(
        meta=ScenarioMeta(identifier, "PLAN", "v1"),
        products=products,
    )


def _manufactured_rows(*, comparison: bool) -> list[ProductRecord]:
    if comparison:
        return [
            ProductRecord(
                "2026-07", "SW", "SW", sales_qty=120, sales_amount=1_320,
                product_cogs=720, sales_fx=1_500,
            ),
            ProductRecord(
                "2026-07", "FS", "FS", unit_basis="LENGTH", sales_length=80,
                sales_amount=960, product_cogs=560, sales_fx=1_500,
            ),
        ]
    return [
        ProductRecord(
            "2026-07", "SW", "SW", sales_qty=100, sales_amount=1_000,
            product_cogs=600, sales_fx=1_450,
        ),
        ProductRecord(
            "2026-07", "FS", "FS", unit_basis="LENGTH", sales_length=100,
            sales_amount=1_000, product_cogs=700, sales_fx=1_450,
        ),
    ]


def _new_business(
    *, comparison: bool, base_revenue: float = 100.0, quantity: float | None = None
) -> ProductRecord:
    return ProductRecord(
        "2026-07",
        "NEW_BUSINESS_SALES",
        "신사업",
        sales_qty=(999 if comparison else 1) if quantity is None else quantity,
        sales_amount=200.0 if comparison else base_revenue,
        product_cogs=140.0 if comparison else 80.0,
        sales_fx=9_999.0 if comparison else 1.0,
        material_applicable_flag=False,
        sales_amount_source="Data!K1669",
        product_cogs_source="Data!K1670",
        source_validation_status="SOURCE_MAPPED",
    )


def test_new_business_uses_revenue_and_gp_rate_only_without_changing_manufactured_effects():
    manufactured = calculate_sales_effects(
        _scenario("base", _manufactured_rows(comparison=False)),
        _scenario("comparison", _manufactured_rows(comparison=True)),
        CONFIG,
    )
    result = calculate_sales_effects(
        _scenario("base", [*_manufactured_rows(comparison=False), _new_business(comparison=False)]),
        _scenario("comparison", [*_manufactured_rows(comparison=True), _new_business(comparison=True)]),
        CONFIG,
    )

    assert result.new_business_revenue_effect == pytest.approx(20.0)
    assert result.new_business_gp_rate_effect == pytest.approx(20.0)
    assert (
        result.new_business_revenue_effect + result.new_business_gp_rate_effect
    ) == pytest.approx((200.0 - 140.0) - (100.0 - 80.0))
    assert result.quantity == pytest.approx(manufactured.quantity + 20.0)
    assert result.displayed_price == pytest.approx(manufactured.displayed_price + 20.0)
    assert result.price == pytest.approx(manufactured.price + 20.0)
    assert result.mix == pytest.approx(manufactured.mix)
    assert result.sales_fx == pytest.approx(manufactured.sales_fx)
    assert {row["product_group"] for row in result.details} == {"SW", "FS"}
    assert {row["pool"] for row in result.pool_details} == {"PCS", "LENGTH"}
    assert result.new_business_details[0]["sales_fx_effect"] == 0.0
    assert result.new_business_details[0]["mix_effect"] == 0.0


def test_new_business_base_revenue_zero_fails_closed():
    with pytest.raises(ValueError, match="신사업: 기준 매출액이 0 이하"):
        calculate_sales_effects(
            _scenario("base", [_new_business(comparison=False, base_revenue=0.0)]),
            _scenario("comparison", [_new_business(comparison=True)]),
            CONFIG,
        )


@pytest.mark.parametrize(
    ("base_quantity", "comparison_quantity"),
    ((0.0, 0.0), (1.0, 999.0), (999.0, 1.0)),
)
def test_new_business_quantity_never_changes_revenue_gp_rate_effects(
    base_quantity: float,
    comparison_quantity: float,
):
    result = calculate_sales_effects(
        _scenario("base", [_new_business(comparison=False, quantity=base_quantity)]),
        _scenario(
            "comparison",
            [_new_business(comparison=True, quantity=comparison_quantity)],
        ),
        CONFIG,
    )

    assert result.new_business_revenue_effect == pytest.approx(20.0)
    assert result.new_business_gp_rate_effect == pytest.approx(20.0)
    assert result.quantity == pytest.approx(20.0)
    assert result.displayed_price == pytest.approx(20.0)
    assert result.price == pytest.approx(20.0)
    assert result.mix == 0.0
    assert result.sales_fx == 0.0


def test_compatibility_sales_path_uses_new_business_revenue_and_cogs_at_zero_quantity():
    row = {
        "product_group": "신사업",
        "baseline_quantity": 0.0,
        "comparison_quantity": 0.0,
        "baseline_amount": 100.0,
        "comparison_amount": 200.0,
        "baseline_cogs": 80.0,
        "comparison_cogs": 140.0,
        "baseline_gross_margin_rate": 0.0,
        "comparison_gross_margin_rate": 0.0,
    }

    result = calculate_sales_effect_rows([row], 1_450.0, 1_500.0)[0]

    assert result.quantity_effect == pytest.approx(20.0)
    assert result.pure_price_effect == pytest.approx(20.0)
    assert result.sales_fx_effect == 0.0
    assert result.total_sales_effect == pytest.approx(40.0)


def test_compatibility_sales_path_fails_closed_without_new_business_cogs():
    with pytest.raises(ValueError, match="매출원가가 필요"):
        calculate_sales_effect_rows(
            [{
                "product_group": "신사업",
                "baseline_quantity": 0.0,
                "comparison_quantity": 0.0,
                "baseline_amount": 100.0,
                "comparison_amount": 200.0,
            }],
            1_450.0,
            1_500.0,
        )


@pytest.mark.parametrize("missing_cogs", (None, "", " "))
def test_compatibility_sales_path_fails_closed_for_blank_new_business_cogs(
    missing_cogs,
):
    with pytest.raises(ValueError, match="유효한.*매출원가"):
        calculate_sales_effect_rows(
            [{
                "product_group": "신사업",
                "baseline_quantity": 0.0,
                "comparison_quantity": 0.0,
                "baseline_amount": 100.0,
                "comparison_amount": 200.0,
                "baseline_cogs": missing_cogs,
                "comparison_cogs": 140.0,
            }],
            1_450.0,
            1_500.0,
        )


def test_sales_evidence_uses_raw_values_and_formulas_for_new_business():
    calculated = calculate_sales_effects(
        _scenario("base", [*_manufactured_rows(comparison=False), _new_business(comparison=False)]),
        _scenario("comparison", [*_manufactured_rows(comparison=True), _new_business(comparison=True)]),
        CONFIG,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "판매효과_근거"
    totals = {
        "quantity_effect": calculated.quantity,
        "mix_effect": calculated.mix,
        "displayed_sales_price_effect": calculated.displayed_price,
        "sales_price_effect": calculated.price,
        "sales_fx_effect": calculated.sales_fx,
        "transport_effect": calculated.transport_effect,
        "tariff_effect": calculated.tariff,
        "total_sales_effect": calculated.quantity + calculated.mix + calculated.price + calculated.sales_fx,
    }
    cells = write_sales_evidence(
        sheet,
        {"sales_analysis": {
            "totals": totals,
            "trace_rows": calculated.details,
            "pool_trace_rows": calculated.pool_details,
            "freight_trace_rows": calculated.freight_details,
            "new_business_trace_rows": calculated.new_business_details,
        }},
        [],
        totals,
        1_450.0,
        1_500.0,
    )

    start, end = cells["new_business_range"]
    assert start == end
    assert [sheet.cell(start, column).value for column in range(8, 12)] == [
        100.0, 200.0, 80.0, 140.0,
    ]
    for column in (12, 13, 14, 15, 16, 18, 20, 21, 22):
        assert str(sheet.cell(start, column).value).startswith("=")
    assert sheet.cell(start, 23).value == "적용하지 않음"
    assert sheet.cell(start, 24).value == "적용하지 않음"
    quantity_row = cells["summary_range"][0]
    price_row = quantity_row + 4
    assert f"SUM(P{start}:P{end})" in sheet.cell(quantity_row, 2).value
    assert f"SUM(R{start}:R{end})" in sheet.cell(price_row - 2, 2).value
    assert sheet.cell(price_row, 2).value == f"=B{price_row - 2}+B{price_row - 1}"
