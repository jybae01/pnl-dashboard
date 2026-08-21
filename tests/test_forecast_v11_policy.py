from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from forecast.bff.forecast_orchestration import ForecastReservation
from forecast.bff.gateway import SupabaseForecastGateway
from forecast.engine import (
    ForecastInput,
    SalesInput,
    calculate_v11_forecast_transport_policy,
)
from forecast.provenance import ResultProvenance
from forecast.workbook import GoldenWorkbook


def _sales() -> dict[str, SalesInput]:
    return {
        "SW400": SalesInput(amount=100),
        "SW440": SalesInput(amount=200),
        "BW400": SalesInput(amount=400),
        "BW440": SalesInput(amount=600),
        "LC": SalesInput(amount=700),
        "LC_MERCHANDISE": SalesInput(amount=900),
        "FS_SW": SalesInput(amount=100),
        "FS_BW": SalesInput(amount=200),
        "FS_TW": SalesInput(amount=300),
        "UF_MBR": SalesInput(amount=1_000),
        "IX": SalesInput(amount=2_000),
    }


def test_v11_forecast_freight_and_tariff_policy_uses_all_authoritative_rates():
    result = calculate_v11_forecast_transport_policy(
        _sales(),
        plan_na_sa_sales=1_000,
        na_sa_sales=2_000,
    )

    assert result["sw_freight"] == pytest.approx(4.5)
    assert result["bw_freight"] == pytest.approx(15)
    assert result["lc_freight"] == pytest.approx(10.5)
    assert result["fs_freight"] == pytest.approx(9)
    assert result["existing_product_freight"] == pytest.approx(39)
    assert result["uf_mbr_freight"] == pytest.approx(100)
    assert result["ix_freight"] == pytest.approx(100)
    assert result["default_customer_freight"] == pytest.approx(239)
    assert result["plan_tariff"] == pytest.approx(85)
    assert result["forecast_tariff"] == pytest.approx(170)
    assert result["tariff_adjustment"] == pytest.approx(85)
    assert result["target_selling_transport"] == pytest.approx(409)


@pytest.mark.parametrize(
    ("product_code", "freight_key"),
    [
        ("SW400", "sw_freight"),
        ("BW400", "bw_freight"),
        ("LC", "lc_freight"),
        ("FS_SW", "fs_freight"),
    ],
)
def test_each_existing_product_group_uses_one_point_five_percent(
    product_code, freight_key,
):
    result = calculate_v11_forecast_transport_policy(
        {product_code: SalesInput(amount=100)},
        plan_na_sa_sales=0,
        na_sa_sales=0,
    )

    assert result[freight_key] == pytest.approx(1.5)
    assert result["existing_product_freight"] == pytest.approx(1.5)
    assert result["default_customer_freight"] == pytest.approx(1.5)
    assert result["target_selling_transport"] == pytest.approx(1.5)


def test_lc_merchandise_revenue_is_not_existing_product_freight():
    sales = _sales()
    baseline = calculate_v11_forecast_transport_policy(
        sales, plan_na_sa_sales=0, na_sa_sales=0
    )
    sales["LC_MERCHANDISE"] = SalesInput(amount=999_999_999)
    changed = calculate_v11_forecast_transport_policy(
        sales, plan_na_sa_sales=0, na_sa_sales=0
    )

    assert changed == baseline


def test_legacy_request_rate_fields_cannot_change_v11_policy():
    first = ForecastInput(
        sales=_sales(),
        uf_mbr_transport_rate=0.03,
        ix_transport_rate=0.99,
        tariff_applicable_rate=0.02,
        tariff_rate=0.97,
        plan_na_sa_sales=1_000,
        na_sa_sales=2_000,
    )
    second = ForecastInput(
        sales=_sales(),
        uf_mbr_transport_rate=0.88,
        ix_transport_rate=0.01,
        tariff_applicable_rate=0.99,
        tariff_rate=0.01,
        plan_na_sa_sales=1_000,
        na_sa_sales=2_000,
    )

    first_policy = calculate_v11_forecast_transport_policy(
        first.sales,
        plan_na_sa_sales=first.plan_na_sa_sales,
        na_sa_sales=first.na_sa_sales,
    )
    second_policy = calculate_v11_forecast_transport_policy(
        second.sales,
        plan_na_sa_sales=second.plan_na_sa_sales,
        na_sa_sales=second.na_sa_sales,
    )

    assert first_policy == second_policy
    assert first_policy["existing_product_freight"] == pytest.approx(39)


def test_v11_forecast_policy_zero_revenue_and_invalid_revenue_contract():
    zero = calculate_v11_forecast_transport_policy(
        {}, plan_na_sa_sales=0, na_sa_sales=0
    )
    assert zero["target_selling_transport"] == 0

    with pytest.raises(ValueError, match="0 이상"):
        calculate_v11_forecast_transport_policy(
            {"SW400": SalesInput(amount=-1)},
            plan_na_sa_sales=0,
            na_sa_sales=0,
        )


def test_authoritative_selling_transport_is_written_to_workbook_audit(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["B1166"] = "★판매관리비 관리 명세서"
    sheet["B1167"] = "판매비"
    sheet["C1168"] = "운반비"
    sheet["K1168"] = 10
    workbook.save(source)

    golden = GoldenWorkbook(source)
    policy_reason = (
        "v1.1 Backend authoritative: 기존제품 운반비 1.5%, UF/MBR 10%, "
        "IX 5%, Forecast 미주매출 관세 8.5%"
    )
    golden.set_input(
        "K1168",
        123,
        "sga_authoritative_default",
        policy_reason,
    )
    golden.save(output)

    generated = load_workbook(output, data_only=False)
    audit = generated["입력반영내역"]
    text = " ".join(
        str(cell.value or "") for row in audit.iter_rows() for cell in row
    )
    assert "K1168" in text
    assert policy_reason in text
    generated.close()


class _Response:
    def __init__(self, data):
        self.data = data


class _Call:
    def __init__(self, data):
        self.data = data

    def execute(self):
        return _Response(self.data)


class _GatewayClient:
    def __init__(self, model_row):
        self.model_row = model_row
        self.rpc_name = None
        self.rpc_params = None

    def rpc(self, name, params):
        self.rpc_name = name
        self.rpc_params = dict(params)
        return _Call([self.model_row])


def test_forecast_gateway_uses_atomic_v11_finalize_rpc_without_direct_update():
    model_id = "22222222-2222-4222-8222-222222222222"
    client = _GatewayClient({
        "id": model_id,
        "workbook_sha256": "a" * 64,
        "regional_sales_monthly": {"7": 1_000_000},
        "tariff_applicable_rate": 0.85,
        "tariff_rate": 0.10,
        "tariff_adjustment_monthly": {},
        "tariff_in_workbook": True,
    })
    gateway = SupabaseForecastGateway(client)
    reservation = ForecastReservation(
        "33333333-3333-4333-8333-333333333333",
        model_id,
        "reserved",
        "44444444-4444-4444-8444-444444444444",
        False,
        "pnl-models",
        "models/base/source.xlsx",
        "b" * 64,
    )

    result = gateway.finalize(
        reservation,
        sha256="a" * 64,
        name="Forecast",
        model_year=2026,
        version="V1",
        file_name="forecast.xlsx",
        period_types={"7": "추정"},
        provenance=ResultProvenance("1.1.0", "mapping", "c" * 64, "1"),
    )

    assert client.rpc_name == "finalize_forecast_generation_v11"
    assert set(client.rpc_params) == {
        "p_generation_id", "p_lease_token", "p_generated_workbook_sha256",
        "p_name", "p_model_year", "p_version", "p_file_name",
        "p_period_types", "p_mapping_version", "p_mapping_hash",
        "p_engine_version",
    }
    assert result["tariff_in_workbook"] is True
