from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook, load_workbook

from forecast.workbook import GoldenWorkbook, MERCHANDISE_EVIDENCE_SHEET_NAME


def _source_workbook(path: Path, *, goods_cogs: float = 800) -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    for address, value in {
        "E105": 100, "F105": 200, "E1660": 80, "F1660": 160,
        "E1733": 300, "F1733": 100, "E1734": 210, "F1734": 70,
        "K105": 1_000, "K114": 2_000, "K1289": goods_cogs,
    }.items():
        data[address] = value
    workbook.save(path)
    workbook.close()


def _record(product: str, *, manual: bool = False) -> dict[str, object]:
    lc = product == "LC"
    return {
        "product_code": product,
        "business_source": "LC 상품 매출원가" if lc else "신사업 상품 매출원가",
        "specification": "4-inch" if lc else "상품",
        "currency": "KRW",
        "mode": "ACTUAL_YTD_DEFAULT" if lc or not manual else "MANUAL_OVERRIDE",
        "calculation_source": "ACTUAL_YTD" if lc or not manual else "MANUAL_OVERRIDE",
        "forecast_month": 7,
        "latest_actual_month": 2,
        "actual_ytd_revenue": 300.0 if lc else 400.0,
        "actual_ytd_cogs": 240.0 if lc else 280.0,
        "actual_ytd_cogs_rate": 0.8 if lc else 0.7,
        "forecast_merchandise_revenue": 1_000.0 if lc else 2_000.0,
        "actual_ytd_derived_cogs": 800.0 if lc else 1_400.0,
        "manual_cogs": 0.0 if manual else None,
        "manual_reason": "명시적 0원 적용" if manual else "",
        "applied_forecast_cogs": 800.0 if lc else (0.0 if manual else 1_400.0),
        "legacy_normalized": False,
        "revenue_source_reference": "Data!E105:F105" if lc else "Data!E1733:F1733",
        "cogs_source_reference": "Data!E1660:F1660" if lc else "Data!E1734:F1734",
        "monthly_rate_source_reference": "Data!E1661:F1661" if lc else "Data!E1736:F1736",
        "forecast_revenue_source_reference": "Data!K105" if lc else "Data!K114",
        "total_forecast_cogs_source_reference": "Data!K1289",
        "canonical_revenue_field": "lc_actual_ytd_merchandise_revenue" if lc else "new_business_actual_ytd_merchandise_revenue",
        "canonical_cogs_field": "lc_actual_ytd_merchandise_cogs" if lc else "new_business_actual_ytd_merchandise_cogs",
        "canonical_rate_field": "lc_actual_ytd_merchandise_cogs_rate" if lc else "new_business_actual_ytd_merchandise_cogs_rate",
        "canonical_forecast_revenue_field": "lc_forecast_merchandise_revenue" if lc else "new_business_forecast_merchandise_revenue",
        "canonical_forecast_cogs_field": "lc_forecast_merchandise_cogs" if lc else "new_business_forecast_merchandise_cogs",
        "source_mapping_version": "forecast-merchandise-v1.1.0",
        "source_mapping_hash": "a" * 64,
        "validation_status": "PASS",
        "actual_cutoff_valid": True,
        "actual_only_numerator": True,
        "actual_only_denominator": True,
        "zero_denominator_check": True,
        "missing_source_check": True,
        "no_forecast_self_reference": True,
        "product_scope_valid": True,
        "mode_valid": True,
        "manual_amount_conflict_free": True,
        "manual_reason_valid": True,
        "no_hardcoded_forecast_cogs": True,
        "golden_source_valid": True,
    }


def test_evidence_formula_trace_manual_zero_reason_and_source_sha(tmp_path):
    source = tmp_path / "golden.xlsx"
    output = tmp_path / "forecast.xlsx"
    _source_workbook(source)
    original_sha = sha256(source.read_bytes()).hexdigest()

    workbook = GoldenWorkbook(source)
    workbook.add_merchandise_cogs_evidence([
        _record("LC"), _record("NEW_BUSINESS", manual=True),
    ])
    workbook.save(output)

    assert sha256(source.read_bytes()).hexdigest() == original_sha
    generated = load_workbook(output, data_only=False)
    assert MERCHANDISE_EVIDENCE_SHEET_NAME in generated.sheetnames
    sheet = generated[MERCHANDISE_EVIDENCE_SHEET_NAME]
    assert sheet["C2"].value == "ACTUAL_YTD_DEFAULT"
    assert sheet["F2"].value == "=SUM(Data!E105:F105)"
    assert sheet["G2"].value == "=SUM(Data!E1660:F1660)"
    assert sheet["H2"].value == "=IFERROR(G2/F2,NA())"
    assert sheet["J2"].value == "=Data!K105"
    assert sheet["L2"].value == "=J2*H2"
    assert sheet["M2"].value == '=IF(C2="MANUAL_OVERRIDE",K2,L2)'
    assert sheet["C3"].value == "MANUAL_OVERRIDE"
    assert sheet["K3"].value == 0
    assert sheet["P3"].value == "명시적 0원 적용"
    assert sheet["M3"].value == '=IF(C3="MANUAL_OVERRIDE",K3,L3)'
    assert sheet["M4"].value == "=SUM(M2:M3)"
    assert sheet["N4"].value == "=Data!K1289"
    assert "forecast-merchandise-v1.1.0" in sheet["AC2"].value
    assert sheet["Q2"].value == sheet["AB2"].value == "PASS"
    generated.close()


def test_chained_evidence_appends_without_expanding_actual_range(tmp_path):
    source = tmp_path / "golden.xlsx"
    first = tmp_path / "forecast-7.xlsx"
    second = tmp_path / "forecast-8.xlsx"
    _source_workbook(source)
    workbook = GoldenWorkbook(source)
    workbook.add_merchandise_cogs_evidence([_record("LC"), _record("NEW_BUSINESS")])
    workbook.save(first)

    records = [_record("LC"), _record("NEW_BUSINESS")]
    for record in records:
        record["forecast_month"] = 8
        record["forecast_revenue_source_reference"] = (
            "Data!L105" if record["product_code"] == "LC" else "Data!L114"
        )
        record["total_forecast_cogs_source_reference"] = "Data!L1289"
    chained = GoldenWorkbook(first)
    chained.add_merchandise_cogs_evidence(records)
    chained.save(second)

    generated = load_workbook(second, data_only=False)
    sheet = generated[MERCHANDISE_EVIDENCE_SHEET_NAME]
    assert sheet.max_row == 7
    assert sheet["F5"].value == "=SUM(Data!E105:F105)"
    assert sheet["J5"].value == "=Data!L105"
    assert sheet["N7"].value == "=Data!L1289"
    generated.close()


def test_manual_zero_reason_is_preserved_in_existing_audit_path(tmp_path):
    source = tmp_path / "golden-zero.xlsx"
    output = tmp_path / "forecast-zero.xlsx"
    _source_workbook(source, goods_cogs=0)
    workbook = GoldenWorkbook(source)
    workbook.set_input("K1289", 0, "cogs.goods", "명시적 0원 적용")
    workbook.save(output)

    generated = load_workbook(output, data_only=False)
    audit = generated["입력반영내역"]
    assert audit["E2"].value == "K1289"
    assert audit["G2"].value == 0
    assert audit["I2"].value == "명시적 0원 적용"
    assert audit["J2"].value == "이동"
    generated.close()
