from __future__ import annotations

import argparse
from collections import Counter
from hashlib import sha256
import json
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from forecast.evidence_reporting import REPORT_SHEETS, audit_reporting_workbook


def _row_with_value(ws: Any, column: str, value: object) -> int | None:
    return next(
        (
            row for row in range(1, ws.max_row + 1)
            if ws[f"{column}{row}"].value == value
        ),
        None,
    )


def _sheet_audit(ws: Any) -> dict[str, Any]:
    formulas = [
        cell.value
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    numeric = [
        cell.coordinate
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
    ]
    widths = {
        column: dimension.width
        for column, dimension in ws.column_dimensions.items()
        if dimension.width is not None
    }
    heights = [
        dimension.height
        for dimension in ws.row_dimensions.values()
        if dimension.height is not None
    ]
    hidden_rows = [
        row for row, dimension in ws.row_dimensions.items() if dimension.hidden
    ]
    formats = Counter(
        cell.number_format
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )
    return {
        "rows": ws.max_row,
        "columns": ws.max_column,
        "formula_count": len(formulas),
        "hard_coded_numeric_count": len(numeric),
        "hard_coded_numeric_cells": numeric[:25],
        "column_widths": widths,
        "row_height_min": min(heights) if heights else None,
        "row_height_max": max(heights) if heights else None,
        "hidden_detail_row_count": len(hidden_rows),
        "outline_level_max": max(
            (dimension.outlineLevel for dimension in ws.row_dimensions.values()),
            default=0,
        ),
        "freeze_panes": str(ws.freeze_panes or ""),
        "gridlines": bool(ws.sheet_view.showGridLines),
        "orientation": ws.page_setup.orientation,
        "fit_to_width": ws.page_setup.fitToWidth,
        "fit_to_height": ws.page_setup.fitToHeight,
        "print_area": str(ws.print_area),
        "print_title_rows": ws.print_title_rows,
        "auto_filter": ws.auto_filter.ref,
        "number_formats": dict(formats),
    }


def audit(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    contract = audit_reporting_workbook(workbook)
    source = workbook["90_원본값"]
    categories = Counter(
        str(source[f"B{row}"].value or "")
        for row in range(6, source.max_row + 1)
        if isinstance(source[f"H{row}"].value, (int, float))
        or isinstance(source[f"J{row}"].value, (int, float))
    )
    effects = workbook["02_손익영향"]
    sales = workbook["03_판매근거"]
    cost = workbook["04_원가근거"]
    effect_formulas = {
        str(effects[f"B{row}"].value): effects[f"D{row}"].value
        for row in range(1, effects.max_row + 1)
        if effects[f"B{row}"].value in {
            "sales_quantity", "sales_mix", "sales_price", "sales_fx", "tariff",
            "material_total", "manufacturing_realized", "inventory_timing",
            "sga_variable", "sga_fixed",
        }
    }
    effects_total = _row_with_value(effects, "C", "Effects Total")
    op_delta = _row_with_value(effects, "C", "OP Delta")
    residual = _row_with_value(effects, "C", "기타 요인")
    sw_sales = _row_with_value(sales, "A", "SW")
    sw_production = _row_with_value(cost, "B", "SW")
    back_production = _row_with_value(cost, "B", "SW+BW+LC")
    new_business_rows = [
        row for row in range(1, sales.max_row + 1)
        if sales[f"A{row}"].value == "신사업"
    ]
    formula_errors = [
        f"{ws.title}!{cell.coordinate}:{cell.value}"
        for ws in workbook.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.startswith("=")
        and any(token in cell.value.upper() for token in (
            "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!",
        ))
    ]
    report = {
        "path": str(path.resolve()),
        "sha256": sha256(path.read_bytes()).hexdigest().upper(),
        "sheet_count": len(workbook.sheetnames),
        "sheet_names": workbook.sheetnames,
        "contract": contract,
        "hard_coded_numeric_by_source_category": dict(categories),
        "hard_coded_derived_duplicates": contract["hard_coded_derived_duplicates"],
        "duplicated_source_values": contract["duplicated_source_values"],
        "derived_source_literal_count": contract["derived_source_literal_count"],
        "formula_errors": formula_errors,
        "calculation": {
            "mode": workbook.calculation.calcMode,
            "full_calc_on_load": workbook.calculation.fullCalcOnLoad,
            "force_full_calc": workbook.calculation.forceFullCalc,
        },
        "unit_legends": {
            name: " | ".join(
                str(cell.value)
                for cell in workbook[name][2]
                if cell.value is not None
            )
            for name in REPORT_SHEETS
        },
        "source_note": source["A3"].value,
        "key_formulas": {
            "official_effect_references": effect_formulas,
            "effects_total": effects[f"D{effects_total}"].value if effects_total else None,
            "op_delta": effects[f"D{op_delta}"].value if op_delta else None,
            "residual": effects[f"D{residual}"].value if residual else None,
            "sw_revenue_delta": sales[f"J{sw_sales}"].value if sw_sales else None,
            "sw_weighted_production_unit_cost": cost[f"I{sw_production}"].value if sw_production else None,
            "back_process_weighted_unit_cost": cost[f"I{back_production}"].value if back_production else None,
        },
        "new_business_quantity_true_blank": all(
            sales[f"E{row}"].value is None
            and sales[f"F{row}"].value is None
            and sales[f"G{row}"].value is None
            for row in new_business_rows
        ),
        "sheets": {name: _sheet_audit(workbook[name]) for name in REPORT_SHEETS},
    }
    if tuple(workbook.sheetnames) != REPORT_SHEETS:
        raise ValueError("sheet contract mismatch")
    if formula_errors:
        raise ValueError(f"formula error tokens: {formula_errors}")
    if not report["new_business_quantity_true_blank"]:
        raise ValueError("new-business quantity cells are not true blank")
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = audit(args.workbook)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
