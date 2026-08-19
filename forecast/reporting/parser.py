from __future__ import annotations

from io import BytesIO
import math
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter, range_boundaries

from ..xlsx_safety import (
    STRICT_REPORTING_XLSX_POLICY,
    XlsxPackageError,
    XlsxSource,
    validate_xlsx_package,
)
from .models import (
    CanonicalMonthlySeries,
    CanonicalNumber,
    CanonicalProductGroup,
    CanonicalSheet,
    DatasetType,
    PnlReportingCanonicalInput,
    PnlReportingValidationResult,
    ValidationErrorCode,
    ValidationIssue,
    ValidationSeverity,
)
from .registry import (
    MANUFACTURING_COGS_ROWS,
    MONTHLY_PNL_ROWS,
    MONTH_KEYS,
    PRODUCT_GROUPS,
    PRODUCT_PNL_ROWS,
    REQUIRED_SHEETS,
    ROW_REGISTRIES,
    SGA_ROWS,
    SHEET_DATA_RANGES,
    SHEET_HEADERS,
    SHEET_MANUFACTURING_COGS,
    SHEET_MONTHLY_PNL,
    SHEET_PRODUCT_PNL,
    SHEET_SGA,
    TEMPLATE_VERSION,
    ProductRowDefinition,
    RowDefinition,
)


MAX_WORKSHEET_ROWS = 10_000
MAX_WORKSHEET_COLUMNS = 256


def parse_pnl_reporting_workbook(
    source: XlsxSource,
    *,
    dataset_type: DatasetType | str,
    reporting_year: int,
    actual_through_month: int | None = None,
    source_filename: str | None = None,
) -> PnlReportingValidationResult:
    """Parse PNL_REPORTING_V1 into a deterministic, base-input-only object."""

    issues: list[ValidationIssue] = []
    resolved_type = _validate_request_metadata(
        dataset_type,
        reporting_year,
        actual_through_month,
        issues,
    )
    if _has_blockers(issues):
        return PnlReportingValidationResult(None, tuple(issues))

    try:
        filename = source_filename or _source_filename(source)
        validate_xlsx_package(
            source,
            file_name=filename,
            policy=STRICT_REPORTING_XLSX_POLICY,
        )
    except (TypeError, XlsxPackageError):
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_TEMPLATE,
                "안전하게 읽을 수 있는 XLSX 파일이 아닙니다.",
                field="file",
            )
        )
        return PnlReportingValidationResult(None, tuple(issues))

    workbook_source: str | Path | BytesIO
    workbook_source = BytesIO(bytes(source)) if isinstance(source, (bytes, bytearray)) else Path(source)
    try:
        workbook = load_workbook(
            workbook_source,
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except Exception:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_TEMPLATE,
                "Workbook 구조를 읽을 수 없습니다.",
                field="file",
            )
        )
        return PnlReportingValidationResult(None, tuple(issues))

    parsed_regular: dict[str, tuple[CanonicalMonthlySeries, ...]] = {}
    parsed_products: tuple[CanonicalProductGroup, ...] = ()
    try:
        _validate_sheets(workbook.sheetnames, issues)
        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            _validate_sheet_shape(worksheet, sheet_name, issues)
            _validate_merged_cells(worksheet, sheet_name, issues)
            _validate_headers(worksheet, sheet_name, issues)
            _warn_unexpected_content(worksheet, sheet_name, issues)

        if SHEET_MONTHLY_PNL in workbook.sheetnames:
            _validate_template_version(workbook[SHEET_MONTHLY_PNL], issues)
            parsed_regular[SHEET_MONTHLY_PNL] = _parse_regular_sheet(
                workbook[SHEET_MONTHLY_PNL],
                MONTHLY_PNL_ROWS,
                key_column=15,
                label_column=1,
                unit_column=2,
                category_column=None,
                month_start_column=3,
                dataset_type=resolved_type,
                actual_through_month=actual_through_month,
                issues=issues,
            )
        if SHEET_MANUFACTURING_COGS in workbook.sheetnames:
            parsed_regular[SHEET_MANUFACTURING_COGS] = _parse_regular_sheet(
                workbook[SHEET_MANUFACTURING_COGS],
                MANUFACTURING_COGS_ROWS,
                key_column=15,
                label_column=1,
                unit_column=2,
                category_column=None,
                month_start_column=3,
                dataset_type=resolved_type,
                actual_through_month=actual_through_month,
                issues=issues,
            )
        if SHEET_SGA in workbook.sheetnames:
            parsed_regular[SHEET_SGA] = _parse_regular_sheet(
                workbook[SHEET_SGA],
                SGA_ROWS,
                key_column=16,
                label_column=1,
                unit_column=3,
                category_column=2,
                month_start_column=4,
                dataset_type=resolved_type,
                actual_through_month=actual_through_month,
                issues=issues,
            )
        if SHEET_PRODUCT_PNL in workbook.sheetnames:
            parsed_products = _parse_product_sheet(
                workbook[SHEET_PRODUCT_PNL],
                dataset_type=resolved_type,
                actual_through_month=actual_through_month,
                issues=issues,
            )
    except Exception:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_TEMPLATE,
                "Workbook 데이터 구조를 안전하게 해석할 수 없습니다.",
            )
        )
    finally:
        workbook.close()

    if _has_blockers(issues):
        return PnlReportingValidationResult(None, tuple(issues))

    canonical = PnlReportingCanonicalInput(
        template_version=TEMPLATE_VERSION,
        dataset_type=resolved_type,
        reporting_year=reporting_year,
        actual_through_month=actual_through_month,
        sheets=(
            CanonicalSheet(
                SHEET_MONTHLY_PNL,
                rows=parsed_regular[SHEET_MONTHLY_PNL],
            ),
            CanonicalSheet(
                SHEET_MANUFACTURING_COGS,
                rows=parsed_regular[SHEET_MANUFACTURING_COGS],
            ),
            CanonicalSheet(
                SHEET_SGA,
                rows=parsed_regular[SHEET_SGA],
            ),
            CanonicalSheet(
                SHEET_PRODUCT_PNL,
                product_groups=parsed_products,
            ),
        ),
    )
    return PnlReportingValidationResult(canonical, tuple(issues))


def _validate_request_metadata(
    dataset_type: DatasetType | str,
    reporting_year: int,
    actual_through_month: int | None,
    issues: list[ValidationIssue],
) -> DatasetType:
    try:
        resolved_type = dataset_type if isinstance(dataset_type, DatasetType) else DatasetType(dataset_type)
    except (TypeError, ValueError):
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_DATASET_TYPE,
                "dataset_type은 PLAN 또는 ACTUAL이어야 합니다.",
                field="dataset_type",
            )
        )
        resolved_type = DatasetType.PLAN

    if (
        not isinstance(reporting_year, int)
        or isinstance(reporting_year, bool)
        or not 1000 <= reporting_year <= 9999
    ):
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_REPORTING_YEAR,
                "reporting_year는 유효한 4자리 연도여야 합니다.",
                field="reporting_year",
            )
        )

    actual_month_is_valid = (
        isinstance(actual_through_month, int)
        and not isinstance(actual_through_month, bool)
        and 1 <= actual_through_month <= 12
    )
    if resolved_type is DatasetType.PLAN and actual_through_month is not None:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_ACTUAL_THROUGH,
                "PLAN에는 actual_through_month를 지정할 수 없습니다.",
                field="actual_through_month",
            )
        )
    if resolved_type is DatasetType.ACTUAL and not actual_month_is_valid:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_ACTUAL_THROUGH,
                "ACTUAL actual_through_month는 1에서 12 사이의 정수여야 합니다.",
                field="actual_through_month",
            )
        )
    return resolved_type


def _validate_sheets(sheet_names: Iterable[str], issues: list[ValidationIssue]) -> None:
    names = tuple(sheet_names)
    for required in REQUIRED_SHEETS:
        if required not in names:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.MISSING_SHEET,
                    f"필수 Sheet '{required}'가 없습니다.",
                    sheet=required,
                )
            )
    for name in names:
        if name not in REQUIRED_SHEETS:
            issues.append(
                _issue(
                    ValidationSeverity.WARNING,
                    ValidationErrorCode.UNEXPECTED_SHEET,
                    f"계약에 포함되지 않은 Sheet '{name}'는 무시됩니다.",
                    sheet=name,
                )
            )
    if all(required in names for required in REQUIRED_SHEETS):
        required_projection = tuple(name for name in names if name in REQUIRED_SHEETS)
        if required_projection != REQUIRED_SHEETS:
            issues.append(
                _issue(
                    ValidationSeverity.WARNING,
                    ValidationErrorCode.SHEET_ORDER_CHANGED,
                    "필수 Sheet 순서가 표준 양식과 다르지만 stable key로 읽습니다.",
                )
            )


def _validate_sheet_shape(worksheet: Any, sheet_name: str, issues: list[ValidationIssue]) -> None:
    if worksheet.max_row > MAX_WORKSHEET_ROWS or worksheet.max_column > MAX_WORKSHEET_COLUMNS:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.INVALID_TEMPLATE,
                "Sheet 크기가 허용 범위를 넘었습니다.",
                sheet=sheet_name,
            )
        )


def _validate_merged_cells(worksheet: Any, sheet_name: str, issues: list[ValidationIssue]) -> None:
    _start_row, end_row, end_column = SHEET_DATA_RANGES[sheet_name]
    for merged_range in worksheet.merged_cells.ranges:
        min_column, min_row, max_column, max_row = range_boundaries(str(merged_range))
        if _ranges_overlap(
            (min_column, min_row, max_column, max_row),
            (1, 1, end_column, end_row),
        ):
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.INVALID_TEMPLATE,
                    "필수 header/data 영역에 merged cell을 사용할 수 없습니다.",
                    sheet=sheet_name,
                    field=str(merged_range),
                )
            )


def _validate_headers(worksheet: Any, sheet_name: str, issues: list[ValidationIssue]) -> None:
    expected = SHEET_HEADERS[sheet_name]
    actual = tuple(worksheet.cell(1, column).value for column in range(1, len(expected) + 1))
    if actual == expected:
        return
    mismatch = next(
        (index for index, (left, right) in enumerate(zip(actual, expected, strict=True), 1) if left != right),
        1,
    )
    issues.append(
        _issue(
            ValidationSeverity.BLOCKING,
            ValidationErrorCode.INVALID_HEADER,
            "Sheet header가 PNL_REPORTING_V1 계약과 다릅니다.",
            sheet=sheet_name,
            field=f"{get_column_letter(mismatch)}1",
        )
    )


def _validate_template_version(worksheet: Any, issues: list[ValidationIssue]) -> None:
    value = worksheet["P2"].value
    if value != TEMPLATE_VERSION:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                ValidationErrorCode.UNSUPPORTED_TEMPLATE_VERSION,
                f"01_월별손익!P2는 '{TEMPLATE_VERSION}'이어야 합니다.",
                sheet=SHEET_MONTHLY_PNL,
                field="P2",
            )
        )


def _warn_unexpected_content(worksheet: Any, sheet_name: str, issues: list[ValidationIssue]) -> None:
    start_row, end_row, end_column = SHEET_DATA_RANGES[sheet_name]
    candidate: Cell | None = None
    if sheet_name == SHEET_MONTHLY_PNL:
        for row in range(start_row + 1, end_row + 1):
            cell = worksheet.cell(row, 16)
            if cell.value is not None:
                candidate = cell
                break
    if candidate is None and worksheet.max_row > end_row:
        for row in worksheet.iter_rows(
            min_row=end_row + 1,
            max_row=min(worksheet.max_row, MAX_WORKSHEET_ROWS),
            min_col=1,
            max_col=min(worksheet.max_column, end_column),
        ):
            candidate = next((cell for cell in row if cell.value is not None), None)
            if candidate is not None:
                break
    if candidate is None and worksheet.max_column > end_column:
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=min(worksheet.max_row, end_row),
            min_col=end_column + 1,
            max_col=min(worksheet.max_column, MAX_WORKSHEET_COLUMNS),
        ):
            candidate = next((cell for cell in row if cell.value is not None), None)
            if candidate is not None:
                break
    if candidate is not None:
        issues.append(
            _issue(
                ValidationSeverity.WARNING,
                ValidationErrorCode.UNEXPECTED_CONTENT,
                "계약 영역 밖의 의미 있는 내용은 canonical input에서 무시됩니다.",
                sheet=sheet_name,
                field=candidate.coordinate,
            )
        )


def _parse_regular_sheet(
    worksheet: Any,
    registry: tuple[RowDefinition, ...],
    *,
    key_column: int,
    label_column: int,
    unit_column: int,
    category_column: int | None,
    month_start_column: int,
    dataset_type: DatasetType,
    actual_through_month: int | None,
    issues: list[ValidationIssue],
) -> tuple[CanonicalMonthlySeries, ...]:
    sheet_name = worksheet.title
    start_row, end_row, _end_column = SHEET_DATA_RANGES[sheet_name]
    expected = {definition.key: definition for definition in registry}
    physical: dict[str, list[int]] = {}

    for row_number in range(start_row, end_row + 1):
        key_cell = worksheet.cell(row_number, key_column)
        key = key_cell.value
        if _is_formula(key_cell):
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.INVALID_VALUE,
                    "stable key에 formula를 사용할 수 없습니다.",
                    sheet=sheet_name,
                    field=key_cell.coordinate,
                )
            )
            continue
        if not isinstance(key, str) or not key:
            continue
        physical.setdefault(key, []).append(row_number)
        if key not in expected:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.UNKNOWN_KEY,
                    f"지원하지 않는 stable key '{key}'입니다.",
                    sheet=sheet_name,
                    row_key=key,
                    field=key_cell.coordinate,
                )
            )

    _validate_key_completeness(expected, physical, sheet_name, issues)
    for key, rows in physical.items():
        if key not in expected and len(rows) > 1:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.DUPLICATE_KEY,
                    f"stable key '{key}'가 중복되었습니다.",
                    sheet=sheet_name,
                    row_key=key,
                )
            )
    _warn_row_order(tuple(expected), physical, sheet_name, issues)

    canonical_rows: list[CanonicalMonthlySeries] = []
    for definition in registry:
        rows = physical.get(definition.key, [])
        if not rows:
            continue
        row_number = rows[0]
        _validate_metadata_cell(
            worksheet.cell(row_number, label_column),
            definition.label,
            ValidationErrorCode.LABEL_MISMATCH,
            "항목명이 PNL_REPORTING_V1 registry와 다릅니다.",
            definition,
            sheet_name,
            issues,
        )
        _validate_metadata_cell(
            worksheet.cell(row_number, unit_column),
            definition.unit,
            ValidationErrorCode.UNIT_MISMATCH,
            "단위가 PNL_REPORTING_V1 registry와 다릅니다.",
            definition,
            sheet_name,
            issues,
        )
        if category_column is not None:
            _validate_metadata_cell(
                worksheet.cell(row_number, category_column),
                definition.category,
                ValidationErrorCode.CATEGORY_MISMATCH,
                "구분이 PNL_REPORTING_V1 registry와 다릅니다.",
                definition,
                sheet_name,
                issues,
            )

        month_cells = tuple(
            worksheet.cell(row_number, month_start_column + offset)
            for offset in range(12)
        )
        if definition.is_input:
            values = _parse_input_cells(
                month_cells,
                dataset_type=dataset_type,
                actual_through_month=actual_through_month,
                sheet_name=sheet_name,
                row_key=definition.key,
                display_label=definition.label,
                product_group_key=None,
                issues=issues,
            )
            canonical_rows.append(CanonicalMonthlySeries(definition.key, values))
        else:
            _inspect_helper_cells(
                month_cells,
                sheet_name=sheet_name,
                row_key=definition.key,
                display_label=definition.label,
                product_group_key=None,
                issues=issues,
            )
    return tuple(canonical_rows)


def _parse_product_sheet(
    worksheet: Any,
    *,
    dataset_type: DatasetType,
    actual_through_month: int | None,
    issues: list[ValidationIssue],
) -> tuple[CanonicalProductGroup, ...]:
    expected = {
        (definition.product_group_key, definition.key): definition
        for definition in PRODUCT_PNL_ROWS
    }
    groups = {definition.key: definition for definition in PRODUCT_GROUPS}
    physical: dict[tuple[str, str], list[int]] = {}
    start_row, end_row, _end_column = SHEET_DATA_RANGES[SHEET_PRODUCT_PNL]

    for row_number in range(start_row, end_row + 1):
        group_cell = worksheet.cell(row_number, 17)
        key_cell = worksheet.cell(row_number, 18)
        group_key = group_cell.value
        row_key = key_cell.value
        if _is_formula(group_cell) or _is_formula(key_cell):
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.INVALID_VALUE,
                    "product stable key에 formula를 사용할 수 없습니다.",
                    sheet=SHEET_PRODUCT_PNL,
                    field=group_cell.coordinate if _is_formula(group_cell) else key_cell.coordinate,
                )
            )
            continue
        if not isinstance(group_key, str) or not isinstance(row_key, str) or not group_key or not row_key:
            continue
        composite = (group_key, row_key)
        physical.setdefault(composite, []).append(row_number)
        if group_key not in groups:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.PRODUCT_GROUP_MISMATCH,
                    f"지원하지 않는 product group key '{group_key}'입니다.",
                    sheet=SHEET_PRODUCT_PNL,
                    row_key=row_key,
                    product_group_key=group_key,
                    field=group_cell.coordinate,
                )
            )
        elif composite not in expected:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.UNKNOWN_KEY,
                    f"지원하지 않는 product metric key '{group_key}/{row_key}'입니다.",
                    sheet=SHEET_PRODUCT_PNL,
                    row_key=row_key,
                    product_group_key=group_key,
                    field=key_cell.coordinate,
                )
            )

    _validate_product_key_completeness(expected, physical, issues)
    for (group_key, row_key), rows in physical.items():
        if (group_key, row_key) not in expected and len(rows) > 1:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.DUPLICATE_KEY,
                    f"product key '{group_key}/{row_key}'가 중복되었습니다.",
                    sheet=SHEET_PRODUCT_PNL,
                    row_key=row_key,
                    product_group_key=group_key,
                )
            )
    _warn_product_row_order(tuple(expected), physical, issues)

    values_by_key: dict[tuple[str, str], CanonicalMonthlySeries] = {}
    for definition in PRODUCT_PNL_ROWS:
        composite = (definition.product_group_key, definition.key)
        rows = physical.get(composite, [])
        if not rows:
            continue
        row_number = rows[0]
        group = groups[definition.product_group_key]
        _validate_product_metadata(
            worksheet,
            row_number,
            group.display,
            group.dimension,
            definition,
            issues,
        )
        month_cells = tuple(worksheet.cell(row_number, 5 + offset) for offset in range(12))
        if definition.is_input:
            values = _parse_input_cells(
                month_cells,
                dataset_type=dataset_type,
                actual_through_month=actual_through_month,
                sheet_name=SHEET_PRODUCT_PNL,
                row_key=definition.key,
                display_label=definition.label,
                product_group_key=definition.product_group_key,
                issues=issues,
            )
            values_by_key[composite] = CanonicalMonthlySeries(definition.key, values)
        else:
            _inspect_helper_cells(
                month_cells,
                sheet_name=SHEET_PRODUCT_PNL,
                row_key=definition.key,
                display_label=definition.label,
                product_group_key=definition.product_group_key,
                issues=issues,
            )

    canonical_groups: list[CanonicalProductGroup] = []
    for group in PRODUCT_GROUPS:
        metrics = tuple(
            values_by_key[(definition.product_group_key, definition.key)]
            for definition in PRODUCT_PNL_ROWS
            if definition.product_group_key == group.key
            and definition.is_input
            and (definition.product_group_key, definition.key) in values_by_key
        )
        canonical_groups.append(CanonicalProductGroup(group.key, metrics))
    return tuple(canonical_groups)


def _validate_key_completeness(
    expected: dict[str, RowDefinition],
    physical: dict[str, list[int]],
    sheet_name: str,
    issues: list[ValidationIssue],
) -> None:
    for key, definition in expected.items():
        rows = physical.get(key, [])
        if not rows:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.MISSING_KEY,
                    f"필수 stable key '{key}'가 없습니다.",
                    sheet=sheet_name,
                    row_key=key,
                    display_label=definition.label,
                )
            )
        elif len(rows) > 1:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.DUPLICATE_KEY,
                    f"stable key '{key}'가 중복되었습니다.",
                    sheet=sheet_name,
                    row_key=key,
                    display_label=definition.label,
                )
            )


def _validate_product_key_completeness(
    expected: dict[tuple[str, str], ProductRowDefinition],
    physical: dict[tuple[str, str], list[int]],
    issues: list[ValidationIssue],
) -> None:
    for composite, definition in expected.items():
        rows = physical.get(composite, [])
        group_key, row_key = composite
        if not rows:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.MISSING_KEY,
                    f"필수 product key '{group_key}/{row_key}'가 없습니다.",
                    sheet=SHEET_PRODUCT_PNL,
                    row_key=row_key,
                    product_group_key=group_key,
                    display_label=definition.label,
                )
            )
        elif len(rows) > 1:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    ValidationErrorCode.DUPLICATE_KEY,
                    f"product key '{group_key}/{row_key}'가 중복되었습니다.",
                    sheet=SHEET_PRODUCT_PNL,
                    row_key=row_key,
                    product_group_key=group_key,
                    display_label=definition.label,
                )
            )


def _warn_row_order(
    expected_order: tuple[str, ...],
    physical: dict[str, list[int]],
    sheet_name: str,
    issues: list[ValidationIssue],
) -> None:
    if all(len(physical.get(key, [])) == 1 for key in expected_order):
        actual_order = tuple(key for key, _rows in sorted(physical.items(), key=lambda item: item[1][0]) if key in expected_order)
        if actual_order != expected_order:
            issues.append(
                _issue(
                    ValidationSeverity.WARNING,
                    ValidationErrorCode.ROW_ORDER_CHANGED,
                    "row 순서가 표준 양식과 다르지만 stable key 순서로 canonicalize합니다.",
                    sheet=sheet_name,
                )
            )


def _warn_product_row_order(
    expected_order: tuple[tuple[str, str], ...],
    physical: dict[tuple[str, str], list[int]],
    issues: list[ValidationIssue],
) -> None:
    if all(len(physical.get(key, [])) == 1 for key in expected_order):
        actual_order = tuple(key for key, _rows in sorted(physical.items(), key=lambda item: item[1][0]) if key in expected_order)
        if actual_order != expected_order:
            issues.append(
                _issue(
                    ValidationSeverity.WARNING,
                    ValidationErrorCode.ROW_ORDER_CHANGED,
                    "product row 순서가 표준 양식과 다르지만 composite key 순서로 canonicalize합니다.",
                    sheet=SHEET_PRODUCT_PNL,
                )
            )


def _validate_metadata_cell(
    cell: Cell,
    expected: Any,
    error_code: ValidationErrorCode,
    message: str,
    definition: RowDefinition,
    sheet_name: str,
    issues: list[ValidationIssue],
) -> None:
    if cell.value != expected or _is_formula(cell) or cell.data_type == "e":
        issues.append(
            _issue(
                ValidationSeverity.BLOCKING,
                error_code,
                message,
                sheet=sheet_name,
                row_key=definition.key,
                display_label=definition.label,
                field=cell.coordinate,
            )
        )


def _validate_product_metadata(
    worksheet: Any,
    row_number: int,
    expected_display: str,
    expected_dimension: str,
    definition: ProductRowDefinition,
    issues: list[ValidationIssue],
) -> None:
    checks = (
        (1, expected_display, ValidationErrorCode.PRODUCT_GROUP_MISMATCH, "제품군 표시명이 registry와 다릅니다."),
        (2, expected_dimension, ValidationErrorCode.PRODUCT_GROUP_MISMATCH, "제품군 규격/dimension이 registry와 다릅니다."),
        (3, definition.label, ValidationErrorCode.LABEL_MISMATCH, "손익 항목명이 registry와 다릅니다."),
        (4, definition.unit, ValidationErrorCode.UNIT_MISMATCH, "제품군 항목 단위가 registry와 다릅니다."),
    )
    for column, expected, code, message in checks:
        cell = worksheet.cell(row_number, column)
        if cell.value != expected or _is_formula(cell) or cell.data_type == "e":
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKING,
                    code,
                    message,
                    sheet=SHEET_PRODUCT_PNL,
                    row_key=definition.key,
                    product_group_key=definition.product_group_key,
                    display_label=definition.label,
                    field=cell.coordinate,
                )
            )


def _parse_input_cells(
    cells: tuple[Cell, ...],
    *,
    dataset_type: DatasetType,
    actual_through_month: int | None,
    sheet_name: str,
    row_key: str,
    display_label: str,
    product_group_key: str | None,
    issues: list[ValidationIssue],
) -> tuple[CanonicalNumber | None, ...]:
    values: list[CanonicalNumber | None] = []
    for month, cell in enumerate(cells, 1):
        value = cell.value
        if value is None:
            if dataset_type is DatasetType.PLAN:
                issues.append(
                    _month_issue(
                        ValidationErrorCode.PLAN_MONTH_MISSING,
                        "PLAN의 applicable 월 값은 비워둘 수 없습니다.",
                        sheet_name,
                        row_key,
                        product_group_key,
                        display_label,
                        month,
                        cell.coordinate,
                    )
                )
            elif month <= int(actual_through_month):
                issues.append(
                    _month_issue(
                        ValidationErrorCode.ACTUAL_MONTH_MISSING,
                        "실적 기준월 이내의 값은 비워둘 수 없습니다.",
                        sheet_name,
                        row_key,
                        product_group_key,
                        display_label,
                        month,
                        cell.coordinate,
                    )
                )
            values.append(None)
            continue

        normalized = _finite_number(cell)
        if dataset_type is DatasetType.ACTUAL and month > int(actual_through_month):
            if normalized is None:
                issues.append(
                    _month_issue(
                        ValidationErrorCode.INVALID_VALUE,
                        "입력 값은 formula/text/boolean/date/error가 아닌 유한한 숫자여야 합니다.",
                        sheet_name,
                        row_key,
                        product_group_key,
                        display_label,
                        month,
                        cell.coordinate,
                    )
                )
            issues.append(
                _month_issue(
                    ValidationErrorCode.ACTUAL_FUTURE_VALUE_PRESENT,
                    "실적 기준월 이후의 값은 blank여야 합니다.",
                    sheet_name,
                    row_key,
                    product_group_key,
                    display_label,
                    month,
                    cell.coordinate,
                )
            )
            values.append(None)
            continue

        if normalized is None:
            issues.append(
                _month_issue(
                    ValidationErrorCode.INVALID_VALUE,
                    "입력 값은 formula/text/boolean/date/error가 아닌 유한한 숫자여야 합니다.",
                    sheet_name,
                    row_key,
                    product_group_key,
                    display_label,
                    month,
                    cell.coordinate,
                )
            )
            values.append(None)
        else:
            values.append(normalized)
    return tuple(values)


def _inspect_helper_cells(
    cells: tuple[Cell, ...],
    *,
    sheet_name: str,
    row_key: str,
    display_label: str,
    product_group_key: str | None,
    issues: list[ValidationIssue],
) -> None:
    for month, cell in enumerate(cells, 1):
        if _is_formula(cell):
            continue
        issues.append(
            _issue(
                ValidationSeverity.WARNING,
                ValidationErrorCode.HELPER_VALUE_CHANGED,
                "helper cell은 authoritative input에서 무시되며 server가 재계산합니다.",
                sheet=sheet_name,
                row_key=row_key,
                product_group_key=product_group_key,
                display_label=display_label,
                month=month,
                field=cell.coordinate,
            )
        )


def _finite_number(cell: Cell) -> CanonicalNumber | None:
    if _is_formula(cell) or cell.data_type == "e":
        return None
    value = cell.value
    if isinstance(value, bool) or type(value) not in {int, float}:
        return None
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if value == 0:
        return 0
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _is_formula(cell: Cell) -> bool:
    return cell.data_type == "f" or (
        isinstance(cell.value, str) and cell.value.startswith("=")
    )


def _month_issue(
    code: ValidationErrorCode,
    message: str,
    sheet_name: str,
    row_key: str,
    product_group_key: str | None,
    display_label: str,
    month: int,
    field: str,
) -> ValidationIssue:
    return _issue(
        ValidationSeverity.BLOCKING,
        code,
        message,
        sheet=sheet_name,
        row_key=row_key,
        product_group_key=product_group_key,
        display_label=display_label,
        month=month,
        field=field,
    )


def _issue(
    severity: ValidationSeverity,
    error_code: ValidationErrorCode,
    safe_message: str,
    *,
    sheet: str | None = None,
    row_key: str | None = None,
    product_group_key: str | None = None,
    display_label: str | None = None,
    month: int | None = None,
    field: str | None = None,
) -> ValidationIssue:
    return ValidationIssue(
        severity=severity,
        error_code=error_code,
        safe_message=safe_message,
        sheet=sheet,
        row_key=row_key,
        product_group_key=product_group_key,
        display_label=display_label,
        month=month,
        field=field,
    )


def _has_blockers(issues: Iterable[ValidationIssue]) -> bool:
    return any(issue.severity is ValidationSeverity.BLOCKING for issue in issues)


def _source_filename(source: XlsxSource) -> str:
    if isinstance(source, (bytes, bytearray)):
        return "PNL_REPORTING_TEMPLATE_V1.xlsx"
    return Path(source).name


def _ranges_overlap(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> bool:
    left_min_column, left_min_row, left_max_column, left_max_row = left
    right_min_column, right_min_row, right_max_column, right_max_row = right
    return not (
        left_max_column < right_min_column
        or right_max_column < left_min_column
        or left_max_row < right_min_row
        or right_max_row < left_min_row
    )
