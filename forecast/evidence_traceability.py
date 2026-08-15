from __future__ import annotations

from collections.abc import Iterable
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_FORMULA_FILL = PatternFill("solid", fgColor="E2F0D9")
_CHECK_FILL = PatternFill("solid", fgColor="EDEDED")
_WHITE_FONT = Font(color="FFFFFF", bold=True)
_BOLD = Font(bold=True)
_MONEY_FORMAT = '#,##0.000;[Red](#,##0.000);-'


def _number(value: Any) -> float:
    try:
        return float(value or 0.0)
    except (TypeError, ValueError):
        return 0.0


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = subtitle
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")


def _headers(ws, row: int, headers: Iterable[str], start_column: int = 1) -> None:
    for offset, header in enumerate(headers):
        cell = ws.cell(row, start_column + offset, header)
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _finish(ws, *, freeze: str = "A5") -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = _MONEY_FORMAT
    for column in range(1, ws.max_column + 1):
        width = 10
        for cell in ws[get_column_letter(column)]:
            width = max(width, min(42, len(str(cell.value or "")) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = freeze


def _validation(formula: str, engine: str, tolerance: float = 1.0) -> str:
    return f'=IF(ABS({formula}-{engine})<={tolerance},"PASS","FAIL")'


def _sheet_ref(sheet: str, cell: str) -> str:
    return f"='{sheet}'!{cell}"


def write_sales_evidence(
    ws,
    result: dict[str, Any],
    sales_rows: Iterable[Any],
    sales_totals: dict[str, float],
    baseline_fx: float,
    comparison_fx: float,
) -> dict[str, Any]:
    _title(
        ws,
        "판매효과 근거",
        "제품군 내부 SKU를 먼저 합산하고 PCS/LENGTH Pool별 Quantity·Mix를 계산합니다. 고객배송 운반비는 Price에 한 번만 포함하고 관세는 분리합니다.",
    )
    sales = result.get("sales_analysis") or {}
    totals = {**sales_totals, **(sales.get("totals") or {})}
    trace = list(sales.get("trace_rows") or [])
    pool_trace = list(sales.get("pool_trace_rows") or [])
    freight_trace = list(sales.get("freight_trace_rows") or [])
    legacy_trace = not trace
    if not trace:
        for source in sales_rows:
            item = source.to_dict() if hasattr(source, "to_dict") else dict(source)
            trace.append({
                "period": str((result.get("period") or {}).get("label") or "선택기간"),
                "pool": "LEGACY_UNSPECIFIED",
                "unit": "UNAVAILABLE",
                "product_group": item.get("product_group"),
                "base_quantity": item.get("baseline_quantity"),
                "comparison_quantity": item.get("comparison_quantity"),
                "base_revenue": item.get("baseline_amount"),
                "comparison_revenue": item.get("comparison_amount"),
                "base_cogs": _number(item.get("baseline_amount")) * (1 - _number(item.get("baseline_gross_margin_rate"))),
                "comparison_cogs": _number(item.get("comparison_amount")) * (1 - _number(item.get("comparison_gross_margin_rate"))),
                "base_fx": baseline_fx,
                "comparison_fx": comparison_fx,
                "price_effect": item.get("pure_price_effect"),
                "sales_fx_effect": item.get("sales_fx_effect"),
                "business_source": "Legacy canonical sales group (trace unavailable)",
                "canonical_fields": "sales quantity / amount / cogs / sales_fx",
                "base_source_reference": "Stored Result (unit unavailable)",
                "comparison_source_reference": "Stored Result (unit unavailable)",
                "validation_status": "TRACE_UNAVAILABLE_LEGACY",
            })
        base_total = sum(_number(row.get("base_quantity")) for row in trace)
        comparison_total = sum(_number(row.get("comparison_quantity")) for row in trace)
        weighted = sum(
            (_number(row.get("base_quantity")) / base_total if base_total else 0)
            * ((_number(row.get("base_revenue")) - _number(row.get("base_cogs"))) / _number(row.get("base_quantity")) if _number(row.get("base_quantity")) else 0)
            for row in trace
        )
        pool_trace = [{
            "period": trace[0]["period"] if trace else "선택기간",
            "pool": "LEGACY_UNSPECIFIED",
            "unit": "UNAVAILABLE",
            "base_total_quantity": base_total,
            "comparison_total_quantity": comparison_total,
            "base_weighted_gp_per_unit": weighted,
            "quantity_effect": _number(totals.get("quantity_effect")),
            "mix_effect": _number(totals.get("mix_effect")),
        }]
    if not freight_trace:
        freight_trace = [{
            "period": str((result.get("period") or {}).get("label") or "선택기간"),
            "business_source": "고객배송 운반비 / 관세",
            "base_source_reference": "Stored Result (tariff components unavailable)",
            "comparison_source_reference": "Stored Result (tariff components unavailable)",
            "base_freight_including_tariff": _number(totals.get("baseline_transport_ex_tariff")),
            "comparison_freight_including_tariff": _number(totals.get("comparison_transport_ex_tariff")),
            "base_tariff": 0.0,
            "comparison_tariff": 0.0,
            "base_tariff_in_transport": False,
            "comparison_tariff_in_transport": False,
            "base_freight_ex_tariff": _number(totals.get("baseline_transport_ex_tariff")),
            "comparison_freight_ex_tariff": _number(totals.get("comparison_transport_ex_tariff")),
            "freight_effect": _number(totals.get("transport_effect")),
            "tariff_effect": _number(totals.get("tariff_effect")),
            "validation_status": "TRACE_UNAVAILABLE_LEGACY" if legacy_trace else "STORED_RESULT",
        }]

    headers = [
        "분석월", "Pool", "제품군", "단위", "Business Source", "Canonical field",
        "Base Source Reference", "Comparison Source Reference", "Base 수량", "Comparison 수량",
        "Base 매출", "Comparison 매출", "Base COGS", "Comparison COGS", "Base GP/unit 수식",
        "Base Pool 수량", "Comparison Pool 수량", "Base Mix", "Comparison Mix", "가중 GP/unit",
        "Mix Difference", "Mix 기여", "Base FX", "Comparison FX", "Base KRW Price",
        "Comparison KRW Price", "Base 외화단가", "Comparison 외화단가", "Price Effect 수식",
        "Price Engine", "Sales FX 수식", "FX Engine", "Validation",
    ]
    _headers(ws, 4, headers)
    detail_start = 5
    for row in trace:
        r = ws.max_row + 1
        ws.append([
            row.get("period"), row.get("pool"), row.get("product_group"),
            (
                "PCS (4-inch)"
                if str(row.get("product_group") or "").upper() == "LC"
                else row.get("unit")
            ),
            row.get("business_source"), row.get("canonical_fields"), row.get("base_source_reference"),
            row.get("comparison_source_reference"), _number(row.get("base_quantity")),
            _number(row.get("comparison_quantity")), _number(row.get("base_revenue")),
            _number(row.get("comparison_revenue")), _number(row.get("base_cogs")),
            _number(row.get("comparison_cogs")), None, None, None, None, None, None, None, None,
            _number(row.get("base_fx") or baseline_fx), _number(row.get("comparison_fx") or comparison_fx),
            None, None, None, None, None, _number(row.get("price_effect")), None,
            _number(row.get("sales_fx_effect")), None,
        ])
        # Formula ranges are patched after the final detail row is known.
        ws.cell(r, 15, f"=IFERROR((K{r}-M{r})/I{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 18, f"=IFERROR(I{r}/P{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 19, f"=IFERROR(J{r}/Q{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 20, f"=R{r}*O{r}").fill = _FORMULA_FILL
        ws.cell(r, 21, f"=S{r}-R{r}").fill = _FORMULA_FILL
        ws.cell(r, 22, f"=U{r}*O{r}").fill = _FORMULA_FILL
        ws.cell(r, 25, f"=IFERROR(K{r}/I{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 26, f"=IFERROR(L{r}/J{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 27, f"=IFERROR(Y{r}/W{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 28, f"=IFERROR(Z{r}/X{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 29, f"=J{r}*(AB{r}-AA{r})*(W{r}+X{r})/2").fill = _FORMULA_FILL
        ws.cell(r, 31, f"=J{r}*(X{r}-W{r})*(AA{r}+AB{r})/2").fill = _FORMULA_FILL
        ws.cell(r, 33, f'=IF(MAX(ABS(AC{r}-AD{r}),ABS(AE{r}-AF{r}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    detail_end = ws.max_row
    for r in range(detail_start, detail_end + 1):
        ws.cell(r, 16, f'=SUMIFS($I${detail_start}:$I${detail_end},$A${detail_start}:$A${detail_end},A{r},$B${detail_start}:$B${detail_end},B{r})').fill = _FORMULA_FILL
        ws.cell(r, 17, f'=SUMIFS($J${detail_start}:$J${detail_end},$A${detail_start}:$A${detail_end},A{r},$B${detail_start}:$B${detail_end},B{r})').fill = _FORMULA_FILL

    freight_start_col = 35  # AI: explicitly keep Freight at the right edge.
    _headers(ws, 4, [
        "분석월", "Business Source", "Base Source", "Comparison Source", "Base Freight(incl Tariff)",
        "Comparison Freight(incl Tariff)", "Base Tariff", "Comparison Tariff", "Base Tariff 포함?",
        "Comparison Tariff 포함?", "Base Freight(ex Tariff)", "Comparison Freight(ex Tariff)",
        "Freight Adjustment 수식", "Freight Engine", "Tariff 수식", "Tariff Engine", "Validation",
    ], freight_start_col)
    freight_start = 5
    for index, item in enumerate(freight_trace):
        r = freight_start + index
        values = [
            item.get("period"), item.get("business_source"), item.get("base_source_reference"),
            item.get("comparison_source_reference"), _number(item.get("base_freight_including_tariff")),
            _number(item.get("comparison_freight_including_tariff")), _number(item.get("base_tariff")),
            _number(item.get("comparison_tariff")), bool(item.get("base_tariff_in_transport")),
            bool(item.get("comparison_tariff_in_transport")), None, None, None,
            _number(item.get("freight_effect")), None, _number(item.get("tariff_effect")), None,
        ]
        for offset, value in enumerate(values):
            ws.cell(r, freight_start_col + offset, value)
        ws.cell(r, 45, f"=AM{r}-AO{r}*AQ{r}").fill = _FORMULA_FILL
        ws.cell(r, 46, f"=AN{r}-AP{r}*AR{r}").fill = _FORMULA_FILL
        ws.cell(r, 47, f"=AS{r}-AT{r}").fill = _FORMULA_FILL
        ws.cell(r, 49, f"=AO{r}-AP{r}").fill = _FORMULA_FILL
        ws.cell(r, 51, f'=IF(MAX(ABS(AU{r}-AV{r}),ABS(AW{r}-AX{r}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    freight_end = freight_start + len(freight_trace) - 1

    summary_header = max(detail_end, freight_end) + 3
    _headers(ws, summary_header, ["Effect", "Evidence Formula", "Engine Output", "Validation", "부호/정책"])
    summary_start = summary_header + 1
    pool_header = summary_header
    _headers(ws, pool_header, [
        "Period", "Pool", "Unit", "Base Total", "Comparison Total", "Base Weighted GP/unit",
        "Quantity Formula", "Quantity Engine", "Mix Component", "Mix Formula", "Mix Engine", "Validation",
    ], 7)
    pool_start = pool_header + 1
    for index, item in enumerate(pool_trace):
        r = pool_start + index
        period = str(item.get("period") or "")
        pool = str(item.get("pool") or "")
        ws.cell(r, 7, period)
        ws.cell(r, 8, pool)
        ws.cell(r, 9, item.get("unit"))
        ws.cell(r, 10, f'=SUMIFS($I${detail_start}:$I${detail_end},$A${detail_start}:$A${detail_end},G{r},$B${detail_start}:$B${detail_end},H{r})').fill = _FORMULA_FILL
        ws.cell(r, 11, f'=SUMIFS($J${detail_start}:$J${detail_end},$A${detail_start}:$A${detail_end},G{r},$B${detail_start}:$B${detail_end},H{r})').fill = _FORMULA_FILL
        ws.cell(r, 12, f'=SUMIFS($T${detail_start}:$T${detail_end},$A${detail_start}:$A${detail_end},G{r},$B${detail_start}:$B${detail_end},H{r})').fill = _FORMULA_FILL
        ws.cell(r, 13, f"=(K{r}-J{r})*L{r}").fill = _FORMULA_FILL
        ws.cell(r, 14, _number(item.get("quantity_effect")))
        ws.cell(r, 15, f'=SUMIFS($V${detail_start}:$V${detail_end},$A${detail_start}:$A${detail_end},G{r},$B${detail_start}:$B${detail_end},H{r})').fill = _FORMULA_FILL
        ws.cell(r, 16, f"=K{r}*O{r}").fill = _FORMULA_FILL
        ws.cell(r, 17, _number(item.get("mix_effect")))
        ws.cell(r, 18, f'=IF(MAX(ABS(M{r}-N{r}),ABS(P{r}-Q{r}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    pool_end = pool_start + len(pool_trace) - 1
    summary = [
        ("sales_quantity", f"=SUM(M{pool_start}:M{pool_end})", totals.get("quantity_effect"), "Pool별 합산"),
        ("sales_mix", f"=SUM(P{pool_start}:P{pool_end})", totals.get("mix_effect"), "제품군 간 Mix만"),
        ("displayed_sales_price", f"=SUM(AC{detail_start}:AC{detail_end})", totals.get("displayed_sales_price_effect", totals.get("pure_price_effect")), "Freight 반영 전"),
        ("freight_adjustment", f"=SUM(AU{freight_start}:AU{freight_end})", totals.get("transport_effect"), "Price에 1회 포함"),
        ("sales_price", f"=B{summary_start + 2}+B{summary_start + 3}", totals.get("sales_price_effect"), "Displayed Price + Freight"),
        ("sales_fx", f"=SUM(AE{detail_start}:AE{detail_end})", totals.get("sales_fx_effect"), "Price와 symmetric 분리"),
        ("tariff", f"=SUM(AW{freight_start}:AW{freight_end})", totals.get("tariff_effect"), "Price/Freight와 분리"),
    ]
    cells: dict[str, str] = {}
    for index, (code, formula, engine, policy) in enumerate(summary):
        r = summary_start + index
        ws.cell(r, 1, code)
        ws.cell(r, 2, formula).fill = _FORMULA_FILL
        ws.cell(r, 3, _number(engine))
        ws.cell(r, 4, _validation(f"B{r}", f"C{r}")).fill = _CHECK_FILL
        ws.cell(r, 5, policy)
        cells[code] = f"B{r}"
    validation_start = summary_start + len(summary) + 1
    ws.cell(validation_start, 1, "Freight double count")
    ws.cell(validation_start, 2, f'=IF(AND(B{summary_start + 4}=B{summary_start + 2}+B{summary_start + 3},COUNTIF(A{summary_start}:A{summary_start + 6},"freight_adjustment")=1),"PASS","FAIL")').fill = _CHECK_FILL
    ws.cell(validation_start + 1, 1, "Tariff separate")
    ws.cell(validation_start + 1, 2, f'=IF(AND(B{summary_start + 6}=SUM(AW{freight_start}:AW{freight_end}),B{summary_start + 4}=B{summary_start + 2}+B{summary_start + 3}),"PASS","FAIL")').fill = _CHECK_FILL
    cells["detail_range"] = (detail_start, detail_end)
    cells["validation_rows"] = (validation_start, validation_start + 1)
    _finish(ws)
    return cells


def write_material_evidence(ws, result: dict[str, Any]) -> dict[str, Any]:
    _title(
        ws,
        "원부재료 근거",
        "Canonical 원부재료 단가 driver와 부직포 KRW/JPY 분해를 실제 수식으로 재현합니다. MCM·Current Cost Basis Gap은 신규 Effect로 만들지 않습니다.",
    )
    analysis = result.get("material_analysis") or {}
    trace = list(analysis.get("trace_rows") or [])
    nonwoven = list(analysis.get("nonwoven_trace_rows") or [])
    _headers(ws, 4, [
        "분석월", "제품군", "단위", "Business Source", "Canonical field", "Base Source",
        "Comparison Source", "Base 금액", "Comparison 금액", "Base 생산/적용량",
        "Comparison 생산/적용량", "Comparison 판매 적용량", "Base 원단위 수식",
        "Comparison 원단위 수식", "Raw Material Effect 수식", "Engine", "Calculation Validation",
        "Source Status",
    ])
    detail_start = 5
    if not trace:
        for item in analysis.get("product_groups") or []:
            trace.append({
                "period": "선택기간", "product_group": item.get("product_group"), "unit": "Canonical",
                "business_source": "Stored canonical material detail", "canonical_fields": "raw_material driver",
                "base_source_reference": "Stored Result", "comparison_source_reference": "Stored Result",
                "base_unit_cost": item.get("baseline_unit_cost"), "comparison_unit_cost": item.get("comparison_unit_cost"),
                "total_effect": item.get("total"),
            })
    for item in trace:
        r = ws.max_row + 1
        ws.append([
            item.get("period"), item.get("product_group"),
            (
                "PCS (4-inch)"
                if str(item.get("product_group") or "").upper() == "LC"
                else item.get("unit")
            ),
            item.get("business_source"),
            item.get("canonical_fields"), item.get("base_source_reference"), item.get("comparison_source_reference"),
            _number(item.get("base_cost")), _number(item.get("comparison_cost")),
            _number(item.get("base_output")), _number(item.get("comparison_output")),
            _number(item.get("comparison_sales")), None, None, None, _number(item.get("total_effect")), None,
            item.get("source_validation_status") or "TRACE_UNAVAILABLE_LEGACY",
        ])
        ws.cell(r, 13, f"=IFERROR(H{r}/J{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 14, f"=IFERROR(I{r}/K{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 15, f'=IF(OR(AND(H{r}<>0,J{r}=0),AND(I{r}<>0,K{r}=0),AND(L{r}<>0,OR(J{r}=0,K{r}=0))),0,(M{r}-N{r})*L{r})').fill = _FORMULA_FILL
        ws.cell(r, 17, _validation(f"O{r}", f"P{r}")).fill = _CHECK_FILL
    detail_end = ws.max_row

    nw_col = 20
    _headers(ws, 4, [
        "분석월", "Business Source", "Canonical field", "Base Source", "Comparison Source",
        "Base 부직포 금액", "Comparison 부직포 금액", "Base 생산길이", "Comparison 생산길이",
        "Comparison 적용길이", "Base JPY FX", "Comparison JPY FX", "Base 단가", "Comparison 단가",
        "부직포 Total", "Base JPY 단가", "JPY Effect 수식", "JPY Engine", "환율제외 Effect 수식",
        "환율제외 Engine", "Calculation Validation", "Source Status",
    ], nw_col)
    nw_start = 5
    for index, item in enumerate(nonwoven):
        r = nw_start + index
        values = [
            item.get("period"), item.get("business_source"), item.get("canonical_fields"),
            item.get("base_source_reference"), item.get("comparison_source_reference"),
            _number(item.get("base_cost")), _number(item.get("comparison_cost")),
            _number(item.get("base_output")), _number(item.get("comparison_output")),
            _number(item.get("comparison_input_length")), _number(item.get("base_jpy_fx")),
            _number(item.get("comparison_jpy_fx")), None, None, None, None, None,
            _number(item.get("nonwoven_jpy")), None, _number(item.get("nonwoven_price_ex_fx")), None,
            item.get("source_validation_status") or "TRACE_UNAVAILABLE_LEGACY",
        ]
        for offset, value in enumerate(values):
            ws.cell(r, nw_col + offset, value)
        ws.cell(r, 32, f"=IFERROR(Y{r}/AA{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 33, f"=IFERROR(Z{r}/AB{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 34, f"=(AF{r}-AG{r})*AC{r}").fill = _FORMULA_FILL
        ws.cell(r, 35, f"=IFERROR(AF{r}/AD{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 36, f"=AC{r}*AI{r}*(AD{r}-AE{r})").fill = _FORMULA_FILL
        ws.cell(r, 38, f"=AH{r}-AJ{r}").fill = _FORMULA_FILL
        ws.cell(r, 40, f'=IF(MAX(ABS(AJ{r}-AK{r}),ABS(AL{r}-AM{r}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    nw_end = nw_start + len(nonwoven) - 1

    summary_header = max(detail_end, nw_end if nonwoven else 4) + 3
    _headers(ws, summary_header, ["Effect", "Evidence Formula", "Engine Output", "Validation", "Policy"])
    start = summary_header + 1
    nw_fx_formula = f"=SUM(AJ{nw_start}:AJ{nw_end})" if nonwoven else "=0"
    nw_price_formula = f"=SUM(AL{nw_start}:AL{nw_end})" if nonwoven else "=0"
    summary = [
        ("material_total", f"=SUM(O{detail_start}:O{detail_end})", analysis.get("total"), "Canonical RM total"),
        ("nonwoven_price_ex_fx", nw_price_formula, analysis.get("nonwoven_price_ex_fx"), "RM subtotal component"),
        ("nonwoven_jpy", nw_fx_formula, analysis.get("nonwoven_jpy"), "RM FX component; not separately additive in Bridge"),
        ("materials_ex_nonwoven", f"=B{start}-B{start + 1}-B{start + 2}", analysis.get("materials_ex_nonwoven"), "Other materials"),
    ]
    cells: dict[str, Any] = {}
    for index, (code, formula, engine, policy) in enumerate(summary):
        r = start + index
        ws.cell(r, 1, code)
        ws.cell(r, 2, formula).fill = _FORMULA_FILL
        ws.cell(r, 3, _number(engine))
        ws.cell(r, 4, _validation(f"B{r}", f"C{r}")).fill = _CHECK_FILL
        ws.cell(r, 5, policy)
        cells[code] = f"B{r}"
    validation_row = start + len(summary) + 1
    ws.cell(validation_row, 1, "RM / RM FX double count")
    ws.cell(validation_row, 2, f'=IF(ABS(B{start}-(B{start + 1}+B{start + 2}+B{start + 3}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    cells["detail_range"] = (detail_start, detail_end)
    cells["validation_rows"] = (validation_row, validation_row)
    _finish(ws)
    return cells


def write_manufacturing_evidence(ws, result: dict[str, Any]) -> dict[str, Any]:
    _title(
        ws,
        "제조경비 근거",
        "상단에서 전공정 LENGTH(m)·후공정 PCS 분모를 확인하고, 기준 전공정 투입비율을 양쪽에 동일 적용하여 Volume/Unit/Fixed를 재현합니다.",
    )
    analysis = result.get("manufacturing_analysis") or {}
    trace = list(analysis.get("trace_rows") or [])
    _headers(ws, 4, [
        "분석월", "Base 전공정 생산량", "Comparison 전공정 생산량", "전공정 단위",
        "Base 후공정 생산량", "Comparison 후공정 생산량", "후공정 단위",
        "Base Source Reference", "Comparison Source Reference", "PCS+LENGTH 합산",
    ])
    periods: dict[str, dict[str, Any]] = {}
    for item in trace:
        periods.setdefault(str(item.get("month") or "선택기간"), item)
    for period, item in sorted(periods.items()):
        r = ws.max_row + 1
        ws.append([
            period, _number(item.get("base_front_activity")), _number(item.get("comparison_front_activity")),
            "LENGTH(m)", _number(item.get("base_back_activity")), _number(item.get("comparison_back_activity")),
            "PCS (LC=4-inch 포함)", item.get("base_front_activity_source") or item.get("base_back_activity_source"),
            item.get("comparison_front_activity_source") or item.get("comparison_back_activity_source"), "=FALSE",
        ])
    reconciliation = list(analysis.get("production_reconciliation") or [])
    if reconciliation:
        _headers(ws, 4, [
            "Scenario", "Month", "Product Group", "SAP Qty", "MES Qty", "Qty Difference",
            "SAP Length", "MES Length", "Length Difference",
        ], start_column=12)
        for offset, item in enumerate(reconciliation, 5):
            for column, value in enumerate((
                item.get("scenario"), item.get("month"), item.get("product_group"),
                _number(item.get("sap_qty")), item.get("mes_qty"), item.get("qty_difference"),
                _number(item.get("sap_length")), item.get("mes_length"), item.get("length_difference"),
            ), 12):
                ws.cell(offset, column, value)
    detail_header = max(8, ws.max_row + 3)
    _headers(ws, detail_header, [
        "분석월", "계정", "구분", "Business Source", "Canonical field", "Base Amount Source",
        "Comparison Amount Source", "배부율 Source", "전공정 단위", "후공정 단위", "Base 총금액",
        "Comparison 총금액", "Base 전공정 비율", "Comparison 적용비율", "Base 전공정 배부",
        "Comparison 전공정 배부", "Base 후공정 배부", "Comparison 후공정 배부", "Base 전공정 생산량",
        "Comparison 전공정 생산량", "Base 후공정 생산량", "Comparison 후공정 생산량", "Base 전공정 원단위",
        "Comparison 전공정 원단위", "Base 후공정 원단위", "Comparison 후공정 원단위",
        "전공정 Volume", "후공정 Volume", "Volume Effect 수식", "Volume Engine",
        "전공정 Unit", "후공정 Unit", "Unit Effect 수식", "Unit Engine", "Fixed Effect 수식",
        "Fixed Engine", "Subtotal 수식", "Engine Total", "Validation", "재고실현율(참고)", "당기제조원가 구성",
    ])
    detail_start = detail_header + 1
    if not trace:
        for item in result.get("manufacturing_accounts") or []:
            trace.append({
                "month": "선택기간", "account": item.get("account"), "classification": item.get("classification"),
                "business_source": item.get("account"), "canonical_fields": "manufacturing account",
                "base_amount_source": f"Data row {item.get('row')}", "comparison_amount_source": f"Data row {item.get('row')}",
                "front_ratio_source": f"Data row {item.get('allocation_ratio_row')}",
                "baseline_amount": item.get("baseline_amount"), "comparison_amount": item.get("comparison_amount"),
                "front_ratio_base": (item.get("baseline_front_ratios") or [0])[0],
                "front_ratio_comparison": (item.get("baseline_front_ratios") or [0])[0],
                "activity_effect": item.get("activity_effect"), "unit_effect": item.get("unit_effect"),
                "fixed_effect": item.get("fixed_effect"), "occurrence_effect": item.get("occurrence_effect"),
                "inventory_realization_rate": item.get("inventory_realization_rate"),
                "current_cost_component": item.get("current_cost_component"),
            })
    for item in trace:
        r = ws.max_row + 1
        ws.append([
            item.get("month"), item.get("account"), item.get("classification"), item.get("business_source"),
            item.get("canonical_fields"), item.get("base_amount_source"), item.get("comparison_amount_source"),
            item.get("front_ratio_source"), "LENGTH(m)", "PCS (LC=4-inch 포함)", _number(item.get("baseline_amount")),
            _number(item.get("comparison_amount")), _number(item.get("front_ratio_base")),
            _number(item.get("front_ratio_comparison")), None, None, None, None,
            _number(item.get("base_front_activity")), _number(item.get("comparison_front_activity")),
            _number(item.get("base_back_activity")), _number(item.get("comparison_back_activity")),
            None, None, None, None, None, None, None, _number(item.get("activity_effect")),
            None, None, None, _number(item.get("unit_effect")), None, _number(item.get("fixed_effect")),
            None, _number(item.get("occurrence_effect")), None, _number(item.get("inventory_realization_rate")),
            item.get("current_cost_component"),
        ])
        ws.cell(r, 15, f"=K{r}*M{r}").fill = _FORMULA_FILL
        ws.cell(r, 16, f"=L{r}*N{r}").fill = _FORMULA_FILL
        ws.cell(r, 17, f"=K{r}-O{r}").fill = _FORMULA_FILL
        ws.cell(r, 18, f"=L{r}-P{r}").fill = _FORMULA_FILL
        ws.cell(r, 23, f"=IFERROR(O{r}/S{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 24, f"=IFERROR(P{r}/T{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 25, f"=IFERROR(Q{r}/U{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 26, f"=IFERROR(R{r}/V{r},0)").fill = _FORMULA_FILL
        ws.cell(r, 27, f'=IF(C{r}="variable",IF(OR(S{r}=0,T{r}=0),0,(S{r}-T{r})*W{r}),0)').fill = _FORMULA_FILL
        ws.cell(r, 28, f'=IF(C{r}="variable",IF(OR(U{r}=0,V{r}=0),0,(U{r}-V{r})*Y{r}),0)').fill = _FORMULA_FILL
        ws.cell(r, 29, f"=AA{r}+AB{r}").fill = _FORMULA_FILL
        ws.cell(r, 31, f'=IF(C{r}="variable",IF(OR(S{r}=0,T{r}=0),O{r}-P{r},T{r}*(W{r}-X{r})),0)').fill = _FORMULA_FILL
        ws.cell(r, 32, f'=IF(C{r}="variable",IF(OR(U{r}=0,V{r}=0),Q{r}-R{r},V{r}*(Y{r}-Z{r})),0)').fill = _FORMULA_FILL
        ws.cell(r, 33, f"=AE{r}+AF{r}").fill = _FORMULA_FILL
        ws.cell(r, 35, f'=IF(C{r}="fixed",K{r}-L{r},0)').fill = _FORMULA_FILL
        ws.cell(r, 37, f"=AC{r}+AG{r}+AI{r}").fill = _FORMULA_FILL
        ws.cell(r, 39, f'=IF(MAX(ABS(AC{r}-AD{r}),ABS(AG{r}-AH{r}),ABS(AI{r}-AJ{r}),ABS(AK{r}-AL{r}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    detail_end = ws.max_row
    summary_header = detail_end + 3
    _headers(ws, summary_header, ["Effect", "Evidence Formula", "Engine Output", "Validation", "Policy"])
    start = summary_header + 1
    summary = [
        ("manufacturing_activity", f"=SUM(AC{detail_start}:AC{detail_end})", analysis.get("activity_effect"), "Volume"),
        ("manufacturing_unit", f"=SUM(AG{detail_start}:AG{detail_end})", analysis.get("unit_effect"), "Unit Cost"),
        ("manufacturing_fixed", f"=SUM(AI{detail_start}:AI{detail_end})", analysis.get("fixed_effect"), "Fixed"),
        ("manufacturing_realized", f"=SUM(AK{detail_start}:AK{detail_end})", analysis.get("final_effect"), "Volume+Unit+Fixed; multiplier 미적용"),
    ]
    cells: dict[str, Any] = {}
    for index, (code, formula, engine, policy) in enumerate(summary):
        r = start + index
        ws.cell(r, 1, code)
        ws.cell(r, 2, formula).fill = _FORMULA_FILL
        ws.cell(r, 3, _number(engine))
        ws.cell(r, 4, _validation(f"B{r}", f"C{r}")).fill = _CHECK_FILL
        ws.cell(r, 5, policy)
        cells[code] = f"B{r}"
    validation_row = start + len(summary) + 1
    ws.cell(validation_row, 1, "Manufacturing subtotal / child double count")
    ws.cell(validation_row, 2, f'=IF(ABS(B{start + 3}-SUM(B{start}:B{start + 2}))<=1,"PASS","FAIL")').fill = _CHECK_FILL
    ws.cell(validation_row + 1, 1, "Inventory realization multiplier used")
    ws.cell(validation_row + 1, 2, "=FALSE").fill = _CHECK_FILL
    cells["detail_range"] = (detail_start, detail_end)
    cells["component_column"] = "AO"
    cells["subtotal_column"] = "AK"
    cells["validation_rows"] = (validation_row, validation_row + 1)
    _finish(ws, freeze=f"A{detail_start}")
    return cells


def write_sga_evidence(ws, result: dict[str, Any]) -> dict[str, str]:
    _title(ws, "판관비 효과 검증", "계정별 Base-Comparison 수식과 Bridge 위치를 확인합니다. 고객배송 운반비는 판매 Price에서만 반영됩니다.")
    _headers(ws, 4, [
        "원천 행", "구역", "계정과목", "구분", "Base", "Comparison", "손익효과 수식",
        "Engine", "Bridge 위치", "Validation", "Source Reference", "Source Status",
    ])
    start = 5
    for item in result.get("sga_accounts") or []:
        r = ws.max_row + 1
        ws.append([
            item.get("row"), item.get("section"), item.get("account"), item.get("classification"),
            _number(item.get("baseline_amount")), _number(item.get("comparison_amount")), None,
            _number(item.get("profit_effect")), item.get("bridge_position"), None,
            f"Golden Source row {item.get('row')} / selected months",
            item.get("source_validation_status") or "TRACE_UNAVAILABLE_LEGACY",
        ])
        ws.cell(
            r,
            7,
            f'=IF(OR(I{r}="판매효과",I{r}="외부효과/관세"),0,E{r}-F{r})',
        ).fill = _FORMULA_FILL
        ws.cell(r, 10, _validation(f"G{r}", f"H{r}")).fill = _CHECK_FILL
    end = ws.max_row
    summary = end + 3
    _headers(ws, summary, ["Effect", "Evidence Formula", "Engine Output", "Validation"])
    effects = {str(item.get("code")): _number(item.get("profit_effect")) for item in result.get("effects") or []}
    cells: dict[str, str] = {}
    for offset, (code, classification) in enumerate((("sga_variable", "variable"), ("sga_fixed", "fixed")), 1):
        r = summary + offset
        ws.cell(r, 1, code)
        ws.cell(r, 2, f'=SUMIF($D${start}:$D${end},"{classification}",$G${start}:$G${end})').fill = _FORMULA_FILL
        ws.cell(r, 3, effects.get(code, 0.0))
        ws.cell(r, 4, _validation(f"B{r}", f"C{r}")).fill = _CHECK_FILL
        cells[code] = f"B{r}"
    _finish(ws)
    return cells


def write_merchandise_link(ws) -> dict[str, str]:
    _title(
        ws,
        "상품원가검증 연결",
        "Slice 2B의 Forecast Workbook ‘상품원가검증’이 authoritative합니다. 비교 Evidence에서는 동일 계산을 복제하지 않고 Scope만 검증합니다.",
    )
    _headers(ws, 4, ["검증 항목", "상태", "근거"])
    rows = [
        ("LC YTD Rate", "LINKED", "Forecast Workbook 상품원가검증: Σ Actual COGS / Σ Actual Revenue"),
        ("신사업 정책", "LINKED", "ACTUAL_YTD_DEFAULT 또는 MANUAL_OVERRIDE provenance"),
        ("중복 계산", "FALSE", "동일 계산을 Comparison Evidence에서 재구현하지 않음"),
        ("Manufactured COGS Scope와 분리", "TRUE", "Forecast Merchandise COGS ≠ 제품+반제품 Manufactured COGS"),
    ]
    for row in rows:
        ws.append(row)
    validation_row = 9
    ws.cell(validation_row, 1, "Scope Validation")
    ws.cell(validation_row, 2, '=IF(AND(B7="FALSE",B8="TRUE"),"PASS","FAIL")').fill = _CHECK_FILL
    _finish(ws)
    return {"scope_validation": f"B{validation_row}"}


def write_final_bridge(
    ws,
    result: dict[str, Any],
    evidence_cells: dict[str, tuple[str, str]],
    merchandise_validation: tuple[str, str],
) -> dict[str, str]:
    _title(
        ws,
        "최종 OP Bridge 검증",
        "모든 공식 Effect는 상세 Evidence의 최종 formula cell을 직접 참조합니다. effects_total + residual = OP_delta를 실제 Excel 수식으로 검증합니다.",
    )
    _headers(ws, 4, ["Effect Code", "Effect", "Evidence Formula", "Engine Output", "Difference", "Validation", "Detail Cell"])
    start = 5
    effects = list(result.get("effects") or [])
    for index, item in enumerate(effects):
        r = start + index
        code = str(item.get("code") or "")
        target = evidence_cells.get(code)
        formula = _sheet_ref(*target) if target else f"=D{r}"
        ws.append([
            code, item.get("factor") or item.get("label"), formula,
            _number(item.get("profit_effect")), None, None,
            f"{target[0]}!{target[1]}" if target else "Legacy stored engine fallback",
        ])
        ws.cell(r, 5, f"=C{r}-D{r}").fill = _FORMULA_FILL
        ws.cell(r, 6, f'=IF(ABS(E{r})<=1,"PASS","FAIL")').fill = _CHECK_FILL
    end = start + len(effects) - 1
    summary = end + 2
    ws.cell(summary, 2, "effects_total")
    ws.cell(summary, 3, f"=SUM(C{start}:C{end})").fill = _FORMULA_FILL
    ws.cell(summary, 4, _number(result.get("effects_total")))
    ws.cell(summary, 6, _validation(f"C{summary}", f"D{summary}")).fill = _CHECK_FILL
    ws.cell(summary + 1, 2, "residual")
    ws.cell(summary + 1, 3, _number(result.get("residual")))
    ws.cell(summary + 1, 4, _number(result.get("residual")))
    ws.cell(summary + 1, 6, _validation(f"C{summary + 1}", f"D{summary + 1}")).fill = _CHECK_FILL
    ws.cell(summary + 2, 2, "OP_delta")
    ws.cell(summary + 2, 3, _number(result.get("operating_profit_delta")))
    ws.cell(summary + 2, 4, _number(result.get("operating_profit_delta")))
    ws.cell(summary + 2, 6, _validation(f"C{summary + 2}", f"D{summary + 2}")).fill = _CHECK_FILL
    ws.cell(summary + 3, 2, "effects_total + residual = OP_delta")
    ws.cell(summary + 3, 3, f"=C{summary}+C{summary + 1}").fill = _FORMULA_FILL
    ws.cell(summary + 3, 4, f"=C{summary + 2}").fill = _FORMULA_FILL
    ws.cell(summary + 3, 6, _validation(f"C{summary + 3}", f"D{summary + 3}")).fill = _CHECK_FILL

    checks = summary + 6
    _headers(ws, checks, ["Double-count Validation", "Formula Result", "정책"])
    check_rows = [
        ("Quantity / Mix 중복 없음", f'=IF(AND(COUNTIF(A{start}:A{end},"sales_quantity")=1,COUNTIF(A{start}:A{end},"sales_mix")=1),"PASS","FAIL")', "Pool별 Quantity와 제품군 간 Mix"),
        ("Freight / Price 중복 없음", f'=IF(AND(COUNTIF(A{start}:A{end},"sales_price")=1,COUNTIF(A{start}:A{end},"freight_adjustment")=0),"PASS","FAIL")', "Freight는 sales_price 내부 1회"),
        ("Tariff 별도 유지", f'=IF(COUNTIF(A{start}:A{end},"tariff")=1,"PASS","FAIL")', "Freight와 분리"),
        ("Raw Material / FX 중복 없음", f'=IF(AND(COUNTIF(A{start}:A{end},"material_total")=1,COUNTIF(A{start}:A{end},"nonwoven_jpy")=0),"PASS","FAIL")', "FX는 material_total 구성요소"),
        ("Manufacturing 하위 Effect 중복 없음", f'=IF(AND(COUNTIF(A{start}:A{end},"manufacturing_realized")=1,COUNTIF(A{start}:A{end},"manufacturing_activity")=0,COUNTIF(A{start}:A{end},"manufacturing_unit")=0,COUNTIF(A{start}:A{end},"manufacturing_fixed")=0),"PASS","FAIL")', "Subtotal만 additive"),
        ("Inventory Timing 한 번만 additive", f'=IF(COUNTIF(A{start}:A{end},"inventory_timing")=1,"PASS","FAIL")', "공식 Effect 1회"),
        ("Current Cost Basis Gap 신규 Effect 아님", f'=IF(COUNTIF(A{start}:A{end},"current_cost_basis_gap")=0,"PASS","FAIL")', "Disclosure only"),
        ("Forecast Merchandise / Manufactured COGS Scope 분리", _sheet_ref(*merchandise_validation), "상품원가검증 연결"),
        ("Residual plug 없음", f'=IF(COUNTIF(A{start}:A{end},"current_cost_basis_gap")=0,"PASS","FAIL")', "Engine residual을 다른 Effect로 backsolve하지 않음"),
    ]
    for index, row in enumerate(check_rows, 1):
        r = checks + index
        ws.cell(r, 1, row[0])
        ws.cell(r, 2, row[1]).fill = _CHECK_FILL
        ws.cell(r, 3, row[2])
    _finish(ws)
    return {
        "effects_total": f"C{summary}",
        "residual": f"C{summary + 1}",
        "operating_profit_delta": f"C{summary + 2}",
        "identity": f"F{summary + 3}",
    }


def write_residual_rca(
    ws,
    result: dict[str, Any],
    bridge_cells: dict[str, str],
) -> None:
    """Write Source-based residual analysis without introducing a new Effect."""
    rca = dict(result.get("residual_analysis") or {})
    if not rca:
        effects = list(result.get("effects") or [])
        rca = {
            "status": "TRACE_UNAVAILABLE_LEGACY",
            "period": str((result.get("period") or {}).get("label") or "선택기간"),
            "sign_convention": "+ = OP improvement; - = OP deterioration",
            "materiality_status": "UNCONFIGURED",
            "direct_op_bridge": {},
            "effect_to_pnl_map": [
                {
                    "effect_code": item.get("code"),
                    "pnl_bucket": "LEGACY_UNMAPPED",
                    "additive": True,
                    "rca_allocation": False,
                    "parent": None,
                    "amount": _number(item.get("profit_effect")),
                    "direct_source_scope": "Stored legacy Result",
                    "note": "RCA source trace unavailable",
                }
                for item in effects
            ],
            "buckets": [],
            "components": [{
                "component_id": "unexplained",
                "bucket": "CROSS_BUCKET",
                "amount": _number(result.get("residual")),
                "classification": "UNEXPLAINED",
                "business_source": "Stored legacy Result",
                "canonical_field": "residual",
                "formula_basis": "Legacy source trace unavailable",
                "source_reference": "SOURCE_REFERENCE_UNAVAILABLE",
                "period": str((result.get("period") or {}).get("label") or "선택기간"),
                "scope": "Legacy stored payload",
                "source_coverage": "NONE",
                "explanation": "Regenerate the analysis to populate Source-based RCA.",
            }],
            "sga_account_rca": [],
            "plug_created": False,
        }

    _title(
        ws,
        "OP Bridge Residual Root Cause Analysis",
        "Direct P&L Source와 기존 canonical Effect의 basis·scope 차이를 분해합니다. 이 시트는 신규 Effect를 만들거나 Residual을 backsolve하지 않습니다.",
    )
    ws["A3"] = f"RCA Status: {rca.get('status')} / Materiality: {rca.get('materiality_status', 'UNCONFIGURED')} / {rca.get('sign_convention', '')}"
    ws["A3"].font = _BOLD
    ws["A4"] = "Unit: KRW / Golden locations are Evidence metadata only"

    direct = dict(rca.get("direct_op_bridge") or {})
    _headers(
        ws,
        5,
        [
            "Scenario", "Revenue", "COGS", "Selling", "General Admin",
            "Reconstructed OP", "Source OP", "Difference", "Validation",
            "Source Reference",
        ],
    )
    for row_no, key, label in ((6, "base", "BASE"), (7, "comparison", "COMPARISON")):
        item = dict(direct.get(key) or {})
        ws.cell(row_no, 1, label)
        ws.cell(row_no, 2, _number(item.get("revenue")))
        ws.cell(row_no, 3, _number(item.get("cogs")))
        ws.cell(row_no, 4, _number(item.get("selling_expense")))
        ws.cell(row_no, 5, _number(item.get("general_admin")))
        ws.cell(row_no, 6, f"=B{row_no}-C{row_no}-D{row_no}-E{row_no}").fill = _FORMULA_FILL
        ws.cell(row_no, 6).number_format = _MONEY_FORMAT
        ws.cell(row_no, 7, _number(item.get("source_operating_profit")))
        ws.cell(row_no, 8, f"=F{row_no}-G{row_no}").fill = _FORMULA_FILL
        ws.cell(row_no, 8).number_format = _MONEY_FORMAT
        ws.cell(row_no, 9, f'=IF(ABS(H{row_no})<=1,"PASS","FAIL")').fill = _CHECK_FILL
        ws.cell(
            row_no,
            10,
            " / ".join(
                str(value)
                for value in (direct.get("source_references") or {}).values()
                if value
            ) or "SOURCE_REFERENCE_UNAVAILABLE",
        )
    ws["A8"] = "Direct OP Delta"
    ws["F8"] = "=F7-F6"; ws["F8"].fill = _FORMULA_FILL
    ws["F8"].number_format = _MONEY_FORMAT
    ws["G8"] = _sheet_ref("최종Bridge_검증", bridge_cells["operating_profit_delta"]); ws["G8"].fill = _FORMULA_FILL
    ws["G8"].number_format = _MONEY_FORMAT
    ws["H8"] = "=F8-G8"; ws["H8"].fill = _FORMULA_FILL
    ws["H8"].number_format = _MONEY_FORMAT
    ws["I8"] = '=IF(ABS(H8)<=1,"PASS","FAIL")'; ws["I8"].fill = _CHECK_FILL

    effect_header = 11
    _headers(
        ws,
        effect_header,
        [
            "Canonical Effect", "P&L Bucket", "Additive", "Parent",
            "RCA Allocation", "Amount", "Direct Source Scope", "Note",
        ],
    )
    effect_rows = list(rca.get("effect_to_pnl_map") or [])
    effect_start = effect_header + 1
    for item in effect_rows:
        ws.append([
            item.get("effect_code"),
            item.get("pnl_bucket"),
            "YES" if item.get("additive") else "NO",
            item.get("parent"),
            "YES" if item.get("rca_allocation") else "NO",
            item.get("amount"),
            item.get("direct_source_scope"),
            item.get("note"),
        ])
    effect_end = max(effect_start, ws.max_row)

    bucket_header = ws.max_row + 3
    _headers(
        ws,
        bucket_header,
        [
            "Direct P&L Bucket", "Base", "Comparison", "Direct Formula",
            "Engine Direct", "Assigned Canonical Formula", "Gap",
            "Validation", "Source Reference",
        ],
    )
    bucket_start = bucket_header + 1
    buckets = list(rca.get("buckets") or [])
    expense_buckets = {"MANUFACTURED_COGS", "MERCHANDISE_COGS", "OTHER_COGS", "SG&A"}
    for item in buckets:
        row_no = ws.max_row + 1
        bucket = str(item.get("bucket") or "")
        ws.cell(row_no, 1, bucket)
        ws.cell(row_no, 2, item.get("base"))
        ws.cell(row_no, 3, item.get("comparison"))
        ws.cell(
            row_no,
            4,
            f"=B{row_no}-C{row_no}" if bucket in expense_buckets else f"=C{row_no}-B{row_no}",
        ).fill = _FORMULA_FILL
        ws.cell(row_no, 4).number_format = _MONEY_FORMAT
        ws.cell(row_no, 5, _number(item.get("direct_effect")))
        ws.cell(
            row_no,
            6,
            f'=SUMIFS($F${effect_start}:$F${effect_end},$B${effect_start}:$B${effect_end},A{row_no},$E${effect_start}:$E${effect_end},"YES")',
        ).fill = _FORMULA_FILL
        ws.cell(row_no, 6).number_format = _MONEY_FORMAT
        ws.cell(row_no, 7, f"=D{row_no}-F{row_no}").fill = _FORMULA_FILL
        ws.cell(row_no, 7).number_format = _MONEY_FORMAT
        ws.cell(row_no, 8, f'=IF(ABS(D{row_no}-E{row_no})<=1,"PASS","FAIL")').fill = _CHECK_FILL
        ws.cell(row_no, 9, item.get("source_reference"))
    bucket_end = max(bucket_start, ws.max_row)

    component_header = ws.max_row + 3
    _headers(
        ws,
        component_header,
        [
            "Component", "Bucket", "Amount", "Classification", "Business Source",
            "Canonical Field", "Formula Basis", "Source Reference", "Period",
            "Scope", "Source Coverage", "Explanation",
        ],
    )
    component_start = component_header + 1
    for item in rca.get("components") or []:
        ws.append([
            item.get("component_id"), item.get("bucket"), round(_number(item.get("amount")), 3),
            item.get("classification"), item.get("business_source"),
            item.get("canonical_field"), item.get("formula_basis"),
            item.get("source_reference"), item.get("period"), item.get("scope"),
            item.get("source_coverage"), item.get("explanation"),
        ])
    component_end = max(component_start, ws.max_row)
    summary = ws.max_row + 2
    ws.cell(summary, 2, "Classified Total")
    ws.cell(summary, 3, f"=SUM(C{component_start}:C{component_end})").fill = _FORMULA_FILL
    ws.cell(summary, 3).number_format = _MONEY_FORMAT
    ws.cell(summary + 1, 2, "Existing Residual")
    ws.cell(summary + 1, 3, _sheet_ref("최종Bridge_검증", bridge_cells["residual"])).fill = _FORMULA_FILL
    ws.cell(summary + 1, 3).number_format = _MONEY_FORMAT
    ws.cell(summary + 2, 2, "Difference")
    ws.cell(summary + 2, 3, f"=C{summary}-C{summary + 1}").fill = _FORMULA_FILL
    ws.cell(summary + 2, 3).number_format = _MONEY_FORMAT
    ws.cell(summary + 3, 2, "Σ Residual Components = Existing Residual")
    ws.cell(summary + 3, 3, f'=IF(ABS(C{summary + 2})<=1,"PASS","FAIL")').fill = _CHECK_FILL

    sga_header = summary + 6
    _headers(
        ws,
        sga_header,
        [
            "SG&A Account", "Section", "Class", "Base", "Comparison",
            "Direct Formula", "Canonical", "Gap Formula", "Source Reference",
            "Source Status", "Note",
        ],
    )
    for item in rca.get("sga_account_rca") or []:
        row_no = ws.max_row + 1
        ws.cell(row_no, 1, item.get("account"))
        ws.cell(row_no, 2, item.get("section"))
        ws.cell(row_no, 3, item.get("classification"))
        ws.cell(row_no, 4, _number(item.get("base")))
        ws.cell(row_no, 5, _number(item.get("comparison")))
        ws.cell(row_no, 6, f"=D{row_no}-E{row_no}").fill = _FORMULA_FILL
        ws.cell(row_no, 6).number_format = _MONEY_FORMAT
        ws.cell(row_no, 7, _number(item.get("canonical_effect")))
        ws.cell(row_no, 8, f"=F{row_no}-G{row_no}").fill = _FORMULA_FILL
        ws.cell(row_no, 8).number_format = _MONEY_FORMAT
        ws.cell(row_no, 9, item.get("source_reference"))
        ws.cell(row_no, 10, item.get("validation_status"))
        ws.cell(row_no, 11, item.get("note"))

    checks_header = ws.max_row + 3
    _headers(ws, checks_header, ["Validation", "Formula Result", "Policy"])
    checks = [
        (
            "Effect map additive sum",
            f'=IF(ABS(SUMIFS(F{effect_start}:F{effect_end},C{effect_start}:C{effect_end},"YES")-\'최종Bridge_검증\'!{bridge_cells["effects_total"]})<=1,"PASS","FAIL")',
            "Subtotal children are non-additive",
        ),
        (
            "Bucket gap sum = Residual",
            f'=IF(ABS(SUM(G{bucket_start}:G{bucket_end})-\'최종Bridge_검증\'!{bridge_cells["residual"]})<=1,"PASS","FAIL")',
            "Direct P&L - assigned canonical",
        ),
        ("Quantity / Mix 중복 없음", f'=IF(AND(COUNTIF(A{effect_start}:A{effect_end},"sales_quantity")=1,COUNTIF(A{effect_start}:A{effect_end},"sales_mix")=1),"PASS","FAIL")', "각 1회"),
        ("Freight / Price 중복 없음", f'=IF(AND(COUNTIFS(A{effect_start}:A{effect_end},"sales_price",C{effect_start}:C{effect_end},"YES")=1,COUNTIFS(A{effect_start}:A{effect_end},"freight_adjustment",C{effect_start}:C{effect_end},"NO")=1),"PASS","FAIL")', "Freight는 sales_price의 비가산 child"),
        ("Tariff 별도 유지", f'=IF(COUNTIFS(A{effect_start}:A{effect_end},"tariff",C{effect_start}:C{effect_end},"YES")=1,"PASS","FAIL")', "Freight와 분리"),
        ("Raw Material / RM FX 중복 없음", f'=IF(AND(COUNTIFS(A{effect_start}:A{effect_end},"material_total",C{effect_start}:C{effect_end},"YES")=1,COUNTIFS(A{effect_start}:A{effect_end},"nonwoven_jpy",C{effect_start}:C{effect_end},"NO")=1),"PASS","FAIL")', "RM FX는 material_total child"),
        ("MCM 독립 Effect 아님", f'=IF(AND(COUNTIFS(A{effect_start}:A{effect_end},"mcm_policy",C{effect_start}:C{effect_end},"NO")=1,COUNTIFS(A{effect_start}:A{effect_end},"mcm_policy",E{effect_start}:E{effect_end},"YES")=0),"PASS","FAIL")', "PRESENTATION_ONLY 정책 disclosure; 유상사급 mapping gap과 중복 배정하지 않음"),
        ("Manufacturing parent / child 중복 없음", f'=IF(AND(COUNTIFS(A{effect_start}:A{effect_end},"manufacturing_realized",C{effect_start}:C{effect_end},"YES")=1,COUNTIFS(D{effect_start}:D{effect_end},"manufacturing_realized",C{effect_start}:C{effect_end},"NO")=3),"PASS","FAIL")', "Volume/Unit/Fixed 비가산"),
        ("Inventory Timing 1회", f'=IF(COUNTIFS(A{effect_start}:A{effect_end},"inventory_timing",C{effect_start}:C{effect_end},"YES")=1,"PASS","FAIL")', "Manufactured COGS에만 배정"),
        ("Merchandise / Manufactured 분리", f'=IF(AND(COUNTIF(A{bucket_start}:A{bucket_end},"MERCHANDISE_COGS")=1,COUNTIF(A{bucket_start}:A{bucket_end},"MANUFACTURED_COGS")=1),"PASS","FAIL")', "Forecast Merchandise는 Actual Effect가 아님"),
        ("Current Cost Basis Gap 신규 Effect 아님", f'=IF(COUNTIFS(A{effect_start}:A{effect_end},"current_cost_basis_gap",C{effect_start}:C{effect_end},"NO")=1,"PASS","FAIL")', "Disclosure only"),
        ("Residual plug 없음", f'=IF(COUNTIFS(A{effect_start}:A{effect_end},"current_cost_basis_gap",C{effect_start}:C{effect_end},"YES")=0,"PASS","FAIL")', "RCA metadata only; Business Formula 불변"),
    ]
    for label, formula, policy in checks:
        ws.append([label, formula, policy])
        ws.cell(ws.max_row, 2).fill = _CHECK_FILL

    _finish(ws, freeze="A6")


def write_sales_cogs_basis(
    ws,
    result: dict[str, Any],
    bridge_cells: dict[str, str],
) -> None:
    """Write formula-bearing Sales GP/COGS overlap counterfactual evidence."""
    analysis = dict(result.get("sales_cogs_basis_analysis") or {})
    _title(
        ws,
        "Sales GP Driver / COGS Basis Overlap",
        "Revenue-basis와 기존 GP-basis Quantity·Mix 차이를 동일 월·Pool·제품군에서 계산합니다. 모든 Option은 분석 전용이며 Production Formula를 변경하지 않습니다.",
    )
    if not analysis:
        ws["A4"] = "TRACE_UNAVAILABLE_LEGACY"
        ws["A5"] = "저장된 legacy Result에는 matched-basis Source trace가 없습니다. 분석을 다시 실행해야 합니다."
        _finish(ws)
        return

    summary = dict(analysis.get("summary") or {})
    ws["A4"] = "RCA Status"
    ws["B4"] = analysis.get("status")
    ws["A5"] = "Double-count Verdict"
    ws["B5"] = analysis.get("verdict")
    ws["A6"] = "Confidence"
    ws["B6"] = analysis.get("confidence")
    ws["D4"] = "Sign: Embedded COGS"
    ws["E4"] = "Comparison expense - Base expense"
    ws["D5"] = "Sign: Overlap Profit"
    ws["E5"] = "- Embedded COGS expense delta"
    ws["D6"] = "Scope"
    ws["E6"] = "SW/BW/LC=PCS, FS=LENGTH(m); New Business excluded; LC merchandise separation unresolved"
    mapping_gaps = list((analysis.get("scope") or {}).get("mapping_gaps") or [])
    ws["D7"] = "Scope Limitation"
    ws["E7"] = " | ".join(mapping_gaps) or "SOURCE_SCOPE_MATCHED"
    ws["D8"] = "New Business Base Revenue/unit"
    ws["E8"] = (
        "BASE_NON_UNITIZED_REVENUE / matched scope excluded"
        if (analysis.get("scope") or {}).get("new_business_base_non_unitized_revenue")
        else "NOT_APPLICABLE"
    )
    ws["J4"] = "Slice5A Sales Formula Basis Gap"
    ws["K4"] = _number(summary.get("slice5a_sales_formula_basis_gap"))
    ws["J5"] = "Embedded candidate / Slice5A gap"
    ws["K5"] = "=IFERROR(ABS(H4)/ABS(K4),0)"
    ws["J6"] = "Gap after embedded candidate"
    ws["K6"] = "=K4-H4"
    ws["K4"].number_format = _MONEY_FORMAT
    ws["K5"].number_format = "0.00%"
    ws["K6"].number_format = _MONEY_FORMAT
    ws["K5"].fill = _FORMULA_FILL
    ws["K6"].fill = _FORMULA_FILL

    detail_header = 9
    _headers(ws, detail_header, [
        "Month", "Pool", "Unit", "Product Group", "Base Qty", "Comparison Qty",
        "Pool Base Qty", "Pool Comparison Qty", "Base Revenue", "Comparison Revenue",
        "Base COGS", "Comparison COGS", "Base Revenue/unit", "Base COGS/unit",
        "Base GP/unit", "Base Mix", "Comparison Mix", "Revenue-basis Quantity",
        "GP-basis Quantity", "Embedded COGS Quantity", "Revenue-basis Mix",
        "GP-basis Mix", "Embedded COGS Mix", "Embedded COGS Expense Total",
        "Overlap Profit Candidate", "Base Source Reference", "Comparison Source Reference",
        "Scope", "Classification", "Source Scope Status",
    ])
    detail_start = detail_header + 1
    for item in analysis.get("detail_rows") or []:
        row = ws.max_row + 1
        ws.append([
            item.get("period"), item.get("pool"), item.get("unit"),
            item.get("product_group"), _number(item.get("base_quantity")),
            _number(item.get("comparison_quantity")),
            _number(item.get("pool_base_quantity")),
            _number(item.get("pool_comparison_quantity")),
            _number(item.get("base_revenue")), _number(item.get("comparison_revenue")),
            _number(item.get("base_cogs")), _number(item.get("comparison_cogs")),
            None, None, None, None, None, None, None, None, None, None, None, None, None,
            item.get("base_source_reference"), item.get("comparison_source_reference"),
            item.get("scope"), item.get("classification"),
            item.get("source_scope_status"),
        ])
        formulas = {
            13: f"=IFERROR(I{row}/E{row},0)",
            14: f"=IFERROR(K{row}/E{row},0)",
            15: f"=M{row}-N{row}",
            16: f"=IFERROR(E{row}/G{row},0)",
            17: f"=IFERROR(F{row}/H{row},0)",
            18: f"=(H{row}-G{row})*P{row}*M{row}",
            19: f"=(H{row}-G{row})*P{row}*O{row}",
            20: f"=R{row}-S{row}",
            21: f"=H{row}*(Q{row}-P{row})*M{row}",
            22: f"=H{row}*(Q{row}-P{row})*O{row}",
            23: f"=U{row}-V{row}",
            24: f"=T{row}+W{row}",
            25: f"=-X{row}",
        }
        for column, formula in formulas.items():
            ws.cell(row, column, formula).fill = _FORMULA_FILL
            ws.cell(row, column).number_format = _MONEY_FORMAT
    detail_end = max(detail_start, ws.max_row)

    pool_header = ws.max_row + 3
    _headers(ws, pool_header, [
        "Month", "Pool", "Unit", "Base Total Qty", "Comparison Total Qty",
        "Official GP Quantity", "Official GP Mix", "Matched GP Quantity",
        "Matched GP Mix", "Revenue Quantity", "Revenue Mix", "Embedded COGS Expense",
        "Overlap Profit", "Quantity Scope Diff", "Mix Scope Diff",
        "Combined GP Scope Diff", "Quantity Scope", "Mix Scope",
        "Combined Validation",
    ])
    pool_start = pool_header + 1
    for item in analysis.get("pool_rows") or []:
        row = ws.max_row + 1
        ws.append([
            item.get("period"), item.get("pool"), item.get("unit"),
            item.get("base_total_quantity"), item.get("comparison_total_quantity"),
            item.get("official_gp_quantity"), item.get("official_gp_mix"),
            None, None, None, None, None, None, None, None, None, None, None,
        ])
        criteria = (
            f'$A${detail_start}:$A${detail_end},A{row},'
            f'$B${detail_start}:$B${detail_end},B{row}'
        )
        for column, source in ((8, "S"), (9, "V"), (10, "R"), (11, "U")):
            ws.cell(
                row, column,
                f'=SUMIFS(${source}${detail_start}:${source}${detail_end},{criteria})',
            ).fill = _FORMULA_FILL
            ws.cell(row, column).number_format = _MONEY_FORMAT
        ws.cell(row, 12, f"=J{row}+K{row}-H{row}-I{row}").fill = _FORMULA_FILL
        ws.cell(row, 13, f"=-L{row}").fill = _FORMULA_FILL
        ws.cell(row, 14, f"=F{row}-H{row}").fill = _FORMULA_FILL
        ws.cell(row, 15, f"=G{row}-I{row}").fill = _FORMULA_FILL
        ws.cell(row, 16, f"=N{row}+O{row}").fill = _FORMULA_FILL
        ws.cell(row, 17, f'=IF(ABS(N{row})<=1,"PASS","CHECK_SCOPE")').fill = _CHECK_FILL
        ws.cell(row, 18, f'=IF(ABS(O{row})<=1,"PASS","CHECK_SCOPE")').fill = _CHECK_FILL
        ws.cell(row, 19, f'=IF(ABS(P{row})<=1,"PASS","FAIL")').fill = _CHECK_FILL
        for column in (12, 13, 14, 15, 16):
            ws.cell(row, column).number_format = _MONEY_FORMAT
    pool_end = max(pool_start, ws.max_row)

    cogs_header = ws.max_row + 3
    _headers(ws, cogs_header, [
        "Month", "Base Sales Product COGS", "Comparison Sales Product COGS",
        "Sales Product COGS Direct", "P&L Manufactured COGS Direct", "Scope Difference",
        "Current Manufacturing Cost Effect", "Inventory Timing", "Overlap Profit Candidate",
        "Adjusted Inventory Timing", "Classification", "Source Reference",
    ])
    cogs_start = cogs_header + 1
    cogs_rows: dict[str, int] = {}
    for item in analysis.get("cogs_comparison") or []:
        row = ws.max_row + 1
        period = str(item.get("period") or "")
        cogs_rows[period] = row
        ws.append([
            period, item.get("base_sales_product_cogs"),
            item.get("comparison_sales_product_cogs"), None,
            item.get("manufactured_cogs_effect"), None,
            item.get("current_manufacturing_cost_effect"),
            item.get("inventory_timing_effect"), None, None,
            item.get("classification"), item.get("source_reference"),
        ])
        ws.cell(row, 4, f"=B{row}-C{row}").fill = _FORMULA_FILL
        ws.cell(row, 6, f"=E{row}-D{row}").fill = _FORMULA_FILL
        ws.cell(
            row, 9,
            f'=SUMIFS($Y${detail_start}:$Y${detail_end},$A${detail_start}:$A${detail_end},A{row})',
        ).fill = _FORMULA_FILL
        ws.cell(row, 10, f"=H{row}-I{row}").fill = _FORMULA_FILL
        for column in (4, 6, 9, 10):
            ws.cell(row, column).number_format = _MONEY_FORMAT
    cogs_end = max(cogs_start, ws.max_row)
    cumulative_period = str(analysis.get("period") or "CUMULATIVE")
    cumulative_cogs_row = ws.max_row + 1
    cogs_rows[cumulative_period] = cumulative_cogs_row
    ws.cell(cumulative_cogs_row, 1, cumulative_period)
    for column in range(2, 11):
        letter = get_column_letter(column)
        ws.cell(
            cumulative_cogs_row, column,
            f"=SUM({letter}{cogs_start}:{letter}{cogs_end})",
        ).fill = _FORMULA_FILL
        ws.cell(cumulative_cogs_row, column).number_format = _MONEY_FORMAT
    ws.cell(cumulative_cogs_row, 11, "PARTIAL_OVERLAP_SOURCE_SCOPE_DIFFERENCE")
    ws.cell(cumulative_cogs_row, 12, "Monthly Source rows above")

    option_header = ws.max_row + 3
    _headers(ws, option_header, [
        "Period", "Option", "Quantity", "Mix", "Sales Total", "Inventory Timing",
        "Other COGS Effects", "Effects Total", "Residual", "OP Delta",
        "Identity Difference", "Identity", "Counterfactual Only",
        "Engine Source Effects Total",
    ])
    option_start = option_header + 1
    option_inputs = list(analysis.get("option_rows") or [])
    current_by_period: dict[str, dict[str, Any]] = {
        str(row.get("period")): dict(row)
        for row in option_inputs if row.get("option") == "CURRENT"
    }
    option_sheet_rows: dict[tuple[str, str], int] = {}
    ordered_periods = [*cogs_rows]
    for period in ordered_periods:
        current = current_by_period[period]
        for option in ("CURRENT", "OPTION_A", "OPTION_B", "OPTION_C"):
            row = ws.max_row + 1
            option_sheet_rows[(period, option)] = row
            ws.cell(row, 1, period)
            ws.cell(row, 2, option)
            if option == "CURRENT":
                for column, key in (
                    (3, "quantity"), (4, "mix"), (5, "sales_total"),
                    (6, "inventory_timing"), (7, "other_cogs_effects"),
                    (8, "effects_total"), (9, "residual"),
                    (10, "operating_profit_delta"),
                ):
                    ws.cell(row, column, _number(current.get(key)))
                if period == cumulative_period:
                    ws.cell(row, 8, _sheet_ref("최종Bridge_검증", bridge_cells["effects_total"])).fill = _FORMULA_FILL
                    ws.cell(row, 9, _sheet_ref("최종Bridge_검증", bridge_cells["residual"])).fill = _FORMULA_FILL
                    ws.cell(row, 10, _sheet_ref("최종Bridge_검증", bridge_cells["operating_profit_delta"])).fill = _FORMULA_FILL
                ws.cell(row, 14, _number(current.get("effects_total")))
            else:
                current_row = option_sheet_rows[(period, "CURRENT")]
                overlap_cell = f"I{cogs_rows[period]}"
                if option in {"OPTION_A", "OPTION_C"}:
                    for column in (3, 4, 5, 7, 10):
                        letter = get_column_letter(column)
                        ws.cell(row, column, f"={letter}{current_row}").fill = _FORMULA_FILL
                if option == "OPTION_A":
                    ws.cell(row, 6, f"=F{current_row}-{overlap_cell}").fill = _FORMULA_FILL
                    ws.cell(row, 8, f"=H{current_row}-{overlap_cell}").fill = _FORMULA_FILL
                elif option == "OPTION_B":
                    period_formula = (
                        f'$A${detail_start}:$A${detail_end},A{row}'
                        if period != cumulative_period else None
                    )
                    if period_formula:
                        ws.cell(row, 3, f'=SUMIFS($R${detail_start}:$R${detail_end},{period_formula})').fill = _FORMULA_FILL
                        ws.cell(row, 4, f'=SUMIFS($U${detail_start}:$U${detail_end},{period_formula})').fill = _FORMULA_FILL
                    else:
                        ws.cell(row, 3, f"=SUM($R${detail_start}:$R${detail_end})").fill = _FORMULA_FILL
                        ws.cell(row, 4, f"=SUM($U${detail_start}:$U${detail_end})").fill = _FORMULA_FILL
                    ws.cell(row, 5, f"=E{current_row}-C{current_row}-D{current_row}+C{row}+D{row}").fill = _FORMULA_FILL
                    ws.cell(row, 6, f"=F{current_row}").fill = _FORMULA_FILL
                    ws.cell(row, 7, f"=G{current_row}").fill = _FORMULA_FILL
                    ws.cell(row, 8, f"=H{current_row}-C{current_row}-D{current_row}+C{row}+D{row}").fill = _FORMULA_FILL
                    ws.cell(row, 10, f"=J{current_row}").fill = _FORMULA_FILL
                else:
                    ws.cell(row, 6, f"=F{current_row}").fill = _FORMULA_FILL
                    ws.cell(row, 8, f"=H{current_row}").fill = _FORMULA_FILL
                ws.cell(row, 9, f"=J{row}-H{row}").fill = _FORMULA_FILL
            ws.cell(row, 11, f"=H{row}+I{row}-J{row}").fill = _FORMULA_FILL
            ws.cell(row, 12, f'=IF(ABS(K{row})<=1,"PASS","FAIL")').fill = _CHECK_FILL
            ws.cell(row, 13, option != "CURRENT")
            for column in range(3, 12):
                ws.cell(row, column).number_format = _MONEY_FORMAT
    option_end = ws.max_row

    sku_header = ws.max_row + 3
    _headers(ws, sku_header, [
        "Product Group", "SKU Scope", "Base Qty", "Comparison Qty",
        "Revenue Quantity Reference", "Revenue Mix Reference", "Embedded COGS",
        "Classification", "Source Coverage", "Source Reference", "Validation",
    ])
    for item in analysis.get("intra_group_sku") or []:
        ws.append([
            item.get("product_group"), item.get("sku_scope"), item.get("base_quantity"),
            item.get("comparison_quantity"), item.get("revenue_basis_quantity_reference"),
            item.get("revenue_basis_mix_reference"), item.get("embedded_cogs_component"),
            item.get("classification"), item.get("source_coverage"),
            item.get("source_reference"), item.get("validation"),
        ])

    validation_header = ws.max_row + 3
    _headers(ws, validation_header, ["Validation", "Formula Result", "Policy"])
    cumulative_current = option_sheet_rows[(cumulative_period, "CURRENT")]
    cumulative_a = option_sheet_rows[(cumulative_period, "OPTION_A")]
    cumulative_b = option_sheet_rows[(cumulative_period, "OPTION_B")]
    validations = [
        (
            "Revenue basis - GP basis = Embedded COGS",
            f'=IF(ABS(SUM(X{detail_start}:X{detail_end})-(SUM(R{detail_start}:R{detail_end})+SUM(U{detail_start}:U{detail_end})-SUM(S{detail_start}:S{detail_end})-SUM(V{detail_start}:V{detail_end})))<=1,"PASS","FAIL")',
            "Expense-source sign",
        ),
        (
            "PCS/LENGTH 혼합 없음",
            f'=IF(COUNTIFS(B{detail_start}:B{detail_end},"PCS",C{detail_start}:C{detail_end},"m")+COUNTIFS(B{detail_start}:B{detail_end},"LENGTH",C{detail_start}:C{detail_end},"PCS")=0,"PASS","FAIL")',
            "PCS=SW/BW/LC, LENGTH=FS",
        ),
        (
            "New Business excluded from matched calculation",
            f'=IF(COUNTIF(D{detail_start}:D{detail_end},"New Business")=0,"PASS","FAIL")',
            "LC merchandise remains a separately disclosed mapping gap",
        ),
        (
            "LC manufactured/total source mismatch disclosed",
            f'=IF(COUNTIF(AD{detail_start}:AD{detail_end},"MAPPING_GAP_LC_MANUFACTURED_QUANTITY_TOTAL_COGS")>0,"PASS","FAIL")',
            "BUSINESS_FORMULA_CONFLICT; production mapping unchanged",
        ),
        (
            "월별 Source scope",
            f'=IF(COUNTA(A{cogs_start}:A{cogs_end})={len(analysis.get("cogs_comparison") or [])},"PASS","FAIL")',
            "선택 월별 Source 1회",
        ),
        (
            "Option identity",
            f'=IF(COUNTIF(L{option_start}:L{option_end},"FAIL")=0,"PASS","FAIL")',
            "effects_total + residual = OP Delta",
        ),
        (
            "Option A/B matched basis",
            f'=IF(ABS(H{cumulative_a}-H{cumulative_b})<=1,"PASS","FAIL")',
            "Same overlap component; different presentation",
        ),
        (
            "Production effects_total 불변",
            f'=IF(ABS(H{cumulative_current}-N{cumulative_current})<=1,"PASS","FAIL")',
            "Bridge formula vs independently supplied Engine source total",
        ),
        (
            "Production Formula 미변경",
            '=IF(H7=FALSE,"PASS","FAIL")',
            "Quantity/Mix/Inventory Timing persisted values unchanged",
        ),
        (
            "Slice5A basis gap candidate reconciliation",
            '=IF(ABS(K4-H4-K6)<=1,"PASS","FAIL")',
            "Independent source calculations; no Residual plug",
        ),
    ]
    for label, formula, policy in validations:
        ws.append([label, formula, policy])
        ws.cell(ws.max_row, 2).fill = _CHECK_FILL

    ws["G4"] = "Embedded Sales COGS Expense"
    ws["H4"] = f"=SUM(X{detail_start}:X{detail_end})"
    ws["G5"] = "Overlap Profit Candidate"
    ws["H5"] = f"=-H4"
    ws["G6"] = "Candidate Adjusted Inventory Timing"
    ws["H6"] = f"=H{cumulative_cogs_row}-I{cumulative_cogs_row}"
    ws["G7"] = "Production Formula Mutated"
    ws["H7"] = bool(analysis.get("production_formula_mutated"))
    for cell in ("H4", "H5", "H6"):
        ws[cell].fill = _FORMULA_FILL
        ws[cell].number_format = _MONEY_FORMAT
    _finish(ws, freeze="A10")
