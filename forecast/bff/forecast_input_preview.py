"""Safe, business-level Forecast input workbook template and preview parser.

This module is deliberately independent from the Forecast engine, model
repositories and production allocator.  The workbook is an interchange format
for the browser only: the parser validates dimensions and numbers, then returns
small source-traceable rows that a later Forecast request adapter may consume.
"""

from __future__ import annotations

import math
import re
import zipfile
from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .auth import AccessCodeSessionService
from .dto import (
    ForecastBusinessProductionPreviewRow,
    ForecastInputIssue,
    ForecastInputPreviewResponse,
    ForecastInputUnitSummary,
    ForecastSalesPreviewRow,
)
from .errors import ApiErrorCode, BffError
from .model_ingestion import _validate_xlsx_package
from .production_allocation import BACK_PROCESS, FRONT_PROCESS, UNIT_LENGTH_M, UNIT_PCS
from ..sales_contract import (
    CANONICAL_FORECAST_SALES_BY_DETAIL,
    CANONICAL_FORECAST_SALES_ITEMS,
    LC_MERCHANDISE_CODE,
)


MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_FILENAME = "forecast_input_template.xlsx"

SALES_HEADERS: tuple[str, ...] = (
    "예상매출월", "구분", "상세 구분", "단위", "수량", "매출액",
)
LEGACY_SALES_HEADERS: tuple[str, ...] = (
    "예상매출월", "제품코드", "제품명", "제품군", "수량", "매출액",
)
PRODUCTION_HEADERS: tuple[str, ...] = (
    "예상생산월", "공정", "제품군", "생산수량", "단위",
)

# Keep this compatibility export because tests and downstream tooling import it.
# The tuple shape remains ``code, detail, category, unit``.
SALES_METADATA: tuple[tuple[str, str, str, str], ...] = tuple(
    (item.code, item.detail, item.category, item.unit)
    for item in CANONICAL_FORECAST_SALES_ITEMS
)
SALES_BY_CODE: Mapping[str, tuple[str, str, str]] = {
    code: (name, group, unit) for code, name, group, unit in SALES_METADATA
}
SALES_BY_DETAIL: Mapping[str, tuple[str, str, str]] = {
    item.detail: (item.code, item.category, item.unit)
    for item in CANONICAL_FORECAST_SALES_ITEMS
}

# Old downloaded templates remain uploadable.  They have one LC row and the
# historic OTHER product-group label; the parser converts them to canonical
# preview rows and orchestration supplies a zero LC merchandise row.
LEGACY_SALES_BY_CODE: Mapping[str, tuple[str, str, str]] = {
    code: value for code, value in SALES_BY_CODE.items() if code != LC_MERCHANDISE_CODE
}
LEGACY_SALES_BY_CODE = {
    **LEGACY_SALES_BY_CODE,
    "LC": ("LC (4인치)", "LC", UNIT_PCS),
    "OTHER": ("기타매출", "신사업", "—"),
}

# These six dimensions are the only accepted business production surface.  The
# values intentionally mirror production_allocation's dimension map without
# importing or invoking the allocator.
PRODUCTION_METADATA: tuple[tuple[str, str, str], ...] = (
    (FRONT_PROCESS, "SW", UNIT_LENGTH_M),
    (FRONT_PROCESS, "BW", UNIT_LENGTH_M),
    (FRONT_PROCESS, "TW", UNIT_LENGTH_M),
    (BACK_PROCESS, "SW", UNIT_PCS),
    (BACK_PROCESS, "BW", UNIT_PCS),
    (BACK_PROCESS, "LC", UNIT_PCS),
)
PRODUCTION_BY_DIMENSION: Mapping[tuple[str, str], str] = {
    (process, group): unit for process, group, unit in PRODUCTION_METADATA
}

MAX_PREVIEW_ROWS = 10_000
MAX_PREVIEW_MONTHS = 6
_FORBIDDEN_OOXML_PREFIXES = (
    "xl/externalLinks/", "xl/embeddings/", "xl/activeX/", "xl/ctrlProps/",
)
_FORBIDDEN_OOXML_ENTRIES = frozenset({"xl/calcChain.xml", "xl/vbaProject.bin"})


class ForecastInputWorkbookError(ValueError):
    """A workbook cannot be safely parsed as the business input format."""

    def __init__(self, code: str, message: str = "Forecast input workbook is invalid") -> None:
        super().__init__(message)
        self.code = code


def build_input_template() -> bytes:
    """Create the exact two-sheet input template as XLSX bytes."""

    workbook = Workbook()
    sales = workbook.active
    sales.title = "판매계획"
    production = workbook.create_sheet("생산계획")
    sales.append(SALES_HEADERS)
    for _code, detail, category, unit in SALES_METADATA:
        # Metadata is pre-filled to make the template self-documenting.  Month
        # and numeric cells remain blank and therefore do not become preview
        # rows until an operator fills them.
        sales.append((None, category, detail, unit, None, None))
    production.append(PRODUCTION_HEADERS)
    for process, group, unit in PRODUCTION_METADATA:
        production.append((None, process, group, None, unit))
    _style_template_sheet(sales, SALES_HEADERS, (12, 14, 18, 10, 14, 18))
    _style_template_sheet(production, PRODUCTION_HEADERS, (12, 14, 14, 14, 10))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_input_workbook(
    path: str | Path,
    *,
    start_month: int,
    end_month: int,
    source_filename: str,
) -> ForecastInputPreviewResponse:
    """Parse and validate a staged workbook without touching Forecast state."""

    _validate_period(start_month, end_month)
    source = Path(path)
    filename = _safe_filename(source_filename)
    _validate_package_for_preview(source, filename)
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except Exception as exc:
        raise ForecastInputWorkbookError("malformed_ooxml") from exc
    try:
        if tuple(workbook.sheetnames) != ("판매계획", "생산계획"):
            raise ForecastInputWorkbookError("wrong_sheets")
        sales_ws = workbook["판매계획"]
        production_ws = workbook["생산계획"]
        sales_headers = _validate_worksheet_shape(
            sales_ws, (SALES_HEADERS, LEGACY_SALES_HEADERS)
        )
        _validate_worksheet_shape(production_ws, (PRODUCTION_HEADERS,))
        sales_rows, sales_issues = _parse_sales_rows(
            sales_ws,
            start_month,
            end_month,
            legacy=sales_headers == LEGACY_SALES_HEADERS,
        )
        production_rows, production_issues = _parse_production_rows(production_ws, start_month, end_month)
    finally:
        workbook.close()

    issues = tuple(sales_issues + production_issues)
    if not sales_rows and not production_rows:
        # A downloaded template with untouched metadata must not be applied:
        # the frontend's zero-fill replacement would otherwise silently erase
        # every selected month.  Keep this a structured preview issue so the
        # browser can render the normal source-traceable error surface.
        issues = issues + (
            ForecastInputIssue(
                "empty_input",
                "판매계획 또는 생산계획에 입력행이 없습니다",
                "ERROR",
                True,
                "판매계획",
                1,
                "입력행",
            ),
        )
    blocking = any(item.blocking for item in issues)
    return ForecastInputPreviewResponse(
        valid=not blocking,
        blocking=blocking,
        source_filename=filename,
        sales_rows=tuple(sales_rows),
        business_production_rows=tuple(production_rows),
        issues=issues,
        sales_summary=_summarize_sales(sales_rows),
        production_summary=_summarize_production(production_rows),
    )


class ForecastInputPreviewService:
    """Admin capability wrapper around the pure template/parser functions."""

    def __init__(self, sessions: AccessCodeSessionService) -> None:
        self._sessions = sessions

    def template(self, session_id: str) -> bytes:
        self._sessions.require_admin(session_id)
        return build_input_template()

    def preview(
        self,
        session_id: str,
        path: str | Path,
        *,
        start_month: int,
        end_month: int,
        source_filename: str,
    ) -> ForecastInputPreviewResponse:
        self._sessions.require_admin(session_id)
        try:
            return parse_input_workbook(
                path,
                start_month=start_month,
                end_month=end_month,
                source_filename=source_filename,
            )
        except ForecastInputWorkbookError as exc:
            raise BffError(
                ApiErrorCode.VALIDATION_ERROR,
                "Forecast input workbook is invalid",
                field_errors={"file": exc.code},
            ) from exc


def _validate_period(start_month: Any, end_month: Any) -> None:
    if (
        isinstance(start_month, bool)
        or isinstance(end_month, bool)
        or not isinstance(start_month, int)
        or not isinstance(end_month, int)
        or not 1 <= start_month <= end_month <= 12
    ):
        raise ForecastInputWorkbookError("invalid_period", "Forecast input period is invalid")
    if end_month - start_month + 1 > MAX_PREVIEW_MONTHS:
        raise ForecastInputWorkbookError("period_too_long", "Forecast input preview period is too long")


def _safe_filename(value: Any) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > 255:
        raise ForecastInputWorkbookError("invalid_filename")
    normalized = value.strip().replace("\\", "/")
    name = normalized.rsplit("/", 1)[-1]
    if normalized != name or name in {"", ".", ".."} or any(ord(ch) < 32 or ord(ch) == 127 for ch in name):
        raise ForecastInputWorkbookError("invalid_filename")
    return name


def _validate_package_for_preview(path: Path, filename: str) -> None:
    if path.suffix.casefold() != ".xlsx" or Path(filename).suffix.casefold() != ".xlsx":
        raise ForecastInputWorkbookError("only_xlsx")
    try:
        # Reuse the staging/package checks used by model ingestion.  This does
        # not open user-controlled XML with openpyxl until all ZIP limits pass.
        _validate_xlsx_package(path, filename)
        with zipfile.ZipFile(path) as archive:
            names = {item.filename.replace("\\", "/") for item in archive.infolist()}
            folded_names = {name.casefold() for name in names}
            if "xl/calcchain.xml" in folded_names or "xl/vbaproject.bin" in folded_names:
                raise ForecastInputWorkbookError("unsafe_ooxml")
            if any(
                name.casefold().startswith(prefix.casefold())
                for name in names for prefix in _FORBIDDEN_OOXML_PREFIXES
            ):
                raise ForecastInputWorkbookError("unsafe_ooxml")
            for item in archive.infolist():
                name = item.filename.replace("\\", "/")
                if name.casefold().endswith((".rels", ".xml")):
                    payload = archive.read(item.filename)
                    normalized_xml = payload.lower()
                    if b"<!doctype" in normalized_xml or b"<!entity" in normalized_xml:
                        raise ForecastInputWorkbookError("unsafe_ooxml")
                    if name.casefold().endswith(".rels"):
                        try:
                            relationships = ET.fromstring(payload)
                        except ET.ParseError as exc:
                            raise ForecastInputWorkbookError("malformed_ooxml") from exc
                        if any(
                            relation.tag.rsplit("}", 1)[-1] == "Relationship"
                            and relation.attrib.get("TargetMode", "").strip().casefold() == "external"
                            for relation in relationships.iter()
                        ):
                            raise ForecastInputWorkbookError("unsafe_ooxml")
    except BffError as exc:
        field = exc.error.field_errors.get("file", "invalid_xlsx")
        raise ForecastInputWorkbookError(str(field)) from exc
    except ForecastInputWorkbookError:
        raise
    except (OSError, KeyError, zipfile.BadZipFile, ValueError) as exc:
        raise ForecastInputWorkbookError("invalid_xlsx") from exc


def _validate_worksheet_shape(
    worksheet: Any,
    accepted_headers: tuple[tuple[str, ...], ...],
) -> tuple[str, ...]:
    try:
        width = len(accepted_headers[0])
        if any(len(headers) != width for headers in accepted_headers):
            raise RuntimeError("accepted headers must have one width")
        if worksheet.max_row > MAX_PREVIEW_ROWS + 1 or worksheet.max_column > width:
            raise ForecastInputWorkbookError("worksheet_too_large")
        actual_headers = tuple(
            _cell_value(cell)
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1, max_col=width))
        )
        if actual_headers not in accepted_headers:
            raise ForecastInputWorkbookError("wrong_headers")
        if getattr(worksheet, "merged_cells", ()):
            raise ForecastInputWorkbookError("unsafe_worksheet")
        return actual_headers
    except StopIteration as exc:
        raise ForecastInputWorkbookError("missing_headers") from exc
    except ForecastInputWorkbookError:
        raise
    except Exception as exc:
        raise ForecastInputWorkbookError("malformed_worksheet") from exc


def _cell_value(cell: Any) -> Any:
    # Formula cells are forbidden even when a cached value exists.  A formula
    # string in a text cell is rejected as well to keep the format data-only.
    if getattr(cell, "data_type", None) == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise ForecastInputWorkbookError("formulas_not_allowed")
    return cell.value


def _iter_data_rows(worksheet: Any, width: int) -> Iterable[tuple[int, tuple[Any, ...]]]:
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, max_col=width), start=2):
        values = tuple(_cell_value(cell) for cell in cells)
        if not any(value is not None and (not isinstance(value, str) or value.strip() != "") for value in values):
            continue
        yield row_number, values


def _parse_sales_rows(
    worksheet: Any,
    start_month: int,
    end_month: int,
    *,
    legacy: bool,
):
    rows: list[ForecastSalesPreviewRow] = []
    issues: list[ForecastInputIssue] = []
    seen: set[tuple[int, str]] = set()
    for row_number, values in _iter_data_rows(worksheet, len(SALES_HEADERS)):
        # Metadata-only rows are placeholders from the downloaded template.
        if _is_placeholder_sales(values, legacy=legacy):
            continue
        month = _parse_month(values[0], "예상매출월", "판매계획", row_number, start_month, end_month, issues)
        quantity = _parse_number(values[4], "수량", "판매계획", row_number, issues)
        amount = _parse_number(values[5], "매출액", "판매계획", row_number, issues)
        if legacy:
            code = _text(values[1])
            name = _text(values[2])
            group = _text(values[3])
            metadata = LEGACY_SALES_BY_CODE.get(code or "")
            metadata_valid = metadata is not None
            if metadata is None:
                _issue(issues, "unknown_product_code", "알 수 없는 판매 제품코드입니다", "판매계획", row_number, "제품코드")
            else:
                expected_name, expected_group, _unit = metadata
                if name != expected_name:
                    metadata_valid = False
                    _issue(issues, "product_name_mismatch", "제품코드에 맞지 않는 제품명입니다", "판매계획", row_number, "제품명")
                if group != expected_group:
                    metadata_valid = False
                    _issue(issues, "product_group_mismatch", "제품코드에 맞지 않는 제품군입니다", "판매계획", row_number, "제품군")
        else:
            group = _text(values[1])
            name = _text(values[2])
            unit = _text(values[3])
            metadata = SALES_BY_DETAIL.get(name or "")
            metadata_valid = metadata is not None
            code = metadata[0] if metadata is not None else ""
            if metadata is None:
                _issue(issues, "unknown_sales_detail", "알 수 없는 판매 상세 구분입니다", "판매계획", row_number, "상세 구분")
            else:
                _code, expected_group, expected_unit = metadata
                if group != expected_group:
                    metadata_valid = False
                    _issue(issues, "sales_category_mismatch", "상세 구분에 맞지 않는 구분입니다", "판매계획", row_number, "구분")
                if unit != expected_unit:
                    metadata_valid = False
                    _issue(issues, "sales_unit_mismatch", "상세 구분에 맞지 않는 단위입니다", "판매계획", row_number, "단위")
        if month is None or quantity is None or amount is None or not metadata_valid:
            continue
        key = (month, code)
        if key in seen:
            _issue(issues, "duplicate_sales_row", "동일 월·상세 구분이 중복되었습니다", "판매계획", row_number, "상세 구분")
            continue
        seen.add(key)
        rows.append(ForecastSalesPreviewRow(month, code, name, group, quantity, amount, "판매계획", row_number))
    return rows, issues


def _parse_production_rows(worksheet: Any, start_month: int, end_month: int):
    rows: list[ForecastBusinessProductionPreviewRow] = []
    issues: list[ForecastInputIssue] = []
    seen: set[tuple[int, str, str]] = set()
    for row_number, values in _iter_data_rows(worksheet, len(PRODUCTION_HEADERS)):
        if _is_placeholder_production(values):
            continue
        month = _parse_month(values[0], "예상생산월", "생산계획", row_number, start_month, end_month, issues)
        process = _text(values[1])
        group = _text(values[2])
        quantity = _parse_number(values[3], "생산수량", "생산계획", row_number, issues)
        unit = _text(values[4])
        expected_unit = PRODUCTION_BY_DIMENSION.get((process or "", group or ""))
        if expected_unit is None:
            _issue(issues, "invalid_production_dimension", "허용되지 않는 공정·제품군입니다", "생산계획", row_number, "공정")
        elif unit != expected_unit:
            _issue(issues, "production_unit_mismatch", "공정·제품군에 맞지 않는 단위입니다", "생산계획", row_number, "단위")
        if month is None or quantity is None or expected_unit is None or unit != expected_unit:
            continue
        key = (month, process, group)
        if key in seen:
            _issue(issues, "duplicate_production_row", "동일 월·공정·제품군이 중복되었습니다", "생산계획", row_number, "제품군")
            continue
        seen.add(key)
        rows.append(ForecastBusinessProductionPreviewRow(month, process, group, quantity, unit, "생산계획", row_number))
    return rows, issues


def _is_placeholder_sales(values: tuple[Any, ...], *, legacy: bool) -> bool:
    if values[0] not in (None, "") or values[4] not in (None, "") or values[5] not in (None, ""):
        return False
    if legacy:
        code, name, group = (_text(values[index]) for index in (1, 2, 3))
        metadata = LEGACY_SALES_BY_CODE.get(code)
        return metadata is not None and metadata[:2] == (name, group)
    category, detail, unit = (_text(values[index]) for index in (1, 2, 3))
    item = CANONICAL_FORECAST_SALES_BY_DETAIL.get(detail)
    return item is not None and (item.category, item.unit) == (category, unit)


def _is_placeholder_production(values: tuple[Any, ...]) -> bool:
    if values[0] not in (None, "") or values[3] not in (None, ""):
        return False
    process, group, unit = (_text(values[index]) for index in (1, 2, 4))
    return PRODUCTION_BY_DIMENSION.get((process, group)) == unit


def _text(value: Any) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        return str(value).strip()
    return value.strip()


def _parse_month(value: Any, field: str, sheet: str, row: int, start: int, end: int, issues: list[ForecastInputIssue]) -> int | None:
    if value in (None, ""):
        _issue(issues, "missing_month", "월이 입력되지 않았습니다", sheet, row, field)
        return None
    # Excel users commonly enter a month as the text ``08``.  Normalize only
    # a strict one/two-digit token; whitespace, signs, decimals and arbitrary
    # text remain invalid.  Quantity/amount cells stay numeric-only.
    if isinstance(value, str) and re.fullmatch(r"[0-9]{1,2}", value):
        number = float(int(value))
    else:
        number = _coerce_finite_number(value)
    if number is None or not number.is_integer():
        _issue(issues, "invalid_month", "월은 1~12의 정수여야 합니다", sheet, row, field)
        return None
    month = int(number)
    if not 1 <= month <= 12:
        _issue(issues, "invalid_month", "월은 1~12의 정수여야 합니다", sheet, row, field)
        return None
    if not start <= month <= end:
        _issue(issues, "month_out_of_period", "요청한 예측 기간 밖의 월입니다", sheet, row, field)
        return None
    return month


def _parse_number(value: Any, field: str, sheet: str, row: int, issues: list[ForecastInputIssue]) -> float | None:
    if value in (None, ""):
        _issue(issues, "missing_numeric", "수치가 입력되지 않았습니다", sheet, row, field)
        return None
    number = _coerce_finite_number(value)
    if number is None:
        _issue(issues, "invalid_numeric", "수치는 유한한 숫자여야 합니다", sheet, row, field)
        return None
    if number < 0:
        _issue(issues, "negative_numeric", "수치는 음수가 될 수 없습니다", sheet, row, field)
        return None
    return number


def _coerce_finite_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            number = float(value)
        except (TypeError, ValueError, OverflowError):
            return None
    else:
        return None
    return number if math.isfinite(number) else None


def _issue(issues: list[ForecastInputIssue], code: str, message: str, sheet: str, row: int, field: str) -> None:
    issues.append(ForecastInputIssue(code, message, "ERROR", True, sheet, row, field))


def _style_template_sheet(worksheet: Any, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    """Apply presentation-only formatting without formulas or validation rules."""

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{chr(64 + len(headers))}{max(1, worksheet.max_row)}"
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for index, (cell, width) in enumerate(zip(worksheet[1], widths), start=1):
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = width
    for row in worksheet.iter_rows(min_row=2, max_col=len(headers)):
        row[0].number_format = "0"
        row[0].alignment = Alignment(horizontal="center")


def _summarize_sales(rows: Iterable[ForecastSalesPreviewRow]) -> tuple[ForecastInputUnitSummary, ...]:
    counts: dict[str, int] = defaultdict(int)
    totals: dict[str, Decimal] = defaultdict(Decimal)
    for row in rows:
        unit = SALES_BY_CODE[row.product_code][2]
        counts[unit] += 1
        totals[unit] += Decimal(str(row.quantity))
    return tuple(
        ForecastInputUnitSummary(unit, counts[unit], _summary_float(totals[unit]))
        for unit in (UNIT_PCS, UNIT_LENGTH_M, "L", "—") if counts[unit]
    )


def _summarize_production(rows: Iterable[ForecastBusinessProductionPreviewRow]) -> tuple[ForecastInputUnitSummary, ...]:
    counts: dict[str, int] = defaultdict(int)
    totals: dict[str, Decimal] = defaultdict(Decimal)
    for row in rows:
        counts[row.unit] += 1
        totals[row.unit] += Decimal(str(row.quantity))
    return tuple(
        ForecastInputUnitSummary(unit, counts[unit], _summary_float(totals[unit]))
        for unit in (UNIT_PCS, UNIT_LENGTH_M) if counts[unit]
    )


def _summary_float(value: Decimal) -> float:
    try:
        result = float(value)
    except (OverflowError, ValueError):
        raise ForecastInputWorkbookError("summary_overflow") from None
    if not math.isfinite(result):
        raise ForecastInputWorkbookError("summary_overflow")
    return result


__all__ = [
    "MIME_XLSX", "TEMPLATE_FILENAME", "SALES_HEADERS", "LEGACY_SALES_HEADERS", "PRODUCTION_HEADERS",
    "SALES_METADATA", "PRODUCTION_METADATA", "MAX_PREVIEW_MONTHS", "ForecastInputIssue",
    "ForecastSalesPreviewRow", "ForecastBusinessProductionPreviewRow",
    "ForecastInputUnitSummary", "ForecastInputPreviewResponse",
    "ForecastInputWorkbookError", "ForecastInputPreviewService",
    "build_input_template", "parse_input_workbook",
]
