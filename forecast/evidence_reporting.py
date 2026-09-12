from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import math
import re
from typing import Any, Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break


REPORT_SHEETS = (
    "01_보고요약",
    "02_손익영향",
    "03_판매근거",
    "04_원가근거",
    "90_원본값",
)

FONT_NAME = "맑은 고딕"
NAVY = "1F4E78"
NAVY_DARK = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_NEUTRAL = "E9EFF5"
LIGHT_TOTAL = "DDEBF7"
LIGHT_SOURCE = "F2F2F2"
WHITE = "FFFFFF"
TEXT = "1F2937"
MUTED = "64748B"
BORDER = "CBD5E1"

MONEY_FORMAT = '#,##0;[Red](#,##0);-'
COUNT_FORMAT = '#,##0;[Red](#,##0);-'
UNIT_COST_FORMAT = '#,##0;[Red](#,##0);-'
FX_FORMAT = '#,##0.00'
PERCENT_FORMAT = '0.0%'

THIN_BORDER = Border(bottom=Side(style="thin", color=BORDER))
TOTAL_BORDER = Border(top=Side(style="thin", color=NAVY))

_SINGLE_CELL_SOURCE_RE = re.compile(
    r"^(?:'[^']+'|[A-Za-z_][\w ]*)!\$?[A-Z]{1,3}\$?\d+$"
)
_SOURCE_CELL_TOKEN_RE = re.compile(
    r"(?:'[^']+'|[A-Za-z_][\w ]*)!\$?[A-Z]{1,3}\$?\d+"
)


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def _records(value: Any) -> list[dict[str, Any]]:
    if value is None:
        return []
    rows: list[dict[str, Any]] = []
    for item in value:
        if isinstance(item, Mapping):
            rows.append(dict(item))
        elif hasattr(item, "to_dict"):
            rows.append(dict(item.to_dict()))
        elif hasattr(item, "__dict__"):
            rows.append(dict(item.__dict__))
    return rows


def _quote_sheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def _sheet_ref(sheet: str, cell: str, *, absolute: bool = False) -> str:
    coordinate = cell
    if absolute:
        letters = "".join(character for character in cell if character.isalpha())
        numbers = "".join(character for character in cell if character.isdigit())
        coordinate = f"${letters}${numbers}"
    return f"={_quote_sheet(sheet)}!{coordinate}"


def _formula_ref(sheet: str, cell: str, *, absolute: bool = True) -> str:
    return _sheet_ref(sheet, cell, absolute=absolute)[1:]


def is_single_cell_source_reference(value: Any) -> bool:
    """Return whether a model source names exactly one workbook cell."""

    return bool(_SINGLE_CELL_SOURCE_RE.fullmatch(str(value or "").strip()))


def _validate_model_source_pair(source: Any, value: Any, *, side: str, key: str) -> None:
    source_text = str(source or "").strip()
    number = _number(value)
    if source_text and not is_single_cell_source_reference(source_text):
        raise ValueError(
            f"MODEL_SOURCE {key} {side} source must be one original workbook cell: "
            f"{source_text!r}"
        )
    if source_text and number is None:
        raise ValueError(
            f"MODEL_SOURCE {key} {side} source has no raw numeric value: {source_text!r}"
        )
    if not source_text and number is not None:
        raise ValueError(
            f"MODEL_SOURCE {key} {side} numeric value has no source cell"
        )


def _source_cells(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    if is_single_cell_source_reference(text):
        return [text]
    cells = _SOURCE_CELL_TOKEN_RE.findall(text)
    if not cells:
        return []
    remainder = _SOURCE_CELL_TOKEN_RE.sub("", text)
    if not re.fullmatch(r"[\s,|]+", remainder):
        return []
    return cells


def _source_reader(source: Any) -> Any:
    if source is None:
        return None
    if hasattr(source, "value"):
        return source.value
    try:
        from .workbook import GoldenWorkbook

        return GoldenWorkbook(source).value
    except (OSError, TypeError, ValueError):
        return None


def _read_source_values(source_cells: Iterable[str], source: Any) -> list[float] | None:
    reader = source if callable(source) else _source_reader(source)
    if reader is None:
        return None
    values: list[float] = []
    for reference in source_cells:
        sheet, cell = reference.split("!", 1)
        # GoldenWorkbook is intentionally scoped to the model Data sheet.
        if sheet.strip("'").replace("''", "'") != "Data":
            return None
        raw_value = reader(cell.replace("$", ""))
        value = _number(raw_value)
        if value is None:
            # The adapter's canonical numeric read treats a genuinely blank
            # source cell as zero; preserve that normalization for readback.
            if raw_value in (None, ""):
                value = 0.0
            else:
                return None
        values.append(value)
    return values


def _source_row_number(reference: str) -> int | None:
    match = _SOURCE_CELL_TOKEN_RE.fullmatch(str(reference or "").strip())
    if not match:
        return None
    digits = re.search(r"\d+$", reference)
    return int(digits.group(0)) if digits else None


def _sga_stored_component_value(
    source: str,
    side: str,
    sga_account_rows: Iterable[Mapping[str, Any]],
) -> float | None:
    row_number = _source_row_number(source)
    if row_number is None:
        return None
    value_key = "baseline_amount" if side == "baseline" else "comparison_amount"
    for item in sga_account_rows:
        try:
            if int(item.get("row") or 0) != row_number:
                continue
        except (TypeError, ValueError):
            continue
        value = _number(item.get(value_key, item.get("amount")))
        if value is not None:
            return value
    return None


def _sga_source_components(
    item: Mapping[str, Any],
    *,
    side: str,
    sga_account_rows: Iterable[Mapping[str, Any]],
    source_workbook: Any = None,
) -> list[dict[str, Any]]:
    """Return one raw component per SGA source cell.

    Current comparison results retain individual source-row amounts in
    ``sga_accounts`` even when ``sga_monthly_trace`` aggregates same-named
    accounts across sections.  Read an original workbook first when one is
    supplied so monthly source cells keep their exact values.  For a stored
    result without the original workbook, use the retained row-level evidence
    and only use the trace aggregate for a genuinely single-cell source.
    """

    source_key = "base_source_reference" if side == "baseline" else "comparison_source_reference"
    aggregate_key = "baseline_amount" if side == "baseline" else "comparison_amount"
    structured = _records(item.get(f"{side}_source_components") or [])
    if structured:
        return [
            {
                "source": str(component.get("source") or ""),
                "value": component.get("value"),
                "role": str(component.get("role") or "amount"),
            }
            for component in structured
            if component.get("source")
        ]

    references = _source_cells(item.get(source_key))
    if not references:
        return []
    components: list[dict[str, Any]] = []
    for reference in references:
        values = _read_source_values([reference], source_workbook)
        value = values[0] if values else None
        if value is None and source_workbook is None and len(references) == 1:
            # A persisted monthly trace already carries the exact value for
            # this one source cell.  Prefer it over sga_accounts, whose
            # amounts are period aggregates and can span several months.
            value = _number(item.get(aggregate_key))
            if value is None and side == "baseline":
                value = _number(item.get("base_amount"))
        if value is None:
            value = _sga_stored_component_value(reference, side, sga_account_rows)
        components.append({"source": reference, "value": value, "role": "amount"})
    if any(component["value"] is None for component in components):
        if len(components) == 1:
            aggregate = _number(item.get(aggregate_key, item.get("base_amount")))
            if aggregate is not None:
                components[0]["value"] = aggregate
        else:
            raise ValueError(
                f"SG&A source components are unavailable for {item.get('account') or 'unknown account'}"
            )
    return components


def validate_source_registry_readback(
    registry: SourceRegistry,
    *,
    baseline_workbook: Any = None,
    comparison_workbook: Any = None,
) -> int:
    """Validate MODEL_SOURCE values against their original workbook cells.

    The check is intentionally optional for stored-result exports that do not
    retain the original workbook paths.  When a workbook is supplied, every
    direct source cell is read back through the same workbook adapter used by
    the analysis and must match the value copied to ``90_원본값``.
    """

    readers = {
        "baseline": _source_reader(baseline_workbook),
        "comparison": _source_reader(comparison_workbook),
    }
    checked = 0
    for row in registry.rows:
        if row.hardcode_class != "MODEL_SOURCE":
            continue
        for side, source, expected in (
            ("baseline", row.baseline_source, row.baseline_value),
            ("comparison", row.comparison_source, row.comparison_value),
        ):
            source_text = str(source or "").strip()
            reader = readers[side]
            if not source_text or reader is None:
                continue
            actual_values = _read_source_values([source_text], reader)
            if not actual_values:
                raise ValueError(
                    f"MODEL_SOURCE {row.key} {side} source cell cannot be read: {source_text!r}"
                )
            actual = actual_values[0]
            tolerance = max(1e-6, abs(expected or 0.0) * 1e-9)
            if expected is None or not math.isclose(actual, expected, rel_tol=0.0, abs_tol=tolerance):
                raise ValueError(
                    f"MODEL_SOURCE {row.key} {side} readback mismatch: "
                    f"{source_text} expected {expected!r}, got {actual!r}"
                )
            checked += 1
    return checked


def _join_sum(parts: Iterable[str]) -> str:
    values = [part for part in parts if part]
    if not values:
        return "=0"
    if len(values) == 1:
        return f"={values[0]}"
    return f"=SUM({','.join(values)})"


def _source_part(reference: Any, index: int) -> str:
    parts = [part.strip() for part in str(reference or "").split("|")]
    if 0 <= index < len(parts) and parts[index]:
        return parts[index]
    return str(reference or "")


@dataclass(frozen=True)
class SourceRow:
    key: str
    category: str
    item: str
    basis: str
    period: str
    unit: str
    baseline_source: str
    baseline_value: float | None
    comparison_source: str
    comparison_value: float | None
    notes: str
    hardcode_class: str
    row: int


class SourceRegistry:
    def __init__(self) -> None:
        self.rows: list[SourceRow] = []
        self.by_key: dict[str, SourceRow] = {}
        self.by_source_signature: dict[tuple[Any, ...], SourceRow] = {}
        self.groups: dict[str, list[str]] = {}

    def add(
        self,
        key: str,
        *,
        category: str,
        item: str,
        basis: str = "",
        period: str = "",
        unit: str = "",
        baseline_source: str = "",
        baseline_value: Any = None,
        comparison_source: str = "",
        comparison_value: Any = None,
        notes: str = "",
        hardcode_class: str = "MODEL_SOURCE",
        group: str | None = None,
    ) -> SourceRow:
        if hardcode_class == "MODEL_SOURCE":
            _validate_model_source_pair(
                baseline_source, baseline_value, side="baseline", key=key
            )
            _validate_model_source_pair(
                comparison_source, comparison_value, side="comparison", key=key
            )
        existing = self.by_key.get(key)
        baseline_number = _number(baseline_value)
        comparison_number = _number(comparison_value)
        if existing is not None:
            if (
                existing.baseline_value != baseline_number
                or existing.comparison_value != comparison_number
                or existing.baseline_source != str(baseline_source or "")
                or existing.comparison_source != str(comparison_source or "")
            ):
                raise ValueError(f"conflicting Evidence source key: {key}")
            if group is not None and key not in self.groups.setdefault(group, []):
                self.groups[group].append(key)
            return existing
        source_signature = (
            str(period or ""),
            str(unit or ""),
            str(baseline_source or ""),
            baseline_number,
            str(comparison_source or ""),
            comparison_number,
        )
        reusable = (
            self.by_source_signature.get(source_signature)
            if hardcode_class == "MODEL_SOURCE"
            and (baseline_source or comparison_source)
            else None
        )
        if reusable is not None:
            self.by_key[key] = reusable
            if group is not None and key not in self.groups.setdefault(group, []):
                self.groups[group].append(key)
            return reusable
        row = SourceRow(
            key=key,
            category=str(category),
            item=str(item),
            basis=str(basis or ""),
            period=str(period or ""),
            unit=str(unit or ""),
            baseline_source=str(baseline_source or ""),
            baseline_value=baseline_number,
            comparison_source=str(comparison_source or ""),
            comparison_value=comparison_number,
            notes=str(notes or ""),
            hardcode_class=str(hardcode_class or ""),
            row=6 + len(self.rows),
        )
        self.rows.append(row)
        self.by_key[key] = row
        if hardcode_class == "MODEL_SOURCE" and (baseline_source or comparison_source):
            self.by_source_signature[source_signature] = row
        if group is not None:
            self.groups.setdefault(group, []).append(key)
        return row

    def ref(self, key: str, side: str) -> str:
        grouped = self.groups.get(key)
        if grouped:
            return _join_sum(
                self._single_ref(item, side) for item in grouped
            )[1:]
        return self._single_ref(key, side)

    def _single_ref(self, key: str, side: str) -> str:
        row = self.by_key[key].row
        column = "H" if side == "baseline" else "J"
        return _formula_ref("90_원본값", f"{column}{row}")

    def refs(self, group: str, side: str) -> list[str]:
        return [self.ref(key, side) for key in self.groups.get(group, ())]


def _selected_periods(result: Mapping[str, Any]) -> list[str]:
    periods: list[str] = []
    for path in (
        ("sales_analysis", "trace_rows", "period"),
        ("sales_analysis", "monthly_effects", "period"),
        ("material_analysis", "trace_rows", "period"),
        ("manufacturing_analysis", "trace_rows", "month"),
        ("inventory_analysis", "selected_monthly_details", "period"),
        (None, "sga_monthly_trace", "period"),
    ):
        parent = result if path[0] is None else result.get(path[0]) or {}
        rows = _records(parent.get(path[1]) or [])
        periods.extend(str(row.get(path[2]) or "") for row in rows if row.get(path[2]))
    return sorted(dict.fromkeys(periods))


def _monthly_fx(
    sales: Mapping[str, Any],
    side: str,
    period: str,
    trace_rows: list[dict[str, Any]],
) -> float | None:
    mapping = sales.get(f"{side}_sales_fx_monthly") or {}
    for key in (period, period[-2:] if len(period) >= 2 else period):
        if key in mapping:
            return _number(mapping[key])
    field = "base_fx" if side == "baseline" else "comparison_fx"
    for row in trace_rows:
        if str(row.get("period") or "") == period and row.get(field) is not None:
            return _number(row.get(field))
    return _number(sales.get(f"{side}_fx_krw_per_usd"))


def _component_identity(component: Mapping[str, Any], index: int) -> tuple[str, str, str]:
    return (
        str(component.get("term_id") or f"component:{index}"),
        str(component.get("role") or "component"),
        str(component.get("source") or ""),
    )


def _component_map(value: Any) -> tuple[list[tuple[tuple[str, str, str], dict[str, Any]]], dict[tuple[str, str, str], dict[str, Any]]]:
    ordered: list[tuple[tuple[str, str, str], dict[str, Any]]] = []
    by_identity: dict[tuple[str, str, str], dict[str, Any]] = {}
    for index, item in enumerate(_records(value), 1):
        identity = _component_identity(item, index)
        # Duplicate source entries with the same semantic role are kept
        # distinct so each arithmetic term remains reconstructible.
        if identity in by_identity:
            identity = (*identity[:2], f"{identity[2]}#{index}")
        by_identity[identity] = item
        ordered.append((identity, item))
    return ordered, by_identity


def _register_raw_component_pairs(
    registry: SourceRegistry,
    *,
    key_prefix: str,
    category: str,
    item: str,
    basis: str,
    period: str,
    unit_by_role: Mapping[str, str] | None,
    baseline_components: Any,
    comparison_components: Any,
    notes: str = "",
) -> tuple[dict[tuple[str, str, str], str], list[tuple[tuple[str, str, str], dict[str, Any], dict[str, Any]]]]:
    baseline_ordered, baseline_by_identity = _component_map(baseline_components)
    comparison_ordered, comparison_by_identity = _component_map(comparison_components)
    identities = [identity for identity, _ in baseline_ordered]
    identities.extend(
        identity for identity, _ in comparison_ordered if identity not in identities
    )
    keys: dict[tuple[str, str, str], str] = {}
    paired: list[tuple[tuple[str, str, str], dict[str, Any], dict[str, Any]]] = []
    for index, identity in enumerate(identities, 1):
        baseline = baseline_by_identity.get(identity, {})
        comparison = comparison_by_identity.get(identity, {})
        baseline_source = str(baseline.get("source") or "")
        comparison_source = str(comparison.get("source") or "")
        baseline_value = baseline.get("value") if baseline else None
        comparison_value = comparison.get("value") if comparison else None
        hardcode_class = "MODEL_SOURCE"
        if (
            (baseline_source and not is_single_cell_source_reference(baseline_source))
            or (comparison_source and not is_single_cell_source_reference(comparison_source))
            or (not baseline_source and baseline_value is not None)
            or (not comparison_source and comparison_value is not None)
        ):
            # Legacy/stored-result traces can contain a descriptive source and
            # an already aggregated value.  Keep those as explicit inputs until
            # a source-cell trace is available; new adapter traces use the
            # MODEL_SOURCE branch above.
            hardcode_class = "REQUEST_INPUT"
        role = str((baseline or comparison).get("role") or "component")
        key = f"{key_prefix}:{index}:{role}"
        registry.add(
            key,
            category=category,
            item=item,
            basis=basis,
            period=period,
            unit=(unit_by_role or {}).get(role, ""),
            baseline_source=baseline_source,
            baseline_value=baseline_value,
            comparison_source=comparison_source,
            comparison_value=comparison_value,
            notes=notes,
            hardcode_class=hardcode_class,
        )
        keys[identity] = key
        paired.append((identity, baseline, comparison))
    return keys, paired


def collect_reporting_sources(
    result: Mapping[str, Any],
    sales_rows: Iterable[Any],
    baseline_fx: float | None,
    comparison_fx: float | None,
    *,
    baseline_workbook: Any = None,
    comparison_workbook: Any = None,
) -> tuple[SourceRegistry, dict[str, Any]]:
    registry = SourceRegistry()
    context: dict[str, Any] = {}
    period_label = str((result.get("period") or {}).get("label") or "선택기간")
    registry.add(
        "policy:reconciliation_tolerance",
        category="정책",
        item="Reconciliation 허용오차",
        unit="원",
        baseline_source="config/analysis_v1.json",
        baseline_value=1.0,
        comparison_source="",
        comparison_value=None,
        notes="Backend absolute tolerance; 보고 시트에서는 천원으로 환산",
        hardcode_class="POLICY_INPUT",
    )
    registry.add(
        "policy:freight_meters_per_pcs",
        category="정책",
        item="FS 운반비 환산 길이",
        unit="m/PCS",
        baseline_source="Analysis v1.1 freight policy",
        baseline_value=45.0,
        comparison_source="",
        comparison_value=None,
        notes="45m = 1 equivalent shipment PCS",
        hardcode_class="POLICY_INPUT",
    )

    direct = ((result.get("residual_analysis") or {}).get("direct_op_bridge") or {})
    direct_refs = direct.get("source_references") or {}
    pnl_rows = _records(result.get("pnl") or [])
    context["pnl_rows"] = pnl_rows
    for item in pnl_rows:
        code = str(item.get("code") or "")
        reference = direct_refs.get(code)
        baseline_cells = _source_cells(reference)
        comparison_cells = _source_cells(reference)
        baseline_values = _read_source_values(baseline_cells, baseline_workbook)
        comparison_values = _read_source_values(comparison_cells, comparison_workbook)
        if len(baseline_cells) == len(comparison_cells) == 1:
            # The normalized single-period result already contains the raw
            # value read from that cell, so a source workbook is optional.
            baseline_values = [_number(item.get("baseline")) or 0.0]
            comparison_values = [_number(item.get("comparison")) or 0.0]
        aggregate_is_reconciled = bool(baseline_cells) and bool(comparison_cells)
        if aggregate_is_reconciled:
            aggregate_is_reconciled = (
                baseline_values is not None
                and comparison_values is not None
                and math.isclose(
                    sum(baseline_values),
                    _number(item.get("baseline")) or 0.0,
                    rel_tol=1e-9,
                    abs_tol=1e-6,
                )
                and math.isclose(
                    sum(comparison_values),
                    _number(item.get("comparison")) or 0.0,
                    rel_tol=1e-9,
                    abs_tol=1e-6,
                )
            )
        if aggregate_is_reconciled:
            pnl_keys: list[str] = []
            for index, (cell, baseline_value, comparison_value) in enumerate(
                zip(baseline_cells, baseline_values or (), comparison_values or (), strict=True),
                1,
            ):
                key = f"pnl:{code}" if index == 1 else f"pnl:{code}:{index}"
                registry.add(
                    key,
                    category="손익",
                    item=str(item.get("item") or code),
                    basis=code,
                    period=period_label,
                    unit="원",
                    baseline_source=cell,
                    baseline_value=baseline_value,
                    comparison_source=comparison_cells[index - 1],
                    comparison_value=comparison_value,
                    notes="원본 모형 P&L line raw component",
                    group=f"pnl:{code}",
                )
                pnl_keys.append(key)
        else:
            # An export of a stored aggregate may not carry the source
            # workbooks.  Keep that limitation explicit as an input rather
            # than placing a comma/range expression in a MODEL_SOURCE row.
            registry.add(
                f"pnl:{code}",
                category="손익",
                item=str(item.get("item") or code),
                basis=code,
                period=period_label,
                unit="원",
                baseline_source=str(reference or "Stored P&L result"),
                baseline_value=item.get("baseline"),
                comparison_source=str(reference or "Stored P&L result"),
                comparison_value=item.get("comparison"),
                notes="구조화된 P&L 원천셀 trace가 없는 저장 결과",
                hardcode_class="REQUEST_INPUT",
            )

    sales = result.get("sales_analysis") or {}
    sales_trace = _records(sales.get("trace_rows") or [])
    if not sales_trace:
        legacy_groups = _records(result.get("sales_groups") or [])
        if not legacy_groups:
            raise ValueError(
                "Evidence sales source trace is unavailable; regenerate the analysis result"
            )
        for item in legacy_groups:
            product_group = str(item.get("product_group") or "")
            if _number(item.get("baseline_cogs")) is None or _number(item.get("comparison_cogs")) is None:
                raise ValueError(
                    f"Evidence legacy raw COGS is unavailable for {product_group or 'unknown product group'}"
                )
            sales_trace.append({
                "period": period_label,
                "pool": "LENGTH" if product_group == "FS" else "PCS",
                "unit": "m" if product_group == "FS" else "PCS",
                "product_group": product_group,
                "base_quantity": item.get("baseline_quantity"),
                "comparison_quantity": item.get("comparison_quantity"),
                "base_revenue": item.get("baseline_amount"),
                "comparison_revenue": item.get("comparison_amount"),
                "base_cogs": item.get("baseline_cogs"),
                "comparison_cogs": item.get("comparison_cogs"),
                "base_source_reference": "Stored Result source identity unavailable",
                "comparison_source_reference": "Stored Result source identity unavailable",
            })
    context["sales_trace"] = sales_trace
    periods = _selected_periods(result) or sorted({str(row.get("period") or period_label) for row in sales_trace})
    context["periods"] = periods
    for period in periods:
        registry.add(
            f"sales_fx:{period}",
            category="입력",
            item="매출환율",
            basis="FX",
            period=period,
            unit="KRW/USD",
            baseline_source=f"Analysis request input: baseline_sales_fx_monthly[{period}]",
            baseline_value=_monthly_fx(sales, "baseline", period, sales_trace) or baseline_fx,
            comparison_source=f"Analysis request input: comparison_sales_fx_monthly[{period}]",
            comparison_value=_monthly_fx(sales, "comparison", period, sales_trace) or comparison_fx,
            notes="실행 시 입력값",
            hardcode_class="REQUEST_INPUT",
            group="sales_fx",
        )

    sales_source_keys: dict[tuple[str, str], dict[str, str]] = {}
    for item in sales_trace:
        period = str(item.get("period") or period_label)
        product = str(item.get("product_group") or "")
        source_key = (period, product)
        keys: dict[str, str] = {}
        for field, label, unit, source_index in (
            ("quantity", "판매수량", str(item.get("unit") or ""), 0),
            ("revenue", "매출", "원", 1),
            ("cogs", "매출원가", "원", 2),
        ):
            key = f"sales:{period}:{product}:{field}"
            keys[field] = key
            registry.add(
                key,
                category="판매",
                item=f"{product} {label}",
                basis=product,
                period=period,
                unit=unit,
                baseline_source=_source_part(item.get("base_source_reference") or "Golden sales source", source_index),
                baseline_value=item.get(f"base_{field}"),
                comparison_source=_source_part(item.get("comparison_source_reference") or "Golden sales source", source_index),
                comparison_value=item.get(f"comparison_{field}"),
                notes=str(item.get("canonical_fields") or ""),
            )
        sales_source_keys[source_key] = keys
    context["sales_source_keys"] = sales_source_keys

    new_business = _records(sales.get("new_business_trace_rows") or [])
    if not new_business:
        new_business = _records(
            ((result.get("sales_cogs_scope_analysis") or {}).get("new_business_rows") or [])
        )
    context["new_business"] = new_business
    new_business_keys: dict[str, dict[str, str]] = {}
    for index, item in enumerate(new_business, 1):
        period = str(item.get("period") or period_label)
        keys: dict[str, str] = {}
        for field, label, source_index in (("revenue", "매출", 0), ("cogs", "매출원가", 1)):
            key = f"new_business:{period}:{index}:{field}"
            keys[field] = key
            registry.add(
                key,
                category="판매",
                item=f"신사업 {label}",
                basis="신사업",
                period=period,
                unit="원",
                baseline_source=_source_part(item.get("base_source_reference") or "Golden new-business source", source_index),
                baseline_value=item.get(f"base_{field}"),
                comparison_source=_source_part(item.get("comparison_source_reference") or "Golden new-business source", source_index),
                comparison_value=item.get(f"comparison_{field}"),
                notes="수량 denominator는 source 미제공 시 true blank",
            )
        new_business_keys[f"{period}:{index}"] = keys
    context["new_business_keys"] = new_business_keys

    freight_rows = _records(sales.get("freight_trace_rows") or [])
    context["freight_rows"] = freight_rows
    freight_keys: dict[str, dict[str, str]] = {}
    for index, item in enumerate(freight_rows, 1):
        period = str(item.get("period") or period_label)
        keys: dict[str, str] = {}
        for field, label in (
            ("freight_including_tariff", "고객배송 운반비(관세 포함)"),
            ("tariff", "관세"),
        ):
            key = f"freight:{period}:{index}:{field}"
            keys[field] = key
            if field == "freight_including_tariff":
                baseline_source = _source_part(
                    item.get("base_source_reference") or "Freight source/input", 0
                )
                comparison_source = _source_part(
                    item.get("comparison_source_reference") or "Freight source/input", 0
                )
                hardcode_class = "MODEL_SOURCE"
            else:
                baseline_source = str(
                    item.get("base_tariff_source_reference")
                    or item.get("base_tariff_calculation_source")
                    or "Scenario metadata tariff input"
                )
                comparison_source = str(
                    item.get("comparison_tariff_source_reference")
                    or item.get("comparison_tariff_calculation_source")
                    or "Scenario metadata tariff input"
                )
                hardcode_class = "REQUEST_INPUT"
            registry.add(
                key,
                category="판매",
                item=label,
                basis="운반비/관세",
                period=period,
                unit="원",
                baseline_source=baseline_source,
                baseline_value=item.get(f"base_{field}"),
                comparison_source=comparison_source,
                comparison_value=item.get(f"comparison_{field}"),
                notes=str(item.get("freight_denominator_policy") or ""),
                hardcode_class=hardcode_class,
            )
        freight_keys[f"{period}:{index}"] = keys
    context["freight_keys"] = freight_keys

    material = result.get("material_analysis") or {}
    material_rows = _records(material.get("trace_rows") or [])
    context["material_rows"] = material_rows
    material_keys: dict[tuple[str, str], dict[str, str]] = {}
    material_formula_groups: dict[tuple[str, str], dict[str, Any]] = {}
    for item in material_rows:
        period = str(item.get("period") or period_label)
        product = str(item.get("product_group") or "")
        group = {
            "front": {},
            "back": {},
            "direct": [],
            "output": [],
            "legacy_cost_key": "",
        }
        baseline_components = item.get("base_material_source_components") or []
        comparison_components = item.get("comparison_material_source_components") or []
        component_keys, paired_components = _register_raw_component_pairs(
            registry,
            key_prefix=f"material:{period}:{product}:source",
            category="원재료",
            item=f"{product} 원재료 원천셀",
            basis=product,
            period=period,
            unit_by_role={
                "front_amount": "원",
                "front_production_basis": str(item.get("unit") or ""),
                "production_quantity": str(item.get("unit") or ""),
                "input_length": "m",
                "adjustment": "배율",
                "back_total_component": "원",
                "direct": "원",
            },
            baseline_components=baseline_components,
            comparison_components=comparison_components,
            notes=str(item.get("canonical_fields") or ""),
        )
        for identity, baseline_component, comparison_component in paired_components:
            component = baseline_component or comparison_component
            role = str(component.get("role") or "component")
            term_id = str(component.get("term_id") or identity[0])
            key = component_keys[identity]
            if role == "back_total_component":
                group["back"].setdefault(term_id, []).append(key)
            elif role == "direct":
                group["direct"].append(key)
            elif role in {
                "front_amount", "front_production_basis", "production_quantity",
                "input_length", "adjustment",
            }:
                group["front"].setdefault(term_id, {})[role] = key

        output_keys, _ = _register_raw_component_pairs(
            registry,
            key_prefix=f"material:{period}:{product}:output",
            category="원재료",
            item=f"{product} 생산기준 원천셀",
            basis=product,
            period=period,
            unit_by_role={"production_basis": str(item.get("unit") or "")},
            baseline_components=item.get("base_output_source_components") or [],
            comparison_components=item.get("comparison_output_source_components") or [],
            notes="material_effects production_basis denominator",
        )
        group["output"] = list(output_keys.values())

        if not baseline_components and not comparison_components:
            # Stored/legacy results may predate structured adapter components.
            # Keep the value explicitly labelled as an input rather than
            # presenting an aggregated source expression as MODEL_SOURCE.
            legacy_key = f"material:{period}:{product}:legacy_cost"
            registry.add(
                legacy_key,
                category="원재료",
                item=f"{product} 원재료 금액(저장 결과)",
                basis=product,
                period=period,
                unit="원",
                baseline_source=str(item.get("base_source_reference") or "Stored material result"),
                baseline_value=item.get("base_cost"),
                comparison_source=str(item.get("comparison_source_reference") or "Stored material result"),
                comparison_value=item.get("comparison_cost"),
                notes="구조화된 원천셀 trace가 없는 저장 결과",
                hardcode_class="REQUEST_INPUT",
            )
            group["legacy_cost_key"] = legacy_key
        material_formula_groups[(period, product)] = group
        material_keys[(period, product)] = {
            "output": group["output"][0] if len(group["output"]) == 1 else "",
        }
    context["material_keys"] = material_keys
    context["material_formula_groups"] = material_formula_groups

    nonwoven_rows = _records(material.get("nonwoven_trace_rows") or [])
    context["nonwoven_rows"] = nonwoven_rows
    nonwoven_keys: dict[str, dict[str, Any]] = {}
    for index, item in enumerate(nonwoven_rows, 1):
        period = str(item.get("period") or period_label)
        base_components = [
            component for component in item.get("base_source_components") or []
            if str(component.get("role") or "") in {"cost", "output", "jpy_fx"}
        ]
        comparison_components = [
            component for component in item.get("comparison_source_components") or []
            if str(component.get("role") or "") in {"cost", "output", "jpy_fx"}
        ]
        component_keys, paired_components = _register_raw_component_pairs(
            registry,
            key_prefix=f"nonwoven:{period}:{index}:source",
            category="원재료",
            item="부직포 원천셀",
            basis="FS",
            period=period,
            unit_by_role={"cost": "원", "output": "m", "jpy_fx": "KRW/JPY"},
            baseline_components=base_components,
            comparison_components=comparison_components,
            notes=str(item.get("canonical_fields") or ""),
        )
        keys: dict[str, Any] = {"cost": [], "output": [], "jpy_fx": []}
        for identity, baseline_component, comparison_component in paired_components:
            component = baseline_component or comparison_component
            role = str(component.get("role") or "")
            if role in keys:
                keys[role].append(component_keys[identity])

        input_components = [
            component for component in item.get("comparison_source_components") or []
            if str(component.get("role") or "") in {"sales_quantity", "input_length"}
        ]
        input_keys, input_pairs = _register_raw_component_pairs(
            registry,
            key_prefix=f"nonwoven:{period}:{index}:input",
            category="원재료",
            item="부직포 비교 적용길이 원천셀",
            basis="FS",
            period=period,
            unit_by_role={"sales_quantity": "PCS", "input_length": "m"},
            baseline_components=[],
            comparison_components=input_components,
            notes="비교 판매수량 × 제품별 입력길이",
        )
        input_terms: dict[str, dict[str, str]] = {}
        for identity, _, comparison_component in input_pairs:
            component = comparison_component
            term_id = str(component.get("term_id") or identity[0])
            role = str(component.get("role") or "")
            if role not in {"sales_quantity", "input_length"}:
                continue
            input_terms.setdefault(term_id, {})[role] = input_keys[identity]
        keys["comparison_input_terms"] = input_terms
        if not input_terms:
            input_key = f"nonwoven:{period}:{index}:comparison_input_length"
            registry.add(
                input_key,
                category="원재료",
                item="부직포 비교 적용길이(저장 결과)",
                basis="FS",
                period=period,
                unit="m",
                comparison_source=str(item.get("comparison_source_reference") or "Stored nonwoven result"),
                comparison_value=item.get("comparison_input_length"),
                notes="구조화된 입력길이 원천셀 trace가 없는 저장 결과",
                hardcode_class="REQUEST_INPUT",
            )
            keys["comparison_input_legacy_key"] = input_key
        nonwoven_keys[f"{period}:{index}"] = keys
    context["nonwoven_keys"] = nonwoven_keys

    _collect_cost_sources(
        result,
        registry,
        context,
        period_label,
        baseline_workbook=baseline_workbook,
        comparison_workbook=comparison_workbook,
    )
    return registry, context


def _collect_cost_sources(
    result: Mapping[str, Any],
    registry: SourceRegistry,
    context: dict[str, Any],
    period_label: str,
    *,
    baseline_workbook: Any = None,
    comparison_workbook: Any = None,
) -> None:
    manufacturing = result.get("manufacturing_analysis") or {}
    manufacturing_rows = _records(manufacturing.get("trace_rows") or [])
    reconciliation = _records(manufacturing.get("production_reconciliation") or [])
    context["manufacturing_rows"] = manufacturing_rows
    activity_periods = {
        str(item.get("month") or period_label)
        for item in manufacturing_rows + reconciliation
    }
    paired_activity: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    for item in reconciliation:
        period = str(item.get("month") or period_label)
        side = str(item.get("scenario") or "").lower()
        product = str(item.get("product_group") or "").upper()
        if side not in {"base", "comparison"}:
            continue
        side_key = "baseline" if side == "base" else "comparison"
        process = "front" if product == "FS" else "back" if product in {"SW", "BW", "LC"} else ""
        if not process:
            continue
        for component in _records(item.get("activity_components") or []):
            operation = str(component.get("operation") or "ADD").upper()
            source_reference = str(component.get("source") or "")
            identity = (period, process, product, operation, source_reference)
            pair = paired_activity.setdefault(identity, {
                "period": period,
                "process": process,
                "product": product,
                "operation": operation,
                "baseline_source": "",
                "baseline_value": None,
                "comparison_source": "",
                "comparison_value": None,
            })
            pair[f"{side_key}_source"] = source_reference
            pair[f"{side_key}_value"] = component.get("value")
    activity_groups: dict[tuple[str, str], dict[str, list[str]]] = {}
    for index, pair in enumerate(paired_activity.values(), 1):
        key = (
            f"manufacturing_activity:{pair['period']}:{pair['process']}:"
            f"{pair['product']}:{pair['operation']}:{index}"
        )
        registry.add(
            key,
            category="제조",
            item=f"{pair['product']} SAP 생산입고",
            basis=pair["process"],
            period=pair["period"],
            unit="m" if pair["process"] == "front" else "PCS",
            baseline_source=pair["baseline_source"],
            baseline_value=pair["baseline_value"],
            comparison_source=pair["comparison_source"],
            comparison_value=pair["comparison_value"],
            notes=f"{pair['operation']} raw activity component",
        )
        if pair["process"] == "front":
            activity_groups.setdefault(
                (pair["period"], "front"), {"ADD": [], "SUBTRACT": []}
            ).setdefault(pair["operation"], []).append(key)
        else:
            if pair["operation"] == "ADD":
                activity_groups.setdefault(
                    (pair["period"], "back"), {"ADD": [], "SUBTRACT": []}
                )["ADD"].append(key)
            activity_groups.setdefault(
                (pair["period"], "outsourcing_back"), {"ADD": [], "SUBTRACT": []}
            ).setdefault(pair["operation"], []).append(key)
    context["manufacturing_activity_groups"] = activity_groups
    context["manufacturing_activity_periods"] = sorted(activity_periods)

    manufacturing_keys: dict[str, dict[str, str]] = {}
    for index, item in enumerate(manufacturing_rows, 1):
        period = str(item.get("month") or period_label)
        account = str(item.get("account") or f"계정 {index}")
        amount_key = f"manufacturing:{period}:{index}:amount"
        ratio_source = str(item.get("front_ratio_source") or "Allocation ratio source")
        ratio_key = f"manufacturing:{period}:ratio:{ratio_source}"
        registry.add(
            amount_key,
            category="제조",
            item=account,
            basis=str(item.get("classification") or ""),
            period=period,
            unit="원",
            baseline_source=str(item.get("base_amount_source") or "Manufacturing source"),
            baseline_value=item.get("baseline_amount"),
            comparison_source=str(item.get("comparison_amount_source") or "Manufacturing source"),
            comparison_value=item.get("comparison_amount"),
            notes=str(item.get("current_cost_component") or ""),
        )
        registry.add(
            ratio_key,
            category="제조",
            item=f"{account} 전공정 배부율",
            basis=str(item.get("classification") or ""),
            period=period,
            unit="%",
            baseline_source=ratio_source,
            baseline_value=item.get("front_ratio_base"),
            comparison_source="V1 policy: baseline ratio reused for comparison",
            comparison_value=item.get("front_ratio_comparison", item.get("front_ratio_base")),
            notes="기준 source ratio를 양쪽에 동일 적용하는 기존 정책 유지",
            hardcode_class="POLICY_INPUT",
        )
        manufacturing_keys[f"{period}:{index}"] = {"amount": amount_key, "ratio": ratio_key}
    context["manufacturing_keys"] = manufacturing_keys

    inventory = result.get("inventory_analysis") or {}
    inventory_source_rows = _records(inventory.get("source_details") or [])
    context["inventory_source_rows"] = inventory_source_rows
    paired_inventory: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for item in inventory_source_rows:
        period = str(item.get("period") or period_label)
        canonical = str(item.get("canonical_field") or "source")
        business_source = str(item.get("business_source") or canonical)
        unit = str(item.get("unit") or "원")
        pair = paired_inventory.setdefault(
            (period, canonical, business_source, unit),
            {
                "period": period,
                "canonical": canonical,
                "business_source": business_source,
                "unit": unit,
                "baseline_source": "",
                "baseline_value": None,
                "comparison_source": "",
                "comparison_value": None,
            },
        )
        side = str(item.get("side") or "").upper()
        if side == "BASE":
            pair["baseline_source"] = str(item.get("source_reference") or "")
            pair["baseline_value"] = item.get("value")
        elif side == "COMPARISON":
            pair["comparison_source"] = str(item.get("source_reference") or "")
            pair["comparison_value"] = item.get("value")
    for index, pair in enumerate(paired_inventory.values(), 1):
        canonical = pair["canonical"]
        key = f"inventory:{pair['period']}:{canonical}:{index}"
        kwargs = dict(
            category="재고/원가",
            item=pair["business_source"],
            basis=canonical,
            period=pair["period"],
            unit=pair["unit"],
            baseline_source=pair["baseline_source"],
            baseline_value=pair["baseline_value"],
            comparison_source=pair["comparison_source"],
            comparison_value=pair["comparison_value"],
            notes="원본 모형 canonical current-cost/COGS line",
        )
        registry.add(key, group=f"inventory:{canonical}:base", **kwargs)
        registry.add(key, group=f"inventory:{canonical}:comparison", **kwargs)

    core_rows = _records(inventory.get("core_overlap_details") or [])
    context["core_rows"] = core_rows
    core_keys: dict[tuple[str, str], str] = {}
    for item in core_rows:
        if not item.get("selected"):
            continue
        period = str(item.get("period") or period_label)
        product = str(item.get("product_group") or "")
        key = f"core:{period}:{product}:cogs"
        baseline_source = str(
            item.get("base_core_cogs_source_reference") or "Core COGS source"
        )
        hardcode_class = (
            "MODEL_SOURCE"
            if is_single_cell_source_reference(baseline_source)
            else "REQUEST_INPUT"
        )
        registry.add(
            key,
            category="재고/원가",
            item=f"{product} Core 제조원가",
            basis=product,
            period=period,
            unit="원",
            baseline_source=baseline_source,
            baseline_value=item.get("base_core_manufactured_cogs"),
            comparison_source=(
                str(item.get("comparison_core_cogs_source_reference") or "")
                if item.get("comparison_core_manufactured_cogs") is not None
                else ""
            ),
            comparison_value=None,
            notes="Quantity/Mix overlap 계산의 기준 COGS source",
            hardcode_class=hardcode_class,
        )
        core_keys[(period, product)] = key
    context["core_keys"] = core_keys

    production_rows = [
        item for item in _records(result.get("production_evidence") or [])
        if str(item.get("production_basis") or "") in {"FS", "SW", "BW", "LC"}
    ]
    context["production_rows"] = production_rows
    production_groups: dict[tuple[str, str], dict[str, list[str]]] = {}
    for item in production_rows:
        product = str(item.get("production_basis") or "")
        components = _records(item.get("source_components") or [])
        if not components:
            components = [{
                "period": ", ".join(str(value) for value in item.get("selected_period") or []) or period_label,
                "kind": "quantity",
                "component": 1,
                "baseline_source": " + ".join(str(value) for value in item.get("baseline_quantity_sources") or []),
                "baseline_value": item.get("baseline_quantity"),
                "comparison_source": " + ".join(str(value) for value in item.get("comparison_quantity_sources") or []),
                "comparison_value": item.get("comparison_quantity"),
            }, {
                "period": ", ".join(str(value) for value in item.get("selected_period") or []) or period_label,
                "kind": "amount",
                "component": 1,
                "baseline_source": " + ".join(str(value) for value in item.get("baseline_amount_sources") or []),
                "baseline_value": item.get("baseline_amount"),
                "comparison_source": " + ".join(str(value) for value in item.get("comparison_amount_sources") or []),
                "comparison_value": item.get("comparison_amount"),
            }]
        key_group = production_groups.setdefault((product, str(item.get("unit") or "")), {"quantity": [], "amount": []})
        for component in components:
            kind = str(component.get("kind") or "")
            if kind not in {"quantity", "amount"}:
                continue
            key = f"production:{product}:{component.get('period')}:{kind}:{component.get('component')}"
            registry.add(
                key,
                category="생산",
                item=f"{product} 생산{'수량' if kind == 'quantity' else '금액'}",
                basis=product,
                period=str(component.get("period") or period_label),
                unit=str(item.get("unit") or "") if kind == "quantity" else "원",
                baseline_source=str(component.get("baseline_source") or "Production source"),
                baseline_value=component.get("baseline_value"),
                comparison_source=str(component.get("comparison_source") or "Production source"),
                comparison_value=component.get("comparison_value"),
                notes=str(item.get("aggregation_basis") or ""),
            )
            key_group[kind].append(key)
    context["production_groups"] = production_groups

    sga_rows = _records(result.get("sga_monthly_trace") or [])
    if not sga_rows:
        sga_rows = _records(result.get("sga_accounts") or [])
    context["sga_rows"] = sga_rows
    sga_account_rows = _records(result.get("sga_accounts") or [])
    sga_keys: dict[str, str] = {}
    sga_key_list: list[str] = []
    sga_formula_groups: dict[str, dict[str, Any]] = {}
    for index, item in enumerate(sga_rows, 1):
        period = str(item.get("period") or period_label)
        account = str(item.get("display_account") or item.get("account") or f"계정 {index}")
        baseline_components = _sga_source_components(
            item,
            side="baseline",
            sga_account_rows=sga_account_rows,
            source_workbook=baseline_workbook,
        )
        comparison_components = _sga_source_components(
            item,
            side="comparison",
            sga_account_rows=sga_account_rows,
            source_workbook=comparison_workbook,
        )
        baseline_by_source = {
            component["source"]: component for component in baseline_components
        }
        comparison_by_source = {
            component["source"]: component for component in comparison_components
        }
        source_order = list(baseline_by_source)
        source_order.extend(
            source for source in comparison_by_source if source not in source_order
        )
        component_keys: dict[str, str] = {}
        for component_index, source in enumerate(source_order, 1):
            baseline_component = baseline_by_source.get(source, {})
            comparison_component = comparison_by_source.get(source, {})
            key = f"sga:{period}:{index}:amount:{component_index}"
            registry.add(
                key,
                category="판관비",
                item=f"{account} 원천셀",
                basis=str(item.get("classification") or ""),
                period=period,
                unit="원",
                baseline_source=str(baseline_component.get("source") or ""),
                baseline_value=baseline_component.get("value"),
                comparison_source=str(comparison_component.get("source") or ""),
                comparison_value=comparison_component.get("value"),
                notes=str(item.get("bridge_position") or ""),
            )
            component_keys[source] = key
        if not source_order:
            # A direct-input/stored result row with no workbook source has no
            # MODEL_SOURCE identity to expose.  Retain the existing explicit
            # input representation for that non-workbook case only.
            key = f"sga:{period}:{index}:amount"
            registry.add(
                key,
                category="판관비",
                item=account,
                basis=str(item.get("classification") or ""),
                period=period,
                unit="원",
                baseline_source=str(item.get("base_source_reference") or item.get("source_reference") or "SG&A input"),
                baseline_value=item.get("base_amount", item.get("baseline_amount")),
                comparison_source=str(item.get("comparison_source_reference") or item.get("source_reference") or "SG&A input"),
                comparison_value=item.get("comparison_amount"),
                notes=str(item.get("bridge_position") or ""),
                hardcode_class="REQUEST_INPUT",
            )
            component_keys["__legacy__"] = key
        group_key = f"{period}:{index}"
        sga_formula_groups[group_key] = {
            "baseline": (
                [
                    component_keys[source]
                    for source in source_order
                    if source in baseline_by_source
                ]
                if source_order else [component_keys["__legacy__"]]
            ),
            "comparison": (
                [
                    component_keys[source]
                    for source in source_order
                    if source in comparison_by_source
                ]
                if source_order else [component_keys["__legacy__"]]
            ),
        }
        first_key = next(iter(component_keys.values()))
        sga_keys[group_key] = first_key
        sga_key_list.append(first_key)
    context["sga_keys"] = sga_keys
    context["sga_key_list"] = sga_key_list
    context["sga_formula_groups"] = sga_formula_groups


def _apply_workbook_font(ws: Any) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.font = Font(
                    name=FONT_NAME,
                    size=cell.font.sz or 10,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=cell.font.color,
                )


def _write_report_title(
    ws: Any,
    title: str,
    subtitle: str,
    *,
    last_column: int,
    unit_legend: str,
) -> None:
    split = max(5, last_column - 3)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY_DARK)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=split)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT_NAME, size=9.5, color=MUTED)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=2, start_column=split + 1, end_row=2, end_column=last_column)
    unit_cell = ws.cell(2, split + 1)
    unit_cell.value = unit_legend
    unit_cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color=NAVY)
    unit_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28


def _section(ws: Any, row: int, label: str, last_column: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    cell = ws.cell(row, 1, label)
    cell.font = Font(name=FONT_NAME, size=11.5, bold=True, color=NAVY_DARK)
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 23
    return row + 1


def _headers(ws: Any, row: int, headers: list[str]) -> None:
    for column, label in enumerate(headers, 1):
        cell = ws.cell(row, column, label)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=TEXT)
        cell.fill = PatternFill("solid", fgColor=LIGHT_NEUTRAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 22


def _style_data_rows(ws: Any, start: int, end: int, last_column: int) -> None:
    if end < start:
        return
    for row in range(start, end + 1):
        ws.row_dimensions[row].height = 19
        for column in range(1, last_column + 1):
            cell = ws.cell(row, column)
            cell.font = Font(name=FONT_NAME, size=9.5, color=TEXT)
            is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
            cell.alignment = Alignment(
                horizontal="right" if column > 3 and (is_formula or not isinstance(cell.value, str)) else "left",
                vertical="center",
                wrap_text=False,
            )
            cell.border = THIN_BORDER


def _style_formula_cells(ws: Any, coordinates: Iterable[str], number_format: str) -> None:
    for coordinate in coordinates:
        cell = ws[coordinate]
        cell.fill = PatternFill("solid", fgColor=WHITE)
        cell.font = Font(name=FONT_NAME, size=9.5, color=TEXT)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.number_format = number_format


def _style_total_row(ws: Any, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.fill = PatternFill("solid", fgColor=LIGHT_TOTAL)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY_DARK)
        cell.border = TOTAL_BORDER
    ws.row_dimensions[row].height = 21


def _set_widths(ws: Any, widths: Mapping[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _group_detail_rows(ws: Any, start: int, end: int) -> None:
    if end < start:
        return
    ws.row_dimensions.group(start, end, outline_level=1, hidden=True)
    ws.sheet_properties.outlinePr.summaryBelow = True


def _finish_sheet(
    ws: Any,
    *,
    last_row: int,
    last_column: int,
    freeze: str,
    print_header_rows: str,
    portrait: bool = False,
    fit_height: int | None = None,
    print_last_column: int | None = None,
) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait" if portrait else "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = (
        fit_height if fit_height is not None else (1 if portrait else 0)
    )
    ws.print_title_rows = print_header_rows
    ws.print_area = (
        f"A1:{get_column_letter(print_last_column or last_column)}{last_row}"
    )
    ws.sheet_view.zoomScale = 90
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    _apply_workbook_font(ws)


def write_source_sheet(ws: Any, registry: SourceRegistry) -> None:
    _write_report_title(
        ws,
        "90_원본값",
        "Baseline/Comparison Workbook 및 분석 실행 입력의 단일 Source of Truth",
        last_column=12,
        unit_legend="원본 금액: 원 / 환율·단가·수량: 고유 단위",
    )
    ws["A3"] = "보고 시트의 모든 원본값은 이 시트를 formula reference하며, derived numeric value는 이 시트에 저장하지 않습니다."
    ws["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color=MUTED)
    ws.merge_cells("A3:L3")
    _headers(ws, 5, [
        "Key", "구분", "항목", "Basis / 제품군", "기간", "단위",
        "기준 Source", "기준 원본값", "비교 Source", "비교 원본값", "비고", "Hard-code Class",
    ])
    for item in registry.rows:
        ws.append([
            item.key, item.category, item.item, item.basis, item.period, item.unit,
            item.baseline_source, item.baseline_value,
            item.comparison_source, item.comparison_value, item.notes, item.hardcode_class,
        ])
        for coordinate in (f"H{item.row}", f"J{item.row}"):
            cell = ws[coordinate]
            cell.fill = PatternFill("solid", fgColor=LIGHT_SOURCE)
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if item.unit == "%":
                cell.number_format = PERCENT_FORMAT
            elif "FX" in item.item or item.unit.startswith("KRW/"):
                cell.number_format = FX_FORMAT
            else:
                cell.number_format = COUNT_FORMAT
        for coordinate in (f"G{item.row}", f"I{item.row}", f"K{item.row}"):
            ws[coordinate].alignment = Alignment(wrap_text=True, vertical="center")
    _style_data_rows(ws, 6, ws.max_row, 12)
    for item in registry.rows:
        for coordinate in (f"G{item.row}", f"I{item.row}", f"K{item.row}"):
            ws[coordinate].alignment = Alignment(wrap_text=True, vertical="center")
        if max(len(item.baseline_source), len(item.comparison_source), len(item.notes)) > 45:
            ws.row_dimensions[item.row].height = 30
    ws.auto_filter.ref = f"A5:L{ws.max_row}"
    _set_widths(ws, {
        "A": 32, "B": 12, "C": 24, "D": 16, "E": 11, "F": 10,
        "G": 30, "H": 14, "I": 30, "J": 14, "K": 28, "L": 16,
    })
    _finish_sheet(
        ws, last_row=ws.max_row, last_column=12, freeze="A6", print_header_rows="1:5"
    )


def _formula_average(refs: list[str]) -> str:
    return _join_sum(refs).replace("=SUM(", "=AVERAGE(", 1) if len(refs) > 1 else _join_sum(refs)


def _activity_formula(
    registry: SourceRegistry,
    groups: Mapping[str, list[str]],
    side: str,
) -> str:
    add_refs = [registry.ref(key, side) for key in groups.get("ADD", ())]
    subtract_refs = [registry.ref(key, side) for key in groups.get("SUBTRACT", ())]
    if not add_refs:
        return '=""'
    add_expression = _join_sum(add_refs)[1:]
    if not subtract_refs:
        return f"={add_expression}"
    subtract_expression = _join_sum(subtract_refs)[1:]
    return f"=MAX({add_expression}-{subtract_expression},0)"


def write_sales_sheet(
    ws: Any,
    registry: SourceRegistry,
    context: Mapping[str, Any],
) -> dict[str, Any]:
    _write_report_title(
        ws,
        "03_판매근거",
        "판매수량·Mix·가격·환율·운반비·신사업의 canonical calculation chain",
        last_column=16,
        unit_legend="단위: 금액 천원 / 단가 원/단위 / 환율 KRW/USD / 수량 PCS·m",
    )
    fx_refs_base = registry.refs("sales_fx", "baseline")
    fx_refs_comparison = registry.refs("sales_fx", "comparison")
    ws["A4"] = "기준 평균 FX"
    ws["B4"] = _formula_average(fx_refs_base)
    ws["D4"] = "비교 평균 FX"
    ws["E4"] = _formula_average(fx_refs_comparison)
    ws["B4"].number_format = FX_FORMAT
    ws["E4"].number_format = FX_FORMAT

    row = _section(ws, 7, "A. 판매수량 / 매출", 12)
    sales_header = row
    _headers(ws, row, [
        "제품군", "단위", "Pool", "기간", "기준 수량", "비교 수량", "수량 차이",
        "기준 매출", "비교 매출", "매출 차이", "기준 COGS", "비교 COGS",
    ])
    row += 1
    sales_rows = list(context.get("sales_trace") or [])
    sales_source_keys = dict(context.get("sales_source_keys") or {})
    sales_table_rows: dict[tuple[str, str], int] = {}
    for item in sales_rows:
        period = str(item.get("period") or "선택기간")
        product = str(item.get("product_group") or "")
        keys = sales_source_keys[(period, product)]
        current = row
        sales_table_rows[(period, product)] = current
        ws.cell(current, 1, product)
        ws.cell(current, 2, "m" if str(item.get("unit") or "").lower() == "m" else item.get("unit"))
        ws.cell(current, 3, item.get("pool"))
        ws.cell(current, 4, period)
        ws.cell(current, 5, f"={registry.ref(keys['quantity'], 'baseline')}")
        ws.cell(current, 6, f"={registry.ref(keys['quantity'], 'comparison')}")
        ws.cell(current, 7, f"=F{current}-E{current}")
        ws.cell(current, 8, f"={registry.ref(keys['revenue'], 'baseline')}/1000")
        ws.cell(current, 9, f"={registry.ref(keys['revenue'], 'comparison')}/1000")
        ws.cell(current, 10, f"=I{current}-H{current}")
        ws.cell(current, 11, f"={registry.ref(keys['cogs'], 'baseline')}/1000")
        ws.cell(current, 12, f"={registry.ref(keys['cogs'], 'comparison')}/1000")
        for column in range(5, 13):
            ws.cell(current, column).number_format = MONEY_FORMAT if column >= 8 else COUNT_FORMAT
        row += 1
    new_business = list(context.get("new_business") or [])
    new_business_keys = dict(context.get("new_business_keys") or {})
    new_business_sales_rows: list[int] = []
    for index, item in enumerate(new_business, 1):
        period = str(item.get("period") or "선택기간")
        keys = new_business_keys[f"{period}:{index}"]
        current = row
        new_business_sales_rows.append(current)
        ws.cell(current, 1, "신사업")
        ws.cell(current, 2, "")
        ws.cell(current, 3, "NEW_BUSINESS")
        ws.cell(current, 4, period)
        # E:G stay physically empty by policy.
        ws.cell(current, 8, f"={registry.ref(keys['revenue'], 'baseline')}/1000")
        ws.cell(current, 9, f"={registry.ref(keys['revenue'], 'comparison')}/1000")
        ws.cell(current, 10, f"=I{current}-H{current}")
        ws.cell(current, 11, f"={registry.ref(keys['cogs'], 'baseline')}/1000")
        ws.cell(current, 12, f"={registry.ref(keys['cogs'], 'comparison')}/1000")
        for column in range(8, 13):
            ws.cell(current, column).number_format = MONEY_FORMAT
        row += 1
    sales_start = sales_header + 1
    sales_end = row - 1
    _style_data_rows(ws, sales_start, sales_end, 12)
    if sales_end >= sales_start:
        ws.auto_filter.ref = f"A{sales_header}:L{sales_end}"

    normal_rows = [
        item for item in sales_rows
        if str(item.get("product_group") or "") != "신사업"
    ]
    pool_items = sorted({
        (str(item.get("period") or "선택기간"), str(item.get("pool") or ""), str(item.get("unit") or ""))
        for item in normal_rows
    })
    driver_header = sales_end + 3
    row = _section(ws, driver_header, "B. 가격 / FX 및 제품군 Mix", 16)
    _headers(ws, row, [
        "기간", "Pool", "제품군", "기준 GP/단위", "기준 Mix", "비교 Mix",
        "가중 GP/단위", "Mix 기여", "기준 판매단가", "비교 판매단가", "단가 차이",
        "기준 외화단가", "비교 외화단가", "Price Effect", "Sales FX", "단위",
    ])
    driver_start = row + 1
    pool_header = driver_start + len(normal_rows) + 2
    pool_start = pool_header + 2
    pool_rows = {item[:2]: pool_start + index for index, item in enumerate(pool_items)}
    driver_rows: list[int] = []
    for item in normal_rows:
        period = str(item.get("period") or "선택기간")
        product = str(item.get("product_group") or "")
        pool = str(item.get("pool") or "")
        source_row = sales_table_rows[(period, product)]
        pool_row = pool_rows[(period, pool)]
        fx_key = f"sales_fx:{period}"
        current = row + 1
        driver_rows.append(current)
        ws.cell(current, 1, period)
        ws.cell(current, 2, pool)
        ws.cell(current, 3, product)
        ws.cell(current, 4, f"=IFERROR((H{source_row}-K{source_row})*1000/E{source_row},0)")
        ws.cell(current, 5, f"=IFERROR(E{source_row}/D{pool_row},0)")
        ws.cell(current, 6, f"=IFERROR(F{source_row}/E{pool_row},0)")
        ws.cell(current, 7, f"=E{current}*D{current}")
        ws.cell(current, 8, f"=(F{current}-E{current})*D{current}")
        ws.cell(current, 9, f"=IFERROR(H{source_row}*1000/E{source_row},0)")
        ws.cell(current, 10, f"=IFERROR(I{source_row}*1000/F{source_row},0)")
        ws.cell(current, 11, f"=J{current}-I{current}")
        ws.cell(current, 12, f"=IFERROR(I{current}/{registry.ref(fx_key, 'baseline')},0)")
        ws.cell(current, 13, f"=IFERROR(J{current}/{registry.ref(fx_key, 'comparison')},0)")
        ws.cell(current, 14, f"=F{source_row}*(M{current}-L{current})*({registry.ref(fx_key, 'baseline')}+{registry.ref(fx_key, 'comparison')})/2/1000")
        ws.cell(current, 15, f"=F{source_row}*({registry.ref(fx_key, 'comparison')}-{registry.ref(fx_key, 'baseline')})*(L{current}+M{current})/2/1000")
        ws.cell(current, 16, "원/PCS" if pool == "PCS" else "원/m")
        for column in range(4, 16):
            ws.cell(current, column).number_format = PERCENT_FORMAT if column in {5, 6} else UNIT_COST_FORMAT
        row += 1
    driver_end = row
    _style_data_rows(ws, driver_start, driver_end, 16)

    pool_section = _section(ws, pool_header, "B-1. Pool별 Quantity / Mix canonical calculation", 9)
    _headers(ws, pool_section, [
        "기간", "Pool", "단위", "기준 총수량", "비교 총수량", "기준 가중 GP/단위",
        "Quantity Effect", "Mix 기여/단위", "Mix Effect",
    ])
    for index, (period, pool, unit) in enumerate(pool_items):
        current = pool_start + index
        ws.cell(current, 1, period)
        ws.cell(current, 2, pool)
        ws.cell(current, 3, unit)
        ws.cell(current, 4, f'=SUMIFS($E${sales_start}:$E${sales_end},$D${sales_start}:$D${sales_end},A{current},$C${sales_start}:$C${sales_end},B{current})')
        ws.cell(current, 5, f'=SUMIFS($F${sales_start}:$F${sales_end},$D${sales_start}:$D${sales_end},A{current},$C${sales_start}:$C${sales_end},B{current})')
        ws.cell(current, 6, f'=SUMIFS($G${driver_start}:$G${driver_end},$A${driver_start}:$A${driver_end},A{current},$B${driver_start}:$B${driver_end},B{current})')
        ws.cell(current, 7, f"=(E{current}-D{current})*F{current}/1000")
        ws.cell(current, 8, f'=SUMIFS($H${driver_start}:$H${driver_end},$A${driver_start}:$A${driver_end},A{current},$B${driver_start}:$B${driver_end},B{current})')
        ws.cell(current, 9, f"=E{current}*H{current}/1000")
        for column in range(4, 10):
            ws.cell(current, column).number_format = MONEY_FORMAT if column in {7, 9} else UNIT_COST_FORMAT
    pool_end = pool_start + len(pool_items) - 1
    _style_data_rows(ws, pool_start, pool_end, 9)

    nb_header = max(pool_end, pool_section) + 3
    row = _section(ws, nb_header, "D. 신사업 Revenue / GP Rate", 16)
    _headers(ws, row, [
        "기간", "제품군", "기준 수량", "비교 수량", "수량 차이",
        "기준 매출", "비교 매출", "매출 차이", "기준 COGS", "비교 COGS",
        "기준 GP", "비교 GP", "기준 GP Rate", "비교 GP Rate",
        "Revenue Effect", "GP Rate Effect",
    ])
    nb_start = row + 1
    nb_effect_rows: list[int] = []
    for index, source_row in enumerate(new_business_sales_rows, 1):
        current = row + index
        nb_effect_rows.append(current)
        ws.cell(current, 1, ws.cell(source_row, 4).value)
        ws.cell(current, 2, "신사업")
        # C:E stay physically empty.
        ws.cell(current, 6, f"=H{source_row}")
        ws.cell(current, 7, f"=I{source_row}")
        ws.cell(current, 8, f"=G{current}-F{current}")
        ws.cell(current, 9, f"=K{source_row}")
        ws.cell(current, 10, f"=L{source_row}")
        ws.cell(current, 11, f"=F{current}-I{current}")
        ws.cell(current, 12, f"=G{current}-J{current}")
        ws.cell(current, 13, f"=IFERROR(K{current}/F{current},0)")
        ws.cell(current, 14, f"=IFERROR(L{current}/G{current},0)")
        ws.cell(current, 15, f"=(G{current}-F{current})*M{current}")
        ws.cell(current, 16, f"=G{current}*(N{current}-M{current})")
        for column in range(6, 17):
            ws.cell(current, column).number_format = PERCENT_FORMAT if column in {13, 14} else MONEY_FORMAT
    nb_end = nb_start + len(nb_effect_rows) - 1
    _style_data_rows(ws, nb_start, nb_end, 16)

    freight_header = max(nb_end, row) + 3
    row = _section(ws, freight_header, "C. 운반비 / 관세", 13)
    _headers(ws, row, [
        "기간", "기준 운반비", "비교 운반비", "기준 관세", "비교 관세",
        "기준 운반비(관세 제외)", "비교 운반비(관세 제외)",
        "기준 환산수량", "비교 환산수량", "기준 운반비 단가", "비교 운반비 단가",
        "Freight Effect", "Tariff Effect",
    ])
    freight_start = row + 1
    freight_keys = dict(context.get("freight_keys") or {})
    freight_rows = list(context.get("freight_rows") or [])
    freight_formula_rows: list[int] = []
    for index, item in enumerate(freight_rows, 1):
        period = str(item.get("period") or "선택기간")
        keys = freight_keys[f"{period}:{index}"]
        current = row + index
        freight_formula_rows.append(current)
        ws.cell(current, 1, period)
        ws.cell(current, 2, f"={registry.ref(keys['freight_including_tariff'], 'baseline')}/1000")
        ws.cell(current, 3, f"={registry.ref(keys['freight_including_tariff'], 'comparison')}/1000")
        ws.cell(current, 4, f"={registry.ref(keys['tariff'], 'baseline')}/1000")
        ws.cell(current, 5, f"={registry.ref(keys['tariff'], 'comparison')}/1000")
        ws.cell(current, 6, f"=B{current}-D{current}")
        ws.cell(current, 7, f"=C{current}-E{current}")
        ws.cell(current, 8, f'=SUMIFS($E${sales_start}:$E${sales_end},$D${sales_start}:$D${sales_end},A{current},$C${sales_start}:$C${sales_end},"PCS")+SUMIFS($E${sales_start}:$E${sales_end},$D${sales_start}:$D${sales_end},A{current},$C${sales_start}:$C${sales_end},"LENGTH")/{registry.ref("policy:freight_meters_per_pcs", "baseline")}')
        ws.cell(current, 9, f'=SUMIFS($F${sales_start}:$F${sales_end},$D${sales_start}:$D${sales_end},A{current},$C${sales_start}:$C${sales_end},"PCS")+SUMIFS($F${sales_start}:$F${sales_end},$D${sales_start}:$D${sales_end},A{current},$C${sales_start}:$C${sales_end},"LENGTH")/{registry.ref("policy:freight_meters_per_pcs", "baseline")}')
        ws.cell(current, 10, f"=IFERROR(F{current}*1000/H{current},0)")
        ws.cell(current, 11, f"=IFERROR(G{current}*1000/I{current},0)")
        ws.cell(current, 12, f"=(J{current}-K{current})*I{current}/1000")
        ws.cell(current, 13, f"=D{current}-E{current}")
        for column in range(2, 14):
            ws.cell(current, column).number_format = UNIT_COST_FORMAT if column in {8, 9, 10, 11} else MONEY_FORMAT
    freight_end = freight_start + len(freight_formula_rows) - 1
    _style_data_rows(ws, freight_start, freight_end, 13)

    summary_header = max(freight_end, row) + 3
    row = _section(ws, summary_header, "판매 Effect canonical cells", 5)
    _headers(ws, row, ["Effect Code", "Effect", "공식 수식", "단위", "정책"])
    summary_start = row + 1
    pool_quantity_range = f"G{pool_start}:G{pool_end}" if pool_end >= pool_start else "G1:G1"
    pool_mix_range = f"I{pool_start}:I{pool_end}" if pool_end >= pool_start else "I1:I1"
    driver_price_range = f"N{driver_start}:N{driver_end}" if driver_end >= driver_start else "N1:N1"
    driver_fx_range = f"O{driver_start}:O{driver_end}" if driver_end >= driver_start else "O1:O1"
    nb_revenue_range = f"O{nb_start}:O{nb_end}" if nb_end >= nb_start else "O1:O1"
    nb_gp_range = f"P{nb_start}:P{nb_end}" if nb_end >= nb_start else "P1:P1"
    freight_effect_range = f"L{freight_start}:L{freight_end}" if freight_end >= freight_start else "L1:L1"
    tariff_range = f"M{freight_start}:M{freight_end}" if freight_end >= freight_start else "M1:M1"
    summary_specs = [
        ("sales_quantity", "판매수량 효과", f"=SUM({pool_quantity_range})+SUM({nb_revenue_range})", "신사업 Revenue Effect 포함"),
        ("sales_mix", "제품 Mix 효과", f"=SUM({pool_mix_range})", "Pool별 canonical Mix"),
        ("displayed_sales_price", "표시 판매가격 효과", f"=SUM({driver_price_range})+SUM({nb_gp_range})", "운반비 제외 price"),
        ("freight_adjustment", "고객배송 운반비 효과", f"=SUM({freight_effect_range})", "sales_price 내부 1회"),
        ("sales_price", "판매가격 효과", f"=C{summary_start + 2}+C{summary_start + 3}", "Displayed Price + Freight"),
        ("sales_fx", "매출환율 효과", f"=SUM({driver_fx_range})", "symmetric FX split"),
        ("tariff", "관세 효과", f"=SUM({tariff_range})", "운반비와 분리"),
    ]
    cells: dict[str, Any] = {"baseline_average_fx": "B4", "comparison_average_fx": "E4"}
    for index, (code, label, formula, policy) in enumerate(summary_specs):
        current = summary_start + index
        ws.cell(current, 1, code)
        ws.cell(current, 2, label)
        ws.cell(current, 3, formula)
        ws.cell(current, 4, "천원")
        ws.cell(current, 5, policy)
        ws.cell(current, 3).number_format = MONEY_FORMAT
        cells[code] = f"C{current}"
    summary_end = summary_start + len(summary_specs) - 1
    _style_data_rows(ws, summary_start, summary_end, 5)
    for current in range(summary_start, summary_end + 1):
        _style_total_row(ws, current, 1, 5)
    _group_detail_rows(ws, summary_header, summary_end)

    _set_widths(ws, {
        "A": 13, "B": 10, "C": 13, "D": 12, "E": 12, "F": 12,
        "G": 12, "H": 13, "I": 13, "J": 13, "K": 13, "L": 13,
        "M": 13, "N": 13, "O": 13, "P": 13,
    })
    _finish_sheet(
        ws, last_row=ws.max_row, last_column=16, freeze=f"A{sales_start}",
        print_header_rows="1:2", fit_height=1,
    )
    cells["_sales_table_rows"] = sales_table_rows
    cells["_pool_rows"] = pool_rows
    return cells


def _material_cost_formula(
    registry: SourceRegistry,
    group: Mapping[str, Any],
    side: str,
) -> str:
    parts: list[str] = []
    for term in group.get("front", {}).values():
        refs = {
            role: registry.ref(key, side)
            for role, key in term.items()
            if key
        }
        required = (
            "front_amount", "front_production_basis", "production_quantity",
            "input_length", "adjustment",
        )
        if all(role in refs for role in required):
            parts.append(
                "IFERROR(("
                f"{refs['front_amount']}/{refs['front_production_basis']}"
                f")*{refs['input_length']}*{refs['adjustment']}*"
                f"{refs['production_quantity']},0)"
            )
    for back_sources in group.get("back", {}).values():
        refs = [registry.ref(key, side) for key in back_sources if key]
        if refs:
            parts.append(_join_sum(refs)[1:])
    parts.extend(
        registry.ref(key, side)
        for key in group.get("direct", ())
        if key
    )
    if not parts:
        legacy_key = str(group.get("legacy_cost_key") or "")
        return f"={registry.ref(legacy_key, side)}/1000" if legacy_key else "=0"
    return f"=SUM({','.join(parts)})/1000"


def _material_output_formula(
    registry: SourceRegistry,
    group: Mapping[str, Any],
    side: str,
) -> str:
    refs = [registry.ref(key, side) for key in group.get("output", ()) if key]
    return _join_sum(refs) if refs else "=0"


def _nonwoven_component_formula(
    registry: SourceRegistry,
    keys: Mapping[str, Any],
    side: str,
    role: str,
) -> str:
    refs = [registry.ref(key, side) for key in keys.get(role, ()) if key]
    return _join_sum(refs) if refs else "=0"


def _nonwoven_input_formula(
    registry: SourceRegistry,
    keys: Mapping[str, Any],
) -> str:
    parts: list[str] = []
    for term in keys.get("comparison_input_terms", {}).values():
        quantity = term.get("sales_quantity")
        input_length = term.get("input_length")
        if quantity and input_length:
            parts.append(
                f"{registry.ref(quantity, 'comparison')}*"
                f"{registry.ref(input_length, 'comparison')}"
            )
    if parts:
        return f"=SUM({','.join(parts)})"
    legacy_key = str(keys.get("comparison_input_legacy_key") or "")
    return f"={registry.ref(legacy_key, 'comparison')}" if legacy_key else "=0"


def _sga_amount_formula(
    registry: SourceRegistry,
    group: Mapping[str, Any],
    side: str,
) -> str:
    keys = [str(key) for key in group.get(side, ()) if key]
    return _join_sum(f"{registry.ref(key, side)}/1000" for key in keys)


def write_cost_sheet(
    ws: Any,
    registry: SourceRegistry,
    context: Mapping[str, Any],
    sales_cells: Mapping[str, Any],
) -> dict[str, str]:
    _write_report_title(
        ws,
        "04_원가근거",
        "원재료·생산수량/가중평균 생산단가·제조경비·재고시차·판관비 근거",
        last_column=24,
        unit_legend="단위: 금액 천원 / 단가 원/PCS·원/m / 수량 PCS·m / 비율 %",
    )
    row = _section(ws, 4, "원가 Effect canonical cells", 5)
    _headers(ws, row, ["Effect Code", "Effect", "공식 수식", "단위", "정책"])
    summary_rows = {
        code: row + index
        for index, code in enumerate((
            "material_total", "manufacturing_realized", "inventory_timing",
            "sga_variable", "sga_fixed",
        ), 1)
    }
    effect_labels = {
        "material_total": "원부재료 총효과",
        "manufacturing_realized": "제조경비 효과",
        "inventory_timing": "재고/원가 반영시차 효과",
        "sga_variable": "변동 판관비 효과",
        "sga_fixed": "고정 판관비 효과",
    }
    effect_policies = {
        "material_total": "JPY child 포함; Bridge additive 1회",
        "manufacturing_realized": "Volume + Unit + Fixed; multiplier 미적용",
        "inventory_timing": "Gross - Core Quantity/Mix overlap",
        "sga_variable": "운반비·관세 제외",
        "sga_fixed": "운반비·관세 제외",
    }
    for code, current in summary_rows.items():
        ws.cell(current, 1, code)
        ws.cell(current, 2, effect_labels[code])
        ws.cell(current, 4, "천원")
        ws.cell(current, 5, effect_policies[code])
        ws.cell(current, 3).number_format = MONEY_FORMAT
        _style_total_row(ws, current, 1, 5)

    row = _section(ws, max(summary_rows.values()) + 3, "A. 원재료", 12)
    material_header = row
    _headers(ws, row, [
        "제품군", "단위", "기간", "기준 금액", "비교 금액", "기준 생산/적용량",
        "비교 생산/적용량", "비교 판매 적용량", "기준 원단위", "비교 원단위",
        "원단위 차이", "Raw Material Effect",
    ])
    row += 1
    material_start = row
    material_rows = list(context.get("material_rows") or [])
    material_formula_groups = dict(context.get("material_formula_groups") or {})
    sales_rows = dict(sales_cells.get("_sales_table_rows") or {})
    for item in material_rows:
        period = str(item.get("period") or "선택기간")
        product = str(item.get("product_group") or "")
        material_group = material_formula_groups.get((period, product), {})
        sales_row = sales_rows.get((period, product))
        current = row
        ws.cell(current, 1, product)
        ws.cell(current, 2, item.get("unit"))
        ws.cell(current, 3, period)
        ws.cell(current, 4, _material_cost_formula(registry, material_group, "baseline"))
        ws.cell(current, 5, _material_cost_formula(registry, material_group, "comparison"))
        ws.cell(current, 6, _material_output_formula(registry, material_group, "baseline"))
        ws.cell(current, 7, _material_output_formula(registry, material_group, "comparison"))
        ws.cell(current, 8, f"={_formula_ref('03_판매근거', f'F{sales_row}')}" if sales_row else "=0")
        ws.cell(current, 9, f"=IFERROR(D{current}*1000/F{current},0)")
        ws.cell(current, 10, f"=IFERROR(E{current}*1000/G{current},0)")
        ws.cell(current, 11, f"=J{current}-I{current}")
        ws.cell(current, 12, f'=IF(OR(AND(D{current}<>0,F{current}=0),AND(E{current}<>0,G{current}=0)),0,(I{current}-J{current})*H{current}/1000)')
        for column in range(4, 13):
            ws.cell(current, column).number_format = MONEY_FORMAT if column in {4, 5, 12} else UNIT_COST_FORMAT
        row += 1
    material_end = row - 1
    _style_data_rows(ws, material_start, material_end, 12)
    material_total_cell = f"L{row}"
    ws.cell(row, 1, "원부재료 총효과")
    ws.cell(row, 12, f"=SUM(L{material_start}:L{material_end})")
    ws.cell(row, 12).number_format = MONEY_FORMAT
    _style_total_row(ws, row, 1, 12)
    row += 3

    row = _section(ws, row, "A-1. 부직포 JPY / 가격 분해 (material_total의 non-additive child)", 14)
    _headers(ws, row, [
        "기간", "기준 금액", "비교 금액", "기준 생산길이", "비교 생산길이",
        "비교 적용길이", "기준 JPY FX", "비교 JPY FX", "기준 단가", "비교 단가",
        "부직포 Total", "기준 JPY 단가", "JPY Effect", "환율 제외 Effect",
    ])
    row += 1
    nonwoven_start = row
    nonwoven_rows = list(context.get("nonwoven_rows") or [])
    nonwoven_keys = dict(context.get("nonwoven_keys") or {})
    for index, item in enumerate(nonwoven_rows, 1):
        period = str(item.get("period") or "선택기간")
        keys = nonwoven_keys[f"{period}:{index}"]
        current = row
        ws.cell(current, 1, period)
        ws.cell(current, 2, f"{_nonwoven_component_formula(registry, keys, 'baseline', 'cost')}/1000")
        ws.cell(current, 3, f"{_nonwoven_component_formula(registry, keys, 'comparison', 'cost')}/1000")
        ws.cell(current, 4, _nonwoven_component_formula(registry, keys, "baseline", "output"))
        ws.cell(current, 5, _nonwoven_component_formula(registry, keys, "comparison", "output"))
        ws.cell(current, 6, _nonwoven_input_formula(registry, keys))
        ws.cell(current, 7, _nonwoven_component_formula(registry, keys, "baseline", "jpy_fx"))
        ws.cell(current, 8, _nonwoven_component_formula(registry, keys, "comparison", "jpy_fx"))
        ws.cell(current, 9, f"=IFERROR(B{current}*1000/D{current},0)")
        ws.cell(current, 10, f"=IFERROR(C{current}*1000/E{current},0)")
        ws.cell(current, 11, f"=(I{current}-J{current})*F{current}/1000")
        ws.cell(current, 12, f"=IFERROR(I{current}/G{current},0)")
        ws.cell(current, 13, f"=F{current}*L{current}*(G{current}-H{current})/1000")
        ws.cell(current, 14, f"=K{current}-M{current}")
        for column in range(2, 15):
            ws.cell(current, column).number_format = FX_FORMAT if column in {7, 8} else MONEY_FORMAT if column in {2, 3, 11, 13, 14} else UNIT_COST_FORMAT
        row += 1
    nonwoven_end = row - 1
    _style_data_rows(ws, nonwoven_start, nonwoven_end, 14)
    jpy_total_cell = f"M{row}"
    price_ex_fx_cell = f"N{row}"
    ws.cell(row, 1, "부직포 분해 합계")
    ws.cell(row, 13, f"=SUM(M{nonwoven_start}:M{nonwoven_end})" if nonwoven_end >= nonwoven_start else "=0")
    ws.cell(row, 14, f"=SUM(N{nonwoven_start}:N{nonwoven_end})" if nonwoven_end >= nonwoven_start else "=0")
    ws.cell(row, 13).number_format = MONEY_FORMAT
    ws.cell(row, 14).number_format = MONEY_FORMAT
    _style_total_row(ws, row, 1, 14)
    nonwoven_summary_row = row
    row += 3

    row = _section(ws, row, "B. 생산수량 / 가중평균 생산단가", 12)
    production_header = row
    _headers(ws, row, [
        "공정", "Basis", "단위", "기준 수량", "비교 수량", "수량 차이",
        "기준 생산금액", "비교 생산금액", "기준 가중평균 단가", "비교 가중평균 단가",
        "단가 차이", "산식/Source",
    ])
    row += 1
    production_start = row
    production_groups = dict(context.get("production_groups") or {})
    production_rows = list(context.get("production_rows") or [])
    production_output_rows: dict[str, int] = {}
    for item in production_rows:
        product = str(item.get("production_basis") or "")
        unit = str(item.get("unit") or "")
        groups = production_groups.get((product, unit), {"quantity": [], "amount": []})
        current = row
        production_output_rows[product] = current
        ws.cell(current, 1, item.get("process"))
        ws.cell(current, 2, product)
        ws.cell(current, 3, unit)
        ws.cell(current, 4, _join_sum(registry.ref(key, "baseline") for key in groups.get("quantity", ())))
        ws.cell(current, 5, _join_sum(registry.ref(key, "comparison") for key in groups.get("quantity", ())))
        ws.cell(current, 6, f"=E{current}-D{current}")
        ws.cell(current, 7, _join_sum(f"{registry.ref(key, 'baseline')}/1000" for key in groups.get("amount", ())))
        ws.cell(current, 8, _join_sum(f"{registry.ref(key, 'comparison')}/1000" for key in groups.get("amount", ())))
        ws.cell(current, 9, f'=IF(D{current}=0,"",G{current}*1000/D{current})')
        ws.cell(current, 10, f'=IF(E{current}=0,"",H{current}*1000/E{current})')
        ws.cell(current, 11, f'=IF(OR(I{current}="",J{current}=""),"",J{current}-I{current})')
        ws.cell(current, 12, item.get("aggregation_basis"))
        for column in range(4, 12):
            ws.cell(current, column).number_format = MONEY_FORMAT if column in {7, 8} else UNIT_COST_FORMAT
        row += 1
    back_row = row
    ws.cell(back_row, 1, "후공정 합계")
    ws.cell(back_row, 2, "SW+BW+LC")
    ws.cell(back_row, 3, "PCS")
    back_component_rows = [production_output_rows[group] for group in ("SW", "BW", "LC") if group in production_output_rows]
    ws.cell(back_row, 4, _join_sum(f"D{current}" for current in back_component_rows))
    ws.cell(back_row, 5, _join_sum(f"E{current}" for current in back_component_rows))
    ws.cell(back_row, 6, f"=E{back_row}-D{back_row}")
    ws.cell(back_row, 7, _join_sum(f"G{current}" for current in back_component_rows))
    ws.cell(back_row, 8, _join_sum(f"H{current}" for current in back_component_rows))
    ws.cell(back_row, 9, f'=IF(D{back_row}=0,"",G{back_row}*1000/D{back_row})')
    ws.cell(back_row, 10, f'=IF(E{back_row}=0,"",H{back_row}*1000/E{back_row})')
    ws.cell(back_row, 11, f'=IF(OR(I{back_row}="",J{back_row}=""),"",J{back_row}-I{back_row})')
    ws.cell(back_row, 12, "Σ SW/BW/LC 생산금액 ÷ Σ SW/BW/LC 생산수량")
    for column in range(4, 12):
        ws.cell(back_row, column).number_format = MONEY_FORMAT if column in {7, 8} else UNIT_COST_FORMAT
    _style_data_rows(ws, production_start, back_row, 12)
    _style_total_row(ws, back_row, 1, 12)
    ws.cell(back_row + 1, 1, "* 생산 수량과 금액: 수불부 기준")
    ws.cell(back_row + 1, 1).font = Font(name=FONT_NAME, size=9, italic=True, color=MUTED)
    ws.merge_cells(start_row=back_row + 1, start_column=1, end_row=back_row + 1, end_column=12)
    row = back_row + 4

    row = _section(ws, row, "C. 제조경비", 24)
    _headers(ws, row, [
        "기간", "기준 전공정 활동", "비교 전공정 활동", "기준 후공정 활동",
        "비교 후공정 활동", "기준 외주 후공정 활동", "비교 외주 후공정 활동",
        "단위", "산식/Source",
    ])
    row += 1
    activity_start = row
    activity_groups = dict(context.get("manufacturing_activity_groups") or {})
    activity_periods = list(context.get("manufacturing_activity_periods") or [])
    activity_output_rows: dict[str, int] = {}
    for period in activity_periods:
        current = row
        activity_output_rows[period] = current
        front_groups = activity_groups.get((period, "front"), {})
        back_groups = activity_groups.get((period, "back"), {})
        outsourcing_groups = activity_groups.get((period, "outsourcing_back"), {})
        ws.cell(current, 1, period)
        ws.cell(current, 2, _activity_formula(registry, front_groups, "baseline"))
        ws.cell(current, 3, _activity_formula(registry, front_groups, "comparison"))
        ws.cell(current, 4, _activity_formula(registry, back_groups, "baseline"))
        ws.cell(current, 5, _activity_formula(registry, back_groups, "comparison"))
        ws.cell(current, 6, _activity_formula(registry, outsourcing_groups, "baseline"))
        ws.cell(current, 7, _activity_formula(registry, outsourcing_groups, "comparison"))
        ws.cell(current, 8, "전공정 m / 후공정 PCS")
        ws.cell(current, 9, "일반 후공정은 MCM 포함; 외주 후공정만 MCM 차감")
        for column in range(2, 8):
            ws.cell(current, column).number_format = COUNT_FORMAT
        row += 1
    activity_end = row - 1
    _style_data_rows(ws, activity_start, activity_end, 9)
    row += 2
    manufacturing_header = row
    _headers(ws, row, [
        "기간", "계정", "구분", "기준 금액", "비교 금액", "손익 차이", "기준 전공정 비율", "비교 적용비율",
        "기준 전공정 배부", "비교 전공정 배부", "기준 후공정 배부", "비교 후공정 배부",
        "기준 전공정 수량", "비교 전공정 수량", "기준 후공정 수량", "비교 후공정 수량",
        "기준 전공정 단가", "비교 전공정 단가", "기준 후공정 단가", "비교 후공정 단가",
        "Volume", "Unit", "Fixed", "제조경비 Effect",
    ])
    row += 1
    manufacturing_start = row
    manufacturing_rows = list(context.get("manufacturing_rows") or [])
    manufacturing_keys = dict(context.get("manufacturing_keys") or {})
    for index, item in enumerate(manufacturing_rows, 1):
        period = str(item.get("month") or "선택기간")
        keys = manufacturing_keys[f"{period}:{index}"]
        activity_row = activity_output_rows[period]
        back_activity_basis = str(item.get("back_activity_basis") or "")
        uses_outsourcing_activity = (
            back_activity_basis == "OUTSOURCING_BACK"
            or str(item.get("account") or "") == "외주가공비"
        )
        base_back_column = "F" if uses_outsourcing_activity else "D"
        comparison_back_column = "G" if uses_outsourcing_activity else "E"
        current = row
        ws.cell(current, 1, period)
        ws.cell(current, 2, item.get("account"))
        ws.cell(current, 3, item.get("classification"))
        ws.cell(current, 4, f"={registry.ref(keys['amount'], 'baseline')}/1000")
        ws.cell(current, 5, f"={registry.ref(keys['amount'], 'comparison')}/1000")
        ws.cell(current, 6, f"=D{current}-E{current}")
        ws.cell(current, 7, f"={registry.ref(keys['ratio'], 'baseline')}")
        ws.cell(current, 8, f"={registry.ref(keys['ratio'], 'comparison')}")
        ws.cell(current, 9, f"=D{current}*G{current}")
        ws.cell(current, 10, f"=E{current}*H{current}")
        ws.cell(current, 11, f"=D{current}-I{current}")
        ws.cell(current, 12, f"=E{current}-J{current}")
        ws.cell(current, 13, f"=B{activity_row}")
        ws.cell(current, 14, f"=C{activity_row}")
        ws.cell(current, 15, f"={base_back_column}{activity_row}")
        ws.cell(current, 16, f"={comparison_back_column}{activity_row}")
        ws.cell(current, 17, f"=IFERROR(I{current}*1000/M{current},0)")
        ws.cell(current, 18, f"=IFERROR(J{current}*1000/N{current},0)")
        ws.cell(current, 19, f"=IFERROR(K{current}*1000/O{current},0)")
        ws.cell(current, 20, f"=IFERROR(L{current}*1000/P{current},0)")
        ws.cell(current, 21, f'=IF(C{current}="variable",IF(OR(M{current}="",N{current}="",M{current}=0,N{current}=0),0,(M{current}-N{current})*Q{current}/1000)+IF(OR(O{current}="",P{current}="",O{current}=0,P{current}=0),0,(O{current}-P{current})*S{current}/1000),0)')
        ws.cell(current, 22, f'=IF(C{current}="variable",IF(OR(M{current}="",N{current}="",M{current}=0,N{current}=0),I{current}-J{current},N{current}*(Q{current}-R{current})/1000)+IF(OR(O{current}="",P{current}="",O{current}=0,P{current}=0),K{current}-L{current},P{current}*(S{current}-T{current})/1000),0)')
        ws.cell(current, 23, f'=IF(C{current}="fixed",D{current}-E{current},0)')
        ws.cell(current, 24, f"=U{current}+V{current}+W{current}")
        for column in range(4, 25):
            ws.cell(current, column).number_format = PERCENT_FORMAT if column in {7, 8} else MONEY_FORMAT if column in {4, 5, 6, 9, 10, 11, 12, 21, 22, 23, 24} else UNIT_COST_FORMAT
        row += 1
    manufacturing_end = row - 1
    _style_data_rows(ws, manufacturing_start, manufacturing_end, 24)
    manufacturing_total_row = row
    ws.cell(row, 1, "제조경비 합계")
    ws.cell(row, 21, f"=SUM(U{manufacturing_start}:U{manufacturing_end})")
    ws.cell(row, 22, f"=SUM(V{manufacturing_start}:V{manufacturing_end})")
    ws.cell(row, 23, f"=SUM(W{manufacturing_start}:W{manufacturing_end})")
    ws.cell(row, 24, f"=SUM(X{manufacturing_start}:X{manufacturing_end})")
    for column in range(21, 25):
        ws.cell(row, column).number_format = MONEY_FORMAT
    ws.cell(row, 2, "Volume")
    ws.cell(row, 3, f"=U{row}")
    ws.cell(row, 4, "Unit")
    ws.cell(row, 5, f"=V{row}")
    ws.cell(row, 6, "Fixed")
    ws.cell(row, 7, f"=W{row}")
    ws.cell(row, 8, "Total")
    ws.cell(row, 9, f"=X{row}")
    for column in (3, 5, 7, 9):
        ws.cell(row, column).number_format = MONEY_FORMAT
    _style_total_row(ws, row, 1, 24)
    row += 3

    row = _section(ws, row, "D. 재고 / 원가 반영시차", 17)
    _headers(ws, row, [
        "공식 계산", "기준", "비교", "손익효과", "정책",
        "기간", "Pool", "제품군", "단위", "기준 수량", "비교 수량",
        "Pool 기준수량", "Pool 비교수량", "기준 Core COGS", "기준 COGS/단위",
        "Quantity Overlap", "Mix Overlap",
    ])
    inventory_header = row
    row += 1
    inventory_rows: dict[str, int] = {}
    for canonical, label in (
        ("finished_goods_cogs", "제품 매출원가"),
        ("semi_finished_goods_cogs", "반제품 매출원가"),
        ("manufactured_cogs", "제품+반제품 매출원가"),
        ("current_manufacturing_cost", "당기투입제조원가"),
        ("manufactured_cogs_effect", "Manufactured COGS Effect"),
        ("current_cost_effect", "Current Manufacturing Cost Effect"),
        ("gross_inventory_timing", "Gross Inventory Timing"),
        ("core_overlap", "Core Manufactured COGS Overlap"),
        ("inventory_timing", "Net Inventory Timing Effect"),
    ):
        inventory_rows[canonical] = row
        ws.cell(row, 1, label)
        row += 1
    finished = inventory_rows["finished_goods_cogs"]
    semi = inventory_rows["semi_finished_goods_cogs"]
    manufactured = inventory_rows["manufactured_cogs"]
    current_cost = inventory_rows["current_manufacturing_cost"]
    mfg_effect = inventory_rows["manufactured_cogs_effect"]
    cost_effect = inventory_rows["current_cost_effect"]
    gross = inventory_rows["gross_inventory_timing"]
    core_total = inventory_rows["core_overlap"]
    inventory_total = inventory_rows["inventory_timing"]
    for canonical, target in (("finished_goods_cogs", finished), ("semi_finished_goods_cogs", semi), ("current_manufacturing_cost", current_cost)):
        ws.cell(target, 2, _join_sum(f"{ref}/1000" for ref in registry.refs(f"inventory:{canonical}:base", "baseline")))
        ws.cell(target, 3, _join_sum(f"{ref}/1000" for ref in registry.refs(f"inventory:{canonical}:comparison", "comparison")))
    ws.cell(manufactured, 2, f"=B{finished}+B{semi}")
    ws.cell(manufactured, 3, f"=C{finished}+C{semi}")
    ws.cell(mfg_effect, 4, f"=B{manufactured}-C{manufactured}")
    ws.cell(cost_effect, 4, f"=B{current_cost}-C{current_cost}")
    ws.cell(gross, 4, f"=D{mfg_effect}-D{cost_effect}")
    ws.cell(gross, 5, "Manufactured COGS Effect - Current Manufacturing Cost Effect")
    for target in (finished, semi, manufactured, current_cost, mfg_effect, cost_effect, gross, core_total, inventory_total):
        for column in (2, 3, 4):
            ws.cell(target, column).number_format = MONEY_FORMAT
    core_detail_start = row + 1
    core_rows = [item for item in context.get("core_rows") or [] if item.get("selected")]
    core_keys = dict(context.get("core_keys") or {})
    pool_rows = dict(sales_cells.get("_pool_rows") or {})
    core_formula_rows: list[int] = []
    for index, item in enumerate(core_rows):
        current = core_detail_start + index
        core_formula_rows.append(current)
        period = str(item.get("period") or "선택기간")
        product = str(item.get("product_group") or "")
        pool = str(item.get("pool") or "")
        sales_row = sales_rows.get((period, product))
        pool_row = pool_rows.get((period, pool))
        core_key = core_keys[(period, product)]
        ws.cell(current, 6, period)
        ws.cell(current, 7, pool)
        ws.cell(current, 8, product)
        ws.cell(current, 9, item.get("unit"))
        ws.cell(current, 10, f"={_formula_ref('03_판매근거', f'E{sales_row}')}" if sales_row else "=0")
        ws.cell(current, 11, f"={_formula_ref('03_판매근거', f'F{sales_row}')}" if sales_row else "=0")
        ws.cell(current, 12, f"={_formula_ref('03_판매근거', f'D{pool_row}')}" if pool_row else "=0")
        ws.cell(current, 13, f"={_formula_ref('03_판매근거', f'E{pool_row}')}" if pool_row else "=0")
        ws.cell(current, 14, f"={registry.ref(core_key, 'baseline')}/1000")
        ws.cell(current, 15, f"=IFERROR(N{current}*1000/J{current},0)")
        ws.cell(current, 16, f"=-((M{current}-L{current})*IFERROR(J{current}/L{current},0)*O{current})/1000")
        ws.cell(current, 17, f"=-(M{current}*(IFERROR(K{current}/M{current},0)-IFERROR(J{current}/L{current},0))*O{current})/1000")
        for column in range(10, 18):
            ws.cell(current, column).number_format = MONEY_FORMAT if column in {14, 16, 17} else UNIT_COST_FORMAT
    core_detail_end = core_detail_start + len(core_formula_rows) - 1
    _style_data_rows(ws, core_detail_start, core_detail_end, 17)
    ws.cell(core_total, 4, f"=SUM(P{core_detail_start}:P{core_detail_end})+SUM(Q{core_detail_start}:Q{core_detail_end})" if core_detail_end >= core_detail_start else "=0")
    ws.cell(core_total, 5, "Sales Quantity/Mix에 이미 포함된 Core COGS")
    ws.cell(inventory_total, 4, f"=D{gross}-D{core_total}")
    ws.cell(inventory_total, 5, "공식 additive Effect = Gross - Core overlap")
    _style_total_row(ws, inventory_total, 1, 5)
    row = max(row, core_detail_end + 1) + 2

    row = _section(ws, row, "E. 판관비", 9)
    sga_header = row
    _headers(ws, row, [
        "기간", "구역", "계정", "구분", "기준 금액", "비교 금액", "손익효과",
        "Bridge 위치", "정책",
    ])
    row += 1
    sga_start = row
    sga_rows = list(context.get("sga_rows") or [])
    sga_keys = dict(context.get("sga_keys") or {})
    sga_key_list = list(context.get("sga_key_list") or [])
    sga_formula_groups = dict(context.get("sga_formula_groups") or {})
    for index, item in enumerate(sga_rows, 1):
        period = str(item.get("period") or "선택기간")
        group_key = f"{period}:{index}"
        group = sga_formula_groups.get(group_key, {})
        key = sga_keys.get(group_key) or sga_key_list[index - 1]
        current = row
        classification = str(item.get("classification") or "")
        bridge_position = str(item.get("bridge_position") or "")
        ws.cell(current, 1, period)
        ws.cell(current, 2, item.get("section"))
        ws.cell(current, 3, item.get("display_account") or item.get("account"))
        ws.cell(current, 4, classification)
        ws.cell(current, 5, _sga_amount_formula(registry, group, "baseline") if group else f"={registry.ref(key, 'baseline')}/1000")
        ws.cell(current, 6, _sga_amount_formula(registry, group, "comparison") if group else f"={registry.ref(key, 'comparison')}/1000")
        ws.cell(current, 7, f'=IF(OR(H{current}="판매효과",H{current}="외부효과/관세",D{current}="transport",D{current}="tariff"),0,E{current}-F{current})')
        ws.cell(current, 8, bridge_position)
        ws.cell(current, 9, "운반비/관세는 판매 Effect에서만 반영")
        for column in (5, 6, 7):
            ws.cell(current, column).number_format = MONEY_FORMAT
        row += 1
    sga_end = row - 1
    _style_data_rows(ws, sga_start, sga_end, 9)
    sga_variable_row = row
    sga_fixed_row = row + 1
    ws.cell(sga_variable_row, 1, "변동 판관비 효과")
    ws.cell(sga_variable_row, 7, f'=SUMIF(D{sga_start}:D{sga_end},"variable",G{sga_start}:G{sga_end})')
    ws.cell(sga_fixed_row, 1, "고정 판관비 효과")
    ws.cell(sga_fixed_row, 7, f'=SUMIF(D{sga_start}:D{sga_end},"fixed",G{sga_start}:G{sga_end})')
    for target in (sga_variable_row, sga_fixed_row):
        ws.cell(target, 7).number_format = MONEY_FORMAT
        _style_total_row(ws, target, 1, 9)

    # Canonical summary references are written last so each Effect has one calculation cell.
    summary_formulas = {
        "material_total": f"={material_total_cell}",
        "manufacturing_realized": f"=X{manufacturing_total_row}",
        "inventory_timing": f"=D{inventory_total}",
        "sga_variable": f"=G{sga_variable_row}",
        "sga_fixed": f"=G{sga_fixed_row}",
    }
    for code, formula in summary_formulas.items():
        ws.cell(summary_rows[code], 3, formula)
        ws.cell(summary_rows[code], 3).number_format = MONEY_FORMAT

    _group_detail_rows(ws, material_start, material_end)
    _group_detail_rows(ws, nonwoven_start, nonwoven_end)
    _group_detail_rows(ws, manufacturing_start, manufacturing_end)
    _group_detail_rows(ws, core_detail_start, core_detail_end)
    _group_detail_rows(ws, sga_start, sga_end)
    ws.row_breaks.append(Break(id=inventory_header - 2))

    _set_widths(ws, {
        "A": 24, "B": 20, "C": 12, "D": 12, "E": 12, "F": 12,
        "G": 12, "H": 13, "I": 13, "J": 13, "K": 13, "L": 13,
        "M": 13, "N": 13, "O": 13, "P": 13, "Q": 13, "R": 13,
        "S": 13, "T": 13, "U": 12, "V": 12, "W": 12, "X": 13,
    })
    _finish_sheet(
        ws, last_row=ws.max_row, last_column=24, freeze=f"A{material_header + 1}",
        print_header_rows="1:2", print_last_column=14,
    )
    return {code: f"C{target}" for code, target in summary_rows.items()} | {
        "material_jpy": jpy_total_cell,
        "material_price_ex_fx": price_ex_fx_cell,
        "production_back_weighted_unit_cost": f"I{back_row}",
    }


def write_effect_sheet(
    ws: Any,
    registry: SourceRegistry,
    sales_cells: Mapping[str, Any],
    cost_cells: Mapping[str, str],
) -> dict[str, str]:
    _write_report_title(
        ws,
        "02_손익영향",
        "공식 Effect는 각 Detail sheet의 단일 canonical calculation cell을 참조합니다.",
        last_column=6,
        unit_legend="단위: 금액 천원",
    )
    row = _section(ws, 4, "손익 요약", 6)
    _headers(ws, row, ["항목", "기준", "비교", "증감", "단위", "근거"])
    pnl_start = row + 1
    baseline_op = registry.ref("pnl:operating_profit", "baseline")
    comparison_op = registry.ref("pnl:operating_profit", "comparison")
    ws.append([
        "영업이익",
        f"={baseline_op}/1000",
        f"={comparison_op}/1000",
        f"=C{pnl_start}-B{pnl_start}",
        "천원",
        "90_원본값 P&L operating_profit",
    ])
    for column in (2, 3, 4):
        ws.cell(pnl_start, column).number_format = MONEY_FORMAT
    ws.cell(pnl_start, 5).alignment = Alignment(horizontal="center", vertical="center")
    _style_total_row(ws, pnl_start, 1, 6)

    row = _section(ws, pnl_start + 3, "손익 영향 요인", 6)
    _headers(ws, row, [
        "그룹", "Effect", "금액", "단위", "Canonical Detail Cell", "", "Effect Code",
    ])
    effect_start = row + 1
    effect_specs = [
        ("매출", "sales_quantity", "판매수량", "03_판매근거", str(sales_cells["sales_quantity"])),
        ("매출", "sales_mix", "Mix", "03_판매근거", str(sales_cells["sales_mix"])),
        ("매출", "sales_price", "판매가격", "03_판매근거", str(sales_cells["sales_price"])),
        ("매출", "sales_fx", "매출환율", "03_판매근거", str(sales_cells["sales_fx"])),
        ("매출", "tariff", "관세", "03_판매근거", str(sales_cells["tariff"])),
        ("원재료", "material_total", "원부재료", "04_원가근거", cost_cells["material_total"]),
        ("제조", "manufacturing_realized", "제조경비", "04_원가근거", cost_cells["manufacturing_realized"]),
        ("제조", "inventory_timing", "재고/원가 반영시차", "04_원가근거", cost_cells["inventory_timing"]),
        ("판관비", "sga_variable", "변동 판관비", "04_원가근거", cost_cells["sga_variable"]),
        ("판관비", "sga_fixed", "고정 판관비", "04_원가근거", cost_cells["sga_fixed"]),
    ]
    effect_cells: dict[str, str] = {}
    for index, (group, code, label, sheet, cell) in enumerate(effect_specs):
        current = effect_start + index
        ws.cell(current, 1, group)
        ws.cell(current, 2, label)
        ws.cell(current, 3, _sheet_ref(sheet, cell))
        ws.cell(current, 4, "천원")
        ws.cell(current, 5, f"{sheet}!{cell}")
        ws.cell(current, 7, code)
        ws.cell(current, 3).number_format = MONEY_FORMAT
        effect_cells[code] = f"C{current}"
    effect_end = effect_start + len(effect_specs) - 1
    _style_data_rows(ws, effect_start, effect_end, 7)

    summary = effect_end + 2
    ws.cell(summary, 3, "Effects Total")
    ws.cell(summary, 4, f"=SUM(C{effect_start}:C{effect_end})")
    ws.cell(summary + 1, 3, "OP Delta")
    ws.cell(summary + 1, 4, f"=D{pnl_start}")
    ws.cell(summary + 2, 3, "기타 요인")
    ws.cell(summary + 2, 4, f"=D{summary + 1}-D{summary}")
    for current in range(summary, summary + 3):
        ws.cell(current, 4).number_format = MONEY_FORMAT
        ws.cell(current, 5, "천원")
        _style_total_row(ws, current, 3, 6)

    reconciliation = summary + 4
    row = _section(ws, reconciliation, "Reconciliation", 6)
    _headers(ws, row, ["검증", "Actual", "Expected", "Difference", "Tolerance", "Status"])
    check = row + 1
    ws.cell(check, 1, "OP Delta = Effects Total + 기타 요인")
    ws.cell(check, 2, f"=D{summary + 1}")
    ws.cell(check, 3, f"=D{summary}+D{summary + 2}")
    ws.cell(check, 4, f"=B{check}-C{check}")
    ws.cell(check, 5, f"={registry.ref('policy:reconciliation_tolerance', 'baseline')}/1000")
    ws.cell(check, 6, f'=IF(ABS(D{check})<=E{check},"PASS","CHECK")')
    for column in range(2, 6):
        ws.cell(check, column).number_format = MONEY_FORMAT
    ws.cell(check, 5).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(check, 6).alignment = Alignment(horizontal="center", vertical="center")
    _style_total_row(ws, check, 1, 6)

    _set_widths(ws, {
        "A": 28, "B": 24, "C": 14, "D": 10, "E": 26, "F": 10, "G": 25,
    })
    ws.column_dimensions["G"].hidden = True
    _finish_sheet(
        ws, last_row=ws.max_row, last_column=7, freeze=f"A{effect_start}",
        print_header_rows="1:10", fit_height=1, print_last_column=6,
    )
    return effect_cells | {
        "baseline_operating_profit": f"B{pnl_start}",
        "comparison_operating_profit": f"C{pnl_start}",
        "operating_profit_delta": f"D{summary + 1}",
        "effects_total": f"D{summary}",
        "residual": f"D{summary + 2}",
        "reconciliation": f"F{check}",
    }


def write_summary_sheet(
    ws: Any,
    result: Mapping[str, Any],
    sales_cells: Mapping[str, Any],
    effect_cells: Mapping[str, str],
) -> None:
    _write_report_title(
        ws,
        "NANOH2O  손익 변동 요인 분석 산출 근거",
        "Backend authoritative Result를 Source → Detail → Summary 수식 흐름으로 재현한 보고용 Evidence Workbook",
        last_column=8,
        unit_legend="단위: 금액 천원 / 환율 KRW/USD",
    )
    baseline = result.get("baseline") or {}
    comparison = result.get("comparison") or {}
    period = result.get("period") or {}
    metadata = [
        ("기준 모형", baseline.get("name"), "비교 모형", comparison.get("name")),
        ("분석기간", period.get("label"), "생성일시", datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")),
        ("기준 평균 FX", _sheet_ref("03_판매근거", str(sales_cells["baseline_average_fx"])), "비교 평균 FX", _sheet_ref("03_판매근거", str(sales_cells["comparison_average_fx"]))),
    ]
    for offset, (left_label, left_value, right_label, right_value) in enumerate(metadata, 4):
        ws.cell(offset, 1, left_label)
        ws.cell(offset, 2, left_value)
        ws.cell(offset, 5, right_label)
        ws.cell(offset, 6, right_value)
        for column in (1, 5):
            ws.cell(offset, column).font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY_DARK)
            ws.cell(offset, column).fill = PatternFill("solid", fgColor=LIGHT_NEUTRAL)
        for column in (2, 6):
            ws.cell(offset, column).alignment = Alignment(horizontal="right" if offset == 6 else "left")
            if offset == 6:
                ws.cell(offset, column).number_format = FX_FORMAT

    row = _section(ws, 8, "손익 요약", 8)
    _headers(ws, row, ["항목", "금액", "단위", "", "항목", "금액", "단위", ""])
    first = row + 1
    summary_items = [
        ("기준 영업이익", "baseline_operating_profit"),
        ("비교 영업이익", "comparison_operating_profit"),
        ("증감", "operating_profit_delta"),
    ]
    for index, (label, key) in enumerate(summary_items):
        current = first + index
        ws.cell(current, 1, label)
        ws.cell(current, 2, _sheet_ref("02_손익영향", effect_cells[key]))
        ws.cell(current, 3, "천원")
        ws.cell(current, 2).number_format = MONEY_FORMAT
        _style_total_row(ws, current, 1, 3)

    row = _section(ws, first + len(summary_items) + 2, "손익 영향 요인", 8)
    _headers(ws, row, ["그룹", "Effect", "금액", "단위", "그룹", "Effect", "금액", "단위"])
    effects = [
        ("매출", "판매수량", "sales_quantity"),
        ("매출", "Mix", "sales_mix"),
        ("매출", "판매가격", "sales_price"),
        ("매출", "매출환율", "sales_fx"),
        ("매출", "관세", "tariff"),
        ("원재료", "원부재료", "material_total"),
        ("제조", "제조경비", "manufacturing_realized"),
        ("제조", "재고/원가 반영시차", "inventory_timing"),
        ("판관비", "변동 판관비", "sga_variable"),
        ("판관비", "고정 판관비", "sga_fixed"),
        ("기타", "기타 요인", "residual"),
    ]
    effect_start = row + 1
    left_count = (len(effects) + 1) // 2
    for index, (group, label, key) in enumerate(effects):
        block = 0 if index < left_count else 4
        current = effect_start + (index if block == 0 else index - left_count)
        ws.cell(current, 1 + block, group)
        ws.cell(current, 2 + block, label)
        ws.cell(current, 3 + block, _sheet_ref("02_손익영향", effect_cells[key]))
        ws.cell(current, 4 + block, "천원")
        ws.cell(current, 3 + block).number_format = MONEY_FORMAT
    effect_end = effect_start + left_count - 1
    _style_data_rows(ws, effect_start, effect_end, 8)

    row = _section(ws, effect_end + 2, "Reconciliation", 8)
    _headers(ws, row, ["항목", "금액/상태", "단위", "", "항목", "금액/상태", "단위", ""])
    reconciliation_items = [
        ("영업이익 증감", "operating_profit_delta", "천원"),
        ("Effects Total", "effects_total", "천원"),
        ("기타 요인", "residual", "천원"),
        ("정합성 확인", "reconciliation", ""),
    ]
    reconciliation_start = row + 1
    for index, (label, key, unit) in enumerate(reconciliation_items):
        current = reconciliation_start + index
        ws.cell(current, 1, label)
        ws.cell(current, 2, _sheet_ref("02_손익영향", effect_cells[key]))
        ws.cell(current, 3, unit)
        if unit:
            ws.cell(current, 2).number_format = MONEY_FORMAT
        _style_total_row(ws, current, 1, 3)
    ws["E10"] = "Formula lineage"
    ws["F10"] = "90_원본값 → 03/04 Detail → 02_손익영향 → 01_보고요약"
    ws.merge_cells("F10:H12")
    ws["F10"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["E10"].font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY_DARK)

    _set_widths(ws, {"A": 18, "B": 18, "C": 10, "D": 3, "E": 18, "F": 22, "G": 12, "H": 10})
    _finish_sheet(
        ws, last_row=ws.max_row, last_column=8, freeze="A9",
        print_header_rows="1:9", fit_height=1,
    )


def audit_reporting_workbook(workbook: Workbook) -> dict[str, Any]:
    if tuple(workbook.sheetnames) != REPORT_SHEETS:
        raise ValueError(f"Evidence Workbook sheet contract mismatch: {workbook.sheetnames}")
    hard_coded_numeric: dict[str, list[str]] = {}
    formula_counts: dict[str, int] = {}
    formula_errors: list[str] = []
    for ws in workbook.worksheets:
        formula_count = 0
        numeric_cells: list[str] = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                    upper = value.upper()
                    if any(token in upper for token in (
                        "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!",
                    )):
                        formula_errors.append(f"{ws.title}!{cell.coordinate}:{value}")
                elif ws.title != "90_원본값" and isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric_cells.append(cell.coordinate)
        formula_counts[ws.title] = formula_count
        if numeric_cells:
            hard_coded_numeric[ws.title] = numeric_cells
    if formula_counts["90_원본값"]:
        raise ValueError("90_원본값 must contain source/input values only")
    if hard_coded_numeric:
        raise ValueError(f"hard-coded derived numeric cells: {hard_coded_numeric}")
    if formula_errors:
        raise ValueError(f"formula error tokens: {formula_errors}")
    if not all(workbook[name].freeze_panes for name in REPORT_SHEETS):
        raise ValueError("freeze panes missing")
    if any(workbook[name].sheet_view.showGridLines for name in REPORT_SHEETS):
        raise ValueError("gridlines must be hidden")
    if any(workbook[name].page_setup.fitToWidth != 1 for name in REPORT_SHEETS):
        raise ValueError("print fit-to-width contract failure")
    source_entries: dict[tuple[str, str], dict[float, list[str]]] = {}
    derived_source_literals: list[str] = []
    source = workbook["90_원본값"]
    for row in range(6, source.max_row + 1):
        hardcode_class = str(source[f"L{row}"].value or "")
        if hardcode_class == "MODEL_SOURCE":
            _validate_model_source_pair(
                source[f"G{row}"].value,
                source[f"H{row}"].value,
                side="baseline",
                key=str(source[f"A{row}"].value or row),
            )
            _validate_model_source_pair(
                source[f"I{row}"].value,
                source[f"J{row}"].value,
                side="comparison",
                key=str(source[f"A{row}"].value or row),
            )
        if hardcode_class not in {"MODEL_SOURCE", "REQUEST_INPUT", "POLICY_INPUT"}:
            for value_column in ("H", "J"):
                value = source[f"{value_column}{row}"].value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    derived_source_literals.append(f"{value_column}{row}:{hardcode_class}")
        for side, source_column, value_column in (
            ("BASE", "G", "H"), ("COMPARISON", "I", "J")
        ):
            source_reference = source[f"{source_column}{row}"].value
            value = source[f"{value_column}{row}"].value
            if not source_reference or not isinstance(value, (int, float)) or isinstance(value, bool):
                continue
            source_text = str(source_reference)
            if not any(marker in source_text for marker in ("Data!", "Analysis request input", "config/")):
                continue
            source_key = (side, source_text)
            source_entries.setdefault(source_key, {}).setdefault(
                float(value), []
            ).append(f"{value_column}{row}")
    source_conflicts = {
        key: values
        for key, values in source_entries.items()
        if len(values) > 1
    }
    if source_conflicts:
        raise ValueError(f"conflicting source/input readbacks: {source_conflicts}")
    hard_coded_derived_count = sum(len(cells) for cells in hard_coded_numeric.values()) + len(derived_source_literals)
    if hard_coded_derived_count:
        raise ValueError(
            f"hard-coded derived values: report={hard_coded_numeric}, source={derived_source_literals}"
        )
    return {
        "sheet_count": len(workbook.sheetnames),
        "sheet_names": list(workbook.sheetnames),
        "formula_counts": formula_counts,
        "formula_count": sum(formula_counts.values()),
        "source_numeric_count": sum(
            1
            for row in workbook["90_원본값"].iter_rows()
            for cell in row
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
        ),
        "hard_coded_derived_duplicates": hard_coded_derived_count,
        "model_source_count": sum(
            1
            for row in range(6, source.max_row + 1)
            if source[f"L{row}"].value == "MODEL_SOURCE"
        ),
        "derived_source_literal_count": len(derived_source_literals),
        "duplicated_source_values": sum(
            sum(len(cells) for cells in values.values()) - 1
            for values in source_entries.values()
            if len(values) == 1
        ),
        "conflicting_source_values": len(source_conflicts),
        "formula_error_count": len(formula_errors),
    }


def build_reporting_workbook(
    *,
    result: Mapping[str, Any],
    sales_rows: Iterable[Any],
    baseline_fx: float | None,
    comparison_fx: float | None,
    baseline_workbook: Any = None,
    comparison_workbook: Any = None,
) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    for name in REPORT_SHEETS:
        workbook.create_sheet(name)
    registry, context = collect_reporting_sources(
        result,
        sales_rows,
        baseline_fx,
        comparison_fx,
        baseline_workbook=baseline_workbook,
        comparison_workbook=comparison_workbook,
    )
    validate_source_registry_readback(
        registry,
        baseline_workbook=baseline_workbook,
        comparison_workbook=comparison_workbook,
    )
    write_source_sheet(workbook["90_원본값"], registry)
    sales_cells = write_sales_sheet(workbook["03_판매근거"], registry, context)
    cost_cells = write_cost_sheet(
        workbook["04_원가근거"], registry, context, sales_cells
    )
    effect_cells = write_effect_sheet(
        workbook["02_손익영향"], registry, sales_cells, cost_cells
    )
    write_summary_sheet(
        workbook["01_보고요약"], result, sales_cells, effect_cells
    )
    audit_reporting_workbook(workbook)
    return workbook
