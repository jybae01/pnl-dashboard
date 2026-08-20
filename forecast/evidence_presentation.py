from __future__ import annotations

from collections.abc import Iterable
from typing import Any

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


ACCOUNTING_FORMAT = '#,##0;[Red](#,##0);-'
QUANTITY_FORMAT = '#,##0.00;[Red](#,##0.00);-'
UNIT_COST_FORMAT = '#,##0.00;[Red](#,##0.00);-'
PERCENT_FORMAT = '0.00%;[Red](0.00%);-'
FX_FORMAT = '#,##0.0000;[Red](#,##0.0000);-'


_KOREAN_LABELS = {
    "Base": "기준(계획)",
    "Comparison": "비교(실적/추정)",
    "Base Source Reference": "기준 원천 위치",
    "Comparison Source Reference": "비교 원천 위치",
    "Base Source": "기준 원천 위치",
    "Comparison Source": "비교 원천 위치",
    "Base Freight(incl Tariff)": "기준 운반비(관세 포함)",
    "Comparison Freight(incl Tariff)": "비교 운반비(관세 포함)",
    "Base Freight(ex Tariff)": "기준 운반비(관세 제외)",
    "Comparison Freight(ex Tariff)": "비교 운반비(관세 제외)",
    "Base Tariff 포함?": "기준 관세 포함 여부",
    "Comparison Tariff 포함?": "비교 관세 포함 여부",
    "Business Source": "원천 항목",
    "Canonical field": "Canonical 필드",
    "Canonical Field": "Canonical 필드",
    "Source Reference": "원천 위치",
    "Calculation Source": "계산 기준",
    "Validation": "정합성 확인",
    "Calculation Validation": "계산 정합성",
    "Trace Validation": "추적 정합성",
    "Source Status": "원천 상태",
    "Engine Output": "Engine 결과",
    "Engine Value": "Engine 결과",
    "Evidence Formula": "근거 수식",
    "Excel Formula": "Excel 수식",
    "Formula Result": "수식 결과",
    "Difference": "차이",
    "Detail Cell": "상세 근거 셀",
    "Policy": "정책",
    "Effect": "효과",
    "Pool": "수량 Pool",
    "Effect Code": "효과 코드",
    "Period": "기간",
    "Product Group": "제품군",
    "Unit": "단위",
    "Selected": "적용 여부",
    "Scenario": "구분",
    "Month": "월",
    "Quantity Formula": "수량효과 수식",
    "Quantity Engine": "수량효과 Engine",
    "Mix Formula": "믹스효과 수식",
    "Mix Engine": "믹스효과 Engine",
    "Mix Component": "믹스 구성요소",
    "Base Total": "기준 총수량",
    "Comparison Total": "비교 총수량",
    "Base Weighted GP/unit": "기준 가중 GP/단위",
    "Base Quantity": "기준 수량",
    "Comparison Quantity": "비교 수량",
    "Pool Base Quantity": "Pool 기준 수량",
    "Pool Comparison Quantity": "Pool 비교 수량",
    "Base Mix": "기준 믹스",
    "Comparison Mix": "비교 믹스",
    "Base Core Manufactured COGS": "기준 핵심 제조품 COGS",
    "Base Core COGS/unit": "기준 핵심 COGS/단위",
    "Embedded Quantity COGS Expense": "수량에 포함된 COGS 비용",
    "Embedded Mix COGS Expense": "믹스에 포함된 COGS 비용",
    "Quantity Overlap OP": "수량 중복분(OP 부호)",
    "Mix Overlap OP": "믹스 중복분(OP 부호)",
    "Total Overlap OP": "총 중복분(OP 부호)",
    "Engine Quantity": "수량 중복 Engine",
    "Engine Mix": "믹스 중복 Engine",
    "Engine Total": "합계 Engine",
    "Base Quantity Source": "기준 수량 원천",
    "Comparison Quantity Source": "비교 수량 원천",
    "Base Core COGS Source": "기준 핵심 COGS 원천",
    "Comparison Core COGS Source": "비교 핵심 COGS 원천",
    "Rolling 3M Period": "최근 3개월 기간",
    "Engine Net": "최종 시차 Engine",
    "Direction": "방향",
    "Direction 수식": "방향 수식",
    "Persistence": "최근 추세",
    "Coverage": "데이터 범위",
    "Confidence": "신뢰도",
    "Materiality": "중요도",
    "Explanation Metadata": "설명 메타데이터",
    "Value": "값",
    "Applied Rule / Notes": "적용 규칙 / 설명",
    "Product Group": "제품군",
    "Specification": "규격",
    "Base Opening Unit Cost": "기준 기초재고 원단위",
    "Comparison Opening Unit Cost": "비교 기초재고 원단위",
    "Delta Formula": "차이 수식",
    "Aligned": "방향 일치",
    "Base Finished Goods COGS": "기준 제품 매출원가",
    "Comparison Finished Goods COGS": "비교 제품 매출원가",
    "Finished Goods COGS": "제품 매출원가",
    "Semi-finished Goods COGS": "반제품 매출원가",
    "Manufactured COGS": "제조품 매출원가",
    "Current Manufacturing Cost": "당기투입제조원가",
    "Manufactured COGS Effect": "제조품 매출원가 효과",
    "Current Manufacturing Cost Effect": "당기투입제조원가 효과",
    "Gross Inventory Timing": "중복 제거 전 재고·원가 반영시차",
    "Core Quantity COGS Overlap": "판매수량에 포함된 제조원가 중복분",
    "Core Mix COGS Overlap": "판매믹스에 포함된 제조원가 중복분",
    "Total Core Manufactured COGS Overlap": "판매효과 포함 제조원가 중복분",
    "Net Inventory Timing Effect": "최종 재고·원가 반영시차",
    "Core overlap policy": "핵심 제조원가 중복 제거 정책",
    "Double-count Validation": "중복계상 확인",
    "Formula Result": "수식 결과",
    "effects_total": "공식 효과 합계",
    "residual": "잔여차이(Residual)",
    "OP_delta": "영업이익 증감",
    "effects_total + residual = OP_delta": "공식 효과 합계 + 잔여차이 = 영업이익 증감",
    "Product Mix 효과": "제품 믹스 효과",
    "제품 Mix 효과": "제품 믹스 효과",
    "Manufacturing subtotal / child double count": "제조경비 합계 / 하위효과 중복 없음",
    "Inventory realization multiplier used": "재고실현율 계산 반영 여부",
    "Freight double count": "운반비 중복계상 없음",
    "Tariff separate": "관세 별도 유지",
    "Freight equivalent-shipment denominator": "운반비 환산 판매수량 검증",
    "RM / RM FX double count": "원부재료 / 환율효과 중복 없음",
    "JPY Source valid": "JPY 원천 유효",
    "Sales quantity source valid": "판매수량 원천 유효",
    "Scope Validation": "범위 정합성",
    "Base/Comparison": "기준/비교",
    "Used Value": "사용 값",
    "Calculation": "계산 기준",
    "Source value (no residual/OP backsolve)": "원천값(Residual/OP 역산 없음)",
    "Core Sales COGS double count removed": "판매 COGS 중복 제거 완료",
    "Core overlap counted in Quantity/Mix": "판매 수량·믹스에 중복분 포함",
    "Core overlap separately additive = FALSE": "중복분 별도 합산 안 함",
    "Gross Inventory Timing additive = FALSE": "중복 제거 전 시차 별도 합산 안 함",
    "Adjustment unitized = FALSE": "조정항목 단위배부 안 함",
    "Merchandise included in overlap = FALSE": "상품 COGS 중복분 제외",
    "row323 separate Effect = FALSE": "row323 별도 효과 없음",
}

_HEADER_PHRASES = (
    ("Comparison Source Reference", "비교 원천 위치"),
    ("Base Source Reference", "기준 원천 위치"),
    ("Comparison Source", "비교 원천"),
    ("Base Source", "기준 원천"),
    ("Business Source", "원천 항목"),
    ("Canonical Field", "Canonical 필드"),
    ("Canonical field", "Canonical 필드"),
    ("Source Reference", "원천 위치"),
    ("Engine Output", "Engine 결과"),
    ("Engine Value", "Engine 결과"),
    ("Engine", "Engine 결과"),
    ("Evidence Formula", "근거 수식"),
    ("Excel Formula", "Excel 수식"),
    ("Calculation Validation", "계산 정합성"),
    ("Trace Validation", "추적 정합성"),
    ("Validation", "정합성"),
    ("Comparison", "비교"),
    ("Base", "기준"),
    ("incl", "관세 포함"),
    ("ex", "관세 제외"),
    ("Front-process", "전공정"),
    ("Back-process", "후공정"),
    ("Front", "전공정"),
    ("Back", "후공정"),
    ("Freight", "운반비"),
    ("Tariff", "관세"),
    ("Price", "판매단가"),
    ("FX", "환율"),
    ("GP/unit", "GP/단위"),
    ("COGS/unit", "COGS/단위"),
    ("/unit", "/단위"),
    ("Volume", "조업도"),
    ("Fixed", "고정비"),
    ("Qty", "수량"),
    ("Length", "길이"),
    ("Amount", "금액"),
    ("Raw Material", "원부재료"),
    ("denominator", "배부기준"),
    ("Component", "구성요소"),
    ("Adjustment", "조정"),
    ("Sales", "판매"),
    ("Total", "합계"),
    ("Rate", "비율"),
    ("Difference", "차이"),
    ("Status", "상태"),
    ("Reference", "근거"),
    ("Quantity", "수량"),
    ("Mix", "믹스"),
    ("Unit Cost", "원단위"),
    ("Business", "업무"),
    ("Calculation", "계산"),
    ("Source", "원천"),
    ("Formula", "수식"),
    ("Effect", "효과"),
    ("Period", "기간"),
    ("Product Group", "제품군"),
    ("Scenario", "구분"),
    ("Unit", "단위"),
    ("Policy", "정책"),
)


def _localize_user_labels(ws) -> int:
    changed = 0
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or value.startswith("="):
                continue
            replacement = _KOREAN_LABELS.get(value)
            if replacement is not None and replacement != value:
                cell.value = replacement
                changed += 1
                continue
            if cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb in {
                "001F4E78", "1F4E78", "FF1F4E78",
            }:
                replacement = value
                for source, target in _HEADER_PHRASES:
                    replacement = replacement.replace(source, target)
                if replacement != value:
                    cell.value = replacement
                    changed += 1
    return changed


def _set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _hide_columns(ws, columns: Iterable[str]) -> None:
    for column in columns:
        dimension = ws.column_dimensions[column]
        dimension.hidden = True
        dimension.outlineLevel = 1


def _format_rows(ws, bounds: Any, formats: dict[int, str]) -> None:
    if not bounds:
        return
    start, end = (int(bounds[0]), int(bounds[1]))
    if start > end:
        return
    for row in range(start, end + 1):
        for column, number_format in formats.items():
            ws.cell(row, column).number_format = number_format


def _finish_user_sheet(
    ws,
    *,
    title_rows: str = "1:4",
    landscape: bool = True,
    fit_width: int = 1,
    zoom: int = 90,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = title_rows
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.outlinePr.summaryRight = True
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="center" if cell.row <= 4 else "top",
                wrap_text=True,
            )
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = max(float(ws.row_dimensions[2].height or 15), 36)


def _polish_sales(ws, cells: dict[str, Any]) -> None:
    ws["A2"] = (
        "제품군 내부 SKU를 먼저 합산하고 PCS/LENGTH Pool별 수량·믹스를 계산합니다. "
        "고객 운반비는 판매단가 효과에 한 번만 포함하며 관세는 분리합니다."
    )
    ws.merge_cells("A2:AG2")
    ws["AI2"] = (
        "월별 단일 고객배송 운반비 Pool을 SW·BW·LC 판매 PCS와 FS 판매길이÷45의 "
        "총 환산 판매수량으로 원단위화합니다. (기준 원단위-비교 원단위)×비교 총 환산 "
        "판매수량을 판매단가 효과에 한 번만 반영합니다."
    )
    ws.merge_cells("AI2:BO2")
    _set_widths(ws, {
        "A": 12, "B": 18, "C": 18, "D": 13, "E": 22, "F": 18,
        "G": 22, "H": 22, "I": 14, "J": 14, "K": 18, "L": 18,
        "M": 18, "N": 18, "O": 16, "P": 14, "Q": 14, "R": 12,
        "S": 12, "T": 16, "U": 12, "V": 16, "W": 12, "X": 12,
        "Y": 16, "Z": 16, "AA": 16, "AB": 16, "AC": 18, "AD": 18,
        "AE": 18, "AF": 18, "AG": 12, "AI": 11, "AJ": 22, "AK": 22,
        "AL": 22, "AM": 18, "AN": 18, "AO": 16, "AP": 16, "AQ": 12,
        "AR": 12, "AS": 18, "AT": 18, "AU": 18, "AV": 18, "AW": 18,
        "AX": 18, "AY": 18, "AZ": 18, "BA": 18, "BB": 18, "BC": 14,
        "BD": 18, "BE": 18, "BF": 20, "BG": 20, "BH": 18, "BI": 18,
        "BJ": 22, "BK": 18, "BL": 18, "BM": 18, "BN": 28, "BO": 12,
    })
    _hide_columns(ws, (
        "E", "F", "G", "H", "AD", "AF", "AG", "AK", "AL", "AM", "AN",
        "AQ", "AR", "BO",
    ))
    _format_rows(ws, cells.get("detail_range"), {
        9: QUANTITY_FORMAT, 10: QUANTITY_FORMAT, 11: ACCOUNTING_FORMAT,
        12: ACCOUNTING_FORMAT, 13: ACCOUNTING_FORMAT, 14: ACCOUNTING_FORMAT,
        15: UNIT_COST_FORMAT, 16: QUANTITY_FORMAT, 17: QUANTITY_FORMAT,
        18: PERCENT_FORMAT, 19: PERCENT_FORMAT, 20: UNIT_COST_FORMAT,
        21: PERCENT_FORMAT, 22: UNIT_COST_FORMAT, 23: FX_FORMAT, 24: FX_FORMAT,
        25: UNIT_COST_FORMAT, 26: UNIT_COST_FORMAT, 27: UNIT_COST_FORMAT,
        28: UNIT_COST_FORMAT, 29: ACCOUNTING_FORMAT, 30: ACCOUNTING_FORMAT,
        31: ACCOUNTING_FORMAT, 32: ACCOUNTING_FORMAT,
    })
    _format_rows(ws, cells.get("pool_range"), {
        10: QUANTITY_FORMAT, 11: QUANTITY_FORMAT, 12: UNIT_COST_FORMAT,
        13: ACCOUNTING_FORMAT, 14: ACCOUNTING_FORMAT, 15: UNIT_COST_FORMAT,
        16: ACCOUNTING_FORMAT, 17: ACCOUNTING_FORMAT,
    })
    _format_rows(ws, cells.get("freight_range"), {
        **{column: ACCOUNTING_FORMAT for column in range(39, 47)},
        **{column: QUANTITY_FORMAT for column in range(47, 60)},
        60: UNIT_COST_FORMAT, 61: UNIT_COST_FORMAT,
        62: ACCOUNTING_FORMAT, 63: ACCOUNTING_FORMAT,
        64: ACCOUNTING_FORMAT, 65: ACCOUNTING_FORMAT,
    })
    _format_rows(ws, cells.get("summary_range"), {2: ACCOUNTING_FORMAT, 3: ACCOUNTING_FORMAT})
    ws.freeze_panes = "A5"
    _finish_user_sheet(ws, fit_width=2, zoom=75)


def _polish_material(ws, cells: dict[str, Any]) -> None:
    ws["A2"] = (
        "왼쪽은 부직포 제외 원부재료, 오른쪽은 부직포·JPY 원천과 계산입니다. "
        "원천값만 직접 기록하고 원단위·가격효과·환율효과·합계는 Excel 수식으로 계산합니다."
    )
    ws.merge_cells("A2:AS2")
    _set_widths(ws, {
        "A": 12, "B": 18, "C": 18, "D": 22, "E": 18, "F": 22,
        "G": 22, "H": 18, "I": 18, "J": 15, "K": 15, "L": 15,
        "M": 16, "N": 16, "O": 18, "P": 18, "Q": 12, "R": 16,
        "T": 12, "U": 22, "V": 18, "W": 22, "X": 22, "Y": 18,
        "Z": 18, "AA": 15, "AB": 15, "AC": 15, "AD": 12, "AE": 12,
        "AF": 16, "AG": 16, "AH": 18, "AI": 16, "AJ": 18, "AK": 18,
        "AL": 18, "AM": 18, "AN": 12, "AO": 16,
    })
    _hide_columns(ws, (
        "E", "F", "G", "P", "Q", "R", "V", "W", "X", "AK", "AM", "AN", "AO",
    ))
    _format_rows(ws, cells.get("detail_range"), {
        8: ACCOUNTING_FORMAT, 9: ACCOUNTING_FORMAT, 10: QUANTITY_FORMAT,
        11: QUANTITY_FORMAT, 12: QUANTITY_FORMAT, 13: UNIT_COST_FORMAT,
        14: UNIT_COST_FORMAT, 15: ACCOUNTING_FORMAT, 16: ACCOUNTING_FORMAT,
    })
    _format_rows(ws, cells.get("nonwoven_range"), {
        25: ACCOUNTING_FORMAT, 26: ACCOUNTING_FORMAT, 27: QUANTITY_FORMAT,
        28: QUANTITY_FORMAT, 29: QUANTITY_FORMAT, 30: FX_FORMAT, 31: FX_FORMAT,
        32: UNIT_COST_FORMAT, 33: UNIT_COST_FORMAT,
        **{column: ACCOUNTING_FORMAT for column in range(34, 40)},
    })
    _format_rows(ws, cells.get("summary_range"), {2: ACCOUNTING_FORMAT, 3: ACCOUNTING_FORMAT})
    ws.freeze_panes = "A5"
    _finish_user_sheet(ws, fit_width=2, zoom=80)


def _polish_manufacturing(ws, cells: dict[str, Any]) -> None:
    ws["A2"] = (
        "상단 생산 기준 → 제조경비 전·후공정 배부 → 생산량 분모 → 원단위 → 조업도·원단위·고정비 효과 순서로 읽습니다. "
        "PCS와 LENGTH(m)는 합산하지 않으며 재고실현율은 효과 계산에 사용하지 않습니다."
    )
    ws.merge_cells("A2:AO2")
    _set_widths(ws, {
        "A": 12, "B": 18, "C": 18, "D": 13, "E": 18, "F": 18,
        "G": 13, "H": 26, "I": 26, "J": 14, "K": 18, "L": 18,
        "M": 12, "N": 12, "O": 18, "P": 18, "Q": 18, "R": 18,
        "S": 16, "T": 16, "U": 16, "V": 16, "W": 16, "X": 16,
        "Y": 16, "Z": 16, "AA": 18, "AB": 18, "AC": 18, "AD": 18,
        "AE": 18, "AF": 18, "AG": 18, "AH": 18, "AI": 18, "AJ": 18,
        "AK": 18, "AL": 18, "AM": 12, "AN": 13, "AO": 20,
    })
    _hide_columns(ws, ("D", "E", "F", "G", "H", "L", "M", "N", "O", "P", "Q", "R", "AD", "AH", "AJ", "AL", "AN", "AO"))
    _format_rows(ws, cells.get("production_summary_range"), {
        2: QUANTITY_FORMAT, 3: QUANTITY_FORMAT, 5: QUANTITY_FORMAT, 6: QUANTITY_FORMAT,
    })
    _format_rows(ws, cells.get("detail_range"), {
        11: ACCOUNTING_FORMAT, 12: ACCOUNTING_FORMAT, 13: PERCENT_FORMAT,
        14: PERCENT_FORMAT, 15: ACCOUNTING_FORMAT, 16: ACCOUNTING_FORMAT,
        17: ACCOUNTING_FORMAT, 18: ACCOUNTING_FORMAT,
        19: QUANTITY_FORMAT, 20: QUANTITY_FORMAT, 21: QUANTITY_FORMAT,
        22: QUANTITY_FORMAT, 23: UNIT_COST_FORMAT, 24: UNIT_COST_FORMAT,
        25: UNIT_COST_FORMAT, 26: UNIT_COST_FORMAT,
        **{column: ACCOUNTING_FORMAT for column in range(27, 39)},
        40: PERCENT_FORMAT,
    })
    _format_rows(ws, cells.get("summary_range"), {2: ACCOUNTING_FORMAT, 3: ACCOUNTING_FORMAT})
    ws.freeze_panes = f"A{cells['detail_range'][0]}"
    _finish_user_sheet(ws)


def _polish_inventory(ws, cells: dict[str, Any]) -> None:
    ws["A2"] = (
        "제조품 매출원가와 당기투입제조원가의 차이에서 판매 수량·믹스에 이미 포함된 핵심 제조원가 중복분을 차감해 "
        "최종 재고·원가 반영시차를 계산합니다. 최근 3개월 추세는 최종 시차 기준입니다."
    )
    ws.merge_cells("A2:G2")
    _set_widths(ws, {
        "A": 28, "B": 18, "C": 18, "D": 20, "E": 18, "F": 14,
        "G": 32, "H": 15, "I": 15, "J": 20, "K": 18, "L": 14,
        "M": 14, "N": 20, "O": 20, "P": 18, "Q": 18, "R": 18,
        "S": 18, "T": 18, "U": 18, "V": 14, "W": 22, "X": 24,
        "Y": 22, "Z": 24,
    })
    _hide_columns(ws, ("S", "T", "U", "V", "W", "X", "Y", "Z"))
    for row in range(5, 18):
        for column in (2, 3, 4, 5):
            ws.cell(row, column).number_format = ACCOUNTING_FORMAT
    _format_rows(ws, cells.get("core_detail_range"), {
        6: QUANTITY_FORMAT, 7: QUANTITY_FORMAT, 8: QUANTITY_FORMAT,
        9: QUANTITY_FORMAT, 10: ACCOUNTING_FORMAT, 11: UNIT_COST_FORMAT,
        12: PERCENT_FORMAT, 13: PERCENT_FORMAT,
        **{column: ACCOUNTING_FORMAT for column in range(14, 22)},
    })
    _format_rows(ws, cells.get("rolling_range"), {
        **{column: ACCOUNTING_FORMAT for column in range(2, 14)},
    })
    _format_rows(ws, cells.get("opening_range"), {
        4: UNIT_COST_FORMAT, 5: UNIT_COST_FORMAT, 6: UNIT_COST_FORMAT,
    })
    ws.freeze_panes = "A5"
    _finish_user_sheet(ws)


def _polish_bridge(ws, cells: dict[str, Any]) -> None:
    ws["A2"] = (
        "상세 근거 시트의 최종 수식 셀을 직접 참조해 공식 효과, Residual, OP 증감을 검증합니다. "
        "중복 제거 전 시차와 핵심 제조원가 중복분은 별도 합산하지 않습니다."
    )
    ws.merge_cells("A2:G2")
    _set_widths(ws, {"A": 18, "B": 28, "C": 20, "D": 18, "E": 18, "F": 14, "G": 24})
    _hide_columns(ws, ("A", "G"))
    _format_rows(ws, cells.get("effect_range"), {
        3: ACCOUNTING_FORMAT, 4: ACCOUNTING_FORMAT, 5: ACCOUNTING_FORMAT,
    })
    _format_rows(ws, cells.get("summary_range"), {
        3: ACCOUNTING_FORMAT, 4: ACCOUNTING_FORMAT,
    })
    ws.freeze_panes = "B5"
    _finish_user_sheet(ws, landscape=False)


def polish_evidence_workbook(
    workbook,
    *,
    sales_cells: dict[str, Any],
    material_cells: dict[str, Any],
    manufacturing_cells: dict[str, Any],
    inventory_cells: dict[str, Any],
    bridge_cells: dict[str, Any],
) -> dict[str, float | int]:
    """Apply presentation-only polish without moving or replacing calculations."""
    key_sheets = {
        "판매효과_근거": lambda ws: _polish_sales(ws, sales_cells),
        "원부재료_근거": lambda ws: _polish_material(ws, material_cells),
        "제조경비_근거": lambda ws: _polish_manufacturing(ws, manufacturing_cells),
        "재고원가반영시차_근거": lambda ws: _polish_inventory(ws, inventory_cells),
        "최종Bridge_검증": lambda ws: _polish_bridge(ws, bridge_cells),
    }
    changed = 0
    for sheet_name, formatter in key_sheets.items():
        ws = workbook[sheet_name]
        changed += _localize_user_labels(ws)
        formatter(ws)
    return {
        "localized_cells": changed,
        "visible_width_total": sum(
            float(ws.column_dimensions[get_column_letter(column)].width or 0)
            for sheet_name in key_sheets
            for ws in (workbook[sheet_name],)
            for column in range(1, ws.max_column + 1)
            if not ws.column_dimensions[get_column_letter(column)].hidden
        ),
    }
